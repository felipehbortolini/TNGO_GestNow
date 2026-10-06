"""Reading an uploaded ``.xlsx``: the bytes of the file become rows of raw cells (D12, ISSUE-018).

The browser never reads the spreadsheet (the SheetJS of the prototype is gone with
its known risk): the file travels to the server, and here it is checked before any
library opens it and then read with openpyxl in its streaming mode.

What is refused, always with ``InvalidDataError`` (422) and a message the person can act on:

* no file, or an empty one;
* a file that is not an ``.xlsx``: another extension, or the right extension over bytes that
  are not an Excel workbook (a PDF renamed, a text file, a corrupted file);
* a file over ``MAX_FILE_BYTES``, or one that would expand past ``MAX_EXPANDED_BYTES`` (a
  compressed bomb is stopped by its declared sizes, before it is opened).

Only the first sheet is read, up to ``MAX_COLUMNS`` columns and ``MAX_SHEET_ROWS`` rows: a
sheet with more data than the limit is cut there, and ``importing`` refuses it as too long.
Formulas read as the value Excel stored last, dates as dates, numbers as numbers, and each
cell keeps its number format (a percent cell holds a fraction: ``0.453`` shown as ``45,3%``).
"""

from __future__ import annotations

import zipfile
import zlib
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from src.core import attachments
from src.core.errors import InvalidDataError

XLSX_EXTENSION = ".xlsx"
WORKBOOK_PART = "xl/workbook.xml"

MEGABYTE = 1024 * 1024
MAX_FILE_BYTES = 5 * MEGABYTE
MAX_EXPANDED_BYTES = 50 * MEGABYTE

# The rows of the sheet that may hold the identification of the model (logo, title, scope, date)
# above the header, and the rows of data after it.
HEADER_SCAN_ROWS = 30
MAX_DATA_ROWS = 5000
MAX_SHEET_ROWS = HEADER_SCAN_ROWS + MAX_DATA_ROWS + 1
MAX_COLUMNS = 60

GENERAL_FORMAT = "General"

NO_FILE_MESSAGE = "Escolha a planilha (.xlsx) antes de conferir."
NOT_A_SPREADSHEET_MESSAGE = (
    "O arquivo «{name}» não é uma planilha do Excel (.xlsx). "
    "Baixe o modelo, preencha e envie em .xlsx."
)
UNREADABLE_MESSAGE = (
    "Não foi possível ler «{name}»: o arquivo está corrompido ou não é uma planilha .xlsx."
)
TOO_BIG_MESSAGE = "O arquivo passa de 5 MB. Divida a importação em planilhas menores."

# What a damaged or hostile workbook makes openpyxl and the zip layer raise: all of them mean
# "this is not a spreadsheet I can read" and none of them is a mistake of the product.
_UNREADABLE = (
    zipfile.BadZipFile,
    InvalidFileException,
    zlib.error,
    EOFError,
    NotImplementedError,
    RuntimeError,
    OSError,
    SyntaxError,
    LookupError,
    ValueError,
    TypeError,
    AttributeError,
)


@dataclass(frozen=True)
class RawCell:
    """What the sheet holds in a cell: the value as openpyxl reads it and its number format."""

    value: Any = None
    number_format: str = GENERAL_FORMAT


@dataclass(frozen=True)
class SheetRow:
    """A row of the sheet: its number (the one Excel shows) and its cells, trailing blanks cut."""

    number: int
    cells: tuple[RawCell, ...]


def read_rows(file_name: str | None, content: bytes) -> tuple[SheetRow, ...]:
    """The rows of the first sheet of the workbook, or the 422 that says why it cannot be read."""
    name = attachments.clean_name(file_name)
    _check_file(name, content)
    try:
        rows = _read_first_sheet(content)
    except _UNREADABLE:
        raise InvalidDataError(UNREADABLE_MESSAGE.format(name=name)) from None
    return rows


def _check_file(name: str, content: bytes) -> None:
    """Refuse what is not an ``.xlsx`` before anything opens it."""
    if not name or not content:
        raise InvalidDataError(NO_FILE_MESSAGE)
    if PurePosixPath(name.lower()).suffix != XLSX_EXTENSION:
        raise InvalidDataError(NOT_A_SPREADSHEET_MESSAGE.format(name=name))
    if len(content) > MAX_FILE_BYTES:
        raise InvalidDataError(TOO_BIG_MESSAGE)
    _check_archive(name, content)


def _check_archive(name: str, content: bytes) -> None:
    """An ``.xlsx`` is a zip with the workbook part; its declared expanded size must stay small."""
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            parts = archive.infolist()
    except zipfile.BadZipFile:
        raise InvalidDataError(NOT_A_SPREADSHEET_MESSAGE.format(name=name)) from None
    if WORKBOOK_PART not in {part.filename for part in parts}:
        raise InvalidDataError(NOT_A_SPREADSHEET_MESSAGE.format(name=name))
    if sum(part.file_size for part in parts) > MAX_EXPANDED_BYTES:
        raise InvalidDataError(TOO_BIG_MESSAGE)


def _read_first_sheet(content: bytes) -> tuple[SheetRow, ...]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet: Any = workbook.worksheets[0]
        rows = sheet.iter_rows(min_row=1, max_row=MAX_SHEET_ROWS, max_col=MAX_COLUMNS)
        return tuple(_sheet_row(number, cells) for number, cells in enumerate(rows, start=1))
    finally:
        workbook.close()


def _sheet_row(number: int, cells: tuple[Any, ...]) -> SheetRow:
    raw = [RawCell(cell.value, cell.number_format or GENERAL_FORMAT) for cell in cells]
    while raw and raw[-1].value is None:
        raw.pop()
    return SheetRow(number=number, cells=tuple(raw))
