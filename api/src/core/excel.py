"""Excel on the server (D12): the one builder every screen uses.

``build_workbook`` turns a ``Document`` (``core.export_document``) into the bytes
of an ``.xlsx`` and ``excel_response`` wraps it as the download a route returns.
The workbook is made with openpyxl; the PDF has nothing to do with it (it is the
browser's print, ``core.printable``).

What the file has:

* The identification on every sheet: the logo, the title, the scope (Portfólio or
  project), the generation date and the context lines of the document.
* A ``Resumo`` sheet, when the document has KPIs: one row per indicator with its
  value, its management reference and, when it has one, its situation in words.
* One sheet per table: the table with the filter on, the header frozen and
  repeated on every printed page, numbers and dates as real cells (so they sort,
  filter and add up), money in full reais and the closing row, if any, below the
  filter. In the Portfólio the first column is ``Projeto``.
* The colors (tones, header, borders, title) come from the tokens of the Design
  System (``core.design_tokens``); no color is written in this module.
* Print setup: landscape, one page wide, A4 or the paper of the document.

Text goes in as text: a cell that starts with ``=`` is never a formula, and the
characters a spreadsheet refuses are dropped.
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from functools import cache
from io import BytesIO
from typing import cast

import azure.functions as func
from openpyxl import Workbook
from openpyxl.cell.cell import Cell as SheetCell
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from src.core import design_tokens
from src.core.export_document import (
    GENERATED_HEADING,
    PRODUCT_NAME,
    SCOPE_HEADING,
    Cell,
    CellValue,
    Column,
    Document,
    Kpi,
    Paper,
    Table,
    Tone,
    ValueKind,
    cents_of,
    date_of,
    format_value,
    number_of,
    resolve_tables,
)
from src.core.responses import file_response

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SUMMARY_SHEET = "Resumo"
KPI_HEADER = "Indicador"
VALUE_HEADER = "Valor"
REFERENCE_HEADER = "Referência"
STATUS_HEADER = "Situação"

# The logo (167 x 163 pixels) shown 36 pixels high in the first row of every sheet.
_LOGO_WIDTH = 37
_LOGO_HEIGHT = 36
_LOGO_ROW_HEIGHT = 32  # points
_FIRST_TEXT_ROW = 2

_MAX_SHEET_NAME = 31
_MIN_WIDTH = 10
_MAX_WIDTH = 60
_WIDTH_PADDING = 4  # room for the arrow of the filter
_WIDTH_SAMPLE = 500  # rows read to fit the width of a column
_SUMMARY_WIDTHS = (38, 20, 52, 22)

# ``paperSize`` of the SpreadsheetML: 8 is A3 and 9 is A4.
_PAPER_CODES = {Paper.A4: 9, Paper.A3: 8}

_MONEY_FORMAT = '"R$" #,##0.00'
_DATE_FORMAT = "DD/MM/YYYY"

_INVALID_SHEET_CHARACTERS = re.compile(r"[\\/?*\[\]:]")
# The control characters a spreadsheet refuses; tab, line feed and carriage return stay.
_ILLEGAL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_MAX_TEXT = 32767
_NOT_A_FILE_NAME_CHARACTER = re.compile(r"[^a-z0-9]+")

# The tone as a pill of the Design System draws it (``.pill--ok``...): background and text tokens.
_TONE_TOKENS: Mapping[Tone, tuple[str, str]] = {
    Tone.OK: ("--ok-50", "--ok-800"),
    Tone.WARN: ("--warn-50", "--warn-800"),
    Tone.ERROR: ("--erro-50", "--erro-700"),
    Tone.INFO: ("--azul-50", "--azul-800"),
    Tone.NEUTRAL: ("--frio-50", "--frio-700"),
}

_HORIZONTAL: Mapping[ValueKind, str] = {
    ValueKind.TEXT: "left",
    ValueKind.INTEGER: "right",
    ValueKind.DECIMAL: "right",
    ValueKind.PERCENT: "right",
    ValueKind.MONEY: "right",
    ValueKind.DATE: "center",
}

_TEXT_COLUMN = Column("")


@dataclass(frozen=True)
class _Styles:
    """The openpyxl styles of one workbook, built once from the tokens and shared by every cell."""

    body: Font
    bold: Font
    title: Font
    muted: Font
    header_font: Font
    header_fill: PatternFill
    total_fill: PatternFill
    line: Border
    header_line: Border
    total_line: Border
    tone_fonts: Mapping[Tone, Font]
    tone_fills: Mapping[Tone, PatternFill]
    alignments: Mapping[ValueKind, Alignment]


def build_workbook(document: Document) -> bytes:
    """The ``.xlsx`` of the document: the ``Resumo`` sheet (KPIs) and one sheet per table."""
    styles = _build_styles()
    tables = resolve_tables(document)
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    taken: set[str] = set()
    if document.kpis or not tables:
        _write_summary(_new_sheet(workbook, SUMMARY_SHEET, taken), document, styles)
    for table in tables:
        _write_table_sheet(_new_sheet(workbook, table.title, taken), document, table, styles)
    workbook.properties.title = document.title
    workbook.properties.creator = PRODUCT_NAME
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def file_name(document: Document) -> str:
    """``mapa-de-controle-2026-10-05.xlsx``: the title without accents, in lower case, and the date."""
    ascii_title = unicodedata.normalize("NFKD", document.title).encode("ascii", "ignore")
    slug = _NOT_A_FILE_NAME_CHARACTER.sub("-", ascii_title.decode("ascii").lower()).strip("-")
    return f"{slug or 'exportacao'}-{document.generated_on.isoformat()}.xlsx"


def excel_response(document: Document) -> func.HttpResponse:
    """The download of the document: the file, not a fragment (``responses.file_response``)."""
    return file_response(
        build_workbook(document),
        filename=file_name(document),
        content_type=XLSX_CONTENT_TYPE,
    )


# ── Styles ───────────────────────────────────────────────────────────────


def _build_styles() -> _Styles:
    family = design_tokens.font_name()
    color = design_tokens.color

    def font(token: str, *, size: int = 10, bold: bool = False) -> Font:
        return Font(name=family, size=size, bold=bold, color=color(token))

    def fill(token: str) -> PatternFill:
        return PatternFill(fill_type="solid", start_color=color(token), end_color=color(token))

    return _Styles(
        body=font("--text-primary"),
        bold=font("--text-primary", bold=True),
        title=font("--brand-title", size=16, bold=True),
        muted=font("--text-secondary"),
        header_font=font("--text-secondary", bold=True),
        header_fill=fill("--neutro-50"),
        total_fill=fill("--neutro-50"),
        line=Border(bottom=Side(style="thin", color=color("--neutro-100"))),
        header_line=Border(bottom=Side(style="thin", color=color("--borda"))),
        total_line=Border(top=Side(style="thin", color=color("--borda"))),
        tone_fonts={tone: font(text, bold=True) for tone, (_, text) in _TONE_TOKENS.items()},
        tone_fills={tone: fill(background) for tone, (background, _) in _TONE_TOKENS.items()},
        alignments={
            kind: Alignment(horizontal=horizontal, vertical="top", wrap_text=kind is ValueKind.TEXT)
            for kind, horizontal in _HORIZONTAL.items()
        },
    )


# ── Sheets ───────────────────────────────────────────────────────────────


def _new_sheet(workbook: Workbook, title: str, taken: set[str]) -> Worksheet:
    name = _sheet_name(title, taken)
    taken.add(name.casefold())
    sheet = cast("Worksheet", workbook.create_sheet(name))
    sheet.sheet_view.showGridLines = False
    return sheet


def _sheet_name(title: str, taken: set[str]) -> str:
    """A name Excel accepts (31 characters, none of ``\\ / ? * [ ] :``) and no sheet has yet."""
    base = _INVALID_SHEET_CHARACTERS.sub(" ", title).strip().strip("'")[:_MAX_SHEET_NAME]
    base = base or "Planilha"
    name, number = base, 1
    while name.casefold() in taken:
        number += 1
        suffix = f" ({number})"
        name = f"{base[: _MAX_SHEET_NAME - len(suffix)]}{suffix}"
    return name


def _write_identification(sheet: Worksheet, document: Document, styles: _Styles) -> int:
    """The logo, the title, the scope, the date and the context; returns the first free row."""
    logo = Image(BytesIO(_logo_bytes()))
    logo.width = _LOGO_WIDTH
    logo.height = _LOGO_HEIGHT
    sheet.add_image(logo, "A1")
    sheet.row_dimensions[1].height = _LOGO_ROW_HEIGHT
    generated = format_value(ValueKind.DATE, document.generated_on)
    lines = [
        (document.title, styles.title),
        (f"{SCOPE_HEADING}: {document.scope_label}", styles.muted),
        (f"{GENERATED_HEADING}: {generated}", styles.muted),
        *((f"{label}: {value}", styles.muted) for label, value in document.context),
    ]
    for row, (text, font) in enumerate(lines, start=_FIRST_TEXT_ROW):
        _put_text(sheet.cell(row=row, column=1), text, font)
    return _FIRST_TEXT_ROW + len(lines) + 1


def _write_summary(sheet: Worksheet, document: Document, styles: _Styles) -> None:
    header_row = _write_identification(sheet, document, styles)
    with_status = any(kpi.status for kpi in document.kpis)
    headers = [
        (KPI_HEADER, ValueKind.TEXT),
        (VALUE_HEADER, ValueKind.DECIMAL),
        (REFERENCE_HEADER, ValueKind.TEXT),
        *([(STATUS_HEADER, ValueKind.TEXT)] if with_status else []),
    ]
    _write_header(sheet, header_row, headers, styles)
    for offset, kpi in enumerate(document.kpis, start=1):
        _write_kpi(sheet, header_row + offset, kpi, styles, with_status=with_status)
    for index, width in enumerate(_SUMMARY_WIDTHS[: len(headers)], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    _setup_printing(sheet, document, header_row=None)


def _write_kpi(sheet: Worksheet, row: int, kpi: Kpi, styles: _Styles, *, with_status: bool) -> None:
    label = sheet.cell(row=row, column=1)
    _write_cell(label, _TEXT_COLUMN, Cell(kpi.label), styles)
    label.font = styles.bold
    value_column = Column(kpi.label, kpi.kind, kpi.digits)
    _write_cell(sheet.cell(row=row, column=2), value_column, Cell(kpi.value, kpi.tone), styles)
    _write_cell(sheet.cell(row=row, column=3), _TEXT_COLUMN, Cell(kpi.reference), styles)
    if with_status:
        _write_cell(sheet.cell(row=row, column=4), _TEXT_COLUMN, Cell(kpi.status, kpi.tone), styles)


def _write_table_sheet(sheet: Worksheet, document: Document, table: Table, styles: _Styles) -> None:
    first_row = _write_identification(sheet, document, styles)
    _put_text(sheet.cell(row=first_row, column=1), table.title, styles.bold)
    header_row = first_row + 1
    _write_header(sheet, header_row, [(c.header, c.kind) for c in table.columns], styles)
    for offset, item in enumerate(table.rows, start=1):
        _write_row(sheet, header_row + offset, table.columns, item.cells, styles)
    last_row = header_row + len(table.rows)
    if table.totals is not None:
        _write_totals(sheet, last_row + 1, table, styles)
    last_column = get_column_letter(len(table.columns))
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{last_row}"
    sheet.freeze_panes = f"A{header_row + 1}"
    for index, column in enumerate(table.columns, start=1):
        width = column.width or _fitted_width(table, index - 1)
        sheet.column_dimensions[get_column_letter(index)].width = width
    _setup_printing(sheet, document, header_row=header_row)


def _write_header(
    sheet: Worksheet, row: int, headers: Sequence[tuple[str, ValueKind]], styles: _Styles
) -> None:
    for index, (text, kind) in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=index)
        _put_text(cell, text, styles.header_font)
        cell.fill = styles.header_fill
        cell.border = styles.header_line
        cell.alignment = styles.alignments[kind]


def _write_row(
    sheet: Worksheet, row: int, columns: Sequence[Column], cells: Sequence[Cell], styles: _Styles
) -> None:
    for index, (column, cell) in enumerate(zip(columns, cells, strict=True), start=1):
        _write_cell(sheet.cell(row=row, column=index), column, cell, styles)


def _write_totals(sheet: Worksheet, row: int, table: Table, styles: _Styles) -> None:
    """The closing row: bold on a light band, with a rule above, unless the cell has its own tone."""
    cells = table.totals or ()
    _write_row(sheet, row, table.columns, cells, styles)
    for index, cell in enumerate(cells, start=1):
        target = sheet.cell(row=row, column=index)
        target.border = styles.total_line
        if cell.tone is None:
            target.font = styles.bold
            target.fill = styles.total_fill


def _write_cell(target: SheetCell, column: Column, cell: Cell, styles: _Styles) -> None:
    """One value in its format (a real number or date for the kinds that have one) and its tone."""
    font = styles.body if cell.tone is None else styles.tone_fonts[cell.tone]
    if column.kind is ValueKind.TEXT:
        _put_text(target, "" if cell.value is None else str(cell.value), font)
    else:
        target.value = _excel_value(column.kind, cell.value)
        target.font = font
        target.number_format = _number_format(column.kind, column.digits)
    target.alignment = styles.alignments[column.kind]
    target.border = styles.line
    if cell.tone is not None:
        target.fill = styles.tone_fills[cell.tone]


def _put_text(cell: SheetCell, text: str, font: Font) -> None:
    """Text as text: ``data_type`` stays ``s`` even when it starts with ``=``."""
    cell.value = _ILLEGAL_CHARACTERS.sub("", text)[:_MAX_TEXT]
    cell.data_type = "s"
    cell.font = font


def _excel_value(kind: ValueKind, value: CellValue) -> int | float | date | None:
    """What the cell stores: percent as a fraction, money as reais, a date as a date."""
    if value is None:
        return None
    if kind is ValueKind.DATE:
        return date_of(value)
    if kind is ValueKind.INTEGER:
        return int(number_of(value).to_integral_value(rounding=ROUND_HALF_UP))
    if kind is ValueKind.MONEY:
        return float(Decimal(cents_of(value)) / 100)
    if kind is ValueKind.PERCENT:
        return float(number_of(value) / 100)
    return float(number_of(value))


def _number_format(kind: ValueKind, digits: int) -> str:
    places = f".{'0' * digits}" if digits > 0 else ""
    formats = {
        ValueKind.INTEGER: "#,##0",
        ValueKind.DECIMAL: f"#,##0{places}",
        ValueKind.PERCENT: f"0{places}%",
        ValueKind.MONEY: _MONEY_FORMAT,
        ValueKind.DATE: _DATE_FORMAT,
    }
    return formats.get(kind, "General")


def _fitted_width(table: Table, position: int) -> int:
    """The width in characters that shows the header and the longest value read of the column."""
    column = table.columns[position]
    cells = [item.cells[position] for item in table.rows[:_WIDTH_SAMPLE]]
    if table.totals is not None:
        cells.append(table.totals[position])
    lengths = [len(column.header)]
    lengths.extend(len(format_value(column.kind, cell.value, column.digits)) for cell in cells)
    return max(_MIN_WIDTH, min(max(lengths) + _WIDTH_PADDING, _MAX_WIDTH))


def _setup_printing(sheet: Worksheet, document: Document, *, header_row: int | None) -> None:
    """Landscape, one page wide, the paper of the document and the table header on every page."""
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = _PAPER_CODES[document.paper]
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if header_row is not None:
        sheet.print_title_rows = f"{header_row}:{header_row}"


@cache
def _logo_bytes() -> bytes:
    return design_tokens.LOGO_FILE.read_bytes()
