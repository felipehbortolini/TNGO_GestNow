"""Rotas do Painel HSE (ISSUE-075, D11, D12, D14): o fragmento, a query e as duas exportações.

O handler é chamado com a requisição de Alpine, o cookie da demonstração identificando a pessoa e
a unidade de trabalho do teste (fixture ``sessao``). O fragmento traz os três blocos de KPIs, os
gráficos como dados (``data-grafico`` + ``data-dados``, sem ``<script>``) e o marcador da
situação; a query escolhe o mês e a referência da pirâmide; o Excel e o imprimível saem pelos
mecanismos genéricos, com a query da tela; no Portfólio as listas são agregadas, sem a coluna
``Projeto``.
"""

from __future__ import annotations

import json
from collections.abc import Mapping
from datetime import date
from decimal import Decimal
from io import BytesIO

import azure.functions as func
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from src.core.auth import DEMO_COOKIE
from src.core.rbac import User
from src.modulos.hse import painel_routes as routes
from src.modulos.hse import service
from tests.apoio_hse import REFERENCIA, Cenario, fechamento_de, horas_de, ocorrencia_de
from tests.html_tags import Tag, find_by_id, parse_tags

ALVO = "painel-conteudo"
JULHO = date(2026, 7, 1)
AGOSTO = date(2026, 8, 1)


def _requisicao(
    usuario: User,
    *,
    params: Mapping[str, str] | None = None,
    alvo: str | None = ALVO,
) -> func.HttpRequest:
    headers = {"Cookie": f"{DEMO_COOKIE}={usuario.id}"}
    if alvo is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = alvo
    return func.HttpRequest(
        method="GET",
        url="/api/hse/painel",
        headers=headers,
        params=dict(params or {}),
        route_params={},
        body=b"",
    )


def _escopo(cenario: Cenario) -> dict[str, str]:
    return {"projeto": str(cenario.projeto.id)}


def _com_dados(sessao: Session, cenario: Cenario) -> None:
    """Dois meses de HHT, uma lesão registrável em julho e um quase acidente em agosto."""
    service.save_hours(
        sessao,
        user=cenario.maria,
        data=horas_de(cenario, month=JULHO, hours=Decimal("100000")),
        reference_date=REFERENCIA,
    )
    service.save_hours(
        sessao,
        user=cenario.maria,
        data=horas_de(cenario, month=AGOSTO, hours=Decimal("100000")),
        reference_date=REFERENCIA,
    )
    service.save_occurrence(
        sessao,
        user=cenario.maria,
        data=ocorrencia_de(
            cenario, kind="Trabalho restrito", role="Operador de moinho", happened_on=JULHO
        ),
        reference_date=REFERENCIA,
    )
    service.save_occurrence(
        sessao,
        user=cenario.maria,
        data=ocorrencia_de(
            cenario, kind="Quase acidente", role="Operador de moinho", happened_on=AGOSTO
        ),
        reference_date=REFERENCIA,
    )
    service.save_closing(
        sessao,
        user=cenario.maria,
        data=fechamento_de(cenario, month=JULHO, deviations=10),
        reference_date=REFERENCIA,
    )


def _corpo(resposta: func.HttpResponse) -> str:
    return resposta.get_body().decode()


def _tags(resposta: func.HttpResponse) -> list[Tag]:
    return parse_tags(_corpo(resposta))


def _dados_do_grafico(resposta: func.HttpResponse, tipo: str) -> dict[str, object]:
    (host,) = [tag for tag in _tags(resposta) if tag.attrs.get("data-grafico") == tipo]
    return json.loads(host.attrs["data-dados"])


# ── O fragmento ──────────────────────────────────────────────────────────────────────────────


def test_o_fragmento_traz_os_blocos_os_graficos_e_a_linha_de_periodo(
    sessao: Session, cenario: Cenario
) -> None:
    _com_dados(sessao, cenario)

    resposta = routes.hse_panel_view(_requisicao(cenario.maria, params=_escopo(cenario)))
    corpo = _corpo(resposta)

    assert resposta.status_code == 200
    assert find_by_id(_tags(resposta), ALVO) is not None
    assert 'data-situacao="conteudo"' in corpo
    assert "No mês selecionado" in corpo
    assert "Acumulado" in corpo
    assert "Indicadores proativos" in corpo
    assert (
        "Mês: Ago/2026 · Acumulado: Jul/2026 a Ago/2026 · base 1.000.000 HHT (NBR 14280)" in corpo
    )
    assert "Taxa de frequência (TF)" in corpo
    assert "Relato de quase acidentes" in corpo
    assert "<script" not in corpo

    piramide = _dados_do_grafico(resposta, "piramide-seguranca")
    assert [item["titulo"] for item in piramide["piramides"]] == [
        "Mês: Ago/2026",
        "Acumulado: Jul/2026 a Ago/2026",
    ]
    assert piramide["piramides"][1]["valores"] == {
        "grave": 0,
        "leve": 1,
        "dano": 0,
        "quase": 1,
        "desvio": 10,
    }
    assert piramide["referencia"]["nome"] == "Bird"

    evolucao = _dados_do_grafico(resposta, "linhas-multiplas")
    assert [periodo["valores"] for periodo in evolucao["periodos"]] == [
        {"tf": 0.0, "trif": 10.0},
        {"tf": 0.0, "trif": 0.0},
    ]
    assert [serie["rotulo"] for serie in evolucao["series"]] == [
        "TF (com afastamento)",
        "TRIF (registráveis)",
    ]


def test_a_query_escolhe_o_mes_e_a_referencia_da_piramide(
    sessao: Session, cenario: Cenario
) -> None:
    _com_dados(sessao, cenario)

    resposta = routes.hse_panel_view(
        _requisicao(
            cenario.maria,
            params={**_escopo(cenario), "mes": "2026-07", "referencia": "heinrich"},
        )
    )
    corpo = _corpo(resposta)

    assert resposta.status_code == 200
    assert "Mês: Jul/2026 · Acumulado: Jul/2026 a Jul/2026" in corpo
    piramide = _dados_do_grafico(resposta, "piramide-seguranca")
    assert piramide["referencia"]["nome"] == "Heinrich"
    assert piramide["referencia"]["nota"] == "níveis 3 e 4 somados"


def test_sem_hht_o_fragmento_traz_o_vazio_de_origem(cenario: Cenario) -> None:
    resposta = routes.hse_panel_view(_requisicao(cenario.maria, params=_escopo(cenario)))
    corpo = _corpo(resposta)

    assert resposta.status_code == 200
    assert 'data-situacao="vazio-origem"' in corpo
    assert "Ainda não há HHT registrado" in corpo
    assert "Taxa de frequência (TF)" not in corpo


def test_sem_o_cabecalho_do_alpine_o_fragmento_manda_para_o_shell(cenario: Cenario) -> None:
    resposta = routes.hse_panel_view(_requisicao(cenario.maria, alvo=None))

    assert resposta.status_code == 302


# ── As exportações ───────────────────────────────────────────────────────────────────────────


def _sheets(resposta: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(resposta.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> tuple[int, list[object]]:
    first_row = int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    width = sheet.max_column
    return first_row, [
        sheet.cell(row=first_row, column=column).value for column in range(1, width + 1)
    ]


def test_o_excel_traz_o_resumo_as_tabelas_e_as_listas(sessao: Session, cenario: Cenario) -> None:
    _com_dados(sessao, cenario)

    resposta = routes.hse_panel_excel(
        _requisicao(cenario.maria, alvo=None, params=_escopo(cenario))
    )

    assert resposta.status_code == 200
    sheets = _sheets(resposta)
    assert "Resumo" in sheets
    assert any(name.startswith("Pirâmide de segurança · acumula") for name in sheets)
    assert "Evolução mensal de TF e TRIF" in sheets
    assert "Ocorrências por empresa" in sheets

    resumo = sheets["Resumo"]
    linhas = {
        resumo.cell(row=row, column=1).value: resumo.cell(row=row, column=2).value
        for row in range(1, resumo.max_row + 1)
    }
    assert linhas["Acumulado · Lesões registráveis (TRIF)"] == 5
    assert linhas["Acumulado · Dias sem afastamento"] == 0

    acumulado = next(
        sheet
        for name, sheet in sheets.items()
        if name.startswith("Pirâmide de segurança · acumula")
    )
    primeira_row, header = _header_row(acumulado)
    assert header == ["Nível", "Ocorrências"]
    assert acumulado.cell(row=primeira_row + 1, column=1).value == "1. Lesão grave"
    assert acumulado.cell(row=primeira_row + 2, column=2).value == 1

    area_row, area_header = _header_row(sheets["Ocorrências por área"])
    assert area_header == ["Área", "Ocorrências"]
    assert sheets["Ocorrências por área"].cell(row=area_row + 1, column=1).value == "Moagem 210"


def test_a_versao_imprimivel_traz_os_graficos_e_o_mesmo_conteudo(
    sessao: Session, cenario: Cenario
) -> None:
    _com_dados(sessao, cenario)

    resposta = routes.hse_panel_printable(
        _requisicao(cenario.maria, alvo="folha-impressao", params=_escopo(cenario))
    )
    corpo = _corpo(resposta)

    assert resposta.status_code == 200
    assert "Painel HSE" in corpo
    assert "Pirâmide de segurança" in corpo
    assert 'data-grafico="piramide-seguranca"' in corpo
    assert 'data-grafico="linhas-multiplas"' in corpo
    assert "Ago/2026" in corpo
    assert "<script" not in corpo


def test_no_portfolio_as_listas_nao_levam_a_coluna_projeto(
    sessao: Session, cenario: Cenario
) -> None:
    _com_dados(sessao, cenario)

    resposta = routes.hse_panel_excel(_requisicao(cenario.maria, alvo=None))

    assert resposta.status_code == 200
    sheets = _sheets(resposta)
    _, empresa_header = _header_row(sheets["Ocorrências por empresa"])
    _, area_header = _header_row(sheets["Ocorrências por área"])
    assert empresa_header == ["Empresa", "Ocorrências"]
    assert area_header == ["Área", "Ocorrências"]

    fragmento = routes.hse_panel_view(_requisicao(cenario.maria))
    assert fragmento.status_code == 200
    assert "Ocorrências por empresa" in _corpo(fragmento)
