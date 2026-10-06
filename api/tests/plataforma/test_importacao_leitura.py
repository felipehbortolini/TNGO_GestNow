"""Importação (ISSUE-018): a leitura do arquivo enviado e a recusa do que não é planilha.

``core.spreadsheet_reader`` confere o arquivo antes de a biblioteca abri-lo e só então o lê: tudo o
que não é um ``.xlsx`` legível é 422 com a mensagem que a pessoa entende (outra extensão, bytes que
não são do Excel, arquivo corrompido, acima de 5 MB, que expandiria demais). A leitura devolve as
linhas como o Excel as numera, com o valor e o formato de número de cada célula.
"""

from __future__ import annotations

import zipfile
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from src.core.errors import InvalidDataError
from src.core.spreadsheet_reader import (
    MAX_EXPANDED_BYTES,
    MAX_FILE_BYTES,
    MAX_SHEET_ROWS,
    WORKBOOK_PART,
    RawCell,
    SheetRow,
    read_rows,
)
from tests.importadores_de_teste import planilha

NAO_E_PLANILHA = "não é uma planilha do Excel (.xlsx)"
NAO_FOI_POSSIVEL_LER = "Não foi possível ler"


def _recusa(nome: str | None, conteudo: bytes) -> str:
    with pytest.raises(InvalidDataError) as erro:
        read_rows(nome, conteudo)
    return str(erro.value)


def _zip(partes: dict[str, bytes]) -> bytes:
    saida = BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as arquivo:
        for nome, conteudo in partes.items():
            arquivo.writestr(nome, conteudo)
    return saida.getvalue()


# ── O que não é planilha é 422 ───────────────────────────────────────────


@pytest.mark.parametrize(
    "nome", ["lista.pdf", "lista.xls", "lista.csv", "lista.xlsm", "lista.xlsx.exe", "lista"]
)
def test_outra_extensao_e_422_mesmo_com_bytes_de_planilha(nome: str) -> None:
    mensagem = _recusa(nome, planilha([("A", None, None)]))

    assert NAO_E_PLANILHA in mensagem
    assert nome in mensagem


def test_a_extensao_vale_em_maiuscula() -> None:
    linhas = read_rows("LISTA.XLSX", planilha([("A", None, None)]))

    assert [linha.number for linha in linhas] == [1, 2]


def test_extensao_xlsx_sobre_texto_pdf_ou_zip_qualquer_e_422() -> None:
    assert NAO_E_PLANILHA in _recusa("lista.xlsx", b"isto nao e uma planilha")
    assert NAO_E_PLANILHA in _recusa("lista.xlsx", b"%PDF-1.7 conteudo")
    assert NAO_E_PLANILHA in _recusa("lista.xlsx", _zip({"a.txt": b"qualquer coisa"}))


def test_planilha_cortada_ao_meio_e_422() -> None:
    inteira = planilha([("A", None, None)])

    assert NAO_E_PLANILHA in _recusa("lista.xlsx", inteira[: len(inteira) // 2])


def test_zip_com_a_parte_do_livro_mas_sem_o_resto_nao_pode_ser_lido() -> None:
    mensagem = _recusa("lista.xlsx", _zip({WORKBOOK_PART: b"<workbook/>"}))

    assert NAO_FOI_POSSIVEL_LER in mensagem
    assert "lista.xlsx" in mensagem


@pytest.mark.parametrize("nome", [None, "", "   "])
def test_sem_arquivo_e_422_pedindo_a_planilha(nome: str | None) -> None:
    assert (
        _recusa(nome, planilha([("A", None, None)]))
        == "Escolha a planilha (.xlsx) antes de conferir."
    )


def test_arquivo_vazio_e_422_pedindo_a_planilha() -> None:
    assert _recusa("lista.xlsx", b"") == "Escolha a planilha (.xlsx) antes de conferir."


def test_um_byte_acima_de_5_mb_e_422_de_tamanho_e_o_limite_exato_ja_e_conferido_como_planilha() -> (
    None
):
    assert _recusa("lista.xlsx", b"x" * (MAX_FILE_BYTES + 1)) == (
        "O arquivo passa de 5 MB. Divida a importação em planilhas menores."
    )
    assert NAO_E_PLANILHA in _recusa("lista.xlsx", b"x" * MAX_FILE_BYTES)


def test_arquivo_que_expandiria_alem_do_limite_e_recusado_sem_ser_aberto() -> None:
    bomba = _zip({WORKBOOK_PART: b"<workbook/>", "xl/dados.bin": b"\0" * (MAX_EXPANDED_BYTES + 1)})

    assert len(bomba) < MAX_FILE_BYTES
    assert _recusa("lista.xlsx", bomba) == (
        "O arquivo passa de 5 MB. Divida a importação em planilhas menores."
    )


def test_o_nome_da_mensagem_perde_a_pasta_e_os_caracteres_de_controle() -> None:
    mensagem = _recusa("C:\\Users\\ana\\Downloads\\lista\r\n.pdf", b"conteudo")

    assert "«lista.pdf»" in mensagem


# ── O que a leitura devolve ──────────────────────────────────────────────


def test_as_linhas_vem_numeradas_como_no_excel_com_valor_e_formato() -> None:
    conteudo = planilha(
        [
            ("Alfa", 0.453, datetime.fromisoformat("2026-10-05T00:00:00")),
            (None, None, None),
            ("Gama", "texto", 7),
        ],
        cabecalho=("Nome", "Avanço", "Início"),
        formatos={"Avanço": "0.0%"},
    )

    linhas = read_rows("lista.xlsx", conteudo)

    assert [linha.number for linha in linhas] == [1, 2, 3, 4]
    assert linhas[0] == SheetRow(1, (RawCell("Nome"), RawCell("Avanço"), RawCell("Início")))
    assert linhas[1].cells[0] == RawCell("Alfa")
    assert linhas[1].cells[1] == RawCell(0.453, "0.0%")
    assert linhas[1].cells[2].value == datetime.fromisoformat("2026-10-05T00:00:00")
    assert linhas[2] == SheetRow(3, ())
    assert linhas[3].cells[1].value == "texto"
    assert linhas[3].cells[2] == RawCell(7)


def test_so_a_primeira_aba_e_lida() -> None:
    livro = Workbook()
    livro.active.title = "Dados"
    livro.active.append(["Nome"])
    livro.active.append(["da primeira"])
    segunda = livro.create_sheet("Outra")
    segunda.append(["Nome"])
    segunda.append(["da segunda"])
    saida = BytesIO()
    livro.save(saida)

    linhas = read_rows("lista.xlsx", saida.getvalue())

    assert [linha.cells[0].value for linha in linhas] == ["Nome", "da primeira"]


def test_celulas_em_branco_no_fim_da_linha_sao_cortadas_e_no_meio_ficam() -> None:
    conteudo = planilha([("Alfa", None, "C", None, None)], cabecalho=("A", "B", "C", "D", "E"))

    linhas = read_rows("lista.xlsx", conteudo)

    assert [celula.value for celula in linhas[1].cells] == ["Alfa", None, "C"]


def test_a_planilha_sem_linhas_devolve_nada() -> None:
    livro = Workbook()
    saida = BytesIO()
    livro.save(saida)

    assert read_rows("lista.xlsx", saida.getvalue()) == ()


def test_a_leitura_para_no_limite_de_linhas_da_folha() -> None:
    conteudo = planilha(
        [(f"L{numero}", None, None) for numero in range(MAX_SHEET_ROWS + 10)],
    )

    linhas = read_rows("lista.xlsx", conteudo)

    assert len(linhas) == MAX_SHEET_ROWS
    assert linhas[-1].number == MAX_SHEET_ROWS
