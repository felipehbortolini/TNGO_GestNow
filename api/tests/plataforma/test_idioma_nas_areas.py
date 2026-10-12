"""O inglês nas telas desta fatia (ISSUE-087, D13): navegação, Início, Central, Governança e Financeiro.

Com o cookie ``gestnow_idioma=en`` a navegação inteira, os fragmentos das
áreas, a folha imprimível e o Excel respondem em inglês; com o cookie ausente
ou em pt os mesmos textos voltam em português (a garantia de não-regressão do
idioma-fonte). Dado cadastrado não é traduzido: o nome do projeto chega como
foi gravado.
"""

from __future__ import annotations

import re
from io import BytesIO

import azure.functions as func
import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from src.blueprints.acesso import switch_language
from src.blueprints.nav import main_nav
from src.core import translations
from src.modulos.central_acoes.routes import list_actions_screen
from src.modulos.financeiro.routes import contingencia_content
from src.modulos.governanca.routes import change_register
from src.modulos.inicio.routes import home_excel, home_printable, home_view
from tests.identidades import criar_projeto, requisicao

COOKIE = translations.LANGUAGE_COOKIE


def _pedido(
    caminho: str,
    *,
    idioma: str | None,
    params: dict[str, str] | None = None,
    alvo: str = "main-nav module-tabs",
) -> func.HttpRequest:
    cookies = {COOKIE: idioma} if idioma else None
    return requisicao(caminho, cookies=cookies, params=params, alvo=alvo)


def _corpo(resposta: func.HttpResponse) -> str:
    return resposta.get_body().decode("utf-8")


# ── A navegação ──────────────────────────────────────────────────────────


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_nav_em_ingles_traz_a_barra_traduzida() -> None:
    resposta = main_nav(
        _pedido("/central-acoes/acoes", idioma="en", params={"caminho": "/central-acoes/acoes"})
    )
    corpo = _corpo(resposta)

    assert resposta.status_code == 200
    assert 'data-idioma="en"' in corpo
    # O mapa de textos do cliente viaja no fragmento (comum, graficos e as áreas).
    assert re.search(r"data-textos='[^']*comum[^']*graficos[^']*'", corpo)
    assert "Loading..." in corpo
    assert "01 Action Center" in corpo
    assert "Modules" in corpo
    assert "Language" in corpo
    assert "Actions · Action Center" in corpo  # o título da página


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_nav_sem_o_cookie_volta_em_portugues() -> None:
    resposta = main_nav(
        _pedido("/central-acoes/acoes", idioma=None, params={"caminho": "/central-acoes/acoes"})
    )
    corpo = _corpo(resposta)

    assert 'data-idioma="pt"' in corpo
    assert "01 Central de Ações" in corpo
    assert "Módulos" in corpo
    assert "Idioma" in corpo
    assert "Ações · Central de Ações" in corpo


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_o_nome_do_projeto_chega_como_foi_gravado(sessao_das_rotas: Session) -> None:
    criar_projeto(sessao_das_rotas, codigo="TN-EN-001", nome="Fábrica Nova")
    resposta = main_nav(_pedido("/", idioma="en", params={"caminho": "/"}))
    corpo = _corpo(resposta)

    assert "TN-EN-001 · Fábrica Nova" in corpo
    assert "Portfolio (1 project)" in corpo  # a contagem é interface e fica em inglês


# ── Os fragmentos das áreas ──────────────────────────────────────────────


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_inicio_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        home_view(_pedido("/api/inicio/resumo", idioma="en", alvo="inicio-conteudo"))
    )
    em_portugues = _corpo(
        home_view(_pedido("/api/inicio/resumo", idioma=None, alvo="inicio-conteudo"))
    )

    assert "Project portfolio" in em_ingles
    assert "No project registered" in em_ingles
    assert "Portfólio de projetos" in em_portugues
    assert "Nenhum projeto cadastrado" in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_central_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        list_actions_screen(_pedido("/api/central-acoes/acoes", idioma="en", alvo="acoes-conteudo"))
    )
    em_portugues = _corpo(
        list_actions_screen(_pedido("/api/central-acoes/acoes", idioma=None, alvo="acoes-conteudo"))
    )

    assert "On time" in em_ingles  # um KPI
    assert "There are no actions yet" in em_ingles  # um estado
    assert "Em dia" in em_portugues
    assert "Ainda não há ações" in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_governanca_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        change_register(_pedido("/api/governanca/mudancas", idioma="en", alvo="mudancas-conteudo"))
    )
    em_portugues = _corpo(
        change_register(_pedido("/api/governanca/mudancas", idioma=None, alvo="mudancas-conteudo"))
    )

    assert "No changes in the scope" in em_ingles
    assert "Nenhuma mudança no escopo" in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_financeiro_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        contingencia_content(
            _pedido("/api/financeiro/contingencia", idioma="en", alvo="contingencia-conteudo")
        )
    )
    em_portugues = _corpo(
        contingencia_content(
            _pedido("/api/financeiro/contingencia", idioma=None, alvo="contingencia-conteudo")
        )
    )

    assert "No portfolio project has reserves established yet." in em_ingles
    assert "Nenhum projeto da carteira tem reservas constituídas." in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_a_folha_imprimivel_em_ingles_tem_o_cabecalho_da_tabela() -> None:
    em_ingles = _corpo(
        home_printable(
            _pedido("/api/inicio/resumo/imprimivel", idioma="en", alvo="folha-impressao")
        )
    )
    em_portugues = _corpo(
        home_printable(
            _pedido("/api/inicio/resumo/imprimivel", idioma=None, alvo="folha-impressao")
        )
    )

    assert "Attention points" in em_ingles
    assert ">Module</th>" in em_ingles
    assert "Pontos de atenção" in em_portugues
    assert ">Módulo</th>" in em_portugues


# ── O Excel ──────────────────────────────────────────────────────────────


def _planilha(resposta: func.HttpResponse):
    return load_workbook(BytesIO(resposta.get_body()))


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_o_excel_em_ingles_tem_abas_e_colunas_traduzidas() -> None:
    pasta = _planilha(home_excel(_pedido("/api/inicio/resumo/excel", idioma="en")))

    assert "Attention points" in pasta.sheetnames
    assert "Modules" in pasta.sheetnames
    aba = pasta["Attention points"]
    cabecalhos = [
        cell.value for linha in aba.iter_rows(min_row=1, max_row=20) for cell in linha if cell.value
    ]
    assert "Module" in cabecalhos
    assert "Level" in cabecalhos


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_o_excel_sem_o_cookie_volta_em_portugues() -> None:
    pasta = _planilha(home_excel(_pedido("/api/inicio/resumo/excel", idioma=None)))

    assert "Pontos de atenção" in pasta.sheetnames
    assert "Módulos" in pasta.sheetnames


# ── A troca de idioma ────────────────────────────────────────────────────


def test_a_rota_do_idioma_lembra_a_escolha(sessao_das_rotas: Session) -> None:
    del sessao_das_rotas
    resposta = switch_language(requisicao("/api/idioma", metodo="POST", corpo={"idioma": "en"}))

    assert resposta.status_code == 200
    assert "gestnow_idioma=en" in resposta.headers["Set-Cookie"]
    assert "HttpOnly" in resposta.headers["Set-Cookie"]


def test_a_rota_do_idioma_recusa_valor_desconhecido(sessao_das_rotas: Session) -> None:
    del sessao_das_rotas
    resposta = switch_language(requisicao("/api/idioma", metodo="POST", corpo={"idioma": "fr"}))

    assert resposta.status_code == 422
    assert resposta.headers.get("Content-Type", "").startswith("text/plain")
