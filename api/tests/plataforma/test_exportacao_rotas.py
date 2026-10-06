"""Rotas do exemplo de exportação (ISSUE-017, D12): o download do Excel e a versão imprimível.

Costura HTTP: o handler é chamado com a ``HttpRequest`` montada no teste. O Excel é um
download (a exceção de download do Padrão): o pedido chega sem o cabeçalho do Alpine, a
resposta é o arquivo e uma recusa vem em texto simples com o status. A versão imprimível
é um fragmento e passa pelo gate do Alpine. As duas rotas resolvem o usuário e o escopo
e conferem a permissão pelo mesmo decorador de toda rota: quem não está no cadastro e o
fornecedor recebem 403 e o Visualizador entra. A mesma descrição alimenta os dois: o que o
Excel traz, a folha traz.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from src.blueprints import exports
from src.core import auth, calendario, config, rbac
from src.core.excel import XLSX_CONTENT_TYPE
from src.core.responses import file_response
from src.modulos.configuracoes.models import Project
from tests.html_tags import Tag, find_all, find_by_id, parse_tags
from tests.identidades import (
    cabecalho_do_principal,
    criar_colaborador,
    criar_empresa,
)

RAIZ = Path(__file__).resolve().parents[3]
HOJE = date(2026, 10, 5)
URL_EXCEL = f"/api/{exports.EXCEL_ROUTE}"
URL_IMPRIMIVEL = f"/api/{exports.PRINTABLE_ROUTE}"
ID_DA_FOLHA = "folha-impressao"
COLUNAS_NO_PORTFOLIO = [
    "Projeto",
    "Atividade",
    "Início",
    "Quantidade",
    "Avanço",
    "Valor",
    "Situação",
]


def _pedido(
    caminho: str,
    *,
    email: str | None = None,
    params: dict[str, str] | None = None,
    cookies: dict[str, str] | None = None,
    alpine: bool = False,
) -> func.HttpRequest:
    """O pedido de um link do navegador; com ``alpine`` é o do botão PDF, que busca o fragmento."""
    cabecalhos: dict[str, str] = {}
    if alpine:
        cabecalhos["X-Alpine-Request"] = "true"
        cabecalhos["X-Alpine-Target"] = ID_DA_FOLHA
    if email is not None:
        cabecalhos[auth.PRINCIPAL_HEADER] = cabecalho_do_principal(email)
    if cookies:
        cabecalhos["Cookie"] = "; ".join(f"{nome}={valor}" for nome, valor in cookies.items())
    return func.HttpRequest(
        method="GET",
        url=caminho,
        headers=cabecalhos,
        params=dict(params or {}),
        route_params={},
        body=b"",
    )


def _baixar(**pedido: object) -> func.HttpResponse:
    return exports.example_excel(_pedido(URL_EXCEL, **pedido))  # type: ignore[arg-type]


def _imprimir(**pedido: object) -> func.HttpResponse:
    return exports.example_printable(_pedido(URL_IMPRIMIVEL, alpine=True, **pedido))  # type: ignore[arg-type]


def _planilha(resposta: func.HttpResponse) -> dict[str, Worksheet]:
    pasta = load_workbook(BytesIO(resposta.get_body()))
    return {aba.title: aba for aba in pasta.worksheets}


def _linha(aba: Worksheet, numero: int, colunas: int) -> list[object]:
    return [aba.cell(row=numero, column=coluna).value for coluna in range(1, colunas + 1)]


def _tags(resposta: func.HttpResponse) -> list[Tag]:
    return parse_tags(resposta.get_body().decode())


def _rotulo(projeto: Project) -> str:
    return f"{projeto.code} · {projeto.name}"


@pytest.fixture
def hoje_fixo(monkeypatch: pytest.MonkeyPatch) -> None:
    """O relógio da plataforma parado em 05/10/2026: o nome do arquivo e as datas não mudam."""
    monkeypatch.setattr(calendario, "today", lambda: HOJE)


# ── O download do Excel ──────────────────────────────────────────────────


@pytest.mark.usefixtures("dois_projetos", "hoje_fixo")
def test_o_download_e_o_arquivo_xlsx_com_o_tipo_e_o_nome() -> None:
    resposta = _baixar()

    assert resposta.status_code == 200
    assert resposta.mimetype == XLSX_CONTENT_TYPE
    assert resposta.headers["Content-Type"] == XLSX_CONTENT_TYPE
    disposicao = resposta.headers["Content-Disposition"]
    assert disposicao.startswith('attachment; filename="exemplo-de-exportacao-2026-10-05.xlsx"')
    assert "filename*=UTF-8''exemplo-de-exportacao-2026-10-05.xlsx" in disposicao
    assert "no-store" in resposta.headers["Cache-Control"]
    assert resposta.get_body()[:2] == b"PK"


@pytest.mark.usefixtures("dois_projetos", "hoje_fixo")
def test_o_download_nao_depende_do_cabecalho_do_alpine() -> None:
    sem_cabecalho = _baixar()
    com_cabecalho = _baixar(alpine=True)

    assert sem_cabecalho.status_code == com_cabecalho.status_code == 200


@pytest.mark.usefixtures("hoje_fixo")
def test_o_excel_traz_cabecalho_indicadores_e_a_tabela_com_a_coluna_projeto(
    dois_projetos: tuple[Project, Project],
) -> None:
    fabrica, caldeira = dois_projetos

    abas = _planilha(_baixar())

    assert list(abas) == ["Resumo", "Atividades de exemplo", "Legenda das situações"]
    resumo = abas["Resumo"]
    assert resumo["A2"].value == "Exemplo de exportação"
    assert resumo["A3"].value == "Escopo: Portfólio"
    assert resumo["A4"].value == "Gerado em: 05/10/2026"
    assert resumo["A5"].value == "Dados: fictícios, só para mostrar o mecanismo"
    assert [resumo.cell(row=numero, column=1).value for numero in range(8, 12)] == [
        "Avanço",
        "Desembolso",
        "Atividades concluídas",
        "Próximo marco",
    ]
    assert resumo["C8"].value == "Meta: 90%"
    atividades = abas["Atividades de exemplo"]
    assert _linha(atividades, 8, 7) == COLUNAS_NO_PORTFOLIO
    assert atividades.auto_filter.ref == "A8:G14"
    projetos = [atividades.cell(row=numero, column=1).value for numero in range(9, 15)]
    assert projetos == [_rotulo(fabrica)] * 3 + [_rotulo(caldeira)] * 3
    assert [atividades.cell(row=numero, column=4).value for numero in (15,)] == [480]
    assert atividades["F15"].value == 1_650_000.0
    assert _linha(abas["Legenda das situações"], 8, 2) == ["Situação", "Significado"]


@pytest.mark.usefixtures("hoje_fixo")
def test_o_download_de_um_projeto_so_traz_o_projeto_e_nao_tem_a_coluna_projeto(
    dois_projetos: tuple[Project, Project],
) -> None:
    fabrica, _ = dois_projetos

    abas = _planilha(_baixar(params={"projeto": str(fabrica.id)}))

    assert abas["Resumo"]["A3"].value == f"Escopo: {_rotulo(fabrica)}"
    atividades = abas["Atividades de exemplo"]
    assert _linha(atividades, 8, 6) == COLUNAS_NO_PORTFOLIO[1:]
    assert atividades.auto_filter.ref == "A8:F11"
    assert atividades["C15"].value is None
    assert atividades["C12"].value == 240


@pytest.mark.usefixtures("hoje_fixo")
def test_o_escopo_lembrado_no_cookie_vale_quando_a_url_nao_o_traz(
    dois_projetos: tuple[Project, Project],
) -> None:
    _, caldeira = dois_projetos

    abas = _planilha(_baixar(cookies={"gestnow_projeto": str(caldeira.id)}))

    assert abas["Resumo"]["A3"].value == f"Escopo: {_rotulo(caldeira)}"


@pytest.mark.usefixtures("rotas_na_transacao_do_teste", "hoje_fixo")
def test_portfolio_sem_projetos_ainda_baixa_um_arquivo_com_a_tabela_vazia() -> None:
    abas = _planilha(_baixar())

    atividades = abas["Atividades de exemplo"]
    assert _linha(atividades, 8, 7) == COLUNAS_NO_PORTFOLIO
    assert atividades.auto_filter.ref == "A8:G8"
    assert atividades.max_row == 8


@pytest.mark.usefixtures("dois_projetos", "hoje_fixo")
def test_o_papel_a3_do_pedido_vai_para_a_planilha() -> None:
    abas = _planilha(_baixar(params={"papel": "a3"}))

    assert {aba.page_setup.paperSize for aba in abas.values()} == {8}


@pytest.mark.usefixtures("dois_projetos", "hoje_fixo")
def test_papel_desconhecido_e_422_em_texto_simples_e_nao_gera_arquivo() -> None:
    resposta = _baixar(params={"papel": "a5"})

    assert resposta.status_code == 422
    assert resposta.headers["Content-Type"].startswith("text/plain")
    assert resposta.get_body().decode() == exports.UNKNOWN_PAPER_MESSAGE


# ── A permissão é a de toda rota ─────────────────────────────────────────


@pytest.mark.usefixtures("hoje_fixo")
def test_quem_nao_esta_no_cadastro_recebe_403_em_texto_simples(sessao_das_rotas: Session) -> None:
    criar_colaborador(sessao_das_rotas, email="admin@example.invalid", perfil="Admin")

    resposta = _baixar(email="intruso@example.invalid")

    assert resposta.status_code == 403
    assert resposta.headers["Content-Type"].startswith("text/plain")
    assert auth.NOT_REGISTERED_MESSAGE in resposta.get_body().decode()
    assert resposta.get_body()[:2] != b"PK"


@pytest.mark.usefixtures("hoje_fixo")
def test_o_fornecedor_nao_tem_permissao_geral_e_recebe_403(sessao_das_rotas: Session) -> None:
    empresa = criar_empresa(sessao_das_rotas, "Contratada Alfa")
    criar_colaborador(
        sessao_das_rotas,
        email="fornecedor@example.invalid",
        vinculo="Fornecedor",
        empresa=empresa,
    )

    resposta = _baixar(email="fornecedor@example.invalid")

    assert resposta.status_code == 403
    assert resposta.get_body().decode() == rbac.SUPPLIER_SCOPE_MESSAGE


@pytest.mark.usefixtures("hoje_fixo")
@pytest.mark.parametrize("perfil", ["Visualizador", "Membro", "Gestor", "Admin"])
def test_todo_perfil_geral_que_le_o_dado_pode_levar_o_excel(
    sessao_das_rotas: Session, perfil: str
) -> None:
    criar_colaborador(sessao_das_rotas, email="pessoa@example.invalid", perfil=perfil)

    resposta = _baixar(email="pessoa@example.invalid")

    assert resposta.status_code == 200


@pytest.mark.usefixtures("hoje_fixo")
def test_o_cliente_visualizador_tambem_leva_o_excel(sessao_das_rotas: Session) -> None:
    criar_colaborador(
        sessao_das_rotas, email="cliente@example.invalid", perfil="Visualizador", vinculo="Cliente"
    )

    assert _baixar(email="cliente@example.invalid").status_code == 200


@pytest.mark.usefixtures("hoje_fixo")
def test_em_producao_o_exemplo_e_recusado_mesmo_para_o_admin(
    sessao_das_rotas: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    criar_colaborador(sessao_das_rotas, email="admin@example.invalid", perfil="Admin")
    monkeypatch.setenv(config.APP_MODE_VARIABLE, config.PRODUCTION)

    excel = _baixar(email="admin@example.invalid")
    folha = _imprimir(email="admin@example.invalid")

    assert excel.status_code == folha.status_code == 403
    assert excel.get_body().decode() == exports.DEMONSTRATION_ONLY_MESSAGE
    assert exports.DEMONSTRATION_ONLY_MESSAGE in folha.get_body().decode()


# ── A versão imprimível ──────────────────────────────────────────────────


def test_a_versao_imprimivel_sem_o_cabecalho_do_alpine_leva_ao_shell() -> None:
    resposta = exports.example_printable(_pedido(URL_IMPRIMIVEL))

    assert resposta.status_code == 302
    assert resposta.headers.get("Location") == "/index.html"


@pytest.mark.usefixtures("hoje_fixo")
def test_a_folha_do_portfolio_traz_os_projetos_o_grafico_e_nenhuma_navegacao(
    dois_projetos: tuple[Project, Project],
) -> None:
    fabrica, caldeira = dois_projetos

    resposta = _imprimir()

    assert resposta.status_code == 200
    assert resposta.mimetype == "text/html"
    tags = _tags(resposta)
    raiz = find_by_id(tags, ID_DA_FOLHA)
    assert raiz is not None
    assert raiz.classes == {"folha", "folha--a4"}
    assert raiz.attrs["data-titulo"] == "Exemplo de exportação - Portfólio - 2026-10-05"
    assert [tag.attrs["data-grafico"] for tag in tags if "data-grafico" in tag.attrs] == [
        "comparativo-barras"
    ]
    celulas = [tag.clean_text() for tag in find_all(tags, "td", "folha__td")]
    assert _rotulo(fabrica) in celulas
    assert _rotulo(caldeira) in celulas
    assert find_all(tags, "a") == []
    assert find_all(tags, "nav") == []


@pytest.mark.usefixtures("dois_projetos", "hoje_fixo")
def test_o_papel_a3_do_pedido_vira_a_classe_da_folha() -> None:
    raiz = find_by_id(_tags(_imprimir(params={"papel": "a3"})), ID_DA_FOLHA)

    assert raiz is not None
    assert raiz.classes == {"folha", "folha--a3"}


@pytest.mark.usefixtures("dois_projetos", "hoje_fixo")
def test_papel_desconhecido_na_folha_e_422_com_o_aviso_e_o_toast() -> None:
    resposta = _imprimir(params={"papel": "a5"})

    assert resposta.status_code == 422
    assert exports.UNKNOWN_PAPER_MESSAGE in resposta.get_body().decode()
    assert resposta.headers.get("X-TN-Toast")
    assert resposta.headers.get("X-TN-Toast-Tipo") == "erro"


@pytest.mark.usefixtures("hoje_fixo")
def test_a_folha_de_quem_nao_esta_no_cadastro_e_403_com_o_fragmento_de_erro(
    sessao_das_rotas: Session,
) -> None:
    criar_colaborador(sessao_das_rotas, email="admin@example.invalid", perfil="Admin")

    resposta = _imprimir(email="intruso@example.invalid")

    assert resposta.status_code == 403
    assert resposta.mimetype == "text/html"
    assert auth.NOT_REGISTERED_MESSAGE in resposta.get_body().decode()
    assert find_by_id(_tags(resposta), ID_DA_FOLHA) is not None
    assert find_all(_tags(resposta), "table") == []


# ── A mesma descrição alimenta o Excel e a folha ─────────────────────────


@pytest.mark.usefixtures("dois_projetos", "hoje_fixo")
def test_o_que_o_excel_traz_a_folha_traz() -> None:
    atividades = _planilha(_baixar())["Atividades de exemplo"]
    folha = _tags(_imprimir())

    colunas_da_folha = [tag.clean_text() for tag in find_all(folha, "th")]
    assert colunas_da_folha[: len(COLUNAS_NO_PORTFOLIO)] == _linha(atividades, 8, 7)
    projetos_da_folha = [
        tag.clean_text()
        for posicao, tag in enumerate(find_all(folha, "td", "folha__td"))
        if posicao % len(COLUNAS_NO_PORTFOLIO) == 0
    ][:6]
    assert projetos_da_folha == [
        atividades.cell(row=numero, column=1).value for numero in range(9, 15)
    ]


# ── O arquivo de download ────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("nome", "simples", "utf8"),
    [
        ("relatório-ç.xlsx", "relatorio-c.xlsx", "relat%C3%B3rio-%C3%A7.xlsx"),
        ('a"b.xlsx', "a_b.xlsx", "a%22b.xlsx"),
        ("plano 2026 (v2).xlsx", "plano 2026 _v2_.xlsx", "plano%202026%20%28v2%29.xlsx"),
        ("中文", "arquivo", "%E4%B8%AD%E6%96%87"),
    ],
)
def test_o_nome_do_arquivo_vai_no_cabecalho_com_e_sem_acento(
    nome: str, simples: str, utf8: str
) -> None:
    resposta = file_response(b"conteudo", filename=nome, content_type="application/octet-stream")

    esperado = f"attachment; filename=\"{simples}\"; filename*=UTF-8''{utf8}"
    assert resposta.headers["Content-Disposition"] == esperado
    assert resposta.get_body() == b"conteudo"


def test_o_arquivo_nao_fica_em_cache_nem_e_reinterpretado_pelo_navegador() -> None:
    resposta = file_response(b"x", filename="a.bin", content_type="application/octet-stream")

    assert resposta.headers["Content-Type"] == "application/octet-stream"
    assert resposta.headers["Cache-Control"] == "private, no-store"
    assert resposta.headers["X-Content-Type-Options"] == "nosniff"


# ── O exemplo está ligado onde precisa estar ─────────────────────────────


def test_a_pagina_de_exemplo_tem_os_botoes_das_duas_rotas_e_carrega_so_o_design_system() -> None:
    pagina = (RAIZ / "app" / "exemplos" / "exportacao.html").read_text(encoding="utf-8")

    assert f'data-tn-excel="{URL_EXCEL}"' in pagina
    assert f'data-tn-pdf="{URL_IMPRIMIVEL}"' in pagina
    assert f'data-tn-pdf="{URL_IMPRIMIVEL}?papel=a3"' in pagina
    assert '<link rel="stylesheet" href="/ds/print.css" media="print" />' in pagina
    assert '<script defer src="/ds/ui.js"></script>' in pagina
    assert "http://" not in pagina
    assert "https://" not in pagina


def test_o_exemplo_esta_registrado_no_azure_e_no_servidor_local() -> None:
    azure = (RAIZ / "api" / "function_app.py").read_text(encoding="utf-8")
    local = (RAIZ / "scripts" / "dev_local.py").read_text(encoding="utf-8")

    assert "from src.blueprints.exports import bp as exports_bp" in azure
    assert "app.register_functions(exports_bp)" in azure
    assert f'r"^{URL_EXCEL}$"' in local
    assert f'r"^{URL_IMPRIMIVEL}$"' in local
