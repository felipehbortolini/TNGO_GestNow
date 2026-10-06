"""Importação em passos (ISSUE-018, HU-139): as quatro rotas, com o importador de teste.

O handler é chamado com a ``HttpRequest`` que o navegador montaria: ``multipart/form-data`` no envio
e na confirmação, o cabeçalho do Alpine nos fragmentos e nenhum cabeçalho no download do modelo.
O que se afirma é o que a pessoa vê e o que fica no banco: a tela dos passos, o modelo baixado, a
conferência com os motivos (e nada gravado), o 422 do que não é planilha e a confirmação (a
gravação, o resultado, o tudo ou nada e a recusa do que não foi conferido).
"""

from __future__ import annotations

import hashlib
from collections.abc import Callable, Iterator, Mapping
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from typing import Any

import azure.functions as func
import pytest
from openpyxl import load_workbook
from sqlalchemy import func as sql
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.blueprints.importing import import_check, import_confirm, import_steps, import_template
from src.core import calendario, importing
from src.core.auth import PRINCIPAL_HEADER
from src.core.excel import XLSX_CONTENT_TYPE
from src.core.models import AuditEntry
from src.core.rbac import Permission
from src.core.responses import cabecalhos_de_toast
from src.core.routing import Access
from tests.html_tags import Tag, find_all, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal, criar_colaborador, criar_empresa
from tests.importadores_de_teste import (
    CHAVE_DE_CLIENTES,
    NOME_RECUSADO,
    PREFIXO,
    clientes_gravados,
    corpo_multipart,
    importador_de_clientes,
    planilha,
)

HOJE = date(2026, 10, 6)
MARIO = "mario.membro@example.invalid"
VERA = "vera.visualizadora@example.invalid"
FABIO = "fabio.fornecedor@example.invalid"

ALVO = "importar-clientes"
ALVO_DO_RESULTADO = f"{ALVO}-resultado"
ALFA = f"{PREFIXO} Alfa"
BETA = f"{PREFIXO} Beta"
GAMA = f"{PREFIXO} Gama"
NAO_E_PLANILHA = "não é uma planilha do Excel (.xlsx)"

Rota = Callable[[func.HttpRequest], func.HttpResponse]


@pytest.fixture(autouse=True)
def importadores(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """O registro só tem o importador de teste, e a data de hoje é fixa."""
    monkeypatch.setattr(importing, "_IMPORTERS", {})
    importing.register(importador_de_clientes())
    monkeypatch.setattr(calendario, "today", lambda: HOJE)
    yield


@pytest.fixture(autouse=True)
def pessoas(sessao_das_rotas: Session) -> Session:
    """Um Membro, um Visualizador e um fornecedor no cadastro, com as rotas na transação do teste."""
    criar_colaborador(sessao_das_rotas, email=MARIO, nome="Mário Membro", perfil="Membro")
    criar_colaborador(
        sessao_das_rotas, email=VERA, nome="Vera Visualizadora", perfil="Visualizador"
    )
    criar_colaborador(
        sessao_das_rotas,
        email=FABIO,
        nome="Fábio Fornecedor",
        vinculo="Fornecedor",
        empresa=criar_empresa(sessao_das_rotas, "Contratada Alfa"),
    )
    return sessao_das_rotas


@dataclass(frozen=True)
class Envio:
    """Como o navegador chama a rota; ``alvo=None`` é o link aberto direto (sem o cabeçalho do Alpine)."""

    metodo: str = "POST"
    email: str = MARIO
    chave: str = CHAVE_DE_CLIENTES
    campos: Mapping[str, str] = field(default_factory=lambda: {"formulario": "importacao"})
    arquivo: tuple[str, bytes] | None = None
    alvo: str | None = ALVO_DO_RESULTADO


def _chamar(rota: Rota, **mudancas: Any) -> func.HttpResponse:
    """Chama a rota com a ``HttpRequest`` do navegador: o principal, o Alpine e o formulário."""
    envio = Envio(**mudancas)
    headers = {PRINCIPAL_HEADER: cabecalho_do_principal(envio.email)}
    corpo = b""
    if envio.metodo == "POST":
        corpo, headers["Content-Type"] = corpo_multipart(envio.campos, envio.arquivo)
    if envio.alvo is not None:
        headers.update({"X-Alpine-Request": "true", "X-Alpine-Target": envio.alvo})
    return rota(
        func.HttpRequest(
            method=envio.metodo,
            url=f"/api/importacao/{envio.chave}",
            headers=headers,
            params={},
            route_params={"chave": envio.chave},
            body=corpo,
        )
    )


def _html(resposta: func.HttpResponse) -> str:
    return resposta.get_body().decode()


def _conferir(
    conteudo: bytes, *, nome: str = "lista.xlsx", email: str = MARIO
) -> func.HttpResponse:
    return _chamar(import_check, email=email, arquivo=(nome, conteudo))


def _campo(tags: list[Tag], nome: str) -> Tag | None:
    return next((tag for tag in find_all(tags, "input") if tag.attrs.get("name") == nome), None)


def _resumo_conferido(conteudo: bytes) -> str:
    """O resumo que a conferência devolve em ``conferido``, lido do fragmento como o navegador o leria."""
    campo = _campo(parse_tags(_html(_conferir(conteudo))), "conferido")
    assert campo is not None
    return campo.attrs["value"]


def _confirmar(
    conteudo: bytes,
    *,
    resumo: str | None,
    ciente: bool = False,
    email: str = MARIO,
) -> func.HttpResponse:
    campos = {"formulario": "importacao"}
    if resumo is not None:
        campos["conferido"] = resumo
    if ciente:
        campos["ciente"] = "sim"
    return _chamar(import_confirm, email=email, campos=campos, arquivo=("lista.xlsx", conteudo))


def _trilha(session: Session) -> int:
    consulta = sql.count(AuditEntry.id)
    return session.scalar(select(consulta).where(AuditEntry.entity == "cliente")) or 0


# Uma linha que entra, uma que entra com aviso (sem sigla) e uma que o erro bloqueia (sem nome).
ENTRA_AVISA_E_ERRA = planilha([(ALFA, "ALF", "Sim"), (BETA, None, "Não"), (None, "SEM", "Sim")])
SO_LINHAS_BOAS = planilha([(ALFA, "ALF", "Sim"), (GAMA, "GAM", "Sim")])


# ── Passo 1 e 2: a tela e o modelo ───────────────────────────────────────


def test_a_tela_dos_passos_traz_o_botao_do_modelo_o_formulario_e_o_lugar_da_conferencia() -> None:
    resposta = _chamar(import_steps, metodo="GET", alvo=ALVO)

    assert resposta.status_code == 200
    tags = parse_tags(_html(resposta))
    assert find_by_id(tags, ALVO) is not None
    (botao,) = [tag for tag in find_all(tags, "button") if "data-tn-excel" in tag.attrs]
    assert botao.attrs["data-tn-excel"] == f"/api/importacao/{CHAVE_DE_CLIENTES}/modelo"
    (formulario,) = find_all(tags, "form")
    assert formulario.attrs["method"] == "post"
    assert formulario.attrs["action"] == f"/api/importacao/{CHAVE_DE_CLIENTES}/conferir"
    assert formulario.attrs["enctype"] == "multipart/form-data"
    assert formulario.attrs["x-target"] == ALVO_DO_RESULTADO
    campo = find_by_id(tags, f"{ALVO}-arquivo")
    assert campo is not None
    assert campo.attrs["type"] == "file"
    assert campo.attrs["name"] == "arquivo"
    assert campo.attrs["accept"] == ".xlsx"
    assert "required" in campo.attrs
    assert find_by_id(tags, ALVO_DO_RESULTADO) is not None
    dica = find_by_id(tags, f"{ALVO}-dica")
    assert dica is not None
    assert dica.clean_text() == "Somente .xlsx, até 5 MB e 5.000 linhas de dados."
    for proibida in ("script", "link", "style"):
        assert find_all(tags, proibida) == []


def test_a_tela_dos_passos_diz_o_que_cada_coluna_pede() -> None:
    tags = parse_tags(_html(_chamar(import_steps, metodo="GET", alvo=ALVO)))

    celulas = [celula.clean_text() for celula in find_all(tags, "td")]
    linhas = [celulas[indice : indice + 4] for indice in (0, 4, 8)]
    assert celulas == [item for linha in linhas for item in linha]
    assert linhas == [
        ["Nome", "Sim", "texto", ALFA],
        ["Sigla", "Não", "texto", "ALF"],
        ["Ativo", "Não", "um de: Sim, Não", "Sim"],
    ]
    assert "Importar planilha: Clientes de teste" in _html(
        _chamar(import_steps, metodo="GET", alvo=ALVO)
    )


def test_o_modelo_e_um_download_xlsx_com_cabecalho_e_instrucoes() -> None:
    resposta = _chamar(import_template, metodo="GET", alvo=None)

    assert resposta.status_code == 200
    assert resposta.headers["Content-Type"] == XLSX_CONTENT_TYPE
    assert (
        "modelo-de-importacao-clientes-de-teste-2026-10-06.xlsx"
        in resposta.headers["Content-Disposition"]
    )
    assert resposta.headers["Cache-Control"] == "private, no-store"
    livro = load_workbook(BytesIO(resposta.get_body()))
    assert livro.sheetnames == ["Dados", "Instruções"]


@pytest.mark.parametrize("email", [VERA, FABIO])
def test_o_modelo_e_recusado_em_texto_a_quem_nao_escreve_no_modulo(email: str) -> None:
    resposta = _chamar(import_template, metodo="GET", email=email, alvo=None)

    assert resposta.status_code == 403
    assert resposta.headers["Content-Type"].startswith("text/plain")


def test_modelo_de_importacao_que_nao_existe_e_422_em_texto() -> None:
    resposta = _chamar(import_template, metodo="GET", chave="nao-existe", alvo=None)

    assert resposta.status_code == 422
    assert _html(resposta) == "A importação pedida não existe."


# ── Passo 2: a conferência não grava ─────────────────────────────────────


def test_a_conferencia_mostra_o_que_cada_linha_vai_fazer_com_o_motivo_e_nao_grava(
    sessao_das_rotas: Session,
) -> None:
    antes = _trilha(sessao_das_rotas)

    resposta = _conferir(ENTRA_AVISA_E_ERRA)

    assert resposta.status_code == 200
    tags = parse_tags(_html(resposta))
    assert find_by_id(tags, ALVO_DO_RESULTADO) is not None
    assert [tag.clean_text() for tag in find_all(tags, "span", "kpi__value")] == ["2", "1", "1"]
    motivos = [tag.clean_text() for tag in find_all(tags, "li", "lista-simples__item")]
    assert motivos == [
        "Linha 3 · aviso Sigla em branco: o cliente entra sem sigla",
        "Linha 4 · erro Nome: obrigatório",
    ]
    situacoes = [
        tag.clean_text()
        for tag in find_all(tags, "span", "pill")
        if tag.clean_text() in ("Entra", "Entra com aviso")
    ]
    assert situacoes == ["Entra", "Entra com aviso"]
    assert clientes_gravados(sessao_das_rotas) == []
    assert _trilha(sessao_das_rotas) == antes


def test_a_conferencia_devolve_o_resumo_do_arquivo_e_o_botao_que_confirma_o_mesmo_formulario() -> (
    None
):
    tags = parse_tags(_html(_conferir(ENTRA_AVISA_E_ERRA)))

    resumo = _campo(tags, "conferido")
    assert resumo is not None
    assert resumo.attrs["type"] == "hidden"
    assert resumo.attrs["value"] == hashlib.sha256(ENTRA_AVISA_E_ERRA).hexdigest()
    (confirmar,) = [tag for tag in find_all(tags, "button") if "formaction" in tag.attrs]
    assert confirmar.attrs["formaction"] == f"/api/importacao/{CHAVE_DE_CLIENTES}/confirmar"
    assert confirmar.attrs["type"] == "submit"
    assert "Importar 2 linha(s)" in confirmar.clean_text()
    ciente = _campo(tags, "ciente")
    assert ciente is not None
    assert ciente.attrs["type"] == "checkbox"
    assert ciente.attrs["value"] == "sim"


def test_sem_linhas_de_erro_a_conferencia_nao_pede_o_estou_ciente() -> None:
    tags = parse_tags(_html(_conferir(SO_LINHAS_BOAS)))

    assert _campo(tags, "ciente") is None
    assert [tag.clean_text() for tag in find_all(tags, "span", "kpi__value")] == ["2", "0", "0"]
    assert find_all(tags, "li", "lista-simples__item") == []


def test_a_conferencia_lista_as_colunas_ausentes_e_nao_oferece_confirmar() -> None:
    conteudo = planilha([("ALF", "Sim")], cabecalho=("Sigla", "Ativo"))

    resposta = _conferir(conteudo)

    assert resposta.status_code == 200
    tags = parse_tags(_html(resposta))
    assert "Colunas obrigatórias ausentes em lista.xlsx" in _html(resposta)
    assert "Nome. Baixe o modelo" in " ".join(_html(resposta).split())
    assert _campo(tags, "conferido") is None
    assert [tag for tag in find_all(tags, "button") if "formaction" in tag.attrs] == []


@pytest.mark.parametrize(
    ("nome", "conteudo"),
    [
        ("lista.pdf", b"%PDF-1.7 conteudo"),
        ("lista.xls", b"\xd0\xcf\x11\xe0 antigo"),
        ("lista.xlsx", b"isto nao e uma planilha"),
        ("lista.csv", b"Nome;Sigla\nAlfa;ALF\n"),
    ],
)
def test_arquivo_que_nao_e_planilha_e_422_com_a_mensagem_no_lugar_da_conferencia(
    sessao_das_rotas: Session, nome: str, conteudo: bytes
) -> None:
    resposta = _conferir(conteudo, nome=nome)

    assert resposta.status_code == 422
    assert resposta.headers["X-TN-Toast-Tipo"] == "erro"
    corpo = " ".join(_html(resposta).split())
    assert NAO_E_PLANILHA in corpo
    assert find_by_id(parse_tags(_html(resposta)), ALVO_DO_RESULTADO) is not None
    assert clientes_gravados(sessao_das_rotas) == []


def test_envio_sem_arquivo_e_422_pedindo_a_planilha() -> None:
    resposta = _chamar(import_check)

    assert resposta.status_code == 422
    assert "Escolha a planilha (.xlsx) antes de conferir." in _html(resposta)
    assert (
        resposta.headers["X-TN-Toast"]
        == cabecalhos_de_toast("Escolha a planilha (.xlsx) antes de conferir.", "erro")[
            "X-TN-Toast"
        ]
    )


# ── Passo 3: a confirmação grava, tudo ou nada ───────────────────────────


def test_a_confirmacao_grava_as_linhas_sem_erro_e_mostra_o_que_ficou_de_fora(
    sessao_das_rotas: Session,
) -> None:
    antes = _trilha(sessao_das_rotas)
    resumo = _resumo_conferido(ENTRA_AVISA_E_ERRA)

    resposta = _confirmar(ENTRA_AVISA_E_ERRA, resumo=resumo, ciente=True)

    assert resposta.status_code == 200
    assert (
        resposta.headers["X-TN-Toast"]
        == cabecalhos_de_toast("2 linhas importadas; 1 com erro não entrou.", "aviso")["X-TN-Toast"]
    )
    assert resposta.headers["X-TN-Toast-Tipo"] == "aviso"
    assert clientes_gravados(sessao_das_rotas) == [ALFA, BETA]
    assert _trilha(sessao_das_rotas) == antes + 2
    corpo = " ".join(_html(resposta).split())
    assert "2 linha(s) de Clientes de teste importada(s)" in corpo
    assert "1 linha(s) entraram com aviso." in corpo
    tags = parse_tags(_html(resposta))
    assert [tag.clean_text() for tag in find_all(tags, "li", "lista-simples__item")] == [
        "Linha 4 · erro Nome: obrigatório"
    ]


def test_a_confirmacao_sem_linhas_de_erro_avisa_ok_e_limpa_o_formulario() -> None:
    resumo = _resumo_conferido(SO_LINHAS_BOAS)

    resposta = _confirmar(SO_LINHAS_BOAS, resumo=resumo)

    assert resposta.status_code == 200
    assert (
        resposta.headers["X-TN-Toast"] == cabecalhos_de_toast("2 linhas importadas.")["X-TN-Toast"]
    )
    assert resposta.headers["X-TN-Toast-Tipo"] == "ok"
    assert "$el.closest('form').reset()" in _html(resposta)


def test_com_linhas_de_erro_a_confirmacao_sem_o_estou_ciente_e_422_e_nada_e_gravado(
    sessao_das_rotas: Session,
) -> None:
    resumo = _resumo_conferido(ENTRA_AVISA_E_ERRA)

    resposta = _confirmar(ENTRA_AVISA_E_ERRA, resumo=resumo, ciente=False)

    assert resposta.status_code == 422
    assert "Confirme que está ciente de que as linhas com erro não serão importadas." in _html(
        resposta
    )
    assert clientes_gravados(sessao_das_rotas) == []


def test_sem_o_resumo_da_conferencia_a_confirmacao_e_422_e_nada_e_gravado(
    sessao_das_rotas: Session,
) -> None:
    resposta = _confirmar(SO_LINHAS_BOAS, resumo=None)

    assert resposta.status_code == 422
    assert "Confira a planilha antes de importar" in _html(resposta)
    assert clientes_gravados(sessao_das_rotas) == []


def test_outro_arquivo_na_confirmacao_e_422_e_nada_e_gravado(sessao_das_rotas: Session) -> None:
    resumo = _resumo_conferido(SO_LINHAS_BOAS)

    resposta = _confirmar(ENTRA_AVISA_E_ERRA, resumo=resumo, ciente=True)

    assert resposta.status_code == 422
    assert "O arquivo mudou depois da conferência." in _html(resposta)
    assert clientes_gravados(sessao_das_rotas) == []


def test_arquivo_que_nao_e_planilha_na_confirmacao_e_422(sessao_das_rotas: Session) -> None:
    conteudo = b"%PDF-1.7 conteudo"
    resposta = _chamar(
        import_confirm,
        campos={"conferido": hashlib.sha256(conteudo).hexdigest()},
        arquivo=("lista.pdf", conteudo),
    )

    assert resposta.status_code == 422
    assert NAO_E_PLANILHA in " ".join(_html(resposta).split())
    assert clientes_gravados(sessao_das_rotas) == []


def test_a_confirmacao_e_tudo_ou_nada_e_a_recusa_diz_a_linha(sessao_das_rotas: Session) -> None:
    conteudo = planilha([(ALFA, "ALF", "Sim"), (NOME_RECUSADO, "REC", "Sim"), (GAMA, "GAM", "Sim")])
    antes = _trilha(sessao_das_rotas)
    resumo = _resumo_conferido(conteudo)

    resposta = _confirmar(conteudo, resumo=resumo)

    assert resposta.status_code == 422
    assert "Linha 3: O cliente recusado não pode ser importado. Nada foi importado." in _html(
        resposta
    )
    assert clientes_gravados(sessao_das_rotas) == []
    assert _trilha(sessao_das_rotas) == antes


def test_cancelar_e_nao_confirmar_nao_deixa_nada_no_banco(sessao_das_rotas: Session) -> None:
    antes = _trilha(sessao_das_rotas)

    _conferir(ENTRA_AVISA_E_ERRA)
    _conferir(SO_LINHAS_BOAS)

    assert clientes_gravados(sessao_das_rotas) == []
    assert _trilha(sessao_das_rotas) == antes


# ── Acesso e endereço ────────────────────────────────────────────────────


@pytest.mark.parametrize("email", [VERA, FABIO])
def test_quem_nao_escreve_no_modulo_recebe_403_nos_tres_passos(
    sessao_das_rotas: Session, email: str
) -> None:
    assert _chamar(import_steps, metodo="GET", email=email, alvo=ALVO).status_code == 403
    assert _conferir(SO_LINHAS_BOAS, email=email).status_code == 403
    assert _confirmar(SO_LINHAS_BOAS, resumo="x", email=email).status_code == 403
    assert clientes_gravados(sessao_das_rotas) == []


@pytest.mark.parametrize("rota", [import_steps, import_check, import_confirm])
def test_importacao_que_nao_existe_e_422_nos_passos(rota: Rota) -> None:
    metodo = "GET" if rota is import_steps else "POST"

    resposta = _chamar(rota, metodo=metodo, chave="nao-existe", alvo=ALVO)

    assert resposta.status_code == 422
    assert "A importação pedida não existe." in _html(resposta)


def test_os_fragmentos_sem_o_cabecalho_do_alpine_voltam_ao_shell() -> None:
    resposta = _chamar(import_check, arquivo=("lista.xlsx", SO_LINHAS_BOAS), alvo=None)

    assert resposta.status_code == 302
    assert resposta.headers["Location"] == "/index.html"


def test_as_rotas_estao_registradas_com_os_caminhos_os_metodos_e_o_acesso_da_spec(
    funcoes_registradas: list,
) -> None:
    rotas = {funcao.get_function_name(): funcao for funcao in funcoes_registradas}

    esperado = {
        "import_steps": ("importacao/{chave}", ["GET"]),
        "import_template": ("importacao/{chave}/modelo", ["GET"]),
        "import_check": ("importacao/{chave}/conferir", ["POST"]),
        "import_confirm": ("importacao/{chave}/confirmar", ["POST"]),
    }
    for nome, (caminho, metodos) in esperado.items():
        gatilho = rotas[nome].get_trigger()
        assert (gatilho.route, [metodo.value for metodo in gatilho.methods]) == (caminho, metodos)
        assert rotas[nome].get_user_function().__dict__["access"] == Access(
            permission=Permission.WRITE
        )
