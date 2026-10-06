"""Excel do servidor (ISSUE-017, D12): o arquivo gerado é aberto com o openpyxl e conferido.

O construtor é um só e toda tela o usa. O teste abre o ``.xlsx`` que ele devolve e
confere o que a pessoa vê ao abrir: as abas, o cabeçalho de cada aba (logo, título,
escopo, data e contexto), os indicadores com a referência de gestão, a tabela com o
filtro ligado e, no Portfólio, a coluna Projeto. As cores das células são as dos
tokens do Design System: aqui elas são lidas do ``app/ds/tokens.css`` de verdade, e
não do arquivo que o servidor leva consigo. Números e datas são células de verdade
(ordenam, filtram e somam), com o formato de cada tipo, e o dinheiro chega em reais
inteiros.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell as SheetCell
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.core import design_tokens, excel
from src.core.export_document import (
    Cell,
    Column,
    Document,
    Kpi,
    Paper,
    Table,
    Tone,
    ValueKind,
    row,
)
from tests.documentos_de_exportacao import CALDEIRA, FABRICA, HOJE
from tests.documentos_de_exportacao import documento as _documento

RAIZ = Path(__file__).resolve().parents[3]
TOKENS_CSS = RAIZ / "app" / "ds" / "tokens.css"
FORMATO_DINHEIRO = '"R$" #,##0.00'

# O que cada tom é no Design System: a pílula (``.pill--<sufixo>`` em tokens.css).
SUFIXO_DO_TOM = {
    Tone.OK: "ok",
    Tone.WARN: "warn",
    Tone.ERROR: "erro",
    Tone.INFO: "fix",
    Tone.NEUTRAL: "neutral",
}


def _abrir(documento: Document) -> Workbook:
    return load_workbook(BytesIO(excel.build_workbook(documento)))


def _tabelas_simples(*titulos: str) -> tuple[Table, ...]:
    return tuple(
        Table(title=titulo, columns=(Column("A"),), rows=(row("x"),)) for titulo in titulos
    )


def _valor(celula: SheetCell) -> object:
    """O valor da célula, com a data de volta como ``date`` (o openpyxl a lê como ``datetime``)."""
    valor = celula.value
    return valor.date() if isinstance(valor, datetime) else valor


def _linha(aba: Worksheet, numero: int, colunas: int) -> list[object]:
    return [_valor(aba.cell(row=numero, column=coluna)) for coluna in range(1, colunas + 1)]


def _tokens() -> design_tokens.DesignTokens:
    return design_tokens.parse_css(TOKENS_CSS.read_text(encoding="utf-8"))


def _cor_do_token(token: str) -> str:
    """O ``RRGGBB`` do token, lido do ``tokens.css`` do Design System."""
    return _tokens().colors[token].removeprefix("#")


def _cores_da_pilula(sufixo: str) -> tuple[str, str]:
    """Fundo e texto da pílula do Design System, como o ``tokens.css`` os declara."""
    css = TOKENS_CSS.read_text(encoding="utf-8")
    regra = re.search(
        rf"\.pill--{sufixo}\s*\{{\s*background:\s*var\((--[\w-]+)\);\s*color:\s*var\((--[\w-]+)\);",
        css,
    )
    assert regra is not None, sufixo
    return _cor_do_token(regra.group(1)), _cor_do_token(regra.group(2))


def _rgb(cor: object) -> str:
    """Os seis dígitos da cor: o openpyxl guarda ``RRGGBB`` como ``00RRGGBB``."""
    return str(getattr(cor, "rgb", ""))[-6:]


# ── As abas e o cabeçalho de cada uma ────────────────────────────────────


def test_o_arquivo_tem_o_resumo_e_uma_aba_por_tabela() -> None:
    pasta = _abrir(_documento())

    assert pasta.sheetnames == ["Resumo", "Atividades", "Legenda"]
    assert pasta.properties.title == "Mapa de controle"
    assert pasta.properties.creator == "Timenow GestNow"


def test_cada_aba_traz_o_logo_o_titulo_o_escopo_a_data_e_o_contexto() -> None:
    pasta = _abrir(_documento())

    for aba in pasta.worksheets:
        assert len(aba._images) == 1, aba.title
        ancora = aba._images[0].anchor._from
        assert (ancora.col, ancora.row) == (0, 0), aba.title
        assert aba["A2"].value == "Mapa de controle", aba.title
        assert aba["A3"].value == "Escopo: Portfólio", aba.title
        assert aba["A4"].value == "Gerado em: 05/10/2026", aba.title
        assert aba["A5"].value == "Período: S40/2026", aba.title


def test_o_escopo_de_um_projeto_aparece_no_cabecalho() -> None:
    pasta = _abrir(_documento(portfolio=False))

    assert pasta["Resumo"]["A3"].value == f"Escopo: {FABRICA}"


def test_o_titulo_e_a_fonte_usam_os_tokens() -> None:
    aba = _abrir(_documento())["Resumo"]

    assert aba["A2"].font.b is True
    assert aba["A2"].font.sz == 16
    assert _rgb(aba["A2"].font.color) == _cor_do_token("--brand-title")
    assert aba["A2"].font.name == _tokens().font
    assert _rgb(aba["A3"].font.color) == _cor_do_token("--text-secondary")
    assert aba.sheet_view.showGridLines is False


# ── Os indicadores, com a referência de gestão ───────────────────────────


def test_o_resumo_traz_cada_indicador_com_a_referencia_e_a_situacao() -> None:
    aba = _abrir(_documento())["Resumo"]

    assert _linha(aba, 7, 4) == ["Indicador", "Valor", "Referência", "Situação"]
    assert _linha(aba, 8, 4) == ["Avanço", 0.875, "Meta: 90%", "Atenção"]
    assert _linha(aba, 9, 4) == [
        "Desembolso",
        1_250_000.0,
        "Orçado: R$ 1.500.000,00",
        "Dentro do orçado",
    ]
    assert _linha(aba, 10, 4) == ["Concluídas", 18, "Previsto: 20", None]
    assert _linha(aba, 11, 4) == ["Próximo marco", date(2026, 10, 19), "Data prevista", "Futuro"]


def test_o_valor_do_indicador_tem_o_formato_do_seu_tipo() -> None:
    aba = _abrir(_documento())["Resumo"]

    assert aba["B8"].number_format == "0.0%"
    assert aba["B9"].number_format == FORMATO_DINHEIRO
    assert aba["B10"].number_format == "#,##0"
    assert aba["B11"].number_format == "DD/MM/YYYY"


def test_o_indicador_com_tom_leva_a_cor_do_design_system_e_o_sem_tom_nao() -> None:
    aba = _abrir(_documento())["Resumo"]
    fundo, texto = _cores_da_pilula("warn")

    assert _rgb(aba["B8"].fill.fgColor) == fundo
    assert _rgb(aba["B8"].font.color) == texto
    assert _rgb(aba["D8"].fill.fgColor) == fundo
    assert aba["B10"].fill.fill_type is None


def test_o_cabecalho_do_resumo_usa_os_tokens() -> None:
    aba = _abrir(_documento())["Resumo"]

    assert _rgb(aba["A7"].fill.fgColor) == _cor_do_token("--neutro-50")
    assert _rgb(aba["A7"].font.color) == _cor_do_token("--text-secondary")
    assert aba["A7"].font.b is True
    assert _rgb(aba["A7"].border.bottom.color) == _cor_do_token("--borda")


def test_sem_nenhuma_situacao_o_resumo_nao_tem_a_coluna_situacao() -> None:
    sem_situacao = (Kpi("Concluídas", 18, "Previsto: 20", ValueKind.INTEGER),)

    aba = _abrir(_documento(kpis=sem_situacao))["Resumo"]

    assert _linha(aba, 7, 4) == ["Indicador", "Valor", "Referência", None]
    assert _linha(aba, 8, 4) == ["Concluídas", 18, "Previsto: 20", None]


def test_indicador_sem_valor_fica_em_branco_e_mantem_a_referencia() -> None:
    sem_dado = (Kpi("Sem dado", None, "Meta: 90%", ValueKind.PERCENT, digits=1),)

    aba = _abrir(_documento(kpis=sem_dado))["Resumo"]

    assert _linha(aba, 8, 3) == ["Sem dado", None, "Meta: 90%"]


def test_documento_sem_indicadores_nem_tabelas_ainda_abre_com_o_resumo() -> None:
    vazio = Document(title="Vazio", scope_label="Portfólio", generated_on=HOJE)

    pasta = _abrir(vazio)

    assert pasta.sheetnames == ["Resumo"]
    assert pasta["Resumo"]["A2"].value == "Vazio"


# ── A tabela: filtro, cabeçalho congelado e a coluna Projeto ─────────────


def test_a_tabela_tem_filtro_cabecalho_congelado_e_o_cabecalho_repetido_ao_imprimir() -> None:
    aba = _abrir(_documento())["Atividades"]

    assert aba["A7"].value == "Atividades"
    assert aba.auto_filter.ref == "A8:G11"
    assert aba.freeze_panes == "A9"
    assert aba.print_title_rows == "$8:$8"


def test_no_portfolio_a_coluna_projeto_vem_primeiro_com_o_projeto_de_cada_linha() -> None:
    aba = _abrir(_documento())["Atividades"]

    assert _linha(aba, 8, 7) == [
        "Projeto",
        "Atividade",
        "Início",
        "Quantidade",
        "Avanço",
        "Valor",
        "Situação",
    ]
    assert [aba.cell(row=numero, column=1).value for numero in (9, 10, 11)] == [
        FABRICA,
        FABRICA,
        CALDEIRA,
    ]


def test_no_portfolio_a_tabela_que_nao_e_por_projeto_nao_ganha_a_coluna() -> None:
    aba = _abrir(_documento())["Legenda"]

    assert _linha(aba, 8, 3) == ["Situação", "Significado", None]
    assert aba.auto_filter.ref == "A8:B10"


def test_fora_do_portfolio_nao_ha_coluna_projeto() -> None:
    aba = _abrir(_documento(portfolio=False))["Atividades"]

    assert _linha(aba, 8, 7) == [
        "Atividade",
        "Início",
        "Quantidade",
        "Avanço",
        "Valor",
        "Situação",
        None,
    ]
    assert aba.auto_filter.ref == "A8:F11"


def test_os_valores_da_tabela_sao_celulas_de_verdade() -> None:
    aba = _abrir(_documento())["Atividades"]

    assert _linha(aba, 9, 7) == [
        FABRICA,
        "Concretagem",
        date(2026, 9, 5),
        120,
        1.0,
        450_000.0,
        "No prazo",
    ]
    assert _linha(aba, 10, 7)[4:6] == [0.625, 275_000.0]
    assert _linha(aba, 11, 7)[4:6] == [0.1, 100_000.0]


def test_cada_tipo_tem_o_seu_formato_e_o_seu_alinhamento() -> None:
    aba = _abrir(_documento())["Atividades"]

    assert aba["C9"].number_format == "DD/MM/YYYY"
    assert aba["D9"].number_format == "#,##0"
    assert aba["E9"].number_format == "0.0%"
    assert aba["F9"].number_format == FORMATO_DINHEIRO
    assert [aba[f"{coluna}9"].alignment.horizontal for coluna in "BCDEFG"] == [
        "left",
        "center",
        "right",
        "right",
        "right",
        "left",
    ]


@pytest.mark.parametrize(
    ("tipo", "digitos", "formato"),
    [
        (ValueKind.DECIMAL, 0, "#,##0"),
        (ValueKind.DECIMAL, 2, "#,##0.00"),
        (ValueKind.DECIMAL, 3, "#,##0.000"),
        (ValueKind.PERCENT, 0, "0%"),
        (ValueKind.PERCENT, 1, "0.0%"),
    ],
)
def test_as_casas_decimais_da_coluna_vao_para_o_formato(
    tipo: ValueKind, digitos: int, formato: str
) -> None:
    tabela = Table(
        title="Casas", columns=(Column("N", tipo, digits=digitos),), rows=(row(Decimal("1.2345")),)
    )

    aba = _abrir(_documento(portfolio=False, tables=(tabela,)))["Casas"]

    assert aba["A9"].number_format == formato


def test_a_linha_de_totais_fica_abaixo_do_filtro_em_destaque() -> None:
    aba = _abrir(_documento())["Atividades"]

    assert _linha(aba, 12, 7) == [None, "Total", None, 240, None, 825_000.0, None]
    assert aba["B12"].font.b is True
    assert _rgb(aba["B12"].fill.fgColor) == _cor_do_token("--neutro-50")
    assert _rgb(aba["B12"].border.top.color) == _cor_do_token("--borda")


def test_tabela_sem_linhas_tem_so_o_cabecalho_com_o_filtro() -> None:
    vazia = Table(title="Vazia", columns=(Column("A"), Column("B")), rows=())

    aba = _abrir(_documento(portfolio=False, tables=(vazia,)))["Vazia"]

    assert _linha(aba, 8, 2) == ["A", "B"]
    assert aba.auto_filter.ref == "A8:B8"
    assert aba.max_row == 8


def test_a_largura_acompanha_o_conteudo_entre_o_minimo_e_o_maximo() -> None:
    texto_longo = "palavra " * 40
    tabela = Table(
        title="Larguras",
        columns=(Column("Id"), Column("Descrição"), Column("Fixa", width=34)),
        rows=(row("1", texto_longo, "x"),),
    )

    aba = _abrir(_documento(portfolio=False, tables=(tabela,)))["Larguras"]

    assert aba.column_dimensions["A"].width == 10
    assert aba.column_dimensions["B"].width == 60
    assert aba.column_dimensions["C"].width == 34


# ── As cores das células vêm dos tokens ──────────────────────────────────


@pytest.mark.parametrize("tom", list(Tone))
def test_a_cor_de_cada_tom_e_a_da_pilula_do_design_system(tom: Tone) -> None:
    tabela = Table(
        title="Tons",
        columns=(Column("Situação"),),
        rows=(row(Cell("Texto", tom)), row("Sem tom")),
    )
    fundo, texto = _cores_da_pilula(SUFIXO_DO_TOM[tom])

    aba = _abrir(_documento(portfolio=False, tables=(tabela,)))["Tons"]

    assert _rgb(aba["A9"].fill.fgColor) == fundo
    assert _rgb(aba["A9"].font.color) == texto
    assert aba["A9"].font.b is True
    assert aba["A10"].fill.fill_type is None
    assert _rgb(aba["A10"].font.color) == _cor_do_token("--text-primary")


def test_nenhuma_cor_esta_escrita_no_codigo_da_exportacao() -> None:
    codigo = Path(excel.__file__).parent
    hexadecimal = re.compile(r"#[0-9A-Fa-f]{6}\b")

    for modulo in ("excel.py", "export_document.py", "printable.py"):
        texto = (codigo / modulo).read_text(encoding="utf-8")
        assert hexadecimal.search(texto) is None, modulo


# ── Texto é texto, e o que a planilha recusa não chega nela ──────────────


def test_texto_que_comeca_com_igual_continua_texto_e_nao_vira_formula() -> None:
    tabela = Table(title="Texto", columns=(Column("T"),), rows=(row("=1+1"), row("+SOMA(A1)")))

    aba = _abrir(_documento(portfolio=False, tables=(tabela,)))["Texto"]

    assert (aba["A9"].value, aba["A9"].data_type) == ("=1+1", "s")
    assert (aba["A10"].value, aba["A10"].data_type) == ("+SOMA(A1)", "s")


def test_caractere_de_controle_sai_e_tabulacao_e_quebra_de_linha_ficam() -> None:
    tabela = Table(title="Controle", columns=(Column("T"),), rows=(row("a\x07b\x1fc\td\ne"),))

    aba = _abrir(_documento(portfolio=False, tables=(tabela,)))["Controle"]

    assert aba["A9"].value == "abc\td\ne"


def test_texto_maior_que_o_limite_da_celula_e_cortado() -> None:
    tabela = Table(title="Longo", columns=(Column("T"),), rows=(row("x" * 40_000),))

    aba = _abrir(_documento(portfolio=False, tables=(tabela,)))["Longo"]

    assert len(aba["A9"].value) == 32_767


@pytest.mark.parametrize(
    ("titulos", "esperado"),
    [
        (
            ("Atividades", "Atividades", "atividades"),
            ["Atividades", "Atividades (2)", "atividades (3)"],
        ),
        (("Resumo",), ["Resumo (2)"]),
        (("x" * 40,), ["x" * 31]),
        (("y" * 40, "y" * 40), ["y" * 31, "y" * 27 + " (2)"]),
        (("A/B:C*D?E[F]G\\H",), ["A B C D E F G H"]),
        (("'entre aspas'",), ["entre aspas"]),
        (("***",), ["Planilha"]),
    ],
)
def test_nome_da_aba_cabe_no_limite_e_nao_se_repete(
    titulos: tuple[str, ...], esperado: list[str]
) -> None:
    documento = _documento(portfolio=False, tables=_tabelas_simples(*titulos))

    pasta = _abrir(documento)

    assert pasta.sheetnames == ["Resumo", *esperado]


# ── A impressão da planilha ──────────────────────────────────────────────


@pytest.mark.parametrize(("papel", "codigo"), [(Paper.A4, 9), (Paper.A3, 8)])
def test_a_planilha_imprime_na_paisagem_e_cabe_na_largura_do_papel(
    papel: Paper, codigo: int
) -> None:
    pasta = _abrir(_documento(paper=papel))

    for aba in pasta.worksheets:
        assert aba.page_setup.orientation == "landscape", aba.title
        assert aba.page_setup.paperSize == codigo, aba.title
        assert aba.page_setup.fitToWidth == 1, aba.title
        assert aba.page_setup.fitToHeight == 0, aba.title
        assert aba.sheet_properties.pageSetUpPr.fitToPage is True, aba.title


def test_o_resumo_nao_repete_cabecalho_ao_imprimir() -> None:
    assert _abrir(_documento())["Resumo"].print_title_rows is None


# ── O arquivo e o nome do download ───────────────────────────────────────


@pytest.mark.parametrize(
    ("titulo", "esperado"),
    [
        ("Mapa de controle", "mapa-de-controle-2026-10-05.xlsx"),
        (
            "Programação Semanal: Relatório (S40)",
            "programacao-semanal-relatorio-s40-2026-10-05.xlsx",
        ),
        ("!!!", "exportacao-2026-10-05.xlsx"),
        ("中文", "exportacao-2026-10-05.xlsx"),
    ],
)
def test_nome_do_arquivo_e_o_titulo_sem_acento_e_a_data(titulo: str, esperado: str) -> None:
    documento = Document(title=titulo, scope_label="Portfólio", generated_on=HOJE)

    assert excel.file_name(documento) == esperado


def test_a_resposta_e_o_arquivo_para_baixar_e_abre() -> None:
    resposta = excel.excel_response(_documento())

    assert resposta.status_code == 200
    assert resposta.mimetype == excel.XLSX_CONTENT_TYPE
    assert resposta.headers["Content-Type"] == excel.XLSX_CONTENT_TYPE
    disposicao = resposta.headers["Content-Disposition"]
    assert disposicao.startswith('attachment; filename="mapa-de-controle-2026-10-05.xlsx"')
    assert "no-store" in resposta.headers["Cache-Control"]
    assert resposta.get_body()[:2] == b"PK"
    assert load_workbook(BytesIO(resposta.get_body())).sheetnames[0] == "Resumo"
