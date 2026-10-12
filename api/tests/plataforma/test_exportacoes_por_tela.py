"""Exportações tela a tela (ISSUE-091, D12): cada rota do inventário responde e traz o que a tela mostra.

O inventário de ``docs/EXPORTACOES.md`` vira dado deste teste: cada tela declara as suas duas
rotas (o Excel e a versão imprimível), o título do documento, o papel, as colunas de cada tabela
e a referência de gestão de cada indicador. O caso parametrizado chama a rota pela mesma porta do
resto (``file_route``/``fragment_route``, o usuário do cadastro e o escopo do request) sobre a
carga de demonstração, com o relógio parado em 12/10/2026, e confere:

* **Excel**: responde 200, abre como ``.xlsx`` de verdade (openpyxl), o título do documento está
  em todas as abas, cada tabela tem o seu cabeçalho na ordem declarada — a verificação é por
  sequência de células, então basta uma coluna sumir para o caso falhar —, o papel do documento
  está no ``pageSetup`` e cada indicador aparece no ``Resumo`` com a referência de gestão.
* **Imprimível**: responde 200 como fragmento, é a folha de impressão (``folha--a4``/``folha--a3``)
  com o cabeçalho do documento, os mesmos indicadores com referência e as mesmas tabelas, e sem
  navegação nenhuma (sem ``#main-nav``, sem ``<nav>``, sem sidebar, sem link).

O teste de cobertura varre as views (``app/_views``), os templates (``api/src/templates``) e a
página do exemplo procurando os botões de exportação: a contagem por arquivo tem de ser
exatamente a declarada, cada endereço literal tem de apontar para uma rota do inventário, toda
rota registrada de exportação tem de estar no inventário e todo caso do inventário tem de estar
parametrizado — botão fora do inventário e rota sem teste reprovam. Um último caso guarda o
vínculo com ``docs/EXPORTACOES.md``: toda rota do inventário tem de estar citada lá.

A demonstração entra na transação do teste (``sessao_das_rotas``), então nada sobrevive ao caso.
As expectativas foram curadas a partir do documento que cada tela monta, com o mesmo relógio do
inventário.
"""

from __future__ import annotations

import json
import re
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import BytesIO
from itertools import pairwise
from pathlib import Path
from typing import Any

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.core import auth, calendario
from src.core.excel import XLSX_CONTENT_TYPE
from src.modulos.central_acoes.models import Minutes
from src.modulos.configuracoes import service as configuracoes
from src.modulos.configuracoes.models import Project
from src.modulos.financeiro.models import Contract
from src.modulos.governanca.models import ChangeRequest
from src.modulos.riscos.models import Risk
from tests.html_tags import find_all, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal
from tests.oraculo import harness

RAIZ = Path(__file__).resolve().parents[3]
NAVEGACAO = RAIZ / "api" / "src" / "core" / "navegacao.json"
DOC = RAIZ / "docs" / "EXPORTACOES.md"
VIEWS = RAIZ / "app" / "_views"
TEMPLATES = RAIZ / "api" / "src" / "templates"
EXEMPLO = RAIZ / "app" / "exemplos" / "exportacao.html"

HOJE = date(2026, 10, 12)
AGORA = datetime(2026, 10, 12, 12, 0, tzinfo=UTC)
PROJETO_1 = "TN-2026-014"
FOLHA = "folha-impressao"
PERIODO_DO_RELATORIO = {"tipo": "Semanal", "periodo": "2026-S38"}

# O ``paperSize`` do SpreadsheetML: 8 é A3 e 9 é A4 (``core.excel``).
PAPEL_XLSX = {"a4": 9, "a3": 8}

# A tela do exemplo não está na lista de navegação: ela só existe na demonstração.
TELAS_FORA_DA_NAVEGACAO = frozenset({"app/exemplos/exportacao.html"})

# O endereço do modelo de importação nomeia a chave no caminho; a rota registrada é uma só.
ROTA_DO_MODELO = "importacao/{chave}/modelo"


@dataclass(frozen=True)
class Exportacao:
    """Uma tela do inventário: as duas rotas, o documento esperado e onde os botões vivem.

    ``parametros`` nomeia o que a rota exige além do escopo (``numero`` do contrato,
    ``codigo`` da ficha, ``id`` da ata, ``tipo``/``periodo`` do relatório); ``chave`` só
    existe para os modelos de importação.
    """

    tela: str
    excel: str
    imprimivel: str | None
    titulo: str
    paper: str
    parametros: tuple[str, ...]
    tabelas: tuple[tuple[str, tuple[str, ...]], ...] = ()
    kpis: tuple[tuple[str, str], ...] = ()
    chave: str | None = None
    # Só a Programação Semanal separa os dois documentos: a planilha e o relatório de impressão
    # têm tabelas e indicadores próprios (``week_document`` x ``week_report_document``); nos
    # demais casos a folha é o mesmo documento do Excel e herda estas expectativas.
    tabelas_imprimivel: tuple[tuple[str, tuple[str, ...]], ...] | None = None
    kpis_imprimivel: tuple[tuple[str, str], ...] | None = None

    def endereco_do_modelo(self) -> str:
        """O endereço concreto deste caso: a chave entra no lugar do ``{chave}`` da rota."""
        return self.excel.format(chave=self.chave) if self.chave else self.excel

    def tabelas_do_papel(self) -> tuple[tuple[str, tuple[str, ...]], ...]:
        return self.tabelas if self.tabelas_imprimivel is None else self.tabelas_imprimivel

    def kpis_do_papel(self) -> tuple[tuple[str, str], ...]:
        return self.kpis if self.kpis_imprimivel is None else self.kpis_imprimivel


INVENTARIO: tuple[Exportacao, ...] = (
    Exportacao(
        tela="app/exemplos/exportacao.html",
        excel="exportacao/exemplo/excel",
        imprimivel="exportacao/exemplo/imprimivel",
        titulo="Exemplo de exportação",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Atividades de exemplo",
                ("Atividade", "Início", "Quantidade", "Avanço", "Valor", "Situação"),
            ),
            ("Legenda das situações", ("Situação", "Significado")),
        ),
        kpis=(
            ("Avanço", "Meta: 90%"),
            ("Desembolso", "Orçado: R$ 1.500.000,00"),
            ("Atividades concluídas", "Previsto: 20"),
            ("Próximo marco", "Data prevista"),
        ),
    ),
    Exportacao(
        tela="configuracoes/parametros",
        excel="configuracoes/parametros/excel",
        imprimivel="configuracoes/parametros/imprimivel",
        titulo="Parâmetros do sistema",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Parâmetros vigentes", ("Grupo", "Parâmetro", "Valor", "Módulo")),
            (
                "Histórico de versões",
                ("Versão", "Grupo", "Vigência", "Alterada por", "Justificativa", "Alterações"),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="configuracoes/colaboradores",
        excel="configuracoes/colaboradores/excel",
        imprimivel="configuracoes/colaboradores/imprimivel",
        titulo="Colaboradores",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Quem tem acesso hoje",
                (
                    "Pessoa",
                    "E-mail",
                    "Função",
                    "Perfil",
                    "Vínculo",
                    "Empresa",
                    "Situação",
                    "Papéis na Programação Semanal",
                ),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="configuracoes/cadastros",
        excel="configuracoes/cadastros/excel",
        imprimivel="configuracoes/cadastros/imprimivel",
        titulo="Cadastros de apoio",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Empresas", ("Nome", "Tipo")),
            ("Pessoas", ("Nome", "Função", "Empresa", "E-mail")),
            (
                "Projetos",
                ("Código", "Nome", "Cliente", "Gerente", "Orçamento", "Início", "Término", "Notas"),
            ),
            ("Sistemas", ("Projeto", "Código", "Nome", "Área")),
            ("Unidades", ("Tipo", "Código", "Nome")),
            ("Locais", ("Projeto", "Código", "Nome")),
            ("Disciplinas", ("Nome",)),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="inicio/home",
        excel="inicio/resumo/excel",
        imprimivel="inicio/resumo/imprimivel",
        titulo="Resumo do projeto",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Pontos de atenção", ("Módulo", "Nível", "Ponto de atenção", "Detalhe")),
            ("Módulos", ("Módulo", "Descrição", "Selos")),
        ),
        kpis=(
            ("01 Central de Ações · Ações atrasadas", "Esperado 0 · de 33 ações em aberto"),
            (
                "02 Planejamento · SPI (desempenho de prazo)",
                "Meta ≥ 1,00 · 0,12 abaixo vs. mês anterior · real 61,8% x previsto 75,8%",
            ),
            (
                "03 Gestão Financeira · CPI (desempenho de custo)",
                "Meta ≥ 1,00 · estável vs. mês anterior · projeção R$ 45,3 mi (+1,5%)",
            ),
            (
                "04 Suprimentos · Pedidos críticos",
                "Esperado 0 · folga negativa em relação ao ROS · 2 de longo prazo",
            ),
            ("05 Gestão de Riscos · Riscos críticos", "Esperado 0 · 1 alto · exposição R$ 3,4 mi"),
            (
                "06 Gestão da Qualidade · Não conformidades abertas",
                "Limite 0 vencidas · 1 vencida · inspeções aprovadas 85,0%",
            ),
            (
                "07 HSE · Dias sem acidente com afastamento",
                "Esperado desde o início do projeto · TF 0,00",
            ),
            (
                "08 Governança · Mudanças aguardando comitê",
                "Esperado 0 · R$ 1,2 mi aprovados (2,7% do orçamento)",
            ),
        ),
    ),
    Exportacao(
        tela="relatorio/gerencial",
        excel="relatorio/gerencial/excel",
        imprimivel=None,
        titulo="Relatório gerencial semanal",
        paper="a4",
        parametros=("tipo", "periodo"),
        tabelas=(
            ("Avanço por área", ("Área", "Peso", "Previsto", "Real", "Desvio (p.p.)")),
            ("Atividades do período", ("Atividade",)),
            ("Atividades do próximo período", ("Atividade",)),
            (
                "Pontos de atenção e riscos",
                ("Nº", "Ponto de atenção", "Risco", "Descrição do risco"),
            ),
            (
                "Valor agregado mês a mês",
                ("Mês", "Planejado (PV)", "Agregado (EV)", "Real (AC)", "CPI", "SPI"),
            ),
            (
                "EAC por pacote",
                ("Pacote", "Orçado", "Comprometido", "Realizado", "Projeção", "Desvio"),
            ),
            ("Pedidos críticos", ("Pedido", "Item", "ROS", "Previsão", "Folga")),
            ("Realizado no período", ("Data", "Marco", "Pedido / pacote", "Item", "Desvio")),
            (
                "Previsto nas próximas 4 semanas",
                ("Data", "Marco", "Pedido / pacote", "Item", "Desvio"),
            ),
            (
                "Principais riscos",
                ("Nº", "Risco", "Residual", "Estratégia", "Dono", "Próx. revisão"),
            ),
            ("Inspeções reprovadas no período", ("Data", "Ponto", "Empresa", "RNC")),
            (
                "Pauta de tratamento",
                ("RNC", "Não conformidade", "Severidade", "Situação", "Motivo"),
            ),
            ("Análises do período", ("Módulo", "Tipo", "Período", "Análise", "Atualizada")),
            ("Desvios negativos e comentários", ("Módulo", "Indicador", "Desvio", "Comentário")),
        ),
        kpis=(
            ("Planejamento · Avanço previsto (acumulado)", "Linha de base · 20/09/2026"),
            ("Planejamento · Avanço real (acumulado)", "Previsto 61,7% · desvio 0,1 p.p."),
            ("Planejamento · SPI físico", "Meta ≥ 1,00 · período anterior 1,05"),
            ("Planejamento · Avanço na semana", "Previsto 2,9 p.p. · desvio -2,9 p.p."),
            ("Planejamento · Término pela tendência", "Linha de base · mar/27"),
            ("Planejamento · Desvio acumulado", "Meta ≥ 0 · faixas \u22122 e \u22125 p.p."),
            ("Planejamento · SPI de quantidades", "Meta ≥ 1,00 · horas ganhas ÷ planejadas"),
            ("Planejamento · Fator de produtividade (FP)", "Meta ≤ 1,00 · HH aprop. ÷ ganhas"),
            ("Planejamento · Aderência semanal", "Meta ≥ 90% · real ÷ LB"),
            ("Planejamento · Pessoas trabalhando", "Meta ≥ 60% · amostragem"),
            ("Planejamento · Utilização da jornada", "Meta ≥ 75% · horas efetivas"),
            ("Planejamento · HH paralisadas", "Limite ≤ 5% · 481 HH na janela"),
            ("Financeiro · Orçamento vigente (BAC)", "EAC revisão vigente"),
            (
                "Financeiro · Realizado acumulado",
                "Previsto R$ 24.084 mil · R$ 5.200 mil no mês · plan. R$ 5.307 mil",
            ),
            ("Financeiro · Valor agregado (EV)", "Previsto R$ 24.084 mil · planejado (PV)"),
            ("Financeiro · CPI", "Meta ≥ 1,00 · mês anterior 0,99 · SPI 0,93"),
            ("Financeiro · Projeção no término", "Orçado R$ 44.600 mil · +R$ 655 mil x orçamento"),
            ("Financeiro · Contingência consumida", "R$ 400 mil de R$ 2.000 mil"),
            (
                "Suprimentos · Aderência ao plano de compras",
                "Meta ≥ 90% · 13 de 14 previstos até o corte",
            ),
            (
                "Suprimentos · Pacotes adjudicados",
                "de 14 · Previsto 14 · R$ 39.000 mil adjudicados",
            ),
            ("Suprimentos · Saving", "Meta ≥ 0% · R$ 2.150 mil sobre a estimativa"),
            ("Suprimentos · Entregas no prazo (OTD)", "Meta ≥ 90% · 2 de 3 entregas até o corte"),
            ("Suprimentos · Pedidos críticos", "Esperado 0 · folga negativa · 1 em atenção"),
            ("Suprimentos · Pedidos emitidos", "Linha de base 14 · até o corte"),
            ("Riscos · Riscos ativos", "Referência de 9 cadastrados · 6 ameaças · 1 oportunidades"),
            ("Riscos · Críticos", "Meta 0 · 1 ameaças · 1 oportunidades"),
            ("Riscos · Altos", "Meta 1 · 1 ameaças · 0 oportunidades"),
            ("Riscos · Exposição (VME)", "Limite ≤ R$ 1,6 mi · ameaças ativas"),
            ("Riscos · Revisão vencida", "Esperado 0 · 4 na pauta de escalonamento"),
            (
                "Riscos · Identificados no período",
                "Referência de 9 cadastrados · 0 encerrados no período",
            ),
            ("Qualidade · RNC em aberto no corte", "Esperado 0 · 0 críticas · 1 vencidas"),
            ("Qualidade · RNC abertas no período", "Esperado 0 · 0 encerradas no período"),
            ("Qualidade · Aprovação em inspeções", "Meta ≥ 95% · 1 reprovadas em 3 inspeções"),
            (
                "Qualidade · Conformidade em auditorias",
                "Meta ≥ 90% · 1 auditorias no período · 1 atrasadas",
            ),
            ("Qualidade · Custo da não qualidade", "Esperado R$ 0,00 · R$ 0,00 em RNC do período"),
            (
                "Qualidade · Disciplina com mais RNC",
                "Referência de 6 RNCs · 3 RNC · 50,0% do total",
            ),
            ("HSE · Taxa de frequência (TF)", "Esperado 0,00 · no mês · acumulado 0,00"),
            ("HSE · Lesões registráveis (TRIF)", "Esperado 0,00 · no mês · acumulado 8,88"),
            ("HSE · Taxa de gravidade (TG)", "Esperado 0,0 · no mês · acumulado 0,0"),
            ("HSE · Alto potencial (HiPo)", "Esperado 0 · no mês · 2 no acumulado"),
            ("HSE · Dias sem afastamento", "Esperado 258 · no corte · desde 05/01/2026"),
            ("HSE · Ocorrências na semana", "Referência de 39 no mês · 8 quase acidentes"),
            ("HSE · DDS realizados", "Meta 100% · no mês 94,1%"),
            ("HSE · Conformidade em inspeções", "Esperado 100% · no mês 87,1%"),
            ("HSE · Observações / 10 mil HHT", "Meta ≥ 40 · abordagens comportamentais"),
            ("HSE · Relato de quase acidentes", "Referência 54,5 · por lesão registrável"),
            ("HSE · Recomendações fechadas", "Esperado 100% · APR e HAZOP"),
            ("HSE · Incidentes ambientais", "Esperado 0 · 1 no mês"),
        ),
    ),
    Exportacao(
        tela="central_acoes/acoes",
        excel="central-acoes/acoes/excel",
        imprimivel="central-acoes/acoes/imprimivel",
        titulo="Central de Ações",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Ações",
                (
                    "Origem",
                    "Referência",
                    "Item",
                    "Assunto",
                    "Descrição",
                    "Grupo",
                    "Solicitante",
                    "Responsável",
                    "Prazo vigente",
                    "Prevista",
                    "Replanejada",
                    "Conclusão",
                    "Status",
                    "Dias de atraso",
                ),
            ),
        ),
        kpis=(
            ("Em dia", "Esperado: 33"),
            ("Atrasadas", "Meta: 0"),
            ("Concluídas", "Previsto: 45"),
            ("Total de ações", "Referência: 52"),
        ),
    ),
    Exportacao(
        tela="central_acoes/atas",
        excel="central-acoes/atas/excel",
        imprimivel="central-acoes/atas/imprimivel",
        titulo="Central de Ações · Atas",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Atas",
                (
                    "Número",
                    "Rev",
                    "Data",
                    "Assunto",
                    "Empresa principal",
                    "Tipo de reunião",
                    "Ações abertas",
                    "Atrasadas",
                ),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="central_acoes/ata",
        excel="central-acoes/ata/excel",
        imprimivel="central-acoes/ata/imprimivel",
        titulo="Ata TN-2026-0028 Rev 0",
        paper="a4",
        parametros=("id",),
        tabelas=(
            ("Empresas executoras", ("Empresa", "Principal")),
            (
                "Lista de Presença",
                ("Nome", "Função", "Empresa", "E-mail", "Ações abertas nesta ata"),
            ),
            (
                "Anotações e Ações",
                (
                    "Item",
                    "Grupo",
                    "Tipo",
                    "Assunto / Descrição",
                    "Responsável",
                    "Prevista",
                    "Replanejada",
                    "Status",
                ),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="central_acoes/dashboard",
        excel="central-acoes/painel/excel",
        imprimivel="central-acoes/painel/imprimivel",
        titulo="Central de Ações: dashboards e KPIs",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Status por origem", ("Origem", "Em dia", "Atrasadas", "Concluídas", "Total")),
            (
                "Desempenho por responsável",
                (
                    "Responsável",
                    "Abertas",
                    "Atrasadas",
                    "% atrasadas",
                    "Concluídas",
                    "Maior atraso (dias)",
                ),
            ),
            ("Previstas x concluídas por mês", ("Mês", "Previstas", "Concluídas")),
        ),
        kpis=(
            ("Total de ações", "Referência: 52"),
            ("Em dia", "Esperado: 33"),
            ("Atrasadas", "Meta: 0"),
            ("Concluídas", "Previsto: 45"),
            ("Concluídas no prazo original", "Esperado: 100%"),
        ),
    ),
    Exportacao(
        tela="planejamento/relato",
        excel="planejamento/relatos/excel",
        imprimivel="planejamento/relatos/imprimivel",
        titulo="Relato do período",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Relatos registrados",
                (
                    "Período",
                    "Atividades do período",
                    "Próximo período",
                    "Pontos de atenção",
                    "Atualizado",
                ),
            ),
            ("Atividades", ("Tipo", "Período", "Campo", "Atividade")),
            (
                "Pontos de atenção e riscos",
                ("Tipo", "Período", "Ponto de atenção", "Natureza do risco", "Risco atrelado"),
            ),
        ),
        kpis=(
            ("Semana anterior", "Esperado: Registrado"),
            ("Mês anterior", "Esperado: Registrado"),
            ("Pontos de atenção (último semanal)", "Referência: 2"),
            ("Relatos registrados", "Previsto: 49"),
        ),
    ),
    Exportacao(
        tela="planejamento/6wla",
        excel="planejamento/6wla/excel",
        imprimivel="planejamento/6wla/imprimivel",
        titulo="6WLA: planejamento de 6 semanas",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Grade de 6 semanas",
                (
                    "ID",
                    "Atividade",
                    "Responsável",
                    "S43 19/10",
                    "S44 26/10",
                    "S45 02/11",
                    "S46 09/11",
                    "S47 16/11",
                    "S48 23/11",
                    "Restrições",
                ),
            ),
            (
                "Restrições",
                (
                    "Atividade",
                    "Tipo",
                    "Restrição",
                    "Responsável",
                    "Necessária até",
                    "Removida em",
                    "Situação",
                ),
            ),
        ),
        kpis=(
            ("Atividades no horizonte", "Referência: S43 (19/10) a S48 (23/11)"),
            ("Prontas nas 2 próximas semanas", "Meta: 7"),
            ("Restrições abertas", "Esperado: 0"),
            ("Restrições vencidas", "Esperado: 0"),
            ("Índice de remoção", "Meta: 100%"),
        ),
    ),
    Exportacao(
        tela="planejamento/eap",
        excel="planejamento/eap/excel",
        imprimivel="planejamento/eap/imprimivel",
        titulo="EAP: estrutura analítica do projeto",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Estrutura analítica do projeto",
                (
                    "Código",
                    "Descrição",
                    "Critério de medição",
                    "Peso",
                    "Previsto",
                    "Real",
                    "Desvio (p.p.)",
                    "Início LB",
                    "Término LB",
                    "Situação",
                    "Item da EAC",
                    "Empresa",
                    "Responsável",
                ),
            ),
            (
                "Revisões da EAP",
                ("Revisão", "Data", "Pacotes", "Alteração", "SM", "Justificativa", "Aprovado por"),
            ),
            (
                "Desdobramentos da revisão vigente",
                (
                    "Data",
                    "Pacote de planejamento",
                    "Pacote de trabalho",
                    "Peso",
                    "Justificativa",
                    "Registrado por",
                ),
            ),
        ),
        kpis=(
            ("Avanço físico real", "Previsto: 65,7%"),
            ("Pacotes de trabalho", "Linha de base: 38"),
            ("Término vencido", "Esperado: 0"),
            ("SMs a incorporar", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="planejamento/curva_s",
        excel="planejamento/curva-fisica/excel",
        imprimivel="planejamento/curva-fisica/imprimivel",
        titulo="Curva S física",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Curva S período a período",
                (
                    "Mês",
                    "Previsto no mês (p.p.)",
                    "Previsto acumulado",
                    "Real no mês (p.p.)",
                    "Real acumulado",
                    "Tendência",
                    "Desvio (p.p.)",
                ),
            ),
        ),
        kpis=(
            ("Previsto acumulado", "Linha de base: 100% em mar/27"),
            ("Real acumulado", "Previsto: 75,8%"),
            ("Desvio", "Meta: ≥ 0"),
            ("SPI", "Meta: ≥ 1,00"),
            ("Término pela tendência", "Linha de base: mar/27"),
        ),
    ),
    Exportacao(
        tela="planejamento/kpis",
        excel="planejamento/kpis/excel",
        imprimivel="planejamento/kpis/imprimivel",
        titulo="KPIs de prazo",
        paper="a4",
        parametros=(),
        tabelas=(
            ("SPI por mês", ("Mês", "SPI")),
            (
                "Avanço por área",
                (
                    "Área",
                    "Peso",
                    "Previsto",
                    "Real",
                    "Desvio (p.p.)",
                    "SPI da área",
                    "Contribuição (p.p.)",
                ),
            ),
        ),
        kpis=(
            ("SPI do projeto", "Meta: ≥ 1,00"),
            ("Previsto x real", "Previsto: 75,8%"),
            ("Desvio acumulado", "Meta: ≥ 0"),
            ("Maior desvio por área", "Limite: ≥ -2"),
            ("Pacotes com término vencido", "Esperado: 0"),
            ("SMs a incorporar", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="planejamento/produtividade",
        excel="planejamento/produtividade/excel",
        imprimivel="planejamento/produtividade/imprimivel",
        titulo="Produtividade: plano de quantidades da linha de base",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Quantidades por grupo",
                (
                    "Grupo",
                    "Unidade",
                    "Itens",
                    "Total da LB",
                    "Previsto até o corte",
                    "Realizado acumulado",
                    "% real (prev.)",
                    "Previsto na semana",
                    "Realizado na semana",
                ),
            ),
            (
                "Plano x realizado por item",
                (
                    "Item",
                    "Empresa",
                    "Grupo",
                    "Unidade",
                    "Total da LB",
                    "Índice HH por unidade",
                    "Realizado acumulado",
                    "Previsto até o corte",
                    "% real (prev.)",
                    "Realizado na semana",
                    "Previsto na semana",
                    "Aderência da semana",
                    "FP",
                    "Término pela tendência",
                    "Situação da LB",
                ),
            ),
        ),
        kpis=(
            ("Avanço por quantidades", "Previsto: 79,4%"),
            ("Aderência da semana", "Meta: ≥ 90%"),
            ("Fator de produtividade", "Meta: ≤ 1,00"),
            ("Apontamentos pendentes", "Esperado: 0"),
            ("Tendência de atraso", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="planejamento/produtividade",
        excel="planejamento/produtividade/horas-efetivas/excel",
        imprimivel="planejamento/produtividade/horas-efetivas/imprimivel",
        titulo="Produtividade: horas efetivas e amostragem do trabalho",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Registros de jornada",
                (
                    "Data",
                    "Empresa",
                    "Área (CWA)",
                    "Encarregado",
                    "Efetivo",
                    "Manhã (chegada · início · término)",
                    "Tarde (chegada · início · término)",
                    "Atraso (min)",
                    "CP (h/dia)",
                    "Utilização",
                    "HH efetivas",
                    "HH improdutivas",
                ),
            ),
            (
                "Capacidade produtiva por área (CWA)",
                (
                    "Área (CWA)",
                    "CP (h/dia)",
                    "Utilização",
                    "Efetivo paralisado (Hhora)",
                    "Máq./equip. parados (Mhora)",
                    "Jornadas",
                ),
            ),
            ("Pareto dos motivos de parada", ("Motivo", "Pessoas", "%", "% acumulado")),
            ("Motivos em trânsito", ("Motivo", "Pessoas", "%", "% acumulado")),
            (
                "Amostragem por empresa e encarregado",
                (
                    "Empresa / encarregado",
                    "Rodadas",
                    "Pessoas observadas",
                    "% trabalhando",
                    "% em trânsito",
                    "% parado",
                ),
            ),
            (
                "Rodadas de observação",
                (
                    "Data",
                    "Hora",
                    "Empresa",
                    "Área (CWA)",
                    "Encarregado",
                    "Trabalhando",
                    "Em trânsito",
                    "Parado",
                    "% trabalhando",
                ),
            ),
            (
                "Registros de paralisação",
                (
                    "Data",
                    "Empresa",
                    "Área (CWA)",
                    "Tipo",
                    "Recurso parado",
                    "Quantidade",
                    "Início",
                    "Término",
                    "Horas",
                    "Hhora / Mhora",
                    "Motivo",
                    "Responsabilidade",
                ),
            ),
            ("Efetivo paralisado por motivo", ("Chave", "Hhora", "Eventos")),
            ("Efetivo paralisado por responsabilidade", ("Chave", "Hhora", "Eventos")),
        ),
        kpis=(
            ("Capacidade produtiva (CP)", "Meta: ≥ 6,60 h"),
            ("Utilização da jornada", "Meta: ≥ 75%"),
            ("Atraso médio de início", "Limite: ≤ 15 min"),
            ("HH efetivas", "Esperado: 0 improdutivas"),
            ("HH improdutivas na frente", "Esperado: 0"),
            ("Trabalhando", "Meta: ≥ 60%"),
            ("Em trânsito", "Esperado: 0 parado"),
            ("Parado", "Esperado: 0"),
            ("Efetivo paralisado", "Esperado: 0"),
            ("Máquinas e equipamentos parados", "Esperado: 0"),
            ("Responsabilidade do cliente", "Esperado: 0"),
            ("Clima", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="planejamento/produtividade",
        excel="planejamento/produtividade/kpis/excel",
        imprimivel="planejamento/produtividade/kpis/imprimivel",
        titulo="Produtividade: KPIs de performance",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Performance por empresa",
                (
                    "Empresa",
                    "Avanço real x prev.",
                    "SPI",
                    "Aderência",
                    "FP",
                    "CP (h/dia)",
                    "% trabalhando",
                    "Hhora parada",
                    "Mhora parada",
                    "Situação",
                ),
            ),
            (
                "Tendência semanal por empresa",
                ("Semana", "Empresa", "FP", "Aderência", "% trabalhando", "CP (h/dia)"),
            ),
        ),
        kpis=(
            ("Avanço por quantidades", "Previsto: 79,4%"),
            ("SPI de quantidades", "Meta: ≥ 0,95"),
            ("Aderência (4 semanas)", "Meta: ≥ 90%"),
            ("Fator de produtividade", "Meta: ≤ 1,00"),
            ("Tendência de atraso", "Esperado: 0"),
            ("Capacidade produtiva (CP)", "Meta: ≥ 6,60 h"),
            ("Pessoas trabalhando", "Meta: ≥ 60%"),
            ("Efetivo paralisado", "Esperado: 0"),
            ("Máquinas e equipamentos parados", "Esperado: 0"),
            ("HH improdutivas na frente", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="planejamento/punch_list",
        excel="planejamento/punch/excel",
        imprimivel="planejamento/punch/imprimivel",
        titulo="Punch list",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Itens da Punch list",
                (
                    "Nº",
                    "Categoria",
                    "Sistema",
                    "Subsistema",
                    "TAG",
                    "Disciplina",
                    "Descrição",
                    "Marco",
                    "Origem",
                    "Responsável",
                    "Empresa",
                    "Abertura",
                    "Prazo",
                    "Idade (dias)",
                    "Situação",
                    "Fechamento",
                ),
            ),
        ),
        kpis=(
            ("Abertos", "Esperado: 0"),
            ("Categoria A abertos", "Esperado: 0"),
            ("Aguardando verificação", "Esperado: 0"),
            ("Vencidos", "Esperado: 0"),
            ("Fechados", "Previsto: 15"),
        ),
    ),
    Exportacao(
        tela="planejamento/punch_list",
        excel="planejamento/punch/painel/excel",
        imprimivel="planejamento/punch/painel/imprimivel",
        titulo="Punch list · painel de completação",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Abertura e fechamento acumulados (burndown)",
                ("Semana", "Abertos (acumulado)", "Fechados (acumulado)", "Saldo em aberto"),
            ),
            (
                "Tempo em aberto por faixa e categoria",
                ("Faixa", "Categoria A", "Categoria B", "Categoria C", "Total"),
            ),
            (
                "Abertos por disciplina e categoria",
                ("Disciplina", "Categoria A", "Categoria B", "Categoria C", "Total"),
            ),
            (
                "Abertos por empresa e categoria",
                ("Empresa", "Categoria A", "Categoria B", "Categoria C", "Total"),
            ),
            (
                "Sistemas e liberação por marco",
                (
                    "Sistema",
                    "Área",
                    "A abertos",
                    "B abertos",
                    "C abertos",
                    "Completação mecânica",
                    "Comissionamento",
                    "Aceite definitivo",
                ),
            ),
        ),
        kpis=(
            ("Sistemas liberados · Completação mecânica", "Meta: 100%"),
            ("Sistemas liberados · Comissionamento", "Meta: 100%"),
            ("Sistemas liberados · Aceite definitivo", "Meta: 100%"),
            ("Sistemas bloqueados", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="programacao_semanal/governanca",
        excel="programacao-semanal/governanca/excel",
        imprimivel="programacao-semanal/governanca/imprimivel",
        titulo="Governança da programação",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Lista de controle",
                (
                    "Contratada",
                    "Semana",
                    "Programações",
                    "Em elaboração",
                    "Validadas",
                    "Publicadas",
                    "Pendências",
                    "Aderência",
                    "Fiscal",
                    "Última alteração",
                ),
            ),
            (
                "Pedidos de alteração",
                (
                    "ID exclusiva",
                    "Contratada",
                    "Semana",
                    "Motivo",
                    "Situação",
                    "Solicitado por",
                    "Solicitado em",
                    "Decidido por",
                    "Decidido em",
                    "Resposta",
                ),
            ),
        ),
        kpis=(
            ("Contratadas ativas", "Contratadas com atividade no histórico"),
            ("Janelas abertas hoje", "Na semana S.42/2026"),
            ("Aguardando fiscal", "Realizado lançado sem aprovação"),
            ("Atividades no histórico", "Todas as semanas"),
        ),
    ),
    Exportacao(
        tela="programacao_semanal/programacao",
        excel="programacao-semanal/semana/excel",
        imprimivel="programacao-semanal/semana/imprimivel",
        titulo="Programação da semana S.42/2026",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Programação",
                (
                    "Item",
                    "ID Exclusiva",
                    "Atividade",
                    "Local",
                    "Empresa",
                    "Encarregado",
                    "Fiscal",
                    "Prev/Real",
                    "2ª",
                    "3ª",
                    "4ª",
                    "5ª",
                    "6ª",
                    "Sáb",
                    "Dom",
                    "Total",
                    "Unidade",
                    "PPC %",
                    "Situação",
                    "Realizado",
                ),
            ),
            (
                "Por contratada",
                (
                    "Contratada",
                    "Atividades",
                    "Previsto",
                    "Realizado",
                    "Aderência %",
                    "PPC médio %",
                    "Pendências",
                ),
            ),
            (
                "Por frente",
                (
                    "Frente",
                    "Atividades",
                    "Previsto",
                    "Realizado",
                    "Aderência %",
                    "PPC médio %",
                    "Pendências",
                ),
            ),
            (
                "Por encarregado",
                (
                    "Encarregado",
                    "Atividades",
                    "Previsto",
                    "Realizado",
                    "Aderência %",
                    "PPC médio %",
                    "Pendências",
                ),
            ),
            (
                "Dia a dia",
                (
                    "Dia",
                    "Previsto",
                    "Realizado dia",
                    "Realizado noite",
                    "Realizado total",
                    "Aderência %",
                ),
            ),
            (
                "Curva S",
                (
                    "Semana",
                    "Período",
                    "Base %",
                    "Previsto %",
                    "Realizado %",
                    "Base acum. %",
                    "Previsto acum. %",
                    "Realizado acum. %",
                ),
            ),
        ),
        kpis=(
            ("Aderência geral", "meta 60%"),
            ("PPC médio", "meta 75%"),
            ("Atividades", "na semana"),
            ("Contratadas", "com programação"),
            ("Frentes", "ativas"),
            ("Realizados pendentes", "aguardando o fiscal"),
            ("Em elaboração", "na semana"),
            ("Validada", "na semana"),
            ("Publicada", "na semana"),
            ("Realizado aprovado", "na semana"),
            ("Sem realizado", "na semana"),
        ),
        tabelas_imprimivel=(
            (
                "Programação",
                (
                    "Item",
                    "ID Exclusiva",
                    "Atividade",
                    "Local",
                    "Empresa",
                    "Encarregado",
                    "Fiscal",
                    "Prev/Real",
                    "2ª",
                    "3ª",
                    "4ª",
                    "5ª",
                    "6ª",
                    "Sáb",
                    "Dom",
                    "Total",
                    "Unidade",
                    "PPC %",
                    "Situação",
                    "Realizado",
                ),
            ),
            (
                "Por contratada",
                (
                    "Contratada",
                    "Atividades",
                    "Previsto",
                    "Realizado",
                    "Aderência %",
                    "PPC médio %",
                    "Pendências",
                ),
            ),
            (
                "Por frente",
                (
                    "Frente",
                    "Atividades",
                    "Previsto",
                    "Realizado",
                    "Aderência %",
                    "PPC médio %",
                    "Pendências",
                ),
            ),
            (
                "Dia a dia",
                (
                    "Dia",
                    "Previsto",
                    "Realizado dia",
                    "Realizado noite",
                    "Realizado total",
                    "Aderência %",
                ),
            ),
        ),
        kpis_imprimivel=(
            ("Aderência", "meta 60%"),
            ("PPC médio", "meta 75%"),
            ("Atividades", "na semana"),
            ("Contratadas", "com programação"),
        ),
    ),
    Exportacao(
        tela="programacao_semanal/dashboard",
        excel="programacao-semanal/dashboard/excel",
        imprimivel="programacao-semanal/dashboard/imprimivel",
        titulo="Dashboard da programação",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Ranking por contratada",
                (
                    "Contratada",
                    "Atividades",
                    "Previsto",
                    "Realizado",
                    "Aderência %",
                    "PPC médio %",
                    "Pendências",
                ),
            ),
            (
                "Ranking por frente",
                (
                    "Frente",
                    "Atividades",
                    "Previsto",
                    "Realizado",
                    "Aderência %",
                    "PPC médio %",
                    "Pendências",
                ),
            ),
            (
                "Ranking por encarregado",
                (
                    "Encarregado",
                    "Atividades",
                    "Previsto",
                    "Realizado",
                    "Aderência %",
                    "PPC médio %",
                    "Pendências",
                ),
            ),
            (
                "O que travou a semana",
                (
                    "ID exclusiva",
                    "Atividade",
                    "Contratada",
                    "Frente",
                    "Encarregado",
                    "Previsto",
                    "Realizado",
                    "PPC %",
                    "Dias parados",
                ),
            ),
            (
                "Mapa de calor — dia por frente",
                ("Frente", "2ª", "3ª", "4ª", "5ª", "6ª", "Sáb", "Dom"),
            ),
            (
                "Produção por dia",
                ("Dia", "Previsto", "Turno dia", "Turno noite", "Realizado", "Aderência %"),
            ),
            (
                "Curva S",
                (
                    "Semana",
                    "Período",
                    "Base %",
                    "Previsto %",
                    "Realizado %",
                    "Base acum. %",
                    "Previsto acum. %",
                    "Realizado acum. %",
                ),
            ),
            ("Previsto contra realizado", ("Semana", "Período", "Previsto %", "Realizado %")),
            (
                "Evolução da aderência por contratada",
                ("Semana", "Bremmer", "Elmec Montagens", "M. Roscoe"),
            ),
        ),
        kpis=(
            ("Aderência", "meta 60%"),
            ("PPC médio", "meta 75%"),
            ("Atividades", "na semana"),
            ("Contratadas", "com programação"),
        ),
    ),
    Exportacao(
        tela="programacao_semanal/importacao",
        excel="programacao-semanal/importacao/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Programação da semana",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Dados",
                (
                    "Item",
                    "Semana",
                    "ID Exclusiva",
                    "Atividade Detalhada",
                    "Local",
                    "Empresa",
                    "Fiscal",
                    "Encarregado",
                    "Prod. Prevista",
                    "Unidade",
                    "Prev/Real",
                    "2ª",
                    "3ª",
                    "4ª",
                    "5ª",
                    "6ª",
                    "Sáb",
                    "Dom",
                ),
            ),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="financeiro/eac",
        excel="financeiro/eac/excel",
        imprimivel="financeiro/eac/imprimivel",
        titulo="EAC: estrutura analítica de custos",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Estrutura analítica de custos",
                (
                    "Código",
                    "Descrição",
                    "Tipo de custo",
                    "Un.",
                    "Qtd.",
                    "Preço unitário",
                    "Valor orçado",
                    "CAPEX/OPEX",
                    "Centro de custo",
                    "Responsável",
                ),
            ),
            (
                "Revisões do orçamento",
                ("Revisão", "Data", "Total", "Variação", "SM", "Justificativa", "Aprovado por"),
            ),
            (
                "Remanejamentos",
                (
                    "Data",
                    "Origem",
                    "Destino",
                    "Valor",
                    "SM",
                    "Situação",
                    "Saldo livre da origem",
                    "Justificativa",
                    "Registrado por",
                ),
            ),
        ),
        kpis=(
            ("Orçamento vigente", "Quantidade x preço unitário dos itens"),
            ("Itens de custo", "3 pacotes"),
        ),
    ),
    Exportacao(
        tela="financeiro/mapa_controle",
        excel="financeiro/mapa-de-controle/excel",
        imprimivel="financeiro/mapa-de-controle/imprimivel",
        titulo="Mapa de controle da EAC",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Mapa de controle",
                (
                    "Código",
                    "Descrição",
                    "Orçado",
                    "Remanejamentos",
                    "Orçado atual",
                    "Comprometido",
                    "Realizado",
                    "Saldo a comprometer",
                    "Projeção no término",
                    "Desvio",
                    "Desvio %",
                ),
            ),
        ),
        kpis=(
            ("Orçado atual", "Linha de base mais remanejamentos"),
            ("Comprometido", "Fechamento do mês do ERP"),
            ("Realizado", "Fechamento do mês do ERP"),
            ("Projeção no término", "Estimativa vigente da equipe ou do fechamento"),
            ("Desvio no término", "Projeção no término menos orçado atual; positivo é sobrecusto"),
        ),
    ),
    Exportacao(
        tela="financeiro/contrato",
        excel="financeiro/contrato/excel",
        imprimivel="financeiro/contrato/imprimivel",
        titulo="Contrato CT-2026-011",
        paper="a4",
        parametros=("numero",),
        tabelas=(
            ("Composição do valor", ("Etapa", "Valor")),
            (
                "Medições",
                (
                    "Nº",
                    "Período",
                    "Valor bruto",
                    "Retenção",
                    "Valor líquido",
                    "Situação",
                    "Motivo da devolução",
                ),
            ),
            ("Aditivos", ("Nº", "Data", "Valor", "Dias", "Motivo", "SM")),
        ),
        kpis=(
            ("Valor atual", "Linha de base: R$ 3.300.000,00"),
            ("Medido aprovado", "Disponível a medir = saldo a faturar - em análise"),
            ("Saldo a faturar", "Valor atual - medido aprovado"),
            ("Término vigente", "Original: 30/10/2026"),
        ),
    ),
    Exportacao(
        tela="financeiro/contrato",
        excel="financeiro/contrato/marcos/excel",
        imprimivel="financeiro/contrato/marcos/imprimivel",
        titulo="Contrato CT-2026-011",
        paper="a4",
        parametros=("numero",),
        tabelas=(
            (
                "Marcos de pagamento",
                (
                    "Nº",
                    "Marco",
                    "Critério de aceite",
                    "%",
                    "Valor",
                    "Retenção",
                    "Prevista",
                    "Atraso (dias)",
                    "Conclusão",
                    "Situação",
                    "Evidência",
                ),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="financeiro/contrato",
        excel="financeiro/contrato/claims/excel",
        imprimivel="financeiro/contrato/claims/imprimivel",
        titulo="Contrato CT-2026-011",
        paper="a4",
        parametros=("numero",),
        tabelas=(
            (
                "Claims",
                (
                    "Código",
                    "Direção",
                    "Tipo",
                    "Causa",
                    "Descrição",
                    "Cláusula",
                    "Evento",
                    "Notificação",
                    "Dias após o evento",
                    "Fora do prazo",
                    "Pleiteado",
                    "Dias pleiteados",
                    "Reconhecido",
                    "Dias reconhecidos",
                    "Situação",
                    "SM",
                    "Risco",
                    "Encerrado em",
                ),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="financeiro/contrato",
        excel="financeiro/contrato/extensoes/excel",
        imprimivel="financeiro/contrato/extensoes/imprimivel",
        titulo="Contrato CT-2026-011",
        paper="a4",
        parametros=("numero",),
        tabelas=(
            (
                "Extensões de prazo",
                (
                    "Código",
                    "Claim",
                    "Evento causador",
                    "Solicitados",
                    "Concedidos",
                    "Classificação",
                    "Método",
                    "Marco afetado",
                    "Solicitada em",
                    "Decidida em",
                    "Situação",
                    "SM",
                ),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="financeiro/contrato",
        excel="financeiro/contrato/avaliacoes/excel",
        imprimivel="financeiro/contrato/avaliacoes/imprimivel",
        titulo="Contrato CT-2026-011",
        paper="a4",
        parametros=("numero",),
        tabelas=(
            (
                "Avaliações",
                (
                    "Período",
                    "Tipo",
                    "Avaliador",
                    "HSE",
                    "Qualidade",
                    "Prazo",
                    "Gestão contratual e comercial",
                    "Recursos e mobilização",
                    "Documentação e comunicação",
                    "Nota",
                    "Classe",
                ),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="financeiro/contratos",
        excel="financeiro/contratos/consolidado/excel",
        imprimivel="financeiro/contratos/consolidado/imprimivel",
        titulo="Contratos: administração contratual",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Contratos",
                (
                    "Nº",
                    "Contratada",
                    "Objeto",
                    "Modalidade",
                    "Valor atual",
                    "Medido",
                    "Saldo a faturar",
                    "Término vigente",
                    "Dias aditados",
                    "Claims abertos",
                    "Última avaliação",
                ),
            ),
            (
                "Claims",
                (
                    "Código",
                    "Contrato",
                    "Direção",
                    "Tipo",
                    "Causa",
                    "Descrição",
                    "Evento",
                    "Notificação",
                    "Dias após o evento",
                    "Fora do prazo",
                    "Pleiteado",
                    "Dias pleiteados",
                    "Reconhecido",
                    "Dias reconhecidos",
                    "Situação",
                ),
            ),
            (
                "Extensões de prazo",
                (
                    "Código",
                    "Contrato",
                    "Claim",
                    "Evento causador",
                    "Solicitados",
                    "Concedidos",
                    "Classificação",
                    "Método",
                    "Marco afetado",
                    "Solicitada em",
                    "Decidida em",
                    "Situação",
                ),
            ),
            (
                "Marcos de pagamento",
                (
                    "Contrato",
                    "Nº",
                    "Marco",
                    "Critério de aceite",
                    "%",
                    "Valor",
                    "Prevista",
                    "Atraso (dias)",
                    "Conclusão",
                    "Situação",
                ),
            ),
            (
                "Avaliações",
                (
                    "Contrato",
                    "Contratada",
                    "Período",
                    "Tipo",
                    "Avaliador",
                    "Nota",
                    "Classe",
                    "Plano de melhoria",
                    "Comentário",
                ),
            ),
        ),
        kpis=(
            ("Contratos", "Contratos do escopo"),
            ("Exposição de claims", "Esperado: R$ 0,00"),
            ("Notificados fora do prazo", "Esperado: 0"),
            ("Marcos atrasados", "Esperado: 0"),
            ("Extensão de prazo", "Limite: ≤ 10%"),
            ("Contratadas C ou D", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="financeiro/contingencia",
        excel="financeiro/contingencia/excel",
        imprimivel="financeiro/contingencia/imprimivel",
        titulo="Contingência e reservas",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Linha de base de custo e reservas", ("Composição", "Valor")),
            ("Saldo real x previsto", ("Mês", "Saldo previsto", "Saldo real")),
            (
                "Extrato de movimentos",
                ("Data", "Reserva", "Movimento", "Descrição", "SM", "Situação", "Valor", "Saldo"),
            ),
            (
                "Ameaças ativas e cobertura",
                ("Risco", "Ameaça", "Severidade", "Situação", "VME", "SMs vinculadas"),
            ),
        ),
        kpis=(
            ("Saldo da contingência", "Previsto: R$ 484.000,00"),
            ("Consumo da contingência", "Limite: 71,8%"),
            ("Pedido em SMs em análise", "Saldo se aprovadas: R$ 1.120.000,00"),
            ("Cobertura da exposição a riscos", "Meta: ≥ 100%"),
            ("Reserva gerencial", "Constituída: R$ 900.000,00"),
        ),
    ),
    Exportacao(
        tela="financeiro/curva_s",
        excel="financeiro/curva-s/excel",
        imprimivel="financeiro/curva-s/imprimivel",
        titulo="Curva S financeira",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Período a período",
                (
                    "Mês",
                    "Planejado no mês",
                    "Planejado acumulado",
                    "Comprometido acumulado",
                    "Realizado no mês",
                    "Realizado acumulado",
                    "Projeção acumulada",
                    "Realizado menos planejado",
                ),
            ),
        ),
        kpis=(
            ("Planejado acumulado", "Orçado: R$ 44.600.000,00"),
            ("Realizado acumulado", "Previsto: R$ 33.806.800,00"),
            ("Comprometido acumulado", "Orçado: R$ 44.600.000,00"),
            ("Projeção no término", "Orçado: R$ 44.600.000,00"),
        ),
    ),
    Exportacao(
        tela="financeiro/kpis",
        excel="financeiro/kpis/excel",
        imprimivel="financeiro/kpis/imprimivel",
        titulo="KPIs de custo",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Valor agregado na data de corte",
                ("Indicador", "Descrição", "Valor", "Índice", "Leitura"),
            ),
            (
                "Período a período",
                ("Mês", "PV", "EV", "AC", "CV", "SV", "CPI", "SPI", "EAC pelo CPI", "TCPI"),
            ),
        ),
        kpis=(
            ("CPI (desempenho de custo)", "Meta: ≥ 1,00"),
            ("SPI de custo (desempenho de prazo)", "Meta: ≥ 1,00"),
            ("CV (variação de custo)", "Meta: ≥ 0"),
            ("VAC (variação no término)", "Meta: ≥ 0"),
            ("Contingência consumida", "Constituída: R$ 2.000.000,00"),
            ("Comprometido", "Orçado: R$ 44.600.000,00"),
        ),
    ),
    Exportacao(
        tela="financeiro/desembolso",
        excel="financeiro/desembolso/excel",
        imprimivel="financeiro/desembolso/imprimivel",
        titulo="Cronograma de desembolso",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Desembolso mensal", ("Mês", "Previsto", "Realizado", "Projetado")),
            (
                "Saldo a pagar por mês",
                (
                    "Código",
                    "Descrição",
                    "Saldo a pagar",
                    "nov/26",
                    "dez/26",
                    "jan/27",
                    "fev/27",
                    "mar/27",
                    "abr/27",
                ),
            ),
        ),
        kpis=(
            ("Realizado até o corte", "Previsto: R$ 33.806.800,00"),
            ("Saldo a desembolsar", "Previsto: R$ 10.793.200,00"),
            ("Próximos 3 meses", "Previsto: R$ 8.607.800,00"),
            ("Mês de corte", "Previsto: R$ 4.504.600,00"),
        ),
    ),
    Exportacao(
        tela="suprimentos/fornecedores",
        excel="suprimentos/fornecedores/excel",
        imprimivel="suprimentos/fornecedores/imprimivel",
        titulo="Fornecedores",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Fornecedores e contratadas",
                (
                    "Empresa",
                    "Tipo",
                    "Categorias",
                    "Situação",
                    "Validade",
                    "Documentos",
                    "Desempenho",
                    "Valor contratado",
                ),
            ),
        ),
        kpis=(
            ("Cadastrados", "Referência: de 20 empresas"),
            ("Qualificados", "Referência: de 17 cadastrados"),
            ("Em qualificação", "Referência: de 17 cadastrados"),
            ("Restritos ou bloqueados", "Esperado: 0"),
            ("Documentos vencidos", "Esperado: 0"),
            ("Qualificação a vencer", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="suprimentos/plano_compras",
        excel="suprimentos/plano-de-compras/excel",
        imprimivel="suprimentos/plano-de-compras/imprimivel",
        titulo="Plano de compras",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Pacotes de compra",
                (
                    "Pacote",
                    "Escopo",
                    "Tipo",
                    "Modalidade",
                    "Disciplina",
                    "Item da EAC",
                    "Comprador",
                    "Estimativa",
                    "Requisição (LB)",
                    "RFx (LB)",
                    "Adjudicação (LB)",
                    "Pedido (LB)",
                    "ROS",
                    "Etapa",
                    "Folga (dias)",
                    "Adjudicado",
                    "Saving",
                ),
            ),
        ),
        kpis=(
            ("Pacotes no plano", "Linha de base: 14"),
            ("Adjudicados", "Previsto até hoje: 14"),
            ("Adjudicação atrasada", "Esperado: 0"),
            ("Itens de longo prazo (LLI)", "Referência: de 14 pacotes"),
            ("Saving sobre estimativa", "Meta: ≥ R$ 0,00"),
        ),
    ),
    Exportacao(
        tela="suprimentos/processos",
        excel="suprimentos/processos/excel",
        imprimivel="suprimentos/processos/imprimivel",
        titulo="Processos de compra (RFx)",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Processos de compra",
                (
                    "Pacote",
                    "Escopo",
                    "Tipo",
                    "RFx",
                    "Etapa",
                    "Propostas",
                    "Melhor proposta",
                    "Alçada",
                    "Adjudicação (LB)",
                    "Adjudicação (real)",
                    "Adjudicação (previsão)",
                ),
            ),
        ),
        kpis=(
            ("Em andamento", "Previsto: 0"),
            ("Aguardando aprovação", "Esperado: 0"),
            ("Propostas válidas por processo", "Meta: ≥ 3"),
            ("Fornecedor único", "Esperado: 0%"),
            ("Ciclo de compra", "Linha de base: 67 dias"),
        ),
    ),
    Exportacao(
        tela="suprimentos/diligenciamento",
        excel="suprimentos/diligenciamento/excel",
        imprimivel="suprimentos/diligenciamento/imprimivel",
        titulo="Diligenciamento e recebimento",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Pedidos em diligenciamento",
                (
                    "Pedido",
                    "Pacote",
                    "Descrição",
                    "Fornecedor",
                    "Valor",
                    "Próximo marco",
                    "Data contratual",
                    "Previsão de entrega",
                    "ROS",
                    "Folga (dias)",
                    "Situação",
                ),
            ),
        ),
        kpis=(
            ("Pedidos em aberto", "Previsto: 2"),
            ("Críticos", "Esperado: 0"),
            ("Em atenção", "Esperado: 0"),
            ("Marcos vencidos", "Esperado: 0"),
            ("Entrega no prazo (OTD)", "Meta: ≥ 90,0%"),
        ),
    ),
    Exportacao(
        tela="suprimentos/mas",
        excel="suprimentos/mas/excel",
        imprimivel="suprimentos/mas/imprimivel",
        titulo="Mapa de Suprimentos (MAS)",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Mapa de Suprimentos",
                (
                    "Pacote",
                    "Escopo",
                    "Tipo",
                    "Disciplina",
                    "LLI",
                    "Item da EAC",
                    "Fornecedor",
                    "Comprador",
                    "Estimativa",
                    "Adjudicado",
                    "Requisição LB",
                    "Requisição previsão ou real",
                    "Requisição situação",
                    "RFx emitida LB",
                    "RFx emitida previsão ou real",
                    "RFx emitida situação",
                    "Propostas LB",
                    "Propostas previsão ou real",
                    "Propostas situação",
                    "Equalização técnica LB",
                    "Equalização técnica previsão ou real",
                    "Equalização técnica situação",
                    "Equalização comercial LB",
                    "Equalização comercial previsão ou real",
                    "Equalização comercial situação",
                    "Aprovação LB",
                    "Aprovação previsão ou real",
                    "Aprovação situação",
                    "Pedido/contrato LB",
                    "Pedido/contrato previsão ou real",
                    "Pedido/contrato situação",
                    "Documentos aprovados LB",
                    "Documentos aprovados previsão ou real",
                    "Documentos aprovados situação",
                    "Fabricação LB",
                    "Fabricação previsão ou real",
                    "Fabricação situação",
                    "Inspeção / FAT LB",
                    "Inspeção / FAT previsão ou real",
                    "Inspeção / FAT situação",
                    "Embarque LB",
                    "Embarque previsão ou real",
                    "Embarque situação",
                    "Entrega LB",
                    "Entrega previsão ou real",
                    "Entrega situação",
                    "Processo",
                    "Pedido ou contrato",
                    "ROS",
                    "Previsão de entrega",
                    "Folga (dias)",
                    "Avanço real (%)",
                    "Avanço LB (%)",
                    "Situação",
                ),
            ),
        ),
        kpis=(
            ("Itens no MAS", "Linha de base: 14"),
            ("Críticos", "Esperado: 0"),
            ("Em atenção", "Esperado: 0"),
            ("Marcos vencidos", "Esperado: 0"),
            ("Realizados com atraso", "Esperado: 0%"),
            ("Avanço de suprimentos", "Linha de base: 98,9%"),
        ),
    ),
    Exportacao(
        tela="suprimentos/painel",
        excel="suprimentos/painel/excel",
        imprimivel="suprimentos/painel/imprimivel",
        titulo="Painel de suprimentos",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Pedidos por folga",
                (
                    "Pacote",
                    "Escopo",
                    "Fornecedor",
                    "LLI",
                    "Próximo marco",
                    "Previsão de entrega",
                    "ROS",
                    "Folga (dias)",
                    "Situação",
                ),
            ),
        ),
        kpis=(
            ("Aderência ao plano de compras", "Meta: ≥ 90%"),
            ("Ciclo de compra", "Linha de base: 67"),
            ("Saving sobre estimativa", "Meta: ≥ R$ 0,00"),
            ("Saving de negociação", "Meta: ≥ R$ 0,00"),
            ("Competitividade", "Meta: ≥ 3"),
            ("Entrega no prazo (OTD)", "Meta: ≥ 90%"),
            ("Pedidos críticos", "Esperado: 0"),
            ("Compras emergenciais", "Esperado: 0%"),
            ("Comprometido x orçado", "Limite: 100% · Orçado: R$ 43.000.000,00"),
        ),
    ),
    Exportacao(
        tela="riscos/registro",
        excel="riscos/registro/excel",
        imprimivel="riscos/registro/imprimivel",
        titulo="Registro de riscos",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Riscos",
                (
                    "Nº",
                    "Risco",
                    "Natureza",
                    "Dono",
                    "Inerente",
                    "Residual",
                    "Estratégia",
                    "Ações",
                    "Próx. revisão",
                    "Situação",
                ),
            ),
        ),
        kpis=(
            ("Críticos (residual)", "Meta: 0"),
            ("Altos (residual)", "Meta: 1"),
            ("Em tratamento", "Esperado: 7"),
            ("Revisão vencida", "Esperado: 0"),
            ("Ativos", "Referência: de 9 registrados"),
        ),
    ),
    Exportacao(
        tela="riscos/matriz",
        excel="riscos/matriz/excel",
        imprimivel="riscos/matriz/imprimivel",
        titulo="Matriz de riscos 5x5",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Distribuição por probabilidade e impacto",
                (
                    "Probabilidade \\ Impacto",
                    "I1 Muito baixo",
                    "I2 Baixo",
                    "I3 Moderado",
                    "I4 Alto",
                    "I5 Muito alto",
                ),
            ),
            (
                "Movimentação inerente para residual",
                ("Nº", "Risco", "Inerente", "Residual", "Observação"),
            ),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="riscos/painel",
        excel="riscos/painel/excel",
        imprimivel="riscos/painel/imprimivel",
        titulo="Painel de riscos",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Exposição por categoria (RBS)",
                ("Categoria", "Score residual somado", "Riscos", "VME"),
            ),
            ("Evolução mensal do score residual", ("Mês", "Score residual somado")),
            ("Pauta de escalonamento", ("Nº", "Risco", "Residual", "Motivo", "Dono")),
        ),
        kpis=(
            ("Críticos (residual)", "Meta: 0"),
            ("Altos (residual)", "Meta: 1"),
            ("Exposição (VME)", "Ameaças ativas"),
            ("Revisão vencida", "Esperado: 0"),
            ("Redução média", "Meta: ≥ 48%"),
        ),
    ),
    Exportacao(
        tela="riscos/ficha",
        excel="riscos/ficha/excel",
        imprimivel="riscos/ficha/imprimivel",
        titulo="Ficha do risco RSK-TN-2026-0001",
        paper="a4",
        parametros=("codigo",),
        tabelas=(
            ("Identificação", ("Campo", "Valor")),
            ("Avaliação", ("Campo", "Valor")),
            ("Plano de resposta", ("Campo", "Valor")),
            ("Ações vinculadas", ("Item", "Ação", "Responsável", "Prevista", "Situação")),
            ("Revisões", ("Data", "Tipo", "Situação apurada", "De", "Para", "Comentário", "Por")),
            ("Histórico", ("Quando", "Quem", "O que mudou")),
        ),
        kpis=(
            ("Inerente", "Sem tratamento"),
            ("Residual", "Com o plano de resposta"),
            ("VME atual", "Probabilidade média x impacto em custo"),
            ("Situação", "Etapa do fluxo"),
            ("Próxima revisão", "Cadência pela severidade"),
        ),
    ),
    Exportacao(
        tela="qualidade/rnc",
        excel="qualidade/rnc/excel",
        imprimivel="qualidade/rnc/imprimivel",
        titulo="Não conformidades (RNC)",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Relatórios de não conformidade",
                (
                    "Nº",
                    "Data",
                    "Não conformidade",
                    "Origem",
                    "Disciplina",
                    "Empresa",
                    "Severidade",
                    "Situação",
                    "Próxima etapa",
                    "Prazo",
                    "Ações na Central",
                    "Disposição",
                    "Custo",
                ),
            ),
        ),
        kpis=(
            ("RNC em aberto", "Esperado: 0"),
            ("Prazo de tratamento vencido", "Esperado: 0"),
            ("Em verificação de eficácia", "Limite: 30 dias"),
            ("Encerradas", "de 6 registradas"),
            ("Custo da não qualidade", "Esperado: R$ 0,00"),
        ),
    ),
    Exportacao(
        tela="qualidade/painel",
        excel="qualidade/painel/excel",
        imprimivel="qualidade/painel/imprimivel",
        titulo="Painel da qualidade",
        paper="a4",
        parametros=(),
        tabelas=(
            ("RNC por mês", ("Mês", "Abertas", "Encerradas", "Em aberto")),
            ("Aprovação em inspeções por mês", ("Mês", "Inspeções", "Aprovação")),
            ("RNC por disciplina", ("Grupo", "RNC", "Percentual", "Acumulado")),
            ("RNC por origem", ("Grupo", "RNC", "Percentual", "Acumulado")),
            (
                "Pauta de tratamento",
                (
                    "Nº",
                    "Não conformidade",
                    "Disciplina",
                    "Empresa",
                    "Severidade",
                    "Situação",
                    "Atraso (dias)",
                    "Motivo",
                ),
            ),
            (
                "Desempenho por empresa",
                (
                    "Empresa",
                    "Inspeções",
                    "Reprovadas",
                    "Aprovação",
                    "RNC",
                    "Custo da não qualidade",
                ),
            ),
        ),
        kpis=(
            ("RNC em aberto", "Esperado: 0"),
            ("Tempo médio de tratamento", "Limite: 15 a 45 dias"),
            ("Aprovação em inspeções", "Meta: 95%"),
            ("Conformidade em auditorias", "Meta: 90%"),
            ("Custo da não qualidade", "Esperado: R$ 0,00"),
        ),
    ),
    Exportacao(
        tela="qualidade/auditorias",
        excel="qualidade/auditorias/excel",
        imprimivel="qualidade/auditorias/imprimivel",
        titulo="Auditorias",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Programa de auditorias",
                ("Nº", "Data", "Escopo", "Auditor líder", "Situação", "Conformidade", "Resultado"),
            ),
        ),
        kpis=(
            ("Aderência ao programa", "Meta: 100%"),
            ("Auditorias atrasadas", "Esperado: 0"),
            ("Conformidade nas auditorias", "Meta: 90%"),
            ("Não conformidades de auditoria", "Esperado: 0"),
            ("Próxima auditoria", "Linha de base: 2026-10-14"),
        ),
    ),
    Exportacao(
        tela="qualidade/inspecoes",
        excel="qualidade/inspecoes/excel",
        imprimivel="qualidade/inspecoes/imprimivel",
        titulo="Inspeções e ITP",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Planos de inspeção e testes (ITP)",
                (
                    "ITP",
                    "Título",
                    "Disciplina",
                    "Empresa",
                    "Pontos",
                    "H · W · R",
                    "Inspeções",
                    "Aprovação",
                    "Cliente",
                ),
            ),
            (
                "Registros de inspeção",
                (
                    "Nº",
                    "Data",
                    "Ponto",
                    "ITP",
                    "Tipo",
                    "Inspetor",
                    "Resultado",
                    "Observação",
                    "RNC",
                ),
            ),
        ),
        kpis=(
            ("Inspeções registradas", "Previsto: 40 pontos"),
            ("Aprovação nas inspeções", "Meta: 95%"),
            ("Reprovadas", "Esperado: 0"),
            ("ITPs aprovados pelo cliente", "de 5"),
            ("Pontos já inspecionados", "de 40"),
        ),
    ),
    Exportacao(
        tela="hse/painel",
        excel="hse/painel/excel",
        imprimivel="hse/painel/imprimivel",
        titulo="Painel HSE",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Pirâmide de segurança · mês", ("Nível", "Ocorrências")),
            ("Pirâmide de segurança · acumulado", ("Nível", "Ocorrências")),
            (
                "Evolução mensal de TF e TRIF",
                ("Mês", "TF (com afastamento)", "TRIF (registráveis)"),
            ),
            ("Ocorrências por área", ("Área", "Ocorrências")),
            ("Ocorrências por empresa", ("Empresa", "Ocorrências")),
        ),
        kpis=(
            ("No mês · Taxa de frequência (TF)", "Esperado: 0,00"),
            ("No mês · Lesões registráveis (TRIF)", "Esperado: 0,00"),
            ("No mês · Taxa de gravidade (TG)", "Esperado: 0,0"),
            ("No mês · Alto potencial (HiPo)", "Esperado: 0"),
            ("Acumulado · Taxa de frequência (TF)", "Esperado: 0,00"),
            ("Acumulado · Lesões registráveis (TRIF)", "Esperado: 0,00"),
            ("Acumulado · Taxa de gravidade (TG)", "Esperado: 0,0"),
            ("Acumulado · Alto potencial (HiPo)", "Esperado: 0"),
            ("Acumulado · Dias sem afastamento", "Esperado: 280"),
            ("Acumulado · Ações HSE no prazo", "Meta: ≥ 90%"),
            ("Proativo · DDS realizados", "Meta: 100%"),
            ("Proativo · Conformidade em inspeções", "Esperado: 100%"),
            ("Proativo · Observações / 10 mil HHT", "Meta: ≥ 40"),
            ("Proativo · Relato de quase acidentes", "Referência: 54,5"),
            ("Proativo · Recomendações fechadas (APR/HAZOP)", "Esperado: 100%"),
            ("Proativo · Incidentes ambientais", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="hse/ocorrencias",
        excel="hse/ocorrencias/excel",
        imprimivel="hse/ocorrencias/imprimivel",
        titulo="Ocorrências HSE",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Ocorrências",
                (
                    "Nº",
                    "Data / hora",
                    "Área",
                    "Empresa",
                    "Tipo / nível",
                    "Potencial",
                    "Situação",
                    "Prazos",
                ),
            ),
        ),
        kpis=(
            ("Ocorrências no projeto", "Referência: de 240 registradas"),
            ("Em aberto", "Esperado: 0"),
            ("Alto potencial (HiPo) em aberto", "Esperado: 0"),
            ("Fora do prazo", "Esperado: 0"),
        ),
    ),
    Exportacao(
        tela="hse/inspecoes",
        excel="hse/inspecoes/excel",
        imprimivel="hse/inspecoes/imprimivel",
        titulo="Inspeções, observações e DDS",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Registros mensais",
                (
                    "Mês",
                    "DDS realizados / programados",
                    "Itens conformes / inspecionados",
                    "Observações",
                    "Desvios",
                ),
            ),
            (
                "Inspeções de segurança",
                ("Data", "Área", "Empresa", "Responsável", "Itens", "Conformes", "Não conformes"),
            ),
            (
                "Observações comportamentais",
                (
                    "Data",
                    "Área",
                    "Tipo",
                    "Descrição",
                    "Empresa",
                    "Responsável",
                    "Observado",
                    "Situação",
                ),
            ),
            ("DDS", ("Data", "Tema", "Empresa", "Responsável", "Participantes")),
        ),
        kpis=(
            ("DDS realizados", "Meta: 100%"),
            ("Conformidade em inspeções", "Esperado: 100%"),
            ("Observações comportamentais", "Meta: ≥ 3152"),
            ("Desvios (nível 5)", "Meta de relato: ≥ 946"),
        ),
    ),
    Exportacao(
        tela="hse/analises_risco",
        excel="hse/analises-risco/excel",
        imprimivel="hse/analises-risco/imprimivel",
        titulo="Análises de risco (APR/HAZOP)",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Estudos", ("Nº", "Tipo", "Área", "Título", "Data", "Participantes", "Recomendações")),
        ),
        kpis=(
            ("Estudos registrados", "Referência: de 5 registrados"),
            ("Recomendações abertas", "Esperado: 0"),
            ("Recomendações atrasadas", "Esperado: 0"),
            ("Recomendações fechadas", "Esperado: 100%"),
        ),
    ),
    Exportacao(
        tela="hse/hht",
        excel="hse/hht/excel",
        imprimivel="hse/hht/imprimivel",
        titulo="Horas trabalhadas (HHT)",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "HHT por mês e empresa",
                ("Mês", "Empresa", "Efetivo médio", "HHT", "Média horas / pessoa"),
            ),
            (
                "Histograma de mão de obra",
                ("Mês", "HHT previsto", "HHT registrado", "Efetivo previsto", "Efetivo registrado"),
            ),
        ),
        kpis=(
            ("HHT total registrado", "Previsto: 788.200"),
            ("Meses com registro", "Esperado: 10"),
            ("Efetivo médio (último mês)", "Previsto: 675"),
            ("Empresas com registro", "Referência: de 4 previstas"),
        ),
    ),
    Exportacao(
        tela="governanca/mudancas",
        excel="governanca/mudancas/excel",
        imprimivel="governanca/mudancas/imprimivel",
        titulo="Registro de mudanças",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Solicitações de mudança",
                (
                    "Nº",
                    "Título",
                    "Solicitada em",
                    "Prioridade",
                    "Tipo · origem",
                    "Custo",
                    "Prazo (dias)",
                    "Situação",
                    "Próxima etapa",
                ),
            ),
            ("Mudanças por situação", ("Situação", "Total")),
            ("Pareto por origem", ("Origem", "Total", "%", "% acumulado")),
            ("Mudanças por tipo", ("Tipo", "Total")),
            (
                "Valor e prazo aprovados acumulados",
                ("Mês", "Aprovadas", "Solicitadas", "Valor acumulado", "Prazo acumulado (dias)"),
            ),
        ),
        kpis=(
            ("Em análise", "Referência: 11"),
            ("Aguardando comitê", "Referência: 11"),
            ("Aprovadas em 2026", "Referência: 5"),
            ("Valor aprovado acumulado", "Orçado: R$ 44,6 mi"),
            ("Impacto de prazo acumulado", "Esperado: 0"),
            ("Tempo médio de decisão", "Referência: 31"),
        ),
    ),
    Exportacao(
        tela="governanca/mudanca",
        excel="governanca/mudanca/excel",
        imprimivel="governanca/mudanca/imprimivel",
        titulo="Solicitação de mudança SM-TN-2026-0001",
        paper="a4",
        parametros=("codigo",),
        tabelas=(
            ("Solicitação", ("Campo", "Valor")),
            ("Análise de impacto", ("Campo", "Valor")),
            ("Decisão", ("Campo", "Valor")),
            ("Implementação", ("Campo", "Valor")),
            ("Histórico", ("Quando", "Quem", "O que aconteceu")),
        ),
        kpis=(
            ("Situação", "Etapa do fluxo"),
            ("Impacto em custo", ""),
            ("Impacto em prazo", ""),
            ("Alçada", ""),
            ("Próxima etapa", ""),
        ),
    ),
    Exportacao(
        tela="governanca/licoes",
        excel="governanca/licoes/excel",
        imprimivel="governanca/licoes/imprimivel",
        titulo="Acervo de lições aprendidas",
        paper="a3",
        parametros=(),
        tabelas=(
            (
                "Lições",
                (
                    "Código",
                    "Título",
                    "Tipo",
                    "Situação",
                    "Fase",
                    "Área",
                    "Disciplina",
                    "Aplicabilidade",
                    "Origem",
                    "Recomendação",
                    "Prazo (dias)",
                    "Custo",
                    "Reusos",
                    "Data",
                ),
            ),
            ("Lições por fase", ("Fase", "Lições", "A repetir", "A evitar", "Publicadas")),
            (
                "Lições por área de conhecimento",
                ("Área", "Lições", "A repetir", "A evitar", "Publicadas"),
            ),
            ("Lições por situação", ("Situação", "Lições")),
            ("Lições mais reusadas", ("Lição", "Título", "Tipo", "Reusos")),
            ("Projetos sem registro de lição no período", ("Projeto",)),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="financeiro/eac",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Itens da EAC (nível 3)",
        chave="eac-itens",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Dados",
                (
                    "Código",
                    "Descrição",
                    "Tipo de custo",
                    "Unidade",
                    "Quantidade",
                    "Preço unitário (R$)",
                    "CAPEX/OPEX",
                    "Centro de custo",
                    "Responsável",
                ),
            ),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="financeiro/mapa_controle",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Custos do ERP (fechamento mensal)",
        chave="custos-erp",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Dados",
                (
                    "Código do item",
                    "Mês (1º dia)",
                    "Comprometido (R$)",
                    "Realizado (R$)",
                    "Projeção no término (R$)",
                ),
            ),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="hse/hht",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Horas trabalhadas (HHT)",
        chave="hht",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Dados", ("Mês (1º dia)", "Empresa", "Efetivo médio", "HHT")),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="hse/inspecoes",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Inspeções, observações e DDS (fechamento mensal)",
        chave="hse-mensal",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Dados",
                (
                    "Mês (1º dia)",
                    "DDS programados",
                    "DDS realizados",
                    "Itens inspecionados",
                    "Itens conformes",
                    "Observações",
                    "Desvios",
                ),
            ),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="planejamento/eap",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Avanço da EAP",
        chave="eap-avanco",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Dados", ("Código do pacote", "Avanço acumulado (%)", "Quantidade executada")),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="planejamento/eap",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Pacotes da EAP (nível 3)",
        chave="eap-pacotes",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Dados",
                (
                    "Código",
                    "Descrição",
                    "Tipo",
                    "Critério de medição",
                    "Modelo de etapas",
                    "Unidade",
                    "Quantidade",
                    "Peso (% do projeto)",
                    "Início LB",
                    "Término LB",
                    "Responsável",
                    "Empresa executante",
                    "Item da EAC",
                ),
            ),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="planejamento/produtividade",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Plano de quantidades da LB",
        chave="produtividade-itens",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Dados",
                (
                    "Empresa",
                    "Grupo",
                    "Tipo",
                    "Disciplina",
                    "Unidade",
                    "Quantidade total",
                    "Índice HH por unidade",
                    "Semana inicial",
                    "Semana final",
                    "Perfil",
                ),
            ),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="planejamento/punch_list",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Itens da Punch list",
        chave="punch-itens",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Dados",
                (
                    "Sistema",
                    "Subsistema",
                    "TAG",
                    "Disciplina",
                    "Categoria",
                    "Marco",
                    "Origem",
                    "Descrição",
                    "Empresa",
                    "Responsável",
                    "Prazo",
                ),
            ),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="suprimentos/plano_compras",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Plano de compras",
        chave="plano-de-compras",
        paper="a4",
        parametros=(),
        tabelas=(
            (
                "Dados",
                (
                    "Escopo",
                    "Tipo",
                    "Modalidade",
                    "Disciplina",
                    "Item da EAC",
                    "Estimativa (R$)",
                    "Comprador",
                    "Requisição",
                    "RFx",
                    "Adjudicação",
                    "Pedido",
                    "ROS",
                    "LLI",
                ),
            ),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
    Exportacao(
        tela="suprimentos/diligenciamento",
        excel="importacao/{chave}/modelo",
        imprimivel=None,
        titulo="Modelo de importação: Diligenciamento e recebimento (ERP)",
        chave="diligenciamento",
        paper="a4",
        parametros=(),
        tabelas=(
            ("Dados", ("Nº do pedido", "Marco", "Previsão", "Realizada")),
            ("Instruções", ("Coluna", "Obrigatória", "Formato", "Exemplo")),
        ),
        kpis=(),
    ),
)

# A contagem exata de botões de exportação por arquivo de tela (view, template ou exemplo):
# ``data-tn-excel``, ``data-tn-pdf``, ``data-exportar`` e o ``href`` de ``/api/.../excel``.
BOTOES: Mapping[str, Mapping[str, int]] = {
    "api/src/templates/central_acoes/acoes.html": {"excel": 1, "pdf": 1},
    "api/src/templates/central_acoes/ata.html": {"excel": 1, "pdf": 1},
    "api/src/templates/central_acoes/atas.html": {"excel": 1, "pdf": 1},
    "api/src/templates/central_acoes/painel.html": {"excel": 1, "pdf": 1},
    "api/src/templates/comum/importacao.html": {"excel": 1},
    "api/src/templates/configuracoes/cadastros.html": {"excel": 1, "pdf": 1},
    "api/src/templates/configuracoes/colaboradores.html": {"excel": 1, "pdf": 1},
    "api/src/templates/configuracoes/parametros.html": {"excel": 1, "pdf": 1},
    "api/src/templates/financeiro/custos_importacao_form.html": {"excel": 1},
    "api/src/templates/financeiro/eac_importacao_form.html": {"excel": 1},
    "api/src/templates/governanca/_licoes_cartoes.html": {"excel": 1, "pdf": 1},
    "api/src/templates/governanca/_mudancas_tabela.html": {"excel": 1, "pdf": 1},
    "api/src/templates/governanca/mudanca.html": {"excel": 1, "pdf": 1},
    "api/src/templates/hse/analises_risco.html": {"excel": 1, "pdf": 1},
    "api/src/templates/hse/hht.html": {"excel": 1, "pdf": 1},
    "api/src/templates/hse/inspecoes.html": {"excel": 1, "pdf": 1},
    "api/src/templates/hse/ocorrencias.html": {"excel": 1, "pdf": 1},
    "api/src/templates/hse/painel.html": {"excel": 1, "pdf": 1},
    "api/src/templates/inicio/home.html": {"excel": 1, "pdf": 1},
    "api/src/templates/planejamento/eap_importacao_form.html": {"excel": 1},
    "api/src/templates/planejamento/eap_pacotes_importacao_form.html": {"excel": 1},
    "api/src/templates/planejamento/produtividade_horas.html": {"excel": 1, "pdf": 1},
    "api/src/templates/planejamento/produtividade_kpis.html": {"excel": 1, "pdf": 1},
    "api/src/templates/planejamento/punch_filtros.html": {"excel": 1, "pdf": 1},
    "api/src/templates/planejamento/punch_painel.html": {"excel": 1, "pdf": 1},
    "api/src/templates/programacao_semanal/_dashboard_painel.html": {"excel": 1, "pdf": 1},
    "api/src/templates/programacao_semanal/_painel.html": {"excel": 1, "pdf": 1},
    "api/src/templates/programacao_semanal/importacao.html": {"excel": 2, "pdf": 1},
    "api/src/templates/qualidade/auditorias.html": {"excel": 1, "pdf": 1},
    "api/src/templates/qualidade/inspecoes.html": {"excel": 1, "pdf": 1},
    "api/src/templates/qualidade/painel.html": {"excel": 1, "pdf": 1},
    "api/src/templates/qualidade/rnc.html": {"excel": 1, "pdf": 1},
    "api/src/templates/relatorio/folhas.html": {"excel": 1},
    "api/src/templates/riscos/ficha.html": {"excel": 3, "pdf": 3},
    "api/src/templates/riscos/matriz.html": {"excel": 1, "pdf": 1},
    "api/src/templates/riscos/painel.html": {"excel": 1, "pdf": 1},
    "api/src/templates/riscos/registro.html": {"excel": 1, "pdf": 1},
    "api/src/templates/suprimentos/_pedido_ficha.html": {"excel": 1, "pdf": 1},
    "api/src/templates/suprimentos/_processo_ficha.html": {"excel": 1, "pdf": 1},
    "api/src/templates/suprimentos/diligenciamento.html": {"excel": 1, "pdf": 1},
    "api/src/templates/suprimentos/fornecedores.html": {"excel": 1, "pdf": 1},
    "api/src/templates/suprimentos/mas.html": {"excel": 1, "pdf": 1},
    "api/src/templates/suprimentos/painel.html": {"excel": 1, "pdf": 1},
    "api/src/templates/suprimentos/plano_compras.html": {"excel": 1, "pdf": 1},
    "api/src/templates/suprimentos/processos.html": {"excel": 1, "pdf": 1},
    "app/_views/financeiro/contingencia.html": {"excel": 1, "pdf": 1},
    "app/_views/financeiro/contrato.html": {"excel": 5, "pdf": 5},
    "app/_views/financeiro/contratos.html": {"excel": 1, "pdf": 1},
    "app/_views/financeiro/curva_s.html": {"excel": 1, "pdf": 1},
    "app/_views/financeiro/desembolso.html": {"excel": 1, "pdf": 1},
    "app/_views/financeiro/eac.html": {"excel": 1, "pdf": 1},
    "app/_views/financeiro/kpis.html": {"excel": 1, "pdf": 1},
    "app/_views/financeiro/mapa_controle.html": {"excel": 1, "pdf": 1},
    "app/_views/planejamento/6wla.html": {"excel": 1, "pdf": 1},
    "app/_views/planejamento/curva_s.html": {"excel": 1, "pdf": 1},
    "app/_views/planejamento/eap.html": {"excel": 1, "pdf": 1},
    "app/_views/planejamento/kpis.html": {"excel": 1, "pdf": 1},
    "app/_views/planejamento/relato.html": {"excel": 1, "pdf": 1},
    "app/exemplos/exportacao.html": {"excel": 1, "pdf": 2},
}


# ── O cenário: a demonstração na transação do teste, com o relógio parado ──


@dataclass(frozen=True)
class Cenario:
    """O que cada caso recebe: a sessão, as rotas registradas e os identificadores da carga."""

    session: Session
    rotas: Mapping[str, Any]
    admin: str
    project_id: int
    contract: str
    risk: str
    change: str
    minutes: int


@pytest.fixture
def cenario(
    sessao_das_rotas: Session,
    funcoes_registradas: list,
    monkeypatch: pytest.MonkeyPatch,
) -> Cenario:
    """Carrega a demonstração na transação do teste e resolve o Admin e os registros nomeados."""
    monkeypatch.setattr(calendario, "today", lambda: HOJE)
    monkeypatch.setattr(calendario, "now", lambda: AGORA)
    harness.load_demonstration(sessao_das_rotas)
    admin = next(
        access
        for access in configuracoes.list_active_access(sessao_das_rotas)
        if access.general_profile == "Admin"
    )
    rotas = {
        funcao.get_trigger().route: funcao.get_user_function() for funcao in funcoes_registradas
    }
    return Cenario(
        session=sessao_das_rotas,
        rotas=rotas,
        admin=admin.email,
        project_id=sessao_das_rotas.scalar(select(Project.id).where(Project.code == PROJETO_1)),
        contract=sessao_das_rotas.scalar(select(Contract.number).order_by(Contract.id).limit(1)),
        risk=sessao_das_rotas.scalar(select(Risk.code).order_by(Risk.id).limit(1)),
        change=sessao_das_rotas.scalar(
            select(ChangeRequest.code).order_by(ChangeRequest.id).limit(1)
        ),
        minutes=sessao_das_rotas.scalar(select(Minutes.id).order_by(Minutes.id).limit(1)),
    )


def _parametros(cenario: Cenario, caso: Exportacao, rota: str) -> dict[str, str]:
    """O endereço da rota como a tela o pede: o escopo e o que aquela tela exige a mais."""
    params = {"projeto": str(cenario.project_id)}
    if "numero" in caso.parametros:
        params["numero"] = cenario.contract
    if "codigo" in caso.parametros:
        params["codigo"] = cenario.risk if rota.startswith("riscos/") else cenario.change
    if "id" in caso.parametros:
        params["id"] = str(cenario.minutes)
    if "tipo" in caso.parametros:
        params |= PERIODO_DO_RELATORIO
    return params


def _chamar(
    cenario: Cenario, caso: Exportacao, rota: str, *, imprimivel: bool
) -> func.HttpResponse:
    """Chama a rota registrada como o navegador a chama; a impressão imita o botão PDF."""
    cabecalhos = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(cenario.admin)}
    if imprimivel:
        cabecalhos |= {"X-Alpine-Request": "true", "X-Alpine-Target": FOLHA}
    registrada = ROTA_DO_MODELO if caso.chave else rota
    pedido = func.HttpRequest(
        method="GET",
        url=f"/api/{rota}",
        headers=cabecalhos,
        params=_parametros(cenario, caso, rota),
        route_params={"chave": caso.chave} if caso.chave else {},
        body=b"",
    )
    return cenario.rotas[registrada](pedido)


def _falha(resposta: func.HttpResponse) -> str:
    return f"{resposta.status_code}: {resposta.get_body()[:200].decode('utf-8', 'replace')}"


# ── O Excel: abre, traz o título, as colunas de cada tabela e a referência de cada KPI ──


def _linhas_com_conteudo(aba: Worksheet) -> list[tuple[int, list[Any]]]:
    """As linhas não vazias da aba, com o número da linha; a cauda de ``None`` sai."""
    linhas: list[tuple[int, list[Any]]] = []
    for numero, linha in enumerate(aba.iter_rows(), start=1):
        valores = [celula.value for celula in linha]
        while valores and valores[-1] is None:
            valores.pop()
        if valores:
            linhas.append((numero, valores))
    return linhas


def _tem_tabela(pasta: Workbook, titulo: str, colunas: Sequence[str]) -> bool:
    """A tabela está na planilha: o título numa linha e, logo abaixo, as colunas na ordem."""
    for aba in pasta.worksheets:
        for (numero, valores), (proximo, seguintes) in pairwise(_linhas_com_conteudo(aba)):
            if valores != [titulo] or proximo != numero + 1:
                continue
            if [str(valor) for valor in seguintes[: len(colunas)]] == list(colunas):
                return True
    return False


def _tem_kpi(aba: Worksheet, rotulo: str, referencia: str) -> bool:
    """O indicador está no ``Resumo`` com o valor e a referência de gestão na mesma linha."""
    for numero, valores in _linhas_com_conteudo(aba):
        if valores[:1] == [rotulo]:
            encontrada = aba.cell(row=numero, column=3).value or ""
            if str(encontrada) == referencia:
                return True
    return False


def _conferir_excel(cenario: Cenario, caso: Exportacao) -> None:
    resposta = _chamar(cenario, caso, caso.excel, imprimivel=False)
    assert resposta.status_code == 200, _falha(resposta)
    assert resposta.mimetype == XLSX_CONTENT_TYPE
    pasta = load_workbook(BytesIO(resposta.get_body()))
    assert pasta.sheetnames, caso.endereco_do_modelo()
    for aba in pasta.worksheets:
        assert aba["A2"].value == caso.titulo, aba.title
        assert aba.page_setup.orientation == "landscape", aba.title
        assert aba.page_setup.paperSize == PAPEL_XLSX[caso.paper], aba.title
    for titulo, colunas in caso.tabelas:
        assert _tem_tabela(pasta, titulo, colunas), (
            f"{caso.endereco_do_modelo()}: tabela «{titulo}» sem as colunas {list(colunas)}"
        )
    if caso.kpis:
        assert "Resumo" in pasta.sheetnames, caso.endereco_do_modelo()
        for rotulo, referencia in caso.kpis:
            assert _tem_kpi(pasta["Resumo"], rotulo, referencia), (
                f"{caso.endereco_do_modelo()}: «{rotulo}» sem a referência «{referencia}»"
            )


# ── A folha: a impressão da mesma descrição, sem navegação ──


def _sequencia(valores: Sequence[str], procurados: Sequence[str]) -> bool:
    """Os valores procurados aparecem como bloco contíguo na lista."""
    tamanho = len(procurados)
    return any(
        valores[inicio : inicio + tamanho] == list(procurados)
        for inicio in range(len(valores) - tamanho + 1)
    )


def _conferir_imprimivel(cenario: Cenario, caso: Exportacao) -> None:
    assert caso.imprimivel is not None
    resposta = _chamar(cenario, caso, caso.imprimivel, imprimivel=True)
    assert resposta.status_code == 200, _falha(resposta)
    assert resposta.mimetype == "text/html"
    tags = parse_tags(resposta.get_body().decode())
    raiz = find_by_id(tags, FOLHA)
    assert raiz is not None
    assert {"folha", f"folha--{caso.paper}"} <= raiz.classes
    assert raiz.attrs["data-titulo"].startswith(f"{caso.titulo} - ")
    assert [tag.clean_text() for tag in find_all(tags, "h1", "folha__titulo")] == [caso.titulo]
    assert find_all(tags, "img", "folha__logo")
    assert find_all(tags, "nav") == []
    assert find_all(tags, "a") == []
    assert find_by_id(tags, "main-nav") is None
    assert not [tag for tag in tags if tag.has_class("sidebar")]
    cartoes = [tag.clean_text() for tag in find_all(tags, "div", "folha__kpi")]
    for rotulo, referencia in caso.kpis_do_papel():
        esperado = f"Referência: {referencia}" if referencia else None
        assert any(
            rotulo in cartao and (esperado is None or esperado in cartao) for cartao in cartoes
        ), f"{caso.imprimivel}: «{rotulo}» sem a referência «{referencia}»"
    cabecalhos = [tag.clean_text() for tag in find_all(tags, "th")]
    for _, colunas in caso.tabelas_do_papel():
        assert _sequencia(cabecalhos, colunas), f"{caso.imprimivel}: colunas {list(colunas)}"


def _identificador(caso: Exportacao) -> str:
    return f"{caso.tela.replace('/', '-')}-{caso.chave}" if caso.chave else caso.excel


@pytest.mark.parametrize("caso", INVENTARIO, ids=_identificador)
def test_excel_e_imprimivel_da_tela(cenario: Cenario, caso: Exportacao) -> None:
    """A tela do inventário: o Excel abre com as colunas e as referências e a folha imprime."""
    _conferir_excel(cenario, caso)
    if caso.imprimivel is not None:
        _conferir_imprimivel(cenario, caso)


# ── A cobertura: os botões das telas e as rotas registradas batem com o inventário ──


TAG_DE_BOTAO = re.compile(r"<(?:a|button)\b[^>]*>", re.S)
ATRIBUTO_EXCEL = re.compile(r'data-tn-excel="([^"]*)"')
ATRIBUTO_PDF = re.compile(r'data-tn-pdf="([^"]*)"')
ATRIBUTO_EXPORTAR = re.compile(r'data-exportar="(excel|pdf)"')
ATRIBUTO_HREF = re.compile(r'href="((?:/api/)[^"]*)"')
APONTA_PARA_EXPORTACAO = re.compile(r"/(?:excel|imprimivel)(?:\?|\"|$)")


def _arquivos_das_telas() -> list[Path]:
    """Todas as telas: as views, os templates e a página do exemplo."""
    return [*sorted(VIEWS.rglob("*.html")), *sorted(TEMPLATES.rglob("*.html")), EXEMPLO]


def _botoes_da_tela(texto: str) -> Iterator[tuple[str, str | None]]:
    """``(tipo, endereço)`` de cada botão de exportação, na ordem em que aparece no arquivo."""
    for achado in TAG_DE_BOTAO.finditer(texto):
        alvo = achado.group(0)
        atributo = ATRIBUTO_EXCEL.search(alvo)
        if atributo:
            yield "excel", atributo.group(1)
            continue
        atributo = ATRIBUTO_PDF.search(alvo)
        if atributo:
            yield "pdf", atributo.group(1)
            continue
        atributo = ATRIBUTO_EXPORTAR.search(alvo)
        if atributo:
            href = ATRIBUTO_HREF.search(alvo)
            yield atributo.group(1), href.group(1) if href else None
            continue
        href = ATRIBUTO_HREF.search(alvo)
        if href and APONTA_PARA_EXPORTACAO.search(href.group(1)):
            tipo = "excel" if "/excel" in href.group(1) else "pdf"
            yield tipo, href.group(1)


def _rota_do_endereco(endereco: str) -> str:
    """O endereço do navegador vira a rota registrada; a chave do modelo vira ``{chave}``."""
    caminho = endereco.split("?")[0].removeprefix("/api/")
    partes = caminho.split("/")
    if len(partes) == 3 and partes[0] == "importacao" and partes[2] == "modelo":
        return ROTA_DO_MODELO
    return caminho


def _rotas_do_inventario() -> set[str]:
    rotas = {caso.excel for caso in INVENTARIO}
    rotas |= {caso.imprimivel for caso in INVENTARIO if caso.imprimivel is not None}
    return rotas


def test_a_contagem_de_botoes_por_arquivo_bate_com_o_inventario() -> None:
    """Cada botão de exportação do app está no inventário: arquivo e quantidade conferem."""
    encontrados: dict[str, dict[str, int]] = {}
    for caminho in _arquivos_das_telas():
        relativo = caminho.relative_to(RAIZ).as_posix()
        for tipo, _ in _botoes_da_tela(caminho.read_text(encoding="utf-8")):
            contagem = encontrados.setdefault(relativo, {})
            contagem[tipo] = contagem.get(tipo, 0) + 1

    assert encontrados == {arquivo: dict(contagem) for arquivo, contagem in BOTOES.items()}


def test_todo_endereco_literal_aponta_para_uma_rota_do_inventario() -> None:
    """O botão com endereço fixo (view ou template) aponta para uma rota registrada."""
    rotas = _rotas_do_inventario()
    for caminho in _arquivos_das_telas():
        for _, endereco in _botoes_da_tela(caminho.read_text(encoding="utf-8")):
            if endereco is None or not endereco.startswith("/api/") or "{{" in endereco:
                continue
            assert _rota_do_endereco(endereco) in rotas, f"{caminho.name}: {endereco}"


def test_toda_rota_de_exportacao_registrada_esta_no_inventario(
    funcoes_registradas: list,
) -> None:
    """Nenhuma rota de exportação existe fora do inventário (e nenhuma inventada sobra)."""
    registradas = {
        funcao.get_trigger().route
        for funcao in funcoes_registradas
        if funcao.get_trigger().route.endswith(("/excel", "/imprimivel", "/modelo"))
    }

    assert registradas == _rotas_do_inventario()


def test_todo_caso_do_inventario_esta_parametrizado() -> None:
    """Rota sem teste reprova: os casos do parametrize são os do inventário, um a um."""
    marcas = [
        marca for marca in test_excel_e_imprimivel_da_tela.pytestmark if marca.name == "parametrize"
    ]
    (marca,) = marcas

    casos = tuple(marca.args[1])

    assert {caso.excel for caso in casos} == {caso.excel for caso in INVENTARIO}
    assert {caso.imprimivel for caso in casos} == {caso.imprimivel for caso in INVENTARIO}


def test_cada_tela_do_inventario_existe_na_navegacao() -> None:
    """A tela de cada caso é um ``modulo/id`` da lista de navegação (ou o exemplo)."""
    navegacao = json.loads(NAVEGACAO.read_text(encoding="utf-8"))
    telas = {
        f"{tela.get('modulo', modulo['id'])}/{tela['id']}"
        for modulo in navegacao["modulos"]
        for tela in modulo["telas"]
    }

    for caso in INVENTARIO:
        assert caso.tela in telas or caso.tela in TELAS_FORA_DA_NAVEGACAO, caso.tela


def test_o_inventario_esta_registrado_em_docs() -> None:
    """O documento e o teste são o mesmo inventário: toda rota do caso está citada em ``docs/``."""
    texto = DOC.read_text(encoding="utf-8")

    for caso in INVENTARIO:
        assert f"`{caso.endereco_do_modelo()}`" in texto, caso.endereco_do_modelo()
        if caso.imprimivel is not None:
            assert f"`{caso.imprimivel}`" in texto, caso.imprimivel
