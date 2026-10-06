"""Importação em passos (ISSUE-018, HU-139): o fluxo da fachada, com importadores de teste.

``core.importing`` é percorrido como um módulo o usaria: o modelo, o envio e a conferência
(``check``), a confirmação (``confirm``). O que se prova:

* o importador de teste percorre os quatro passos, do modelo baixado à gravação com a trilha;
* linha com erro não grava, linha com aviso grava, e a conferência diz o motivo de cada uma;
* cancelar na conferência não grava nada (a conferência nunca grava);
* a confirmação grava tudo ou nada: a recusa de uma linha desfaz as anteriores;
* a confirmação só vale para o arquivo conferido, e o acesso é do módulo do importador.
"""

from __future__ import annotations

import re
from collections.abc import Callable, Iterator
from dataclasses import replace
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from src.core import auth, importing
from src.core.errors import AccessDeniedError, InvalidDataError
from src.core.excel import build_workbook, file_name
from src.core.export_document import ValueKind
from src.core.importing import (
    Confirmation,
    ImportColumn,
    ImportContext,
    Importer,
    InvalidImporterError,
    RowCheck,
    Upload,
    digest_of,
)
from src.core.models import AuditEntry, Client
from src.core.rbac import Permission
from src.core.scope import Scope
from tests.identidades import criar_colaborador, criar_empresa, requisicao
from tests.importadores_de_teste import (
    CABECALHO_DE_TIPOS,
    CHAVE_DE_CLIENTES,
    NOME_RECUSADO,
    PREFIXO,
    clientes_gravados,
    importador_de_clientes,
    importador_de_tipos,
    planilha,
)

HOJE = date(2026, 10, 6)
MARIO = "mario.membro@example.invalid"
VERA = "vera.visualizadora@example.invalid"
GIL = "gil.gestor@example.invalid"
FABIO = "fabio.fornecedor@example.invalid"

ALFA = f"{PREFIXO} Alfa"
BETA = f"{PREFIXO} Beta"
GAMA = f"{PREFIXO} Gama"


@pytest.fixture(autouse=True)
def registro_isolado(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """Cada teste começa com o registro de importadores vazio, e o desfaz ao terminar."""
    monkeypatch.setattr(importing, "_IMPORTERS", {})
    yield


@pytest.fixture(autouse=True)
def pessoas(db_session: Session) -> None:
    """Um Membro, um Visualizador, um Gestor e um fornecedor no cadastro."""
    criar_colaborador(db_session, email=MARIO, nome="Mário Membro", perfil="Membro")
    criar_colaborador(db_session, email=VERA, nome="Vera Visualizadora", perfil="Visualizador")
    criar_colaborador(db_session, email=GIL, nome="Gil Gestor", perfil="Gestor")
    criar_colaborador(
        db_session,
        email=FABIO,
        nome="Fábio Fornecedor",
        vinculo="Fornecedor",
        empresa=criar_empresa(db_session, "Contratada Alfa"),
    )


def _contexto(session: Session, email: str = MARIO) -> ImportContext:
    usuario = auth.resolve_user(session, requisicao(email=email))
    return ImportContext(
        user=usuario, scope=Scope(project_id=None, source="padrao"), reference_date=HOJE
    )


def _conferir(
    session: Session, importador: Importer, conteudo: bytes, *, email: str = MARIO
) -> importing.ImportPreview:
    return importing.check(
        session,
        importer=importador,
        upload=Upload("lista.xlsx", conteudo),
        context=_contexto(session, email),
    )


def _confirmar(
    session: Session,
    importador: Importer,
    conteudo: bytes,
    *,
    ciente: bool = False,
    conferido: str | None = "do-arquivo",
) -> importing.ImportResult:
    """Confirma o arquivo; ``conferido="do-arquivo"`` manda o resumo do próprio arquivo."""
    upload = Upload("lista.xlsx", conteudo)
    resumo = digest_of(upload) if conferido == "do-arquivo" else conferido
    return importing.confirm(
        session,
        importer=importador,
        upload=upload,
        context=_contexto(session),
        confirmation=Confirmation(checked_digest=resumo, acknowledged=ciente),
    )


def _trilha_de_clientes(session: Session) -> int:
    consulta = select(func.count()).select_from(AuditEntry).where(AuditEntry.entity == "cliente")
    return session.scalar(consulta) or 0


def _linha_do_cabecalho(folha: Any) -> int:
    """Onde o modelo pôs o cabeçalho: a primeira linha do filtro da folha (``A7:C7``)."""
    return int(re.search(r"\d+", folha.auto_filter.ref).group())


def _mensagem(chamada: Callable[[], object]) -> str:
    with pytest.raises(InvalidDataError) as erro:
        chamada()
    return str(erro.value)


# ── O fluxo inteiro ──────────────────────────────────────────────────────


def test_um_importador_de_teste_percorre_modelo_envio_conferencia_e_confirmacao(
    db_session: Session,
) -> None:
    importador = importador_de_clientes()
    importing.register(importador)
    contexto = _contexto(db_session)

    # 1. Modelo: o Excel gerado pelo construtor da exportação, com o cabeçalho pronto.
    modelo = build_workbook(
        importing.model_document(db_session, importer=importador, context=contexto)
    )
    livro = load_workbook(BytesIO(modelo))
    assert livro.sheetnames == ["Dados", "Instruções"]
    dados = livro["Dados"]
    cabecalho = _linha_do_cabecalho(dados)
    assert [celula.value for celula in dados[cabecalho]][:3] == ["Nome", "Sigla", "Ativo"]
    assert dados.max_row == cabecalho

    # 2. Envio: a pessoa preenche o modelo, na própria planilha, abaixo do cabeçalho.
    preenchido = [(ALFA, "ALF", "Sim"), (BETA, None, "Não"), (GAMA, "GAM", None)]
    for numero, linha in enumerate(preenchido, start=cabecalho + 1):
        for coluna, valor in enumerate(linha, start=1):
            dados.cell(row=numero, column=coluna, value=valor)
    saida = BytesIO()
    livro.save(saida)
    enviado = saida.getvalue()

    # 3. Conferência: as linhas com o número da planilha, nada gravado.
    antes = _trilha_de_clientes(db_session)
    conferencia = _conferir(db_session, importador, enviado)
    assert [linha.number for linha in conferencia.rows] == [
        cabecalho + 1,
        cabecalho + 2,
        cabecalho + 3,
    ]
    assert [linha.errors for linha in conferencia.rows] == [(), (), ()]
    assert [bool(linha.warnings) for linha in conferencia.rows] == [False, True, False]
    assert conferencia.can_confirm
    assert conferencia.digest == digest_of(Upload("lista.xlsx", enviado))
    assert clientes_gravados(db_session) == []
    assert _trilha_de_clientes(db_session) == antes

    # 4. Confirmação: grava tudo, pela fachada, com a trilha.
    resultado = _confirmar(db_session, importador, enviado, conferido=conferencia.digest)
    assert [linha.number for linha in resultado.written] == [
        cabecalho + 1,
        cabecalho + 2,
        cabecalho + 3,
    ]
    assert resultado.skipped == ()
    assert resultado.with_warning == 1
    assert clientes_gravados(db_session) == [ALFA, BETA, GAMA]
    ativos = dict(
        db_session.execute(
            select(Client.name, Client.active).where(Client.name.like(f"{PREFIXO}%"))
        )
    )
    assert ativos == {ALFA: True, BETA: False, GAMA: True}
    assert _trilha_de_clientes(db_session) == antes + 3
    autores = set(
        db_session.scalars(
            select(AuditEntry.user_id).where(
                AuditEntry.entity == "cliente", AuditEntry.action == "criado"
            )
        )
    )
    assert contexto.user.id in autores


def test_o_modelo_traz_so_o_cabecalho_e_as_instrucoes_de_cada_coluna(db_session: Session) -> None:
    importador = importador_de_tipos([])
    documento = importing.model_document(
        db_session, importer=importador, context=_contexto(db_session)
    )

    livro = load_workbook(BytesIO(build_workbook(documento)))

    dados = livro["Dados"]
    cabecalho = _linha_do_cabecalho(dados)
    assert [celula.value for celula in dados[cabecalho]][: len(CABECALHO_DE_TIPOS)] == list(
        CABECALHO_DE_TIPOS
    )
    assert dados.max_row == cabecalho, "o exemplo não vai na aba Dados: seria importado por engano"
    instrucoes = livro["Instruções"]
    inicio = _linha_do_cabecalho(instrucoes)
    linhas = [
        tuple(celula.value for celula in linha)
        for linha in instrucoes.iter_rows(min_row=inicio, max_row=instrucoes.max_row, max_col=4)
    ]
    assert linhas == [
        ("Coluna", "Obrigatória", "Formato", "Exemplo"),
        ("Código", "Sim", "texto", "A-1"),
        ("Quantidade", "Não", "número inteiro", "12"),
        ("Peso", "Não", "número (ex.: 12,5)", "12,5"),
        ("Avanço", "Não", "percentual (ex.: 45,5)", "45,3"),
        ("Valor", "Não", "valor em reais (ex.: 1234,56)", "1234,56"),
        ("Início", "Não", "dd/mm/aaaa", "05/10/2026"),
        ("Situação", "Não", "um de: Aberta, Fechada", "Aberta"),
    ]
    assert documento.kpis == ()
    assert "Resumo" not in livro.sheetnames


def test_o_modelo_se_chama_modelo_de_importacao_com_o_titulo_e_a_data(db_session: Session) -> None:
    documento = importing.model_document(
        db_session, importer=importador_de_clientes(), context=_contexto(db_session)
    )

    assert documento.title == "Modelo de importação: Clientes de teste"
    assert documento.generated_on == HOJE
    assert documento.scope_label == "Portfólio"
    assert file_name(documento) == "modelo-de-importacao-clientes-de-teste-2026-10-06.xlsx"


# ── Erro bloqueia a linha, aviso não ─────────────────────────────────────


def test_linha_com_erro_nao_grava_linha_com_aviso_grava_e_a_conferencia_diz_o_motivo(
    db_session: Session,
) -> None:
    importador = importador_de_clientes()
    db_session.add(Client(name=f"{PREFIXO} Existente", active=True))
    db_session.flush()
    conteudo = planilha(
        [
            (f"{PREFIXO} Nova", "NOV", "Sim"),
            (f"{PREFIXO} Sem Sigla", None, None),
            (f"{PREFIXO} EXISTENTE", "EX", "Sim"),
            (None, "SEM", "Sim"),
            (f"{PREFIXO} Zeta", "ZET", "Talvez"),
            (f"{PREFIXO} Nova", "N2", None),
        ]
    )

    conferencia = _conferir(db_session, importador, conteudo)

    motivos = {linha.number: (linha.errors, linha.warnings) for linha in conferencia.rows}
    assert motivos == {
        2: ((), ()),
        3: ((), ("Sigla em branco: o cliente entra sem sigla",)),
        4: (("Nome: já existe um cliente com este nome",), ()),
        5: (("Nome: obrigatório",), ()),
        6: (("Ativo: valor fora da lista (Sim, Não)",), ()),
        7: (("repete a linha 2 (mesmo valor em Nome)",), ()),
    }
    assert [linha.number for linha in conferencia.writable_rows] == [2, 3]
    assert [linha.number for linha in conferencia.error_rows] == [4, 5, 6, 7]
    assert [linha.number for linha in conferencia.warning_rows] == [3]
    assert conferencia.can_confirm

    resultado = _confirmar(db_session, importador, conteudo, ciente=True)

    assert clientes_gravados(db_session) == [
        f"{PREFIXO} Existente",
        f"{PREFIXO} Nova",
        f"{PREFIXO} Sem Sigla",
    ]
    assert [linha.number for linha in resultado.written] == [2, 3]
    assert [linha.number for linha in resultado.skipped] == [4, 5, 6, 7]


def test_com_linhas_de_erro_a_confirmacao_exige_o_estou_ciente(db_session: Session) -> None:
    importador = importador_de_clientes()
    conteudo = planilha([(ALFA, "ALF", "Sim"), (None, "SEM", "Sim")])

    mensagem = _mensagem(lambda: _confirmar(db_session, importador, conteudo, ciente=False))

    assert mensagem == "Confirme que está ciente de que as linhas com erro não serão importadas."
    assert clientes_gravados(db_session) == []


def test_o_repetido_na_planilha_ignora_caixa_e_acento(db_session: Session) -> None:
    conteudo = planilha(
        [(ALFA, "A", "Sim"), (ALFA.upper().replace("Ç", "C").replace("Ã", "A"), "B", "Sim")]
    )

    conferencia = _conferir(db_session, importador_de_clientes(), conteudo)

    assert conferencia.rows[1].errors == ("repete a linha 2 (mesmo valor em Nome)",)


def test_planilha_so_com_linhas_de_erro_nao_tem_nada_a_importar(db_session: Session) -> None:
    importador = importador_de_clientes()
    conteudo = planilha([(None, "A", "Sim"), (None, "B", "Sim")])

    conferencia = _conferir(db_session, importador, conteudo)

    assert not conferencia.can_confirm
    assert _mensagem(lambda: _confirmar(db_session, importador, conteudo, ciente=True)) == (
        "Nada a importar: nenhuma linha passou na conferência."
    )


# ── Cancelar não grava; confirmar é tudo ou nada ─────────────────────────


def test_cancelar_na_conferencia_nao_grava_nada(db_session: Session) -> None:
    importador = importador_de_clientes()
    conteudo = planilha([(ALFA, "ALF", "Sim"), (BETA, None, "Não")])
    antes = _trilha_de_clientes(db_session)

    conferencia = _conferir(db_session, importador, conteudo)
    # A pessoa cancela: nenhum passo seguinte acontece. Conferir de novo também não deixa marca.
    _conferir(db_session, importador, conteudo)

    assert conferencia.can_confirm
    assert clientes_gravados(db_session) == []
    assert _trilha_de_clientes(db_session) == antes
    assert not db_session.new
    assert not db_session.dirty
    assert not db_session.deleted


def test_a_confirmacao_grava_tudo_ou_nada(db_session: Session) -> None:
    importador = importador_de_clientes()
    conteudo = planilha([(ALFA, "ALF", "Sim"), (NOME_RECUSADO, "REC", "Sim"), (GAMA, "GAM", "Sim")])
    antes = _trilha_de_clientes(db_session)

    mensagem = _mensagem(lambda: _confirmar(db_session, importador, conteudo))

    assert mensagem == "Linha 3: O cliente recusado não pode ser importado. Nada foi importado."
    assert clientes_gravados(db_session) == []
    assert _trilha_de_clientes(db_session) == antes

    # A transação segue utilizável: sem a linha recusada, a mesma importação entra inteira.
    sem_a_recusada = planilha([(ALFA, "ALF", "Sim"), (GAMA, "GAM", "Sim")])
    _confirmar(db_session, importador, sem_a_recusada)
    assert clientes_gravados(db_session) == [ALFA, GAMA]
    assert _trilha_de_clientes(db_session) == antes + 2


def test_a_recusa_no_meio_desfaz_as_linhas_ja_gravadas_mesmo_dentro_da_transacao_do_chamador(
    db_session: Session,
) -> None:
    gravadas: list[str] = []

    def gravar(session: Session, *, values: Any, **_contexto: object) -> None:
        if len(gravadas) == 2:
            raise InvalidDataError("recusada")
        gravadas.append(values["name"])
        session.add(Client(name=values["name"], active=True))
        session.flush()

    importador = replace(importador_de_clientes(), save=gravar)
    conteudo = planilha(
        [
            (ALFA, "A", "Sim"),
            (BETA, "B", "Sim"),
            (GAMA, "G", "Sim"),
            (f"{PREFIXO} Delta", "D", "Sim"),
        ]
    )

    mensagem = _mensagem(lambda: _confirmar(db_session, importador, conteudo))

    assert gravadas == [ALFA, BETA]
    assert mensagem == "Linha 4: recusada Nada foi importado."
    assert clientes_gravados(db_session) == []


# ── A confirmação é do arquivo que foi conferido ─────────────────────────


def test_sem_conferir_a_confirmacao_e_recusada_e_nada_e_gravado(db_session: Session) -> None:
    importador = importador_de_clientes()
    conteudo = planilha([(ALFA, "ALF", "Sim")])

    for resumo in (None, ""):
        assert _mensagem(
            lambda resumo=resumo: _confirmar(db_session, importador, conteudo, conferido=resumo)
        ) == ("Confira a planilha antes de importar: nada é gravado sem a conferência.")
    assert clientes_gravados(db_session) == []


def test_outro_arquivo_depois_da_conferencia_e_recusado(db_session: Session) -> None:
    importador = importador_de_clientes()
    conferido = planilha([(ALFA, "ALF", "Sim")])
    trocado = planilha([(ALFA, "ALF", "Sim"), (BETA, "BET", "Sim")])
    resumo = _conferir(db_session, importador, conferido).digest

    mensagem = _mensagem(lambda: _confirmar(db_session, importador, trocado, conferido=resumo))

    assert mensagem == "O arquivo mudou depois da conferência. Envie-o e confira de novo."
    assert clientes_gravados(db_session) == []


# ── O arquivo e o cabeçalho ──────────────────────────────────────────────


def test_arquivo_que_nao_e_planilha_e_422_na_conferencia_e_na_confirmacao(
    db_session: Session,
) -> None:
    importador = importador_de_clientes()
    texto = Upload("lista.xlsx", b"isto nao e uma planilha")

    contexto = _contexto(db_session)
    with pytest.raises(InvalidDataError, match=r"não é uma planilha do Excel"):
        importing.check(db_session, importer=importador, upload=texto, context=contexto)
    with pytest.raises(InvalidDataError, match=r"não é uma planilha do Excel"):
        importing.confirm(
            db_session,
            importer=importador,
            upload=texto,
            context=contexto,
            confirmation=Confirmation(checked_digest=digest_of(texto), acknowledged=True),
        )
    assert clientes_gravados(db_session) == []


def test_colunas_obrigatorias_ausentes_nada_e_conferido_nem_confirmado(db_session: Session) -> None:
    importador = importador_de_clientes()
    conteudo = planilha([("ALF", "Sim")], cabecalho=("Sigla", "Ativo"))

    conferencia = _conferir(db_session, importador, conteudo)

    assert conferencia.missing == ("Nome",)
    assert conferencia.rows == ()
    assert not conferencia.can_confirm
    assert _mensagem(lambda: _confirmar(db_session, importador, conteudo, ciente=True)) == (
        "Colunas obrigatórias ausentes: Nome. Baixe o modelo e ajuste a planilha."
    )


def test_planilha_que_nao_e_o_modelo_lista_todas_as_colunas_como_ausentes(
    db_session: Session,
) -> None:
    conteudo = planilha([("a", "b")], cabecalho=("Foo", "Bar"))

    conferencia = _conferir(db_session, importador_de_clientes(), conteudo)

    assert conferencia.missing == ("Nome", "Sigla", "Ativo")
    assert conferencia.ignored == ()
    assert conferencia.rows == ()


def test_coluna_opcional_ausente_vale_em_branco(db_session: Session) -> None:
    conteudo = planilha([(ALFA,)], cabecalho=("Nome",))

    conferencia = _conferir(db_session, importador_de_clientes(), conteudo)

    assert conferencia.missing == ()
    assert conferencia.rows[0].values == {"name": ALFA, "acronym": None, "active": None}
    assert conferencia.rows[0].warnings == ("Sigla em branco: o cliente entra sem sigla",)


def test_cabecalho_com_outra_caixa_acento_e_espacos_e_reconhecido_e_o_que_sobra_e_ignorado(
    db_session: Session,
) -> None:
    conteudo = planilha(
        [("ALF", ALFA, "Sim", "obs")], cabecalho=("  sigla ", "NOME", "ativo", "Observação")
    )

    conferencia = _conferir(db_session, importador_de_clientes(), conteudo)

    assert conferencia.missing == ()
    assert conferencia.ignored == ("Observação",)
    assert conferencia.rows[0].values == {"name": ALFA, "acronym": "ALF", "active": "Sim"}


def test_o_cabecalho_vem_depois_das_linhas_de_identificacao_e_a_numeracao_e_a_do_excel(
    db_session: Session,
) -> None:
    conteudo = planilha([(ALFA, "A", "Sim"), (None, None, None), (BETA, "B", "Sim")], acima=6)

    conferencia = _conferir(db_session, importador_de_clientes(), conteudo)

    assert [linha.number for linha in conferencia.rows] == [8, 10]


def test_linha_vazia_e_espacamento_e_nao_e_conferida(db_session: Session) -> None:
    conteudo = planilha(
        [(None, None, None), ("  ", None, None), (ALFA, "A", "Sim"), (None, None, None)]
    )

    conferencia = _conferir(db_session, importador_de_clientes(), conteudo)

    assert [linha.number for linha in conferencia.rows] == [4]


def test_a_planilha_so_com_o_cabecalho_nao_tem_nada_a_importar(db_session: Session) -> None:
    conferencia = _conferir(db_session, importador_de_clientes(), planilha([]))

    assert conferencia.rows == ()
    assert conferencia.missing == ()
    assert not conferencia.can_confirm


def test_o_limite_de_5000_linhas_passa_e_uma_a_mais_e_422(db_session: Session) -> None:
    destino: list[Any] = []
    importador = importador_de_tipos(destino)

    no_limite = _conferir(
        db_session, importador, planilha([(f"C{n}",) for n in range(5000)], cabecalho=("Código",))
    )
    assert len(no_limite.rows) == 5000

    mensagem = _mensagem(
        lambda: _conferir(
            db_session,
            importador,
            planilha([(f"C{n}",) for n in range(5001)], cabecalho=("Código",)),
        )
    )
    assert (
        mensagem
        == "A planilha passa de 5.000 linhas de dados. Divida a importação em planilhas menores."
    )
    assert destino == []


# ── Cada tipo de coluna ──────────────────────────────────────────────────


def test_cada_tipo_de_coluna_e_convertido_conferido_e_gravado_como_o_prototipo(
    db_session: Session,
) -> None:
    destino: list[Any] = []
    importador = importador_de_tipos(destino)
    inicio = date(2026, 10, 5)
    formatos = {"Avanço": "0.0%", "Valor": '"R$" #,##0.00', "Início": "DD/MM/YYYY"}
    conteudo = planilha(
        [
            ("A-1", 12, 12.5, 0.453, 1234.56, inicio, "aberta"),
            ("A-2", "7", "1.234,5", "50%", "R$ 99,90", "31/12/2026", " FECHADA "),
            ("A-3", 1.5, "abc", "x", "dez", "31/02/2026", "Talvez"),
            ("A-4", -2, None, None, None, None, None),
            ("A-5", None, None, 1.5, None, None, None),
        ],
        cabecalho=CABECALHO_DE_TIPOS,
        formatos=formatos,
    )

    conferencia = _conferir(db_session, importador, conteudo)

    primeira, segunda, com_erros, negativa, acima = conferencia.rows
    assert primeira.values == {
        "code": "A-1",
        "quantity": 12,
        "weight": Decimal("12.5"),
        "progress": Decimal("45.3"),
        "value": 123456,
        "start": inicio,
        "status": "Aberta",
    }
    assert primeira.shown == ("A-1", "12", "12,50", "45,3%", "R$ 1.234,56", "05/10/2026", "Aberta")
    assert segunda.values["quantity"] == 7
    assert segunda.values["weight"] == Decimal("1234.5")
    assert segunda.values["progress"] == Decimal(50)
    assert segunda.values["value"] == 9990
    assert segunda.values["start"] == date(2026, 12, 31)
    assert segunda.values["status"] == "Fechada"
    assert com_erros.errors == (
        "Quantidade: número inteiro inválido",
        "Peso: número inválido",
        "Avanço: número inválido",
        "Valor: valor inválido",
        "Início: data inválida (use dd/mm/aaaa)",
        "Situação: valor fora da lista (Aberta, Fechada)",
    )
    assert com_erros.shown[2] == "abc"
    assert negativa.errors == ("Quantidade: não pode ser negativa",)
    assert acima.writable
    assert acima.warnings == ("Avanço acima de 100%",)
    assert acima.values["progress"] == Decimal("150")
    assert destino == [], "a conferência não grava"

    _confirmar(db_session, importador, conteudo, ciente=True)

    assert [linha["code"] for linha in destino] == ["A-1", "A-2", "A-5"]
    assert destino[2]["quantity"] is None
    assert destino[2]["start"] is None


def test_a_validacao_do_modulo_so_ve_linhas_ja_convertidas(db_session: Session) -> None:
    recebidas: list[Any] = []

    def validar(_session: Session, *, values: Any, **_contexto: object) -> RowCheck:
        recebidas.append(dict(values))
        return importing.VALID

    importador = replace(importador_de_tipos([]), validate=validar)
    conteudo = planilha(
        [("A-1", 3, None, None, None, None, None), ("A-2", "x", None, None, None, None, None)],
        cabecalho=CABECALHO_DE_TIPOS,
    )

    _conferir(db_session, importador, conteudo)

    assert [linha["code"] for linha in recebidas] == ["A-1"], (
        "linha com erro de tipo nem chega ao módulo"
    )
    assert recebidas[0]["quantity"] == 3


# ── Acesso ───────────────────────────────────────────────────────────────


@pytest.mark.parametrize("email", [VERA, FABIO])
def test_quem_nao_escreve_no_modulo_nao_baixa_nao_confere_nem_confirma(
    db_session: Session, email: str
) -> None:
    importador = importador_de_clientes()
    conteudo = planilha([(ALFA, "ALF", "Sim")])
    contexto = _contexto(db_session, email)

    with pytest.raises(AccessDeniedError):
        importing.model_document(db_session, importer=importador, context=contexto)
    with pytest.raises(AccessDeniedError):
        _conferir(db_session, importador, conteudo, email=email)
    with pytest.raises(AccessDeniedError):
        importing.confirm(
            db_session,
            importer=importador,
            upload=Upload("lista.xlsx", conteudo),
            context=contexto,
            confirmation=Confirmation(checked_digest=None),
        )
    assert clientes_gravados(db_session) == []


def test_o_importador_pode_pedir_mais_que_gravar(db_session: Session) -> None:
    importador = importador_de_tipos([], permissao=Permission.MANAGE)
    conteudo = planilha([("A-1",)], cabecalho=("Código",))

    with pytest.raises(AccessDeniedError):
        _conferir(db_session, importador, conteudo, email=MARIO)

    assert len(_conferir(db_session, importador, conteudo, email=GIL).rows) == 1


# ── O registro de importadores ───────────────────────────────────────────


def test_registrar_o_mesmo_importador_de_novo_nao_faz_nada_e_outro_na_mesma_chave_falha() -> None:
    importador = importador_de_clientes()

    importing.register(importador)
    importing.register(importador_de_clientes())

    assert importing.find(CHAVE_DE_CLIENTES) == importador
    assert importing.registered() == [importador]
    with pytest.raises(ValueError, match=CHAVE_DE_CLIENTES):
        importing.register(replace(importador, title="Outro título"))


def test_chave_desconhecida_e_422() -> None:
    assert importing.find("nao-existe") is None
    assert _mensagem(lambda: importing.require("nao-existe")) == "A importação pedida não existe."


@pytest.mark.parametrize(
    "mudanca",
    [
        {"key": "Chave Com Espaço"},
        {"key": "com_sublinhado"},
        {"key": ""},
        {"columns": ()},
        {"columns": (ImportColumn("a", "A"), ImportColumn("a", "B"))},
        {"columns": (ImportColumn("a", "Nome"), ImportColumn("b", "  nome "))},
        {"unique": ("campo_que_nao_existe",)},
    ],
)
def test_importador_que_quebra_o_contrato_falha_na_hora(mudanca: dict[str, Any]) -> None:
    with pytest.raises(InvalidImporterError):
        replace(importador_de_clientes(), **mudanca)


def test_a_coluna_diz_o_formato_que_pede() -> None:
    assert ImportColumn("a", "A").format_hint == "texto"
    assert ImportColumn("a", "A", ValueKind.MONEY).format_hint == "valor em reais (ex.: 1234,56)"
    assert ImportColumn("a", "A", options=("X", "Y")).format_hint == "um de: X, Y"
