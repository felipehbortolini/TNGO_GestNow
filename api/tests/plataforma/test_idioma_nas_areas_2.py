"""O inglês nas telas da parte 2 (ISSUE-088, D13): as oito áreas restantes.

Com o cookie ``gestnow_idioma=en`` os fragmentos das áreas, o relatório gerencial
e o seu Excel respondem em inglês; com o cookie ausente os mesmos textos voltam em
português (a garantia de não-regressão do idioma-fonte). Dado cadastrado não é
traduzido: o nome do projeto chega como foi gravado. O rótulo curto da semana ISO
segue o idioma (``S38``/``W38``).
"""

from __future__ import annotations

import html
from io import BytesIO

import azure.functions as func
import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from src.blueprints.nav import main_nav
from src.core import calendario, translations
from src.modulos.configuracoes.routes import parameters_screen
from src.modulos.hse.ocorrencias_routes import occurrence_list_screen
from src.modulos.planejamento.routes import curva_fisica_content
from src.modulos.planejamento.validation import WEEKLY
from src.modulos.programacao_semanal.routes import matrix
from src.modulos.qualidade.routes import rnc_list_screen
from src.modulos.relatorio.routes import report_screen, report_sheets_excel
from src.modulos.riscos.routes import risk_register_screen
from src.modulos.suprimentos.routes import supplier_list
from tests.identidades import criar_projeto, requisicao

COOKIE = translations.LANGUAGE_COOKIE


def _pedido(
    caminho: str,
    *,
    idioma: str | None,
    alvo: str,
    params: dict[str, str] | None = None,
) -> func.HttpRequest:
    cookies = {COOKIE: idioma} if idioma else None
    return requisicao(caminho, cookies=cookies, alvo=alvo, params=params)


def _corpo(resposta: func.HttpResponse) -> str:
    return html.unescape(resposta.get_body().decode("utf-8"))


def _periodo_corrente() -> str:
    return calendario.period_of(WEEKLY, calendario.today())


# ── O motor e o rótulo da semana ─────────────────────────────────────────


def test_o_rotulo_da_semana_iso_segue_o_idioma() -> None:
    assert translations.semana_label(38, "pt") == "S38"
    assert translations.semana_label(38, "en") == "W38"


# ── Os fragmentos das oito áreas ─────────────────────────────────────────


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_planejamento_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        curva_fisica_content(
            _pedido("/api/planejamento/curva-fisica", idioma="en", alvo="curva-fisica-conteudo")
        )
    )
    em_portugues = _corpo(
        curva_fisica_content(
            _pedido("/api/planejamento/curva-fisica", idioma=None, alvo="curva-fisica-conteudo")
        )
    )

    assert "No portfolio project has an S-Curve registered." in em_ingles
    assert "Nenhum projeto da carteira tem Curva S cadastrada." in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_programacao_semanal_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        matrix(
            _pedido(
                "/api/programacao-semanal/programacoes",
                idioma="en",
                alvo="prog-resumo prog-tabela",
            )
        )
    )
    em_portugues = _corpo(
        matrix(
            _pedido(
                "/api/programacao-semanal/programacoes",
                idioma=None,
                alvo="prog-resumo prog-tabela",
            )
        )
    )

    assert "Adherence" in em_ingles
    assert "Awaiting inspector" in em_ingles
    assert "Aderência" in em_portugues
    assert "Aguardando fiscal" in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_suprimentos_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        supplier_list(
            _pedido("/api/suprimentos/fornecedores", idioma="en", alvo="fornecedores-conteudo")
        )
    )
    em_portugues = _corpo(
        supplier_list(
            _pedido("/api/suprimentos/fornecedores", idioma=None, alvo="fornecedores-conteudo")
        )
    )

    assert "Registered" in em_ingles
    assert "There are no qualified suppliers yet" in em_ingles
    assert "Cadastrados" in em_portugues
    assert "Ainda não há fornecedores qualificados" in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_riscos_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        risk_register_screen(_pedido("/api/riscos/registro", idioma="en", alvo="registro-conteudo"))
    )
    em_portugues = _corpo(
        risk_register_screen(_pedido("/api/riscos/registro", idioma=None, alvo="registro-conteudo"))
    )

    assert "New risk" in em_ingles
    assert "There are no risks registered yet" in em_ingles
    assert "Novo risco" in em_portugues
    assert "Ainda não há riscos registrados" in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_qualidade_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        rnc_list_screen(_pedido("/api/qualidade/rnc", idioma="en", alvo="rnc-conteudo"))
    )
    em_portugues = _corpo(
        rnc_list_screen(_pedido("/api/qualidade/rnc", idioma=None, alvo="rnc-conteudo"))
    )

    assert "New NCR" in em_ingles
    assert "Status: open" in em_ingles
    assert "Nova RNC" in em_portugues
    assert "Situação: em aberto" in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_hse_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        occurrence_list_screen(
            _pedido("/api/hse/ocorrencias", idioma="en", alvo="ocorrencias-conteudo")
        )
    )
    em_portugues = _corpo(
        occurrence_list_screen(
            _pedido("/api/hse/ocorrencias", idioma=None, alvo="ocorrencias-conteudo")
        )
    )

    assert "Record the first incident with date, area, type and the potential (P x I)." in em_ingles
    assert (
        "Registre a primeira ocorrência com data, área, tipo e o potencial (P x I)." in em_portugues
    )


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_configuracoes_em_ingles_e_em_portugues() -> None:
    em_ingles = _corpo(
        parameters_screen(
            _pedido("/api/configuracoes/parametros", idioma="en", alvo="parametros-conteudo")
        )
    )
    em_portugues = _corpo(
        parameters_screen(
            _pedido("/api/configuracoes/parametros", idioma=None, alvo="parametros-conteudo")
        )
    )

    assert "Current version" in em_ingles
    assert "Version history" in em_ingles
    assert "Versão vigente" in em_portugues
    assert "Histórico de versões" in em_portugues


# ── O relatório gerencial: o fragmento, o Excel e a impressão ────────────


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_o_relatorio_gerencial_em_ingles_e_em_portugues() -> None:
    params = {"tipo": WEEKLY, "periodo": _periodo_corrente()}
    em_ingles = _corpo(
        report_screen(
            _pedido(
                "/api/relatorio/gerencial", idioma="en", alvo="gerencial-conteudo", params=params
            )
        )
    )
    em_portugues = _corpo(
        report_screen(
            _pedido(
                "/api/relatorio/gerencial", idioma=None, alvo="gerencial-conteudo", params=params
            )
        )
    )

    # A tela é a impressão: o cabeçalho e as folhas saem no idioma.
    assert "Management report weekly" in em_ingles
    assert "Issued on" in em_ingles
    assert "Project portfolio" in em_ingles
    assert "Relatório gerencial semanal" in em_portugues
    assert "Emitido em" in em_portugues
    assert "Carteira de projetos" in em_portugues


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_o_excel_do_relatorio_gerencial_em_ingles_e_em_portugues() -> None:
    params = {"tipo": WEEKLY, "periodo": _periodo_corrente()}
    em_ingles = load_workbook(
        BytesIO(
            report_sheets_excel(
                _pedido("/api/relatorio/gerencial/excel", idioma="en", alvo="x", params=params)
            ).get_body()
        )
    )
    em_portugues = load_workbook(
        BytesIO(
            report_sheets_excel(
                _pedido("/api/relatorio/gerencial/excel", idioma=None, alvo="x", params=params)
            ).get_body()
        )
    )

    assert "Summary" in em_ingles.sheetnames
    assert "Period analyses" in em_ingles.sheetnames
    assert "Resumo" in em_portugues.sheetnames
    assert "Análises do período" in em_portugues.sheetnames


# ── Dado cadastrado não é traduzido ──────────────────────────────────────


@pytest.mark.usefixtures("rotas_na_transacao_do_teste")
def test_o_nome_do_projeto_chega_como_foi_gravado(sessao_das_rotas: Session) -> None:
    criar_projeto(sessao_das_rotas, codigo="TN-EN-088", nome="Fábrica Nova")
    resposta = main_nav(
        _pedido("/api/nav", idioma="en", alvo="main-nav module-tabs", params={"caminho": "/"})
    )
    corpo = _corpo(resposta)

    assert "TN-EN-088 · Fábrica Nova" in corpo
    assert "Portfolio (1 project)" in corpo
