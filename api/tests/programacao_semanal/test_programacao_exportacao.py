"""Os documentos da semana: o Excel com as abas do app e o relatório imprimível (ISSUE-055)."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from src.core import calendario
from src.core.excel import build_workbook
from src.core.printable import build_sheet
from src.core.scope import Scope
from src.modulos.configuracoes import service as configuracoes
from src.modulos.programacao_semanal import calculations, export, service
from src.modulos.programacao_semanal.service import NewActivity
from tests.cenario_programacao import WEEK, Scenario

TODAY = date(2026, 8, 3)

SCHEDULE_HEADERS = (
    "Item",
    "ID Exclusiva",
    "Atividade",
    "Local",
    "Empresa",
    "Encarregado",
    "Fiscal",
    "Prev/Real",
    *calculations.DAY_LABELS,
    "Total",
    "Unidade",
    "PPC %",
    "Situação",
    "Realizado",
)


def _store(
    scenario: Scenario,
    unique_id: str,
    *,
    week: str = WEEK,
    planned: float = 70.0,
    done: float = 70.0,
    company_id: int | None = None,
) -> None:
    service.store_activity(
        scenario.session,
        author_id=scenario.author_id,
        new=NewActivity(
            project_id=scenario.project_id,
            company_id=company_id or scenario.company_a,
            week=week,
            unique_id=unique_id,
            item=1,
            description=f"Serviço {unique_id}",
            planned_headline=planned,
            author_person_id=scenario.planner.person_id,
            at=calendario.now(),
            location_id=scenario.location_id,
            unit_id=scenario.unit_id,
            foreman_id=scenario.foreman_a,
            inspector_id=scenario.fiscal.person_id,
            planned_days=(planned / 7,) * 7,
            day_shift=(done / 7,) * 7,
        ),
    )


def _scope(scenario: Scenario) -> Scope:
    return Scope(project_id=scenario.project_id, source="url")


def _week_document(scenario: Scenario, *, portfolio: bool = False, week: str = WEEK):
    caller = scenario.caller(scenario.planner)
    views = service.list_activities(scenario.session, caller=caller, week=week)
    scope = Scope(project_id=None, source="url") if portfolio else _scope(scenario)
    return export.week_document(
        week,
        views,
        service.week_summary(views),
        service.week_curve(scenario.session, caller=caller, week=week),
        service.parameters_of(scenario.session, scenario.project_id),
        scope=scope,
        projects=configuracoes.list_projects(scenario.session),
        today=TODAY,
    )


def _report_document(scenario: Scenario):
    caller = scenario.caller(scenario.planner)
    views = service.list_activities(scenario.session, caller=caller, week=WEEK)
    return export.week_report_document(
        WEEK,
        views,
        service.week_summary(views),
        service.parameters_of(scenario.session, scenario.project_id),
        scope=_scope(scenario),
        projects=configuracoes.list_projects(scenario.session),
        today=TODAY,
    )


# ── A curva S ────────────────────────────────────────────────────────────


def test_the_curve_gives_equal_weight_per_activity_over_the_horizon() -> None:
    points = calculations.curve_points(
        [
            calculations.CurveActivity(
                week="S.30/2026", identity="A", planned_total=100.0, done_total=50.0
            ),
            calculations.CurveActivity(
                week="S.30/2026", identity="B", planned_total=50.0, done_total=25.0
            ),
            calculations.CurveActivity(
                week="S.31/2026", identity="A", planned_total=100.0, done_total=100.0
            ),
        ]
    )

    first, second = points
    assert first.week == "S.30/2026"
    assert first.base == 50.0
    assert first.planned == 75.0
    assert first.done == 37.5
    assert (first.base_accumulated, first.planned_accumulated, first.done_accumulated) == (
        50.0,
        75.0,
        37.5,
    )
    assert first.period == "20/07 a 26/07/2026"
    assert second.planned == 25.0
    assert second.done == 25.0
    assert (second.base_accumulated, second.planned_accumulated, second.done_accumulated) == (
        100.0,
        100.0,
        62.5,
    )


def test_the_curve_is_empty_without_weeks() -> None:
    assert calculations.curve_points([]) == []


def test_the_week_curve_reads_the_previous_weeks_of_the_scenario(scenario: Scenario) -> None:
    _store(scenario, "ANT-1", week="S.30/2026")
    _store(scenario, "ATU-1", week=WEEK)

    points = service.week_curve(
        scenario.session, caller=scenario.caller(scenario.planner), week=WEEK
    )

    assert [point.week for point in points] == ["S.30/2026", WEEK]
    assert points[1].planned == 50.0


# ── O Excel da semana ────────────────────────────────────────────────────


def test_the_workbook_has_the_sheets_of_the_app(scenario: Scenario) -> None:
    _store(scenario, "IMP-1")

    book = load_workbook(BytesIO(build_workbook(_week_document(scenario))))

    assert book.sheetnames == [
        "Resumo",
        "Programação",
        "Por contratada",
        "Por frente",
        "Por encarregado",
        "Dia a dia",
        "Curva S",
    ]


def test_the_schedule_sheet_has_the_columns_of_the_app_and_two_rows_per_activity(
    scenario: Scenario,
) -> None:
    _store(scenario, "IMP-1")

    book = load_workbook(BytesIO(build_workbook(_week_document(scenario))))
    sheet = book["Programação"]
    rows = list(sheet.iter_rows(values_only=True))
    header = next(row for row in rows if "Prev/Real" in row)
    assert tuple(header[: len(SCHEDULE_HEADERS)]) == SCHEDULE_HEADERS

    header_index = rows.index(header)
    planned = rows[header_index + 1]
    done = rows[header_index + 2]
    assert planned[2] == "Serviço IMP-1"
    assert planned[7] == "Previsto"
    assert planned[8:15] == (10.0, 10.0, 10.0, 10.0, 10.0, 10.0, 10.0)
    assert done[7] == "Realizado"
    assert done[8:15] == (10.0, 10.0, 10.0, 10.0, 10.0, 10.0, 10.0)


def test_the_summary_sheet_carries_the_indicators_of_the_app(scenario: Scenario) -> None:
    _store(scenario, "IMP-1")

    book = load_workbook(BytesIO(build_workbook(_week_document(scenario))))
    summary = book["Resumo"]
    labels = [row[0] for row in summary.iter_rows(values_only=True) if row[0]]

    for label in (
        "Aderência geral",
        "PPC médio",
        "Atividades",
        "Contratadas",
        "Realizados pendentes",
        "Em elaboração",
    ):
        assert label in labels


def test_the_workbook_in_the_portfolio_opens_with_the_project_column(scenario: Scenario) -> None:
    _store(scenario, "IMP-1")

    book = load_workbook(BytesIO(build_workbook(_week_document(scenario, portfolio=True))))
    sheet = book["Programação"]
    rows = list(sheet.iter_rows(values_only=True))
    header = next(row for row in rows if "Prev/Real" in row)

    assert header[0] == "Projeto"
    assert "PS-1 · Projeto da programação" in " ".join(
        str(value) for value in rows[rows.index(header) + 1]
    )


# ── O relatório imprimível ───────────────────────────────────────────────


def test_the_report_has_the_indicators_and_the_tables_of_the_sheet(scenario: Scenario) -> None:
    _store(scenario, "IMP-1")

    document = _report_document(scenario)
    sheet = build_sheet(document)

    assert document.title == f"Programação da semana {WEEK}"
    assert [kpi.label for kpi in document.kpis] == [
        "Aderência",
        "PPC médio",
        "Atividades",
        "Contratadas",
    ]
    assert [table.title for table in document.tables] == [
        "Programação",
        "Por contratada",
        "Por frente",
        "Dia a dia",
    ]
    assert sheet.scope_label == "PS-1 · Projeto da programação"
    assert sheet.paper == "a4"
    assert f"Programação da semana {WEEK}" in sheet.file_title


def test_the_report_is_rebuilt_from_the_same_views_as_the_excel(scenario: Scenario) -> None:
    _store(scenario, "IMP-1")

    excel_rows = _week_document(scenario).tables[0].rows
    report_rows = _report_document(scenario).tables[0].rows

    assert excel_rows == report_rows
