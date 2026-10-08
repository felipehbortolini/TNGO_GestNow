"""As rotas da importação e impressão da semana: o modelo, a conferência e a confirmação (D10, D12)."""

from __future__ import annotations

import html
from collections.abc import Mapping
from io import BytesIO

import azure.functions as func
import pytest
from openpyxl import load_workbook

from src.core.auth import PRINCIPAL_HEADER
from src.core.excel import XLSX_CONTENT_TYPE
from src.modulos.programacao_semanal import repository, routes
from src.modulos.programacao_semanal.models import Activity
from tests.cenario_programacao import Scenario
from tests.html_tags import find_all, parse_tags
from tests.identidades import cabecalho_do_principal
from tests.importadores_de_teste import corpo_multipart
from tests.programacao_semanal.test_programacao_importacao import (
    HEADERS,
    prev_row,
    real_row,
    sheet,
)

pytestmark = pytest.mark.usefixtures("sessao_das_rotas")

PLANNER = "ps-plan@example.invalid"
FISCAL = "ps-fisc@example.invalid"
SUPPLIER_A = "ps-forn-a@example.invalid"
VIEWER = "ps-view@example.invalid"
IMPORT_TARGET = "importacao-painel"
CHECK_TARGET = "importacao-conferencia"


def _body(response: func.HttpResponse) -> str:
    return html.unescape(response.get_body().decode())


def _request(
    scenario: Scenario,
    *,
    method: str = "POST",
    email: str = PLANNER,
    alvo: str | None = CHECK_TARGET,
    params: Mapping[str, str] | None = None,
    campos: Mapping[str, str] | None = None,
    arquivo: tuple[str, bytes] | None = None,
) -> func.HttpRequest:
    headers = {PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    body = b""
    if method == "POST":
        body, headers["Content-Type"] = corpo_multipart(campos or {}, arquivo)
    if alvo is not None:
        headers.update({"X-Alpine-Request": "true", "X-Alpine-Target": alvo})
    query = {"projeto": str(scenario.project_id), **(params or {})}
    return func.HttpRequest(
        method=method,
        url="/api/programacao-semanal/importacao",
        headers=headers,
        params=query,
        route_params={},
        body=body,
    )


def _digest(response: func.HttpResponse) -> str:
    tags = parse_tags(_body(response))
    campo = next(tag for tag in find_all(tags, "input") if tag.attrs.get("name") == "conferido")
    return campo.attrs["value"]


def _stored(scenario: Scenario) -> list[Activity]:
    return repository.activities(
        scenario.session, project_id=scenario.project_id, week=None, company_id=None
    )


def _confirm(
    scenario: Scenario,
    content: bytes,
    *,
    email: str = PLANNER,
    conferido: str | None = None,
    ciente: bool = False,
) -> func.HttpResponse:
    campos: dict[str, str] = {}
    if conferido is not None:
        campos["conferido"] = conferido
    if ciente:
        campos["ciente"] = "sim"
    return routes.week_import_confirm(
        _request(scenario, email=email, campos=campos, arquivo=("semana.xlsx", content))
    )


# ── A tela ───────────────────────────────────────────────────────────────


def test_the_screen_serves_the_steps_the_registers_and_the_outputs(scenario: Scenario) -> None:
    response = routes.week_import(
        _request(scenario, method="GET", alvo=IMPORT_TARGET, params={"semana": "S.31/2026"})
    )

    body = _body(response)
    assert response.status_code == 200
    assert f'id="{IMPORT_TARGET}"' in body
    assert 'data-tn-excel="/api/programacao-semanal/importacao/modelo"' in body
    assert 'data-tn-excel="/api/programacao-semanal/semana/excel?semana=S.31%2F2026"' in body
    assert 'data-tn-pdf="/api/programacao-semanal/semana/imprimivel?semana=S.31%2F2026"' in body
    assert 'action="/api/programacao-semanal/importacao/conferir"' in body
    assert 'name="arquivo"' in body
    assert 'id="importacao-conferencia"' in body
    for register in ("Frente 1", "Empresa A", "ps-enc-a", "ps-fisc"):
        assert register in body


def test_the_screen_of_a_viewer_explains_the_restriction(scenario: Scenario) -> None:
    response = routes.week_import(
        _request(scenario, method="GET", email=VIEWER, alvo=IMPORT_TARGET)
    )

    body = _body(response)
    assert response.status_code == 200
    assert "Importação restrita" in body
    assert "/importacao/conferir" not in body


def test_the_screen_in_the_portfolio_asks_for_a_project(scenario: Scenario) -> None:
    response = routes.week_import(
        _request(scenario, method="GET", alvo=IMPORT_TARGET, params={"projeto": "portfolio"})
    )

    body = _body(response)
    assert response.status_code == 200
    assert "Escolha um projeto para importar" in body
    assert "Baixar o modelo" not in body


# ── O modelo ─────────────────────────────────────────────────────────────


def test_the_model_is_a_download_without_the_alpine_header(scenario: Scenario) -> None:
    response = routes.week_import_model(_request(scenario, method="GET", alvo=None))

    assert response.status_code == 200
    assert response.headers["Content-Type"] == XLSX_CONTENT_TYPE
    assert "modelo-de-importacao-programacao-da-semana" in response.headers["Content-Disposition"]
    book = load_workbook(BytesIO(response.get_body()))
    assert book.sheetnames == ["Dados", "Instruções"]
    header_row = next(row for row in book["Dados"].iter_rows(values_only=True) if "Semana" in row)
    assert tuple(header_row[: len(HEADERS)]) == HEADERS


def test_the_model_is_refused_for_whoever_does_not_import(scenario: Scenario) -> None:
    response = routes.week_import_model(_request(scenario, method="GET", email=VIEWER, alvo=None))

    assert response.status_code == 403
    assert response.headers["Content-Type"].startswith("text/plain")


# ── A conferência ────────────────────────────────────────────────────────


def test_the_check_route_lists_the_invalid_lines_and_writes_nothing(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1"), prev_row("IMP-2", location="Frente 9"))

    response = routes.week_import_check(_request(scenario, arquivo=("semana.xlsx", content)))

    body = _body(response)
    assert response.status_code == 200
    assert f'id="{CHECK_TARGET}"' in body
    assert "Linha 3" in body
    assert "não está no cadastro de locais" in body
    assert _stored(scenario) == []
    assert _digest(response)


def test_the_check_route_refuses_what_is_not_a_spreadsheet(scenario: Scenario) -> None:
    response = routes.week_import_check(
        _request(scenario, arquivo=("semana.pdf", b"%PDF-1.7 conteudo"))
    )

    assert response.status_code == 422
    assert "não é uma planilha do Excel (.xlsx)" in _body(response)
    assert _stored(scenario) == []


def test_the_check_route_in_the_portfolio_asks_for_a_project(scenario: Scenario) -> None:
    response = routes.week_import_check(
        _request(
            scenario,
            arquivo=("semana.xlsx", sheet(prev_row())),
            params={"projeto": "portfolio"},
        )
    )

    assert response.status_code == 422
    assert "Escolha um projeto" in _body(response)


@pytest.mark.parametrize("email", [SUPPLIER_A, VIEWER, FISCAL])
def test_the_check_route_refuses_whoever_does_not_import(scenario: Scenario, email: str) -> None:
    response = routes.week_import_check(
        _request(scenario, email=email, arquivo=("semana.xlsx", sheet(prev_row())))
    )

    assert response.status_code == 403
    assert _stored(scenario) == []


# ── A confirmação ────────────────────────────────────────────────────────


def test_the_confirm_route_writes_the_checked_lines(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1"), real_row("IMP-1"))
    checked = routes.week_import_check(_request(scenario, arquivo=("semana.xlsx", content)))

    response = _confirm(scenario, content, conferido=_digest(checked))

    body = _body(response)
    assert response.status_code == 200
    assert "linha(s) de Programação da semana importada(s)" in body
    assert "Abrir a programação" in body
    assert [activity.unique_id for activity in _stored(scenario)] == ["IMP-1"]


def test_the_confirm_route_without_the_check_is_422(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1"))

    response = _confirm(scenario, content)

    assert response.status_code == 422
    assert "Confira a planilha antes de importar" in _body(response)
    assert _stored(scenario) == []


def test_the_confirm_route_with_another_file_is_422(scenario: Scenario) -> None:
    checked = routes.week_import_check(
        _request(scenario, arquivo=("semana.xlsx", sheet(prev_row())))
    )

    response = _confirm(scenario, sheet(prev_row("IMP-2")), conferido=_digest(checked))

    assert response.status_code == 422
    assert "O arquivo mudou depois da conferência" in _body(response)
    assert _stored(scenario) == []


def test_the_confirm_route_with_an_error_line_needs_the_acknowledgement(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1"), prev_row("IMP-2", location="Frente 9"))
    checked = routes.week_import_check(_request(scenario, arquivo=("semana.xlsx", content)))

    response = _confirm(scenario, content, conferido=_digest(checked), ciente=False)

    assert response.status_code == 422
    assert "confirme que está ciente" in _body(response).lower()
    assert _stored(scenario) == []

    accepted = _confirm(scenario, content, conferido=_digest(checked), ciente=True)

    assert accepted.status_code == 200
    assert [activity.unique_id for activity in _stored(scenario)] == ["IMP-1"]


def test_the_week_excel_and_printable_answer_with_their_mechanisms(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1"))
    checked = routes.week_import_check(_request(scenario, arquivo=("semana.xlsx", content)))
    _confirm(scenario, content, conferido=_digest(checked))

    excel = routes.week_excel(
        _request(
            scenario,
            method="GET",
            alvo=None,
            params={"semana": "S.31/2026"},
        )
    )
    printable = routes.week_printable(
        _request(scenario, method="GET", alvo="folha-impressao", params={"semana": "S.31/2026"})
    )

    assert excel.status_code == 200
    assert excel.headers["Content-Type"] == XLSX_CONTENT_TYPE
    assert printable.status_code == 200
    printable_body = _body(printable)
    assert "Programação da semana S.31/2026" in printable_body
    assert "Aderência" in printable_body
    assert "Por contratada" in printable_body
    assert "<a " not in printable_body
    assert "<button" not in printable_body
