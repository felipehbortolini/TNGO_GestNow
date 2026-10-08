"""Rotas do dashboard: o painel, os filtros, o recorte do fornecedor e as exportações (ISSUE-056)."""

from __future__ import annotations

import html
from io import BytesIO

import azure.functions as func
import pytest
from openpyxl import load_workbook

from src.core import calendario
from src.core.excel import XLSX_CONTENT_TYPE
from src.modulos.programacao_semanal import calculations, routes, service
from src.modulos.programacao_semanal.service import NewActivity
from tests.cenario_programacao import WEEK, Scenario
from tests.identidades import requisicao

pytestmark = pytest.mark.usefixtures("sessao_das_rotas")

PLANNER = "ps-plan@example.invalid"
SUPPLIER_A = "ps-forn-a@example.invalid"
SUPPLIER_B = "ps-forn-b@example.invalid"
EMPTY_WEEK = "S.30/2026"


def _body(response: func.HttpResponse) -> str:
    return html.unescape(response.get_body().decode())


def _get(scenario: Scenario, email: str, **params: str) -> dict:
    return {"email": email, "params": {"projeto": str(scenario.project_id), **params}}


def _store(
    scenario: Scenario,
    unique_id: str,
    *,
    week: str = WEEK,
    company_id: int | None = None,
    planned: float = 70.0,
    done: float = 70.0,
) -> None:
    service.store_activity(
        scenario.session,
        author_id=scenario.author_id,
        new=NewActivity(
            project_id=scenario.project_id,
            company_id=company_id or scenario.company_a,
            week=week,
            unique_id=unique_id,
            item=1,
            description=f"Serviço {unique_id}",
            planned_headline=planned,
            author_person_id=scenario.planner.person_id,
            at=calendario.now(),
            location_id=scenario.location_id,
            unit_id=scenario.unit_id,
            foreman_id=scenario.foreman_a,
            inspector_id=scenario.fiscal.person_id,
            planned_days=(planned / calculations.DAYS,) * calculations.DAYS,
            day_shift=(done / calculations.DAYS,) * calculations.DAYS,
            situation=calculations.SITUATION_PUBLISHED,
            approval=calculations.APPROVAL_APPROVED,
            approved_at=calendario.now(),
        ),
    )


def _panel(scenario: Scenario, email: str, **params: str) -> func.HttpResponse:
    return routes.dashboard(
        requisicao(
            "/api/programacao-semanal/dashboard",
            alvo="dash-painel",
            **_get(scenario, email, **params),
        )
    )


# ── O painel e a barra ───────────────────────────────────────────────────


def test_o_painel_serve_os_indicadores_os_graficos_e_as_exportacoes(scenario: Scenario) -> None:
    _store(scenario, "D-1")

    response = _panel(scenario, PLANNER, semana=WEEK)

    body = _body(response)
    assert response.status_code == 200
    assert 'id="dash-painel"' in body
    for label in (
        "Aderência",
        "PPC médio",
        "Evolução da aderência por contratada",
        "Curva S do avanço",
        "Previsto contra realizado",
        "Situação da programação",
        "Aprovação do realizado",
        "Faixa de PPC",
        "Ranking por contratada",
        "Ranking por frente",
        "Ranking por encarregado",
        "Mapa de calor — dia por frente",
        "Produção por dia",
        "O que travou a semana",
    ):
        assert label in body
    for kind in (
        "curva-s-linha",
        "curva-s-barra-linha",
        "kpi-status",
        "tabela-heatmap",
        "comparativo-barras",
    ):
        assert f'data-grafico="{kind}"' in body
    assert 'data-tn-excel="/api/programacao-semanal/dashboard/excel?' in body
    assert 'data-tn-pdf="/api/programacao-semanal/dashboard/imprimivel?' in body


def test_o_painel_marca_a_serie_da_aderencia_como_razao(scenario: Scenario) -> None:
    """O drill depende do par (num, den): sem ``tipo: razao`` a linha não é desenhada."""
    _store(scenario, "D-1")

    body = _body(_panel(scenario, PLANNER, semana=WEEK))

    assert '"tipo": "razao"' in body
    assert '"num":' in body and '"den":' in body


def test_a_barra_serve_os_filtros_do_app(scenario: Scenario) -> None:
    response = routes.dashboard_filters(
        requisicao(
            "/api/programacao-semanal/dashboard/filtros",
            alvo="dash-filtros",
            **_get(scenario, PLANNER),
        )
    )

    body = _body(response)
    assert response.status_code == 200
    assert 'id="dash-filtros"' in body
    assert 'name="semana"' in body
    assert 'name="empresa"' in body
    assert 'name="local"' in body
    assert 'name="encarregado"' in body
    assert 'x-target="dash-painel"' in body


def test_o_painel_do_fornecedor_nao_mostra_outra_empresa(scenario: Scenario) -> None:
    _store(scenario, "A-1")
    _store(scenario, "B-1", company_id=scenario.company_b)

    body = _body(_panel(scenario, SUPPLIER_A, semana=WEEK))

    assert "Empresa A" in body
    assert "Empresa B" not in body


def test_o_painel_do_fornecedor_nao_oferece_a_outra_empresa_no_filtro(
    scenario: Scenario,
) -> None:
    _store(scenario, "B-1", company_id=scenario.company_b)

    body = _body(
        routes.dashboard_filters(
            requisicao(
                "/api/programacao-semanal/dashboard/filtros",
                alvo="dash-filtros",
                **_get(scenario, SUPPLIER_A),
            )
        )
    )

    # O fornecedor não escolhe contratada: o corte é dele e o seletor nem aparece.
    assert 'name="empresa"' not in body
    assert "Empresa B" not in body
    assert 'name="local"' in body and 'name="encarregado"' in body


def test_o_painel_no_portfolio_abre_a_coluna_projeto_nos_gargalos(scenario: Scenario) -> None:
    _store(scenario, "D-1", planned=70.0, done=0.0)

    response = routes.dashboard(
        requisicao(
            "/api/programacao-semanal/dashboard",
            alvo="dash-painel",
            email=PLANNER,
            params={"projeto": "portfolio", "semana": WEEK},
        )
    )

    body = _body(response)
    assert response.status_code == 200
    assert '<th scope="col">Projeto</th>' in body


def test_a_semana_sem_programacao_pede_outra_semana(scenario: Scenario) -> None:
    body = _body(_panel(scenario, PLANNER, semana=EMPTY_WEEK))

    assert "Ainda não há programação nesta semana" in body


def test_o_filtro_sem_resultado_diz_que_e_o_filtro(scenario: Scenario) -> None:
    _store(scenario, "A-1")

    body = _body(_panel(scenario, PLANNER, semana=WEEK, empresa=str(scenario.company_b)))

    assert "Nenhum resultado para os filtros aplicados" in body


# ── As exportações ───────────────────────────────────────────────────────


def test_as_exportacoes_respondem_com_os_mecanismos_da_plataforma(scenario: Scenario) -> None:
    _store(scenario, "D-1")

    excel = routes.dashboard_excel(
        requisicao(
            "/api/programacao-semanal/dashboard/excel",
            email=PLANNER,
            params={"projeto": str(scenario.project_id), "semana": WEEK},
        )
    )
    printable = routes.dashboard_printable(
        requisicao(
            "/api/programacao-semanal/dashboard/imprimivel",
            alvo="folha-impressao",
            **_get(scenario, PLANNER, semana=WEEK),
        )
    )

    assert excel.status_code == 200
    assert excel.headers.get("Content-Type") == XLSX_CONTENT_TYPE
    assert printable.status_code == 200
    body = _body(printable)
    assert "Dashboard da programação" in body
    assert 'data-grafico="curva-s-linha"' in body
    assert "Ranking por contratada" in body


def test_o_excel_do_dashboard_tem_o_resumo_e_as_tabelas_dos_graficos(scenario: Scenario) -> None:
    _store(scenario, "D-1")

    excel = routes.dashboard_excel(
        requisicao(
            "/api/programacao-semanal/dashboard/excel",
            email=PLANNER,
            params={"projeto": str(scenario.project_id), "semana": WEEK},
        )
    )
    book = load_workbook(BytesIO(excel.get_body()))

    assert "Resumo" in book.sheetnames
    assert "Ranking por contratada" in book.sheetnames
    assert "Mapa de calor — dia por frente" in book.sheetnames
    assert any(name.startswith("Evolução da aderência") for name in book.sheetnames)
