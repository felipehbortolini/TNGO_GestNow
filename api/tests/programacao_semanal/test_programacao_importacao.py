"""Importação da semana: o modelo, a conferência linha a linha e a confirmação (D10, D12)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from src.core import calendario, importing
from src.core.errors import AccessDeniedError, InvalidDataError
from src.core.excel import build_workbook
from src.core.rbac import Bond, GeneralProfile, ScheduleRole, User
from src.core.scope import Scope
from src.modulos.configuracoes.models import Person
from src.modulos.programacao_semanal import importers, repository, service
from src.modulos.programacao_semanal.models import Activity, ActivityDay
from src.modulos.programacao_semanal.service import NewActivity
from tests.cenario_programacao import WEEK, Scenario
from tests.identidades import criar_colaborador
from tests.importadores_de_teste import planilha

HEADERS = (
    importers.ITEM_HEADER,
    importers.WEEK_HEADER,
    importers.UNIQUE_ID_HEADER,
    importers.DESCRIPTION_HEADER,
    importers.LOCATION_HEADER,
    importers.COMPANY_HEADER,
    importers.INSPECTOR_HEADER,
    importers.FOREMAN_HEADER,
    importers.HEADLINE_HEADER,
    importers.UNIT_HEADER,
    importers.KIND_HEADER,
    *importers.DAY_HEADERS,
)

SEVEN_TENS = (10, 10, 10, 10, 10, 10, 10)


def prev_row(
    unique_id: str = "IMP-1",
    *,
    week: str = WEEK,
    description: str = "Armação das paredes",
    location: str = "Frente 1",
    company: str = "Empresa A",
    inspector: str = "ps-fisc",
    foreman: str = "ps-enc-a",
    headline: Any = 70,
    unit: str = "m",
    days: tuple[Any, ...] = SEVEN_TENS,
) -> list[Any]:
    """Uma linha ``Prev.`` da planilha do app, com os cadastros do cenário."""
    return [
        "01",
        week,
        unique_id,
        description,
        location,
        company,
        inspector,
        foreman,
        headline,
        unit,
        importers.KIND_PREV,
        *days,
    ]


def real_row(
    unique_id: str = "IMP-1",
    *,
    week: str = WEEK,
    days: tuple[Any, ...] = (7, 7, 7, 7, 7, 7, 7),
) -> list[Any]:
    """Uma linha ``Real.`` da planilha do app: só a semana, a ID e os sete dias."""
    return ["", week, unique_id, "", "", "", "", "", "", "", importers.KIND_REAL, *days]


def sheet(*rows: list[Any]) -> bytes:
    return planilha(rows, cabecalho=HEADERS)


def context(
    scenario: Scenario, *, user: User | None = None, project: bool = True
) -> importing.ImportContext:
    return importing.ImportContext(
        user=user or scenario.planner,
        scope=Scope(project_id=scenario.project_id if project else None, source="url"),
        reference_date=calendario.today(),
    )


def check(
    scenario: Scenario, content: bytes, *, user: User | None = None, project: bool = True
) -> importing.ImportPreview:
    return importing.check(
        scenario.session,
        importer=importers.programacao_importer(),
        upload=importing.Upload(name="semana.xlsx", content=content),
        context=context(scenario, user=user, project=project),
    )


def confirm(
    scenario: Scenario,
    content: bytes,
    *,
    digest: str | None = None,
    acknowledged: bool = True,
    user: User | None = None,
) -> importing.ImportResult:
    upload = importing.Upload(name="semana.xlsx", content=content)
    return importing.confirm(
        scenario.session,
        importer=importers.programacao_importer(),
        upload=upload,
        context=context(scenario, user=user),
        confirmation=importing.Confirmation(
            checked_digest=importing.digest_of(upload) if digest is None else digest,
            acknowledged=acknowledged,
        ),
    )


def activities(scenario: Scenario) -> list[Activity]:
    return repository.activities(
        scenario.session, project_id=scenario.project_id, week=None, company_id=None
    )


def days_of(scenario: Scenario, activity: Activity) -> list[ActivityDay]:
    return repository.days_by_activity(scenario.session, [activity.id])[activity.id]


def store_approved(scenario: Scenario, unique_id: str) -> Activity:
    """An activity already stored with the done work approved: the frozen case."""
    return service.store_activity(
        scenario.session,
        author_id=scenario.author_id,
        new=NewActivity(
            project_id=scenario.project_id,
            company_id=scenario.company_a,
            week=WEEK,
            unique_id=unique_id,
            item=1,
            description="Já aprovada",
            planned_headline=70.0,
            author_person_id=scenario.planner.person_id,
            at=calendario.now(),
            location_id=scenario.location_id,
            unit_id=scenario.unit_id,
            foreman_id=scenario.foreman_a,
            inspector_id=scenario.fiscal.person_id,
            planned_days=SEVEN_TENS,
            day_shift=(5.0,) * 7,
            situation="publicada",
            approval="aprovado",
            approved_at=calendario.now(),
        ),
    )


# ── O modelo ─────────────────────────────────────────────────────────────


def test_the_model_has_the_columns_of_the_app_and_is_read_back(scenario: Scenario) -> None:
    document = importing.model_document(
        scenario.session,
        importer=importers.programacao_importer(),
        context=context(scenario),
    )
    content = build_workbook(document)

    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ["Dados", "Instruções"]
    header_row = next(row for row in book["Dados"].iter_rows(values_only=True) if "Semana" in row)
    assert tuple(header_row[: len(HEADERS)]) == HEADERS
    instrucoes = [
        row[0] for row in book["Instruções"].iter_rows(values_only=True) if row[0] in HEADERS
    ]
    assert instrucoes == list(HEADERS)

    book["Dados"].append(prev_row())
    book.save(buffer := BytesIO())
    preview = check(scenario, buffer.getvalue())
    (line,) = preview.rows
    assert line.writable
    assert line.values[importers.FIELD_WEEK] == WEEK
    assert line.values[importers.FIELD_KIND] == importers.KIND_PREV
    assert line.values[importers.FIELD_COMPANY] == "Empresa A"


# ── A conferência ────────────────────────────────────────────────────────


def test_the_check_lists_the_invalid_lines_with_the_reason_and_writes_nothing(
    scenario: Scenario,
) -> None:
    content = sheet(
        prev_row(),
        prev_row("IMP-2", location="Frente 9"),
        prev_row("IMP-3", week="S.99/2026"),
        prev_row("IMP-4", days=(5, 5, 5, 5, 5, 5, 5)),
        prev_row("IMP-1"),
    )

    preview = check(scenario, content)

    assert [line.number for line in preview.rows] == [2, 3, 4, 5, 6]
    assert [line.writable for line in preview.rows] == [True, False, False, False, False]
    problems = {line.number: " ".join(line.problems) for line in preview.error_rows}
    assert "“Frente 9” não está no cadastro de locais." in problems[3]
    assert "“S.99/2026” não é uma semana válida" in problems[4]
    assert "A soma dos dias (35) não bate com a produção prevista (70)." in problems[5]
    assert "repete a linha 2" in problems[6]
    assert activities(scenario) == []


def test_a_blank_required_column_blocks_the_line(scenario: Scenario) -> None:
    preview = check(scenario, sheet(prev_row("IMP-1", headline="")))

    (line,) = preview.rows
    assert not line.writable
    assert "Produção prevista não pode ficar em branco." in line.problems


def test_an_id_already_in_the_week_blocks_the_line(scenario: Scenario) -> None:
    store_approved(scenario, "IMP-1")

    preview = check(scenario, sheet(prev_row("IMP-1")))

    (line,) = preview.rows
    assert not line.writable
    assert f"A ID “IMP-1” já existe na semana {WEEK}." in line.problems


def test_a_kind_out_of_the_list_blocks_the_line(scenario: Scenario) -> None:
    row = prev_row()
    row[10] = "Total"

    preview = check(scenario, sheet(row))

    (line,) = preview.rows
    assert not line.writable
    assert "valor fora da lista (Prev., Real.)" in " ".join(line.problems)


def test_a_lowercase_company_matches_the_register(scenario: Scenario) -> None:
    preview = check(scenario, sheet(prev_row(company="empresa a")))

    (line,) = preview.rows
    assert line.writable


def test_an_unknown_column_is_ignored(scenario: Scenario) -> None:
    content = planilha([prev_row()], cabecalho=(*HEADERS, "Observação"))

    preview = check(scenario, content)

    assert preview.ignored == ("Observação",)
    assert preview.rows[0].writable


def test_the_header_scan_finds_the_header_below_instructions(scenario: Scenario) -> None:
    content = planilha([prev_row()], cabecalho=HEADERS, acima=5)

    preview = check(scenario, content)

    (line,) = preview.rows
    assert line.writable
    assert line.number == 7


# ── A confirmação ────────────────────────────────────────────────────────


def test_the_confirmation_writes_only_the_lines_without_error(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1"), prev_row("IMP-2", location="Frente 9"))

    result = confirm(scenario, content, acknowledged=True)

    stored = activities(scenario)
    assert [activity.unique_id for activity in stored] == ["IMP-1"]
    assert [line.values[importers.FIELD_UNIQUE_ID] for line in result.written] == ["IMP-1"]
    assert [line.values[importers.FIELD_UNIQUE_ID] for line in result.skipped] == ["IMP-2"]
    assert stored[0].situation == "em_elaboracao"
    assert stored[0].company_id == scenario.company_a
    assert days_of(scenario, stored[0])[0].planned == 10.0


def test_the_real_line_completes_the_day_shift_of_the_activity(scenario: Scenario) -> None:
    result = confirm(scenario, sheet(prev_row("IMP-1"), real_row("IMP-1")))

    (activity,) = activities(scenario)
    stored = days_of(scenario, activity)
    assert [day.done_day_shift for day in stored] == [7.0] * 7
    assert [day.done_night_shift for day in stored] == [0.0] * 7
    assert activity.approval == "pendente"
    assert len(result.written) == 2


def test_a_real_row_with_custom_days_keeps_the_order(scenario: Scenario) -> None:
    confirm(scenario, sheet(prev_row(), real_row(days=(1, 2, 3, 4, 5, 6, 7))))

    (activity,) = activities(scenario)
    stored = days_of(scenario, activity)
    assert [day.done_day_shift for day in stored] == [1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0]


def test_the_real_line_without_its_prev_is_ignored(scenario: Scenario) -> None:
    result = confirm(scenario, sheet(real_row("SEM-PREV")))

    assert activities(scenario) == []
    assert len(result.written) == 1


def test_the_real_line_of_an_approved_activity_is_an_error(scenario: Scenario) -> None:
    store_approved(scenario, "IMP-1")

    preview = check(scenario, sheet(real_row("IMP-1")))

    (line,) = preview.rows
    assert not line.writable
    assert importers.FROZEN in line.problems


def test_the_confirmation_requires_the_digest_of_the_checked_file(scenario: Scenario) -> None:
    content = sheet(prev_row())

    with pytest.raises(InvalidDataError, match="Confira a planilha"):
        confirm(scenario, content, digest="")
    with pytest.raises(InvalidDataError, match="O arquivo mudou"):
        confirm(scenario, content, digest="0" * 64)
    assert activities(scenario) == []


def test_the_confirmation_with_an_error_line_needs_the_acknowledgement(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1"), prev_row("IMP-2", location="Frente 9"))

    with pytest.raises(InvalidDataError, match="ciente"):
        confirm(scenario, content, acknowledged=False)
    assert activities(scenario) == []


def test_the_confirmation_is_all_or_nothing(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1"), prev_row("IMP-1", week="S.031/2026"))

    assert [line.writable for line in check(scenario, content).rows] == [True, True]

    with pytest.raises(InvalidDataError, match="Linha 3") as refused:
        confirm(scenario, content)

    assert "já existe" in str(refused.value)
    assert activities(scenario) == []


def test_the_planned_days_of_a_fractional_row_are_kept(scenario: Scenario) -> None:
    content = sheet(prev_row("IMP-1", headline=0.5, days=(0.5, 0, 0, 0, 0, 0, 0)))

    result = confirm(scenario, content)

    assert len(result.written) == 1
    (activity,) = activities(scenario)
    assert days_of(scenario, activity)[0].planned == 0.5


# ── Acesso e escopo ──────────────────────────────────────────────────────


def test_a_viewer_cannot_check_the_sheet(scenario: Scenario) -> None:
    with pytest.raises(AccessDeniedError):
        check(scenario, sheet(prev_row()), user=scenario.viewer)


def test_the_planner_of_another_project_cannot_check(scenario: Scenario) -> None:
    other = importing.ImportContext(
        user=scenario.planner,
        scope=Scope(project_id=scenario.other_project_id, source="url"),
        reference_date=calendario.today(),
    )

    with pytest.raises(AccessDeniedError):
        importing.check(
            scenario.session,
            importer=importers.programacao_importer(),
            upload=importing.Upload(name="semana.xlsx", content=sheet(prev_row())),
            context=other,
        )


def test_a_viewer_with_the_planner_role_imports(scenario: Scenario) -> None:
    collaborator = criar_colaborador(
        scenario.session,
        email="ps-view-plan@example.invalid",
        perfil="Visualizador",
        papeis=[(scenario.project_id, "Planejador")],
    )
    person = scenario.session.get(Person, collaborator.person_id)
    user = User(
        id=collaborator.id,
        person_id=person.id,
        name=person.name,
        email=person.email,
        general_profile=GeneralProfile(collaborator.general_profile),
        bond=Bond(collaborator.bond),
        schedule_roles=frozenset({(scenario.project_id, ScheduleRole.PLANNER)}),
    )

    preview = check(scenario, sheet(prev_row()), user=user)

    assert preview.writable_rows


def test_the_portfolio_asks_for_a_project(scenario: Scenario) -> None:
    with pytest.raises(InvalidDataError, match="Escolha um projeto"):
        check(scenario, sheet(prev_row()), project=False)


# ── O registro do importador ─────────────────────────────────────────────


def test_the_importer_is_registered_under_its_key() -> None:
    importers.register_importers()

    registered = importing.find(importers.KEY)
    assert registered is not None
    assert registered.module == "programacao_semanal"
    assert registered.title == importers.TITLE


def test_registering_again_keeps_the_importer_already_registered() -> None:
    importers.register_importers()
    before = len(importing.registered())

    importers.register_importers()

    assert len(importing.registered()) == before
