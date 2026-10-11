"""As rotas do Início (ISSUE-079): o fragmento no alvo, os perfis, o Excel e o imprimível.

A rota roda na transação do teste (``sessao_das_rotas``), em modo demonstração, e a requisição
é montada como o shell a envia: cabeçalho do Alpine AJAX, cookie do seletor de perfil e o
escopo no parâmetro ``projeto``. O Visualizador vê o resumo; o Fornecedor, que não alcança o
módulo (D7), recebe 403. A data de referência é a de hoje (``calendario.today()``).
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.core import calendario
from src.core.auth import DEMO_COOKIE
from src.core.export_document import ValueKind, format_value
from src.modulos.inicio import routes
from tests.identidades import requisicao
from tests.inicio.conftest import Cenario, Demonstracao

HOME_TARGET = "inicio-conteudo"
CABECALHO_AUSENTE = "cabeçalho não encontrado: {title}"


def _requisicao(
    caminho: str,
    *,
    colaborador_id: int,
    params: dict[str, str] | None = None,
    alvo: str = HOME_TARGET,
):
    return requisicao(
        caminho,
        cookies={DEMO_COOKIE: str(colaborador_id)},
        params=params,
        alvo=alvo,
    )


def _header_row(sheet: Worksheet, title: str) -> int:
    for row in sheet.iter_rows(min_row=1, max_row=30):
        if any(cell.value == title for cell in row):
            return row[0].row
    raise AssertionError(CABECALHO_AUSENTE.format(title=title))


def test_a_rota_do_resumo_responde_o_fragmento_no_alvo(
    demonstracao: Demonstracao, sessao_das_rotas
) -> None:
    del sessao_das_rotas
    resposta = routes.home_view(
        _requisicao(
            "/api/inicio/resumo",
            colaborador_id=demonstracao.admin.id,
            params={"projeto": str(demonstracao.project_id)},
        )
    )

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert 'id="inicio-conteudo"' in corpo
    assert "TN-2026-014 · Construção de uma nova fábrica" in corpo
    assert f"Data de referência {format_value(ValueKind.DATE, calendario.today())}" in corpo
    assert "Orçamento R$ 44,6 mi" in corpo
    assert "Ações atrasadas" in corpo
    assert "Não conformidades abertas" in corpo
    assert "Pontos de atenção" in corpo
    assert "CLM-TN-2026-0003 notificado fora do prazo contratual" in corpo
    assert "Sistema 210 Moagem bloqueado por item A" in corpo
    assert "Módulos" in corpo
    assert corpo.count("home__modulo--") == 8, "os oito cards dos módulos"
    assert 'data-tn-excel="/api/inicio/resumo/excel"' in corpo
    assert 'data-tn-pdf="/api/inicio/resumo/imprimivel"' in corpo
    assert 'href="/central-acoes/acoes?status=atrasada"' in corpo


def test_o_portfolio_da_rota_traz_a_etiqueta_do_portfolio(
    demonstracao: Demonstracao, sessao_das_rotas
) -> None:
    del sessao_das_rotas
    resposta = routes.home_view(
        _requisicao("/api/inicio/resumo", colaborador_id=demonstracao.admin.id)
    )

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert "Portfólio de projetos · 3 projetos" in corpo
    assert "Orçamento R$ 70,2 mi" in corpo
    assert "TN-2026-014 · Contratar consultoria credenciada" in corpo


def test_os_perfis_do_projeto_veem_o_resumo(cenario: Cenario) -> None:
    for perfil in ("Visualizador", "Membro", "Gestor"):
        resposta = routes.home_view(
            _requisicao(
                "/api/inicio/resumo",
                colaborador_id=cenario.colaboradores[perfil],
                params={"projeto": str(cenario.project.id)},
            )
        )
        assert resposta.status_code == 200, perfil
        assert "Indicadores-chave por módulo" in resposta.get_body().decode(), perfil


def test_o_fornecedor_nao_alcanca_o_inicio(cenario: Cenario) -> None:
    resposta = routes.home_view(
        _requisicao(
            "/api/inicio/resumo",
            colaborador_id=cenario.colaboradores["Fornecedor"],
            params={"projeto": str(cenario.project.id)},
        )
    )

    assert resposta.status_code == 403
    assert "Programação Semanal" in resposta.get_body().decode()


def test_o_excel_do_resumo_traz_o_resumo_a_atencao_e_os_modulos(
    demonstracao: Demonstracao, sessao_das_rotas
) -> None:
    del sessao_das_rotas
    resposta = routes.home_excel(
        _requisicao(
            "/api/inicio/resumo/excel",
            colaborador_id=demonstracao.admin.id,
            params={"projeto": str(demonstracao.project_id)},
        )
    )

    assert resposta.status_code == 200
    assert resposta.headers.get("Content-Type", "").startswith("application/vnd.openxmlformats")
    pasta = load_workbook(BytesIO(resposta.get_body()))
    assert set(pasta.sheetnames) >= {"Resumo", "Pontos de atenção", "Módulos"}

    resumo = pasta["Resumo"]
    indicadores = [cell.value for row in resumo.iter_rows() for cell in row]
    assert "01 Central de Ações · Ações atrasadas" in indicadores
    assert "07 HSE · Dias sem acidente com afastamento" in indicadores

    atencao = pasta["Pontos de atenção"]
    header = _header_row(atencao, "Ponto de atenção")
    colunas = [cell.value for cell in atencao[header]]
    assert colunas == ["Módulo", "Nível", "Ponto de atenção", "Detalhe"]
    assert atencao.page_setup.orientation == "landscape"

    modulos = pasta["Módulos"]
    header = _header_row(modulos, "Módulo")
    assert [cell.value for cell in modulos[header]] == ["Módulo", "Descrição", "Selos"]


def test_o_imprimivel_do_resumo_sai_em_a4(demonstracao: Demonstracao, sessao_das_rotas) -> None:
    del sessao_das_rotas
    resposta = routes.home_printable(
        _requisicao(
            "/api/inicio/resumo/imprimivel",
            colaborador_id=demonstracao.admin.id,
            params={"projeto": str(demonstracao.project_id)},
            alvo="folha-impressao",
        )
    )

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert "folha--a4" in corpo
    assert "Resumo do projeto" in corpo
    assert "01 Central de Ações · Ações atrasadas" in corpo
    assert "Pontos de atenção" in corpo
