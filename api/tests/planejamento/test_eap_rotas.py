"""Rotas da EAP (ISSUE-036, D14): fragmento, dicionário, Portfólio e exportações.

A costura é HTTP: a rota recebe a ``HttpRequest`` do navegador e a resposta é lida como a pessoa
a veria. O cenário repete o projeto 1 da demonstração em pequena escala: uma área, uma subárea,
três pacotes (um de planejamento), a Rev 0 vigente, o desdobramento 1.1.3 > 1.1.2 e os vínculos
com a EAC. O Portfólio é somente leitura; a carga de demonstração é conferida pelos números do
critério: 110 itens, 38 pacotes no projeto 1, Rev 2 vigente, 1 desdobramento e 2 vencidos.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from urllib.parse import unquote

import azure.functions as func
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.carga import DEMO_ANCHOR, run_demonstration
from src.core.auth import _user_of
from src.core.scope import Scope
from src.modulos.configuracoes import service as configuracoes
from src.modulos.configuracoes.models import Company, Person, Project, Unit
from src.modulos.financeiro.models import EacItem
from src.modulos.planejamento import routes, service
from src.modulos.planejamento.models import EapBreakdown, EapItem, EapItemStep, EapRevision
from tests.apoio_6wla import Envio, pedido
from tests.html_tags import find_by_id, parse_tags
from tests.identidades import criar_colaborador, criar_empresa, criar_projeto

D = Decimal
BASE = "/api/planejamento/eap"
EMAIL_MEMBRO = "membro-eap-rotas@example.invalid"
EMAIL_VISUALIZADOR = "visualizador-eap-rotas@example.invalid"
ALVO_CONTEUDO = "eap-conteudo"
ALVO_FORMULARIO = "eap-formulario"


@dataclass(frozen=True)
class CenarioEapRotas:
    """O projeto com EAP do teste das rotas e os itens criados."""

    session: Session
    projeto: Project
    empresa: Company
    responsavel: Person
    itens: dict[str, EapItem]

    @property
    def params(self) -> dict[str, str]:
        return {"projeto": str(self.projeto.id)}


def _eac(
    session: Session, projeto: Project, code: str, description: str, parent: EacItem | None = None
) -> EacItem:
    item = EacItem(
        project_id=projeto.id,
        parent_id=parent.id if parent is not None else None,
        code=code,
        description=description,
        level=len(code.split(".")),
    )
    session.add(item)
    session.flush()
    return item


def _item(
    session: Session,
    projeto: Project,
    code: str,
    description: str,
    parent: EapItem | None = None,
    **fields: object,
) -> EapItem:
    item = EapItem(
        project_id=projeto.id,
        parent_id=parent.id if parent is not None else None,
        code=code,
        description=description,
        level=len(code.split(".")),
        **fields,
    )
    session.add(item)
    session.flush()
    return item


@pytest.fixture
def cenario_eap_rotas(sessao: Session) -> CenarioEapRotas:
    """Uma EAP pequena com Rev 0, desdobramento e os vínculos do dicionário."""
    projeto = criar_projeto(sessao, codigo="TN-EAP-ROTA-001", nome="Fábrica de rotas")
    empresa = criar_empresa(sessao, "Montadora de rotas")
    membro = criar_colaborador(sessao, email=EMAIL_MEMBRO, nome="Maria Rota", perfil="Membro")
    criar_colaborador(sessao, email=EMAIL_VISUALIZADOR, perfil="Visualizador")
    unidade = Unit(kind="medida", code="un", name="Unidade")
    sessao.add(unidade)
    sessao.flush()
    responsavel = sessao.get(Person, membro.person_id)
    assert responsavel is not None

    raiz_eac = _eac(sessao, projeto, "1", "Obra")
    sub_eac = _eac(sessao, projeto, "1.1", "Fundações", raiz_eac)
    _eac(sessao, projeto, "1.1.1", "Estacas", sub_eac)
    eac_blocos = _eac(sessao, projeto, "1.1.2", "Blocos", sub_eac)
    _eac(sessao, projeto, "1.1.3", "Partida", sub_eac)

    obra = _item(sessao, projeto, "1", "Obra civil")
    fundacoes = _item(sessao, projeto, "1.1", "Fundações", obra)
    estacas = _item(
        sessao,
        projeto,
        "1.1.1",
        "Estacas",
        fundacoes,
        kind="Trabalho",
        criterion="Unidades",
        weight=D("55"),
        predicted=D("90"),
        start=date(2026, 2, 1),
        finish=date(2026, 12, 1),
        quantity=D("100"),
        executed=D("90"),
        unit_id=unidade.id,
        company_id=empresa.id,
        owner_id=responsavel.id,
        deliverable="Estacas concretadas",
        acceptance="Laudo de integridade",
        eac_item_id=eac_blocos.id,
    )
    blocos = _item(
        sessao,
        projeto,
        "1.1.2",
        "Blocos",
        fundacoes,
        kind="Trabalho",
        criterion="Marco 0/100",
        weight=D("40"),
        predicted=D("50"),
        start=date(2026, 7, 1),
        finish=date(2026, 12, 15),
        state="Não iniciado",
        eac_item_id=eac_blocos.id,
    )
    partida = _item(
        sessao,
        projeto,
        "1.1.3",
        "Partida assistida",
        fundacoes,
        kind="Planejamento",
        weight=D("5"),
        predicted=D("0"),
        start=date(2026, 8, 1),
        finish=date(2027, 1, 1),
    )
    sessao.add(
        EapRevision(
            project_id=projeto.id,
            approved_by_id=responsavel.id,
            revision=0,
            revision_date=date(2026, 1, 5),
            alteration="Linha de base",
            justification="Linha de base aprovada no gate de investimento.",
        )
    )
    sessao.add(
        EapItemStep(item_id=blocos.id, order=1, name="Concretagem", weight=D("100"), pct=D("0"))
    )
    sessao.add(
        EapBreakdown(
            project_id=projeto.id,
            source_item_id=partida.id,
            target_item_id=blocos.id,
            by_id=responsavel.id,
            revision=0,
            applied_on=date(2026, 2, 2),
            weight=D("5"),
            justification="Plano de partida detalhado.",
        )
    )
    sessao.flush()
    return CenarioEapRotas(
        session=sessao,
        projeto=projeto,
        empresa=empresa,
        responsavel=responsavel,
        itens={"1.1.1": estacas, "1.1.2": blocos, "1.1.3": partida},
    )


def _texto(resposta: func.HttpResponse) -> str:
    return unquote(resposta.get_body().decode())


def _toast(resposta: func.HttpResponse) -> str:
    return unquote(resposta.headers["X-TN-Toast"])


def _corpo_do_dicionario(cenario: CenarioEapRotas, **campos: str) -> dict[str, str]:
    item = cenario.itens["1.1.1"]
    base = {
        "descricao": item.description,
        "entregavel": item.deliverable or "",
        "aceitacao": item.acceptance or "",
        "empresa": str(cenario.empresa.id),
        "responsavel": str(cenario.responsavel.id),
        "eac": "",
        "versao": str(item.version),
    }
    base.update(campos)
    return base


# ── O fragmento da tela ──────────────────────────────────────────────────


def test_o_fragmento_traz_kpis_arvore_revisoes_e_desdobramentos(
    cenario_eap_rotas: CenarioEapRotas,
) -> None:
    resposta = routes.eap_content(
        pedido(
            BASE,
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_eap_rotas.params, alvo=ALVO_CONTEUDO),
        )
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert find_by_id(parse_tags(texto), ALVO_CONTEUDO) is not None
    assert "Estrutura analítica do projeto" in texto
    assert "Estacas" in texto
    assert "sem medição" in texto
    assert "Rev 0" in texto
    assert "Vigente" in texto
    assert "1.1.3 · Partida assistida" in texto
    assert "Não iniciado" in texto
    assert 'data-carteira="nao"' in texto


@pytest.mark.usefixtures("cenario_eap_rotas")
def test_o_portfolio_e_somente_leitura_e_traz_a_coluna_projeto() -> None:
    resposta = routes.eap_content(pedido(BASE, email=EMAIL_MEMBRO, envio=Envio(alvo=ALVO_CONTEUDO)))
    texto = _texto(resposta)

    assert resposta.status_code == 200
    assert 'data-carteira="sim"' in texto
    assert "Portfólio" in texto
    assert "TN-EAP-ROTA-001" in texto
    assert "/dicionario" not in texto


def test_o_dicionario_recusa_o_portfolio(cenario_eap_rotas: CenarioEapRotas) -> None:
    item = cenario_eap_rotas.itens["1.1.1"]
    formulario = routes.eap_dictionary_form(
        pedido(
            f"{BASE}/itens/{item.id}/dicionario",
            email=EMAIL_MEMBRO,
            envio=Envio(alvo=ALVO_FORMULARIO, rota={"item_id": str(item.id)}),
        )
    )
    gravacao = routes.eap_dictionary_save(
        pedido(
            f"{BASE}/itens/{item.id}/dicionario",
            email=EMAIL_MEMBRO,
            envio=Envio(
                metodo="POST",
                campos=_corpo_do_dicionario(cenario_eap_rotas),
                alvo=ALVO_FORMULARIO,
                rota={"item_id": str(item.id)},
            ),
        )
    )

    assert formulario.status_code == 403
    assert gravacao.status_code == 403


def test_sem_o_cabecalho_do_alpine_a_tela_manda_para_o_shell(
    cenario_eap_rotas: CenarioEapRotas,
) -> None:
    resposta = routes.eap_content(
        pedido(
            BASE,
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_eap_rotas.params, alpine=False),
        )
    )

    assert resposta.status_code == 302


# ── O formulário e a gravação do dicionário ──────────────────────────────


def test_o_formulario_do_dicionario_mostra_os_valores_e_as_listas(
    cenario_eap_rotas: CenarioEapRotas,
) -> None:
    item = cenario_eap_rotas.itens["1.1.1"]
    resposta = routes.eap_dictionary_form(
        pedido(
            f"{BASE}/itens/{item.id}/dicionario",
            email=EMAIL_MEMBRO,
            envio=Envio(
                params=cenario_eap_rotas.params,
                alvo=ALVO_FORMULARIO,
                rota={"item_id": str(item.id)},
            ),
        )
    )

    texto = _texto(resposta)
    tags = parse_tags(texto)
    assert resposta.status_code == 200
    assert find_by_id(tags, ALVO_FORMULARIO) is not None
    names = {tag.attrs.get("name") for tag in tags if tag.name in {"input", "textarea", "select"}}
    assert {"descricao", "entregavel", "aceitacao", "empresa", "responsavel", "eac", "versao"} <= (
        names
    )
    assert "Estacas" in texto
    assert "1.1.2" in texto
    assert "Peso, datas e critério" in texto


def test_gravar_o_dicionario_devolve_o_conteudo_e_fecha_o_formulario(
    cenario_eap_rotas: CenarioEapRotas,
) -> None:
    item = cenario_eap_rotas.itens["1.1.1"]
    campos = _corpo_do_dicionario(
        cenario_eap_rotas,
        descricao="Estacas executadas",
        entregavel="Estacas concretadas e ensaiadas",
        aceitacao="Laudo de integridade aprovado",
        eac="1.1.2",
    )

    resposta = routes.eap_dictionary_save(
        pedido(
            f"{BASE}/itens/{item.id}/dicionario",
            email=EMAIL_MEMBRO,
            envio=Envio(
                metodo="POST",
                campos=campos,
                params=cenario_eap_rotas.params,
                rota={"item_id": str(item.id)},
                alvo=f"{ALVO_FORMULARIO} {ALVO_CONTEUDO}",
            ),
        )
    )

    texto = _texto(resposta)
    tags = parse_tags(texto)
    assert resposta.status_code == 200
    assert _toast(resposta) == "Pacote 1.1.1 atualizado."
    assert find_by_id(tags, ALVO_CONTEUDO) is not None
    formulario = find_by_id(tags, ALVO_FORMULARIO)
    assert formulario is not None
    assert formulario.clean_text() == ""
    assert "Estacas executadas" in texto
    cenario_eap_rotas.session.refresh(item)
    assert (item.description, item.deliverable, item.acceptance) == (
        "Estacas executadas",
        "Estacas concretadas e ensaiadas",
        "Laudo de integridade aprovado",
    )
    assert item.eac_item_id is not None
    assert item.weight == D("55.00")


def test_gravar_com_descricao_curta_volta_422_com_a_mensagem_no_campo(
    cenario_eap_rotas: CenarioEapRotas,
) -> None:
    item = cenario_eap_rotas.itens["1.1.1"]
    resposta = routes.eap_dictionary_save(
        pedido(
            f"{BASE}/itens/{item.id}/dicionario",
            email=EMAIL_MEMBRO,
            envio=Envio(
                metodo="POST",
                campos=_corpo_do_dicionario(cenario_eap_rotas, descricao="ab"),
                params=cenario_eap_rotas.params,
                rota={"item_id": str(item.id)},
                alvo=f"{ALVO_FORMULARIO} {ALVO_CONTEUDO}",
            ),
        )
    )

    texto = _texto(resposta)
    assert resposta.status_code == 422
    assert find_by_id(parse_tags(texto), ALVO_FORMULARIO) is not None
    assert "Descreva o pacote" in texto
    cenario_eap_rotas.session.refresh(item)
    assert item.description == "Estacas"


def test_gravar_com_versao_vencida_volta_409(cenario_eap_rotas: CenarioEapRotas) -> None:
    item = cenario_eap_rotas.itens["1.1.1"]
    resposta = routes.eap_dictionary_save(
        pedido(
            f"{BASE}/itens/{item.id}/dicionario",
            email=EMAIL_MEMBRO,
            envio=Envio(
                metodo="POST",
                campos=_corpo_do_dicionario(cenario_eap_rotas, versao="999"),
                params=cenario_eap_rotas.params,
                rota={"item_id": str(item.id)},
                alvo=f"{ALVO_FORMULARIO} {ALVO_CONTEUDO}",
            ),
        )
    )

    assert resposta.status_code == 409


# ── Permissão, Excel e imprimível ────────────────────────────────────────


def test_o_visualizador_le_a_tela_mas_recebe_403_no_formulario_e_na_gravacao(
    cenario_eap_rotas: CenarioEapRotas,
) -> None:
    item = cenario_eap_rotas.itens["1.1.1"]
    leitura = routes.eap_content(
        pedido(
            BASE,
            email=EMAIL_VISUALIZADOR,
            envio=Envio(params=cenario_eap_rotas.params, alvo=ALVO_CONTEUDO),
        )
    )
    formulario = routes.eap_dictionary_form(
        pedido(
            f"{BASE}/itens/{item.id}/dicionario",
            email=EMAIL_VISUALIZADOR,
            envio=Envio(
                params=cenario_eap_rotas.params,
                alvo=ALVO_FORMULARIO,
                rota={"item_id": str(item.id)},
            ),
        )
    )
    gravacao = routes.eap_dictionary_save(
        pedido(
            f"{BASE}/itens/{item.id}/dicionario",
            email=EMAIL_VISUALIZADOR,
            envio=Envio(
                metodo="POST",
                campos=_corpo_do_dicionario(cenario_eap_rotas),
                params=cenario_eap_rotas.params,
                rota={"item_id": str(item.id)},
                alvo=ALVO_FORMULARIO,
            ),
        )
    )

    assert leitura.status_code == 200
    assert formulario.status_code == 403
    assert gravacao.status_code == 403


def test_o_excel_e_o_imprimivel_trazem_o_mesmo_conteudo_da_tela(
    cenario_eap_rotas: CenarioEapRotas,
) -> None:
    excel = routes.eap_excel(
        pedido(
            f"{BASE}/excel",
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_eap_rotas.params, alpine=False),
        )
    )
    imprimivel = routes.eap_printable(
        pedido(
            f"{BASE}/imprimivel",
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_eap_rotas.params, alvo="folha-impressao"),
        )
    )

    assert excel.status_code == 200
    assert excel.get_body()[:2] == b"PK"
    sheets = load_workbook(BytesIO(excel.get_body())).worksheets
    assert [sheet.title for sheet in sheets] == [
        "Resumo",
        "Estrutura analítica do projeto",
        "Revisões da EAP",
        "Desdobramentos da revisão vigen",
    ]
    texto = _texto(imprimivel)
    assert imprimivel.status_code == 200
    assert "EAP: estrutura analítica do projeto" in texto
    assert "Revisões da EAP" in texto
    assert "Desdobramentos da revisão vigente" in texto


@pytest.mark.usefixtures("cenario_eap_rotas")
def test_o_excel_do_portfolio_traz_a_coluna_projeto() -> None:
    excel = routes.eap_excel(pedido(f"{BASE}/excel", email=EMAIL_MEMBRO, envio=Envio(alpine=False)))

    sheets = load_workbook(BytesIO(excel.get_body())).worksheets
    tree = next(sheet for sheet in sheets if sheet.title == "Estrutura analítica do projeto")
    header_row = int(tree.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))

    assert excel.status_code == 200
    assert tree.cell(row=header_row, column=1).value == "Projeto"


# ── A carga com os números do critério ───────────────────────────────────


def test_a_carga_de_demonstracao_traz_os_numeros_do_criterio(sessao: Session) -> None:
    run_demonstration(sessao, reference_date=DEMO_ANCHOR)

    projeto = sessao.scalars(select(Project).where(Project.code == "TN-2026-014")).one()
    access = configuracoes.find_access_by_email(sessao, "leonardo.gomes@exemplo.com")
    assert access is not None
    visao = service.eap_overview(
        sessao,
        user=_user_of(access),
        scope=Scope(project_id=projeto.id, source="url"),
        reference_date=DEMO_ANCHOR,
    )

    assert visao.current_revision is not None
    assert (visao.current_revision.revision, visao.current_revision.package_count) == (2, 38)
    assert [(item.source_code, item.target_code) for item in visao.breakdowns] == [
        ("5.2.1", "5.2.2")
    ]
    assert visao.indicators.overdue == 2
    assert [node.code for node in visao.nodes if node.is_work_package and node.overdue] == [
        "2.1.1",
        "3.2.1",
    ]
