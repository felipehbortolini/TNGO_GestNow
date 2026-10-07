"""Rotas dos KPIs de planejamento (ISSUE-040): fragmento, Portfólio, vazio e exportações.

A costura é HTTP: a rota recebe a ``HttpRequest`` do navegador e a resposta é lida como a pessoa
a veria. O cenário é o da fachada em escala menor: cinco áreas da EAP com os pesos e as médias
do protótipo (previsto 65,7 e real 61,8 no corte de set/2026, o SPI de 0,94 do oráculo), a
carteira com o segundo projeto em 100/100 e o projeto sem curva. O Portfólio pondera 75/25 pelos
orçamentos e traz a coluna Projeto; sem curva a tela mostra o vazio de origem.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from urllib.parse import unquote

import azure.functions as func
import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from src.modulos.configuracoes.models import Project
from src.modulos.governanca.models import ChangeImpact, ChangeRequest
from src.modulos.planejamento import routes
from src.modulos.planejamento.models import (
    CurvaFisicaMonth,
    CurvaFisicaRevision,
    EapItem,
    EapMeasurement,
    EapRevision,
    EapRevisionItem,
)
from tests.apoio_6wla import Envio, pedido
from tests.html_tags import find_by_id, parse_tags
from tests.identidades import criar_colaborador, criar_projeto

D = Decimal
BASE = "/api/planejamento/kpis"
EMAIL_MEMBRO = "membro-kpis-rotas@example.invalid"
ALVO_CONTEUDO = "kpis-conteudo"

JUL = date(2026, 7, 1)
AUG = date(2026, 8, 1)
SEP = date(2026, 9, 1)
OCT = date(2026, 10, 1)

MESES = (JUL, AUG, SEP, OCT)

AREAS = (
    ("1", "Engenharia", "15", "98", "96", date(2026, 12, 31)),
    ("2", "Suprimentos", "30", "80", "78", date(2026, 12, 31)),
    ("3", "Obras civis", "25", "70", "66", date(2026, 12, 31)),
    ("4", "Montagem eletromecânica", "25", "38", "30", date(2026, 8, 31)),
    ("5", "Comissionamento", "5", "0", "0", date(2026, 12, 31)),
)


@dataclass(frozen=True)
class CenarioKpisRotas:
    """O projeto dos KPIs, o da carteira e o projeto sem curva, na transação do teste."""

    session: Session
    projeto: Project
    outro: Project
    sem_curva: Project
    itens: dict[str, EapItem]

    @property
    def params(self) -> dict[str, str]:
        return {"projeto": str(self.projeto.id)}


def _item(
    session: Session,
    project: Project,
    code: str,
    description: str,
    parent: EapItem | None = None,
    **fields: object,
) -> EapItem:
    item = EapItem(
        project_id=project.id,
        parent_id=parent.id if parent is not None else None,
        code=code,
        description=description,
        level=len(code.split(".")),
        **fields,
    )
    session.add(item)
    session.flush()
    return item


def _medicao(session: Session, autor_id: int, item: EapItem, dia: date, para: str) -> None:
    session.add(
        EapMeasurement(
            item_id=item.id,
            author_id=autor_id,
            measured_on=dia,
            from_pct=D("0"),
            to_pct=D(para),
            note="Boletim",
        )
    )
    session.flush()


def _congelar(
    session: Session, projeto: Project, itens: tuple[EapItem, ...], *, baseline: tuple[str, ...]
) -> None:
    """A revisão 0 com os pesos congelados e a linha de base da curva, congelada à mão."""
    revisao = EapRevision(
        project_id=projeto.id,
        approved_by_id=projeto.manager_id,
        revision=0,
        revision_date=date(2026, 1, 5),
        alteration="Linha de base",
        justification="Linha de base aprovada no gate de investimento.",
    )
    session.add(revisao)
    session.flush()
    for item in itens:
        session.add(
            EapRevisionItem(revision_id=revisao.id, item_id=item.id, weight=item.weight or D(0))
        )
    curva = CurvaFisicaRevision(project_id=projeto.id, eap_revision_id=revisao.id)
    session.add(curva)
    session.flush()
    session.add_all(
        CurvaFisicaMonth(revision_id=curva.id, month=month, baseline_pct=D(value))
        for month, value in zip(MESES, baseline, strict=True)
    )
    session.flush()


def _montar_projeto(session: Session, projeto: Project) -> dict[str, EapItem]:
    """As cinco áreas do cenário, cada uma com o seu pacote de trabalho."""
    itens: dict[str, EapItem] = {}
    for code, nome, peso, previsto, real, termino in AREAS:
        area = _item(session, projeto, code, nome)
        pacote = _item(
            session,
            projeto,
            f"{code}.1.1",
            f"Pacote de {nome}",
            area,
            kind="Trabalho",
            criterion="Percentual estimado",
            weight=D(peso),
            predicted=D(previsto),
            estimated_pct=D(real),
            start=date(2026, 1, 1),
            finish=termino,
        )
        if D(real) > 0:
            _medicao(session, projeto.manager_id, pacote, date(2026, 8, 15), real)
        itens[code] = pacote
    return itens


def _montar_outro(session: Session, projeto: Project) -> EapItem:
    """O projeto da carteira: uma área e um pacote único em 100/100."""
    area = _item(session, projeto, "1", "Engenharia")
    pacote = _item(
        session,
        projeto,
        "1.1.1",
        "Pacote único",
        area,
        kind="Trabalho",
        criterion="Percentual estimado",
        weight=D("100"),
        predicted=D("100"),
        estimated_pct=D("100"),
        start=date(2026, 1, 1),
        finish=date(2026, 12, 31),
    )
    _medicao(session, projeto.manager_id, pacote, date(2026, 8, 15), "100")
    return pacote


def _mudanca(session: Session, projeto: Project) -> ChangeRequest:
    """A SM aprovada com impacto em escopo que nenhuma revisão incorporou."""
    change = ChangeRequest(
        project_id=projeto.id,
        requester_id=projeto.manager_id,
        code="SM-TN-2026-0400",
        title="Inclusão do sistema de tratamento de efluentes",
        kind="Escopo",
        origin="Cliente",
        priority="Normal",
        request_date=date(2026, 9, 1),
        description="Descrição da mudança de teste.",
        situation="Aprovada",
    )
    session.add(change)
    session.flush()
    session.add(
        ChangeImpact(
            change_id=change.id,
            analyst_id=projeto.manager_id,
            analysis_date=date(2026, 9, 10),
            cost_cents=0,
            term_days=0,
            scope="Nova unidade de tratamento",
            quality="Sem impacto",
            risks="Sem impacto",
            safety="Sem impacto",
            contract="Sem impacto",
            affects_contract_milestone=False,
        )
    )
    session.flush()
    return change


@pytest.fixture
def cenario_kpis_rotas(sessao: Session) -> CenarioKpisRotas:
    """O cenário do protótipo na transação do teste: 65,7/61,8 no corte e a carteira 75/25."""
    projeto = criar_projeto(sessao, codigo="TN-KPI-ROTAS-001", nome="Fábrica dos KPIs")
    outro = criar_projeto(sessao, codigo="TN-KPI-ROTAS-002", nome="Caldeira dos KPIs")
    sem_curva = criar_projeto(sessao, codigo="TN-KPI-ROTAS-003", nome="Projeto sem curva")
    projeto.budget_cents = 75_000_000
    outro.budget_cents = 25_000_000
    criar_colaborador(sessao, email=EMAIL_MEMBRO, nome="Maria Rota", perfil="Membro")
    sessao.flush()
    itens = _montar_projeto(sessao, projeto)
    _congelar(sessao, projeto, tuple(itens.values()), baseline=("50.0", "60.0", "65.7", "100.0"))
    item_outro = _montar_outro(sessao, outro)
    _congelar(sessao, outro, (item_outro,), baseline=("80.0", "90.0", "100.0", "100.0"))
    _mudanca(sessao, projeto)
    return CenarioKpisRotas(
        session=sessao, projeto=projeto, outro=outro, sem_curva=sem_curva, itens=itens
    )


def _texto(resposta: func.HttpResponse) -> str:
    return unquote(resposta.get_body().decode())


def _graficos(texto: str) -> dict[str, dict[str, object]]:
    """O ``data-dados`` de cada gráfico do fragmento, lido como a biblioteca o leria."""
    tags = parse_tags(texto)
    return {
        tag.attrs["data-grafico"]: json.loads(tag.attrs["data-dados"])
        for tag in tags
        if "data-grafico" in tag.attrs
    }


# ── O fragmento da tela ──────────────────────────────────────────────────


def test_o_fragmento_traz_kpis_graficos_tabela_e_rodape(
    cenario_kpis_rotas: CenarioKpisRotas,
) -> None:
    resposta = routes.planning_kpis_content(
        pedido(
            BASE,
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_kpis_rotas.params, alvo=ALVO_CONTEUDO),
        )
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert find_by_id(parse_tags(texto), ALVO_CONTEUDO) is not None
    assert 'data-situacao="conteudo"' in texto
    assert 'data-carteira="nao"' in texto
    assert "SPI do projeto" in texto
    assert "Meta: ≥ 1,00" in texto
    assert "0,94" in texto
    assert "-0,09 vs. mês anterior" in texto
    assert "Previsto x real" in texto
    assert "Previsto: 65,7%" in texto
    assert "61,8%" in texto
    assert "acumulado em set/26" in texto
    assert "Desvio acumulado" in texto
    assert "Maior desvio por área" in texto
    assert "-8,0" in texto
    assert "Limite: ≥ -2" in texto
    assert "Montagem eletromecânica: real 30% x previsto 38%" in texto
    assert "Pacotes com término vencido" in texto
    assert "SMs a incorporar" in texto
    assert "SPI por mês" in texto
    assert "Previsto x real por área" in texto
    assert "Avanço por área" in texto
    assert "Pesos somam 100%; a média ponderada fecha com a Curva S" in texto
    assert "Projeto (ponderado)" in texto


def test_o_fragmento_traz_os_dois_graficos_com_dados_parseaveis(
    cenario_kpis_rotas: CenarioKpisRotas,
) -> None:
    resposta = routes.planning_kpis_content(
        pedido(
            BASE,
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_kpis_rotas.params, alvo=ALVO_CONTEUDO),
        )
    )

    graficos = _graficos(_texto(resposta))
    assert set(graficos) == {"linhas-multiplas", "comparativo-barras"}
    assert [item["valores"]["spi"] for item in graficos["linhas-multiplas"]["periodos"]] == [
        0.0,
        1.03,
        0.94,
    ]
    painel = graficos["comparativo-barras"]["paineis"][0]
    assert [categoria["id"] for categoria in painel["categorias"]] == ["1", "2", "3", "4", "5"]
    assert [linha["rotulo"] for linha in painel["linhas"]] == ["Previsto", "Real"]


@pytest.mark.usefixtures("cenario_kpis_rotas")
def test_sem_o_cabecalho_do_alpine_a_tela_manda_para_o_shell() -> None:
    resposta = routes.planning_kpis_content(
        pedido(BASE, email=EMAIL_MEMBRO, envio=Envio(alpine=False))
    )

    assert resposta.status_code == 302


@pytest.mark.usefixtures("cenario_kpis_rotas")
def test_o_portfolio_traz_a_coluna_projeto_e_o_spi_ponderado() -> None:
    resposta = routes.planning_kpis_content(
        pedido(BASE, email=EMAIL_MEMBRO, envio=Envio(alvo=ALVO_CONTEUDO))
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert 'data-carteira="sim"' in texto
    assert "SPI ponderado da carteira" in texto
    assert "0,96" in texto
    assert "Maior desvio por projeto" in texto
    assert "TN-KPI-ROTAS-001: real 62% x previsto 66%" in texto
    assert "Previsto x real por projeto" in texto
    assert "Avanço por projeto" in texto
    assert '<th scope="col">Projeto</th>' in texto
    assert "TN-KPI-ROTAS-001 · Fábrica dos KPIs" in texto
    assert "TN-KPI-ROTAS-002 · Caldeira dos KPIs" in texto
    assert "75,0%" in texto
    assert "25,0%" in texto
    assert "Carteira (ponderado)" in texto
    assert (
        "Pesos da ponderação da carteira (somam 100%); a média ponderada fecha com a Curva S da carteira"
        in texto
    )


def test_o_projeto_sem_curva_mostra_o_vazio_de_origem(
    cenario_kpis_rotas: CenarioKpisRotas,
) -> None:
    resposta = routes.planning_kpis_content(
        pedido(
            BASE,
            email=EMAIL_MEMBRO,
            envio=Envio(
                params={"projeto": str(cenario_kpis_rotas.sem_curva.id)}, alvo=ALVO_CONTEUDO
            ),
        )
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert 'data-situacao="vazio-origem"' in texto
    assert "Sem Curva S cadastrada para este projeto." in texto


def test_o_portfolio_sem_curva_mostra_o_vazio_de_origem(sessao: Session) -> None:
    criar_colaborador(
        sessao, email="visualizador-kpis-vazio@example.invalid", perfil="Visualizador"
    )

    resposta = routes.planning_kpis_content(
        pedido(
            BASE, email="visualizador-kpis-vazio@example.invalid", envio=Envio(alvo=ALVO_CONTEUDO)
        )
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert 'data-situacao="vazio-origem"' in texto
    assert "Sem Curva S cadastrada nos projetos da carteira." in texto


# ── Excel e imprimível ───────────────────────────────────────────────────


def test_o_excel_e_o_imprimivel_trazem_o_mesmo_conteudo_da_tela(
    cenario_kpis_rotas: CenarioKpisRotas,
) -> None:
    excel = routes.planning_kpis_excel(
        pedido(f"{BASE}/excel", email=EMAIL_MEMBRO, envio=Envio(params=cenario_kpis_rotas.params))
    )
    imprimivel = routes.planning_kpis_printable(
        pedido(
            f"{BASE}/imprimivel",
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_kpis_rotas.params, alvo="folha-impressao"),
        )
    )

    assert excel.status_code == 200
    assert excel.get_body()[:2] == b"PK"
    sheets = load_workbook(BytesIO(excel.get_body())).worksheets
    assert [sheet.title for sheet in sheets] == ["Resumo", "SPI por mês", "Avanço por área"]
    texto = _texto(imprimivel)
    assert imprimivel.status_code == 200
    assert "KPIs de prazo" in texto
    assert "SPI por mês" in texto
    assert "Avanço por área" in texto
    assert "Projeto (ponderado)" in texto
    assert 'data-grafico="linhas-multiplas"' in texto
    assert 'data-grafico="comparativo-barras"' in texto


def test_o_projeto_sem_curva_exporta_o_documento_vazio(
    cenario_kpis_rotas: CenarioKpisRotas,
) -> None:
    params = {"projeto": str(cenario_kpis_rotas.sem_curva.id)}
    excel = routes.planning_kpis_excel(
        pedido(f"{BASE}/excel", email=EMAIL_MEMBRO, envio=Envio(params=params))
    )
    imprimivel = routes.planning_kpis_printable(
        pedido(
            f"{BASE}/imprimivel",
            email=EMAIL_MEMBRO,
            envio=Envio(params=params, alvo="folha-impressao"),
        )
    )

    assert excel.status_code == 200
    assert excel.get_body()[:2] == b"PK"
    assert imprimivel.status_code == 200
    assert "KPIs de prazo" in _texto(imprimivel)


@pytest.mark.usefixtures("cenario_kpis_rotas")
def test_o_excel_do_portfolio_traz_a_coluna_projeto() -> None:
    excel = routes.planning_kpis_excel(pedido(f"{BASE}/excel", email=EMAIL_MEMBRO, envio=Envio()))

    sheets = load_workbook(BytesIO(excel.get_body())).worksheets
    avancos = next(sheet for sheet in sheets if sheet.title == "Avanço por projeto")
    header_row = int(avancos.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))

    assert excel.status_code == 200
    assert avancos.cell(row=header_row, column=1).value == "Projeto"
