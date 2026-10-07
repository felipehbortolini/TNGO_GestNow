"""Rotas da Curva S física (ISSUE-039, D14): fragmento, Portfólio e exportações.

A costura é HTTP: a rota recebe a ``HttpRequest`` do navegador e a resposta é lida como a pessoa
a veria. O cenário repete o projeto do teste da fachada em escala menor: duas curvas congeladas
(75/25 pelos orçamentos), a referência em 25/09/2026 e o projeto sem curva. O Portfólio é a média
ponderada e traz a composição com a coluna ``Projeto``, também nas exportações.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from urllib.parse import unquote

import azure.functions as func
import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from src.modulos.configuracoes.models import Project, Unit
from src.modulos.planejamento import routes, service
from src.modulos.planejamento.models import EapItem, EapMeasurement, EapRevision, EapRevisionItem
from tests.apoio_6wla import Envio, pedido
from tests.html_tags import find_by_id, parse_tags
from tests.identidades import criar_colaborador, criar_projeto

D = Decimal
BASE = "/api/planejamento/curva-fisica"
EMAIL_MEMBRO = "membro-curva-rotas@example.invalid"
ALVO_CONTEUDO = "curva-fisica-conteudo"


@dataclass(frozen=True)
class CenarioCurvaRotas:
    """Os dois projetos com curva congelada e o projeto sem curva."""

    session: Session
    projeto: Project
    outro: Project
    sem_curva: Project
    itens: dict[str, EapItem]

    @property
    def params(self) -> dict[str, str]:
        return {"projeto": str(self.projeto.id)}


def _item(
    session: Session,
    projeto: Project,
    code: str,
    description: str,
    parent: EapItem | None = None,
    **fields: object,
) -> EapItem:
    item = EapItem(
        project_id=projeto.id,
        parent_id=parent.id if parent is not None else None,
        code=code,
        description=description,
        level=len(code.split(".")),
        **fields,
    )
    session.add(item)
    session.flush()
    return item


def _curva(
    session: Session,
    projeto: Project,
    unidade: Unit,
    autor_id: int,
    medicoes: tuple[tuple[str, date, str], ...],
) -> dict[str, EapItem]:
    area = _item(session, projeto, "1", "Obra civil")
    sub = _item(session, projeto, "1.1", "Fundações", area)
    estacas = _item(
        session,
        projeto,
        "1.1.1",
        "Estacas",
        sub,
        kind="Trabalho",
        criterion="Unidades",
        weight=D("60"),
        predicted=D("100"),
        start=date(2026, 1, 1),
        finish=date(2026, 6, 30),
        quantity=D("100"),
        executed=D("100"),
        unit_id=unidade.id,
    )
    blocos = _item(
        session,
        projeto,
        "1.1.2",
        "Blocos",
        sub,
        kind="Trabalho",
        criterion="Unidades",
        weight=D("40"),
        predicted=D("0"),
        start=date(2026, 7, 1),
        finish=date(2026, 12, 31),
        quantity=D("100"),
        executed=D("50"),
        unit_id=unidade.id,
    )
    itens = {"1.1.1": estacas, "1.1.2": blocos}
    revisao = EapRevision(
        project_id=projeto.id,
        approved_by_id=projeto.manager_id,
        revision=0,
        revision_date=date(2026, 1, 5),
        alteration="Linha de base",
        justification="Linha de base aprovada no gate de investimento.",
    )
    session.add(revisao)
    session.flush()
    for item in itens.values():
        session.add(
            EapRevisionItem(revision_id=revisao.id, item_id=item.id, weight=item.weight or D(0))
        )
    for code, dia, para in medicoes:
        session.add(
            EapMeasurement(
                item_id=itens[code].id,
                author_id=autor_id,
                measured_on=dia,
                from_pct=D("0"),
                to_pct=D(para),
                note="Boletim",
            )
        )
    session.flush()
    service.freeze_curve(
        session, project_id=projeto.id, revision_id=revisao.id, reference_date=date(2026, 9, 25)
    )
    return itens


def _curva_simples(session: Session, projeto: Project, unidade: Unit, autor_id: int) -> EapItem:
    area = _item(session, projeto, "1", "Obra civil")
    sub = _item(session, projeto, "1.1", "Fundações", area)
    item = _item(
        session,
        projeto,
        "1.1.1",
        "Estacas",
        sub,
        kind="Trabalho",
        criterion="Unidades",
        weight=D("100"),
        predicted=D("100"),
        start=date(2026, 3, 1),
        finish=date(2026, 11, 30),
        quantity=D("100"),
        executed=D("100"),
        unit_id=unidade.id,
    )
    revisao = EapRevision(
        project_id=projeto.id,
        approved_by_id=projeto.manager_id,
        revision=0,
        revision_date=date(2026, 1, 5),
        alteration="Linha de base",
        justification="Linha de base aprovada no gate de investimento.",
    )
    session.add(revisao)
    session.flush()
    session.add(
        EapRevisionItem(revision_id=revisao.id, item_id=item.id, weight=item.weight or D(0))
    )
    session.add(
        EapMeasurement(
            item_id=item.id,
            author_id=autor_id,
            measured_on=date(2026, 6, 15),
            from_pct=D("0"),
            to_pct=D("100"),
            note="Boletim",
        )
    )
    session.flush()
    service.freeze_curve(
        session, project_id=projeto.id, revision_id=revisao.id, reference_date=date(2026, 9, 25)
    )
    return item


@pytest.fixture
def cenario_curva_rotas(sessao: Session) -> CenarioCurvaRotas:
    """Duas curvas congeladas e um projeto sem curva, na transação do teste."""
    projeto = criar_projeto(sessao, codigo="TN-CURVA-ROTA-001", nome="Fábrica de rotas")
    outro = criar_projeto(sessao, codigo="TN-CURVA-ROTA-002", nome="Caldeira de rotas")
    sem_curva = criar_projeto(sessao, codigo="TN-CURVA-ROTA-003", nome="Sem curva")
    projeto.budget_cents = 75_000_000
    outro.budget_cents = 25_000_000
    unidade = Unit(kind="medida", code="un", name="Unidade")
    sessao.add(unidade)
    membro = criar_colaborador(sessao, email=EMAIL_MEMBRO, nome="Maria Rota", perfil="Membro")
    sessao.flush()
    itens = _curva(
        sessao,
        projeto,
        unidade,
        membro.person_id,
        (
            ("1.1.1", date(2026, 2, 10), "50"),
            ("1.1.1", date(2026, 4, 20), "100"),
            ("1.1.2", date(2026, 8, 5), "50"),
        ),
    )
    _curva_simples(sessao, outro, unidade, membro.person_id)
    return CenarioCurvaRotas(
        session=sessao,
        projeto=projeto,
        outro=outro,
        sem_curva=sem_curva,
        itens=itens,
    )


def _texto(resposta: func.HttpResponse) -> str:
    return unquote(resposta.get_body().decode())


# ── O fragmento da tela ──────────────────────────────────────────────────


def test_o_fragmento_traz_kpis_grafico_tabela_e_o_corte(
    cenario_curva_rotas: CenarioCurvaRotas,
) -> None:
    resposta = routes.curva_fisica_content(
        pedido(
            BASE,
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_curva_rotas.params, alvo=ALVO_CONTEUDO),
        )
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert find_by_id(parse_tags(texto), ALVO_CONTEUDO) is not None
    assert "Previsto acumulado" in texto
    assert "Real acumulado" in texto
    assert "Desvio" in texto
    assert "SPI" in texto
    assert "Término pela tendência" in texto
    assert "78,8%" in texto
    assert "80,0%" in texto
    assert "+1,2" in texto
    assert 'href="/planejamento/kpis?projeto=' in texto
    assert 'data-grafico="curva-s-linha"' in texto
    assert "Período a período" in texto
    assert "Corte em set/26" in texto
    assert "corte" in texto
    assert 'data-carteira="nao"' in texto
    assert "Composição da carteira" not in texto


@pytest.mark.usefixtures("cenario_curva_rotas")
def test_sem_o_cabecalho_do_alpine_a_tela_manda_para_o_shell() -> None:
    resposta = routes.curva_fisica_content(
        pedido(BASE, email=EMAIL_MEMBRO, envio=Envio(alpine=False))
    )

    assert resposta.status_code == 302


@pytest.mark.usefixtures("cenario_curva_rotas")
def test_o_portfolio_traz_a_composicao_com_a_coluna_projeto() -> None:
    resposta = routes.curva_fisica_content(
        pedido(BASE, email=EMAIL_MEMBRO, envio=Envio(alvo=ALVO_CONTEUDO))
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert 'data-carteira="sim"' in texto
    assert "Composição da carteira" in texto
    assert '<th scope="col">Projeto</th>' in texto
    assert "TN-CURVA-ROTA-001 · Fábrica de rotas" in texto
    assert "TN-CURVA-ROTA-002 · Caldeira de rotas" in texto
    assert "75,00%" in texto
    assert "25,00%" in texto
    assert "75,9%" in texto
    assert "+24,1" in texto
    assert "+6,03" in texto
    assert "média ponderada pelos pesos dos projetos na carteira" in texto


def test_o_projeto_sem_curva_mostra_o_vazio_de_origem(
    cenario_curva_rotas: CenarioCurvaRotas,
) -> None:
    resposta = routes.curva_fisica_content(
        pedido(
            BASE,
            email=EMAIL_MEMBRO,
            envio=Envio(
                params={"projeto": str(cenario_curva_rotas.sem_curva.id)}, alvo=ALVO_CONTEUDO
            ),
        )
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert 'data-situacao="vazio-origem"' in texto
    assert "Este projeto ainda não tem Curva S cadastrada." in texto


# ── Excel e imprimível ───────────────────────────────────────────────────


def test_o_excel_e_o_imprimivel_trazem_o_mesmo_conteudo_da_tela(
    cenario_curva_rotas: CenarioCurvaRotas,
) -> None:
    excel = routes.curva_fisica_excel(pedido(f"{BASE}/excel", email=EMAIL_MEMBRO, envio=Envio()))
    imprimivel = routes.curva_fisica_printable(
        pedido(
            f"{BASE}/imprimivel",
            email=EMAIL_MEMBRO,
            envio=Envio(params=cenario_curva_rotas.params, alvo="folha-impressao"),
        )
    )

    assert excel.status_code == 200
    assert excel.get_body()[:2] == b"PK"
    sheets = load_workbook(BytesIO(excel.get_body())).worksheets
    assert [sheet.title for sheet in sheets] == [
        "Resumo",
        "Curva S período a período",
        "Composição da carteira",
    ]
    compose = next(sheet for sheet in sheets if sheet.title == "Composição da carteira")
    header_row = int(compose.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    assert compose.cell(row=header_row, column=1).value == "Projeto"
    texto = _texto(imprimivel)
    assert imprimivel.status_code == 200
    assert "Curva S física" in texto
    assert "Curva S período a período" in texto
    assert "Corte" in texto
    assert 'data-grafico="curva-s-linha"' in texto


def test_o_visualizador_le_a_tela(cenario_curva_rotas: CenarioCurvaRotas) -> None:
    criar_colaborador(
        cenario_curva_rotas.session,
        email="visualizador-curva-rotas@example.invalid",
        perfil="Visualizador",
    )
    resposta = routes.curva_fisica_content(
        pedido(
            BASE,
            email="visualizador-curva-rotas@example.invalid",
            envio=Envio(params=cenario_curva_rotas.params, alvo=ALVO_CONTEUDO),
        )
    )

    assert resposta.status_code == 200
    assert "Previsto acumulado" in _texto(resposta)


@pytest.mark.usefixtures("cenario_curva_rotas")
def test_o_projeto_desconhecido_cai_no_portfolio() -> None:
    resposta = routes.curva_fisica_content(
        pedido(
            BASE,
            email=EMAIL_MEMBRO,
            envio=Envio(params={"projeto": "999999"}, alvo=ALVO_CONTEUDO),
        )
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert 'data-carteira="sim"' in texto


def test_o_portfolio_sem_curva_mostra_o_vazio_de_origem(sessao: Session) -> None:
    criar_colaborador(sessao, email="visualizador-vazio@example.invalid", perfil="Visualizador")

    resposta = routes.curva_fisica_content(
        pedido(BASE, email="visualizador-vazio@example.invalid", envio=Envio(alvo=ALVO_CONTEUDO))
    )

    texto = _texto(resposta)
    assert resposta.status_code == 200
    assert 'data-situacao="vazio-origem"' in texto
    assert "Nenhum projeto da carteira tem Curva S cadastrada." in texto
