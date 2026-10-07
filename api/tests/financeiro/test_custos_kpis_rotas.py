"""Rotas dos KPIs de custo (ISSUE-042, HU-092, D12, D14).

Costura HTTP como em ``test_contingencia_rotas.py``. A carga de demonstração é a base do escopo,
com a âncora injetada no relógio: no corte de set/26 o projeto 1 tem CPI 0,96, SPI de custo 0,94,
CV de menos R$ 1,1 mi, VAC de menos R$ 655 mil, contingência consumida de 20,0% e comprometido de
92,2% do orçamento. O fragmento traz os cartões com a variação do mês anterior, o gráfico de CPI e
SPI (``linhas-multiplas``) e a tabela de valor agregado com a EAC pelo CPI e o TCPI — no Portfólio,
o desempenho por projeto com a coluna Projeto. O Excel e o imprimível levam o mesmo conteúdo; o
Visualizador lê e o fornecedor não.
"""

from __future__ import annotations

import json
from io import BytesIO

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.carga import DEMO_ANCHOR, run_demonstration
from src.core import auth, calendario
from src.modulos.configuracoes.models import Project
from src.modulos.financeiro import routes
from tests.financeiro.conftest import MEMBER_EMAIL, EacScenario
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal

ADMIN_EMAIL = "leonardo.gomes@exemplo.com"
VIEWER_EMAIL = "visualizador@cliente.demo"
SUPPLIER_EMAIL = "carlos.nunes@exemplo.com"
PROJECT_CODE = "TN-2026-014"
PATH = "/api/financeiro/kpis"
CONTENT_TARGET = "kpis-conteudo"
FOLHA_TARGET = "folha-impressao"

KPI_LABELS = (
    "CPI (desempenho de custo)",
    "SPI de custo (desempenho de prazo)",
    "CV (variação de custo)",
    "VAC (variação no término)",
    "Contingência consumida",
    "Comprometido",
)
VALUE_HEADERS = (
    "Mês",
    "PV",
    "EV",
    "AC",
    "CV",
    "SV",
    "CPI",
    "SPI",
    "EAC pelo CPI",
    "TCPI",
)
PROJECT_HEADERS = (
    "Projeto",
    "BAC",
    "EV",
    "AC",
    "CPI",
    "SPI",
    "Projeção no término",
    "VAC",
    "Contingência consumida",
)
RESUMO_HEADERS = ("Indicador", "Descrição", "Valor", "Índice", "Leitura")


@pytest.fixture(autouse=True)
def data_de_referencia(monkeypatch: pytest.MonkeyPatch) -> None:
    """O relógio parado na âncora da demonstração: os números do mock não mudam."""
    monkeypatch.setattr(calendario, "today", lambda: DEMO_ANCHOR)


@pytest.fixture
def carga(rotas_na_transacao: Session) -> Session:
    """A base de demonstração na transação do teste, com as rotas lendo a mesma sessão."""
    run_demonstration(rotas_na_transacao, reference_date=DEMO_ANCHOR)
    return rotas_na_transacao


def _request(
    path: str,
    *,
    email: str = ADMIN_EMAIL,
    target: str | None = CONTENT_TARGET,
    **params: str,
) -> func.HttpRequest:
    """A request como o navegador a envia; ``target=None`` é um link puro (sem Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params={},
        body=b"",
    )


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _table_text(response: func.HttpResponse, attr: str) -> str:
    """O texto de dentro da tabela marcada com o atributo (``data-valor``...)."""
    table = next(tag for tag in _tags(response) if tag.name == "table" and attr in tag.attrs)
    return table.clean_text()


def _projeto(session: Session) -> Project:
    return session.scalars(select(Project).where(Project.code == PROJECT_CODE)).one()


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> int:
    return int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))


# ── Leitura do fragmento ─────────────────────────────────────────────────


def test_os_kpis_do_projeto_trazem_cartoes_e_os_dois_graficos(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.custos_kpis_content(_request(PATH, projeto=str(project.id)))

    assert response.status_code == 200
    body = _body(response)
    assert find_by_id(_tags(response), CONTENT_TARGET) is not None
    assert 'data-situacao="conteudo"' in body
    assert 'data-carteira="nao"' in body
    for rotulo in KPI_LABELS:
        assert rotulo in body
    assert "0,96" in body, "o CPI do corte"
    assert "0,94" in body, "o SPI de custo do corte"
    assert "-R$ 1.137.200,00" in body, "o CV do corte"
    assert "-R$ 655.000,00" in body, "o VAC do corte"
    assert "20,0%" in body, "o consumo da contingência do projeto 1"
    assert "92,2%" in body, "o comprometido do corte"
    assert "Meta: ≥ 1,00" in body
    assert "-0,02 vs. mês anterior" in body, "a variação do CPI no rodapé"
    assert "+0,01 vs. mês anterior" in body, "a variação do SPI no rodapé"
    assert "EV menos AC · mês anterior -R$ 366.200,00" in body
    assert "R$ 400.000,00 de R$ 2.000.000,00" in body
    assert "R$ 41.130.000,00 do orçamento" in body
    assert "Corte em set/26 · EV = avanço físico real x orçamento (BAC)" in body

    graficos = {
        dados["titulo"]: dados
        for dados in (
            json.loads(tag.attrs["data-dados"])
            for tag in _tags(response)
            if tag.attrs.get("data-grafico") == "linhas-multiplas"
        )
    }
    assert list(graficos) == ["CPI e SPI por mês", "Valor agregado acumulado"]

    indice = graficos["CPI e SPI por mês"]
    assert [serie["id"] for serie in indice["series"]] == ["cpi", "spi"]
    assert indice["referencias"] == [{"valor": 1, "rotulo": "Meta (1,00)", "papel": "linha-base"}]
    assert len(indice["periodos"]) == 9, "de jan/26 ao corte"
    assert indice["periodos"][-1]["valores"] == {"cpi": 0.96, "spi": 0.94}
    assert indice["base_zero"] is False

    agregado = graficos["Valor agregado acumulado"]
    assert [serie["id"] for serie in agregado["series"]] == ["pv", "ev", "ac"]
    assert [serie["papel"] for serie in agregado["series"]] == [
        "linha-base",
        "realizado",
        "previsto",
    ]
    assert agregado["periodos"][-1]["valores"] == {
        "pv": 29302.2,
        "ev": 27562.8,
        "ac": 28700.0,
    }


def test_os_kpis_do_projeto_trazem_o_resumo_e_a_tabela_periodo(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.custos_kpis_content(_request(PATH, projeto=str(project.id)))

    resumo = _table_text(response, "data-resumo")
    for coluna in RESUMO_HEADERS:
        assert coluna in resumo
    for sigla in ("BAC", "PV", "EV", "AC", "CV", "SV", "EAC (CPI)", "Projeção", "VAC", "TCPI"):
        assert sigla in resumo
    assert "Projeção no término (bottom-up)" in resumo
    assert "R$ 44.600.000,00" in resumo, "o BAC"
    assert "R$ 45.255.000,00" in resumo, "a projeção no término bottom-up"
    assert "1,07" in resumo, "o TCPI"
    assert "Gastou mais do que agregou" in resumo
    assert "Agregou menos do que o planejado" in resumo
    assert "BAC menos projeção no término" in resumo
    assert "Precisa render mais do que até agora para fechar no BAC" in resumo

    valor = _table_text(response, "data-valor")
    for coluna in VALUE_HEADERS:
        assert coluna in valor
    assert "R$ 29.302.200,00" in valor, "o PV do corte"
    assert "R$ 27.562.800,00" in valor, "o EV do corte"
    assert "R$ 28.700.000,00" in valor, "o AC do corte"
    assert "R$ 46.440.129,45" in valor, "a EAC pelo CPI"
    assert "1,07" in valor, "o TCPI do corte"
    assert "Projeto" not in valor, (
        "a coluna Projeto só entra no Portfólio no desempenho por projeto"
    )


@pytest.mark.usefixtures("carga")
def test_o_portfolio_traz_a_soma_e_o_desempenho_por_projeto() -> None:
    response = routes.custos_kpis_content(_request(PATH, projeto="portfolio"))

    assert response.status_code == 200
    body = _body(response)
    assert 'data-carteira="sim"' in body
    assert "0,97" in body, "o CPI da soma dos projetos"
    assert "carteira: soma do EV, PV e AC dos projetos" in body
    assert "Desempenho de custo por projeto" in body

    projetos = _table_text(response, "data-projetos")
    for coluna in PROJECT_HEADERS:
        assert coluna in projetos
    assert "TN-2026-014" in projetos and "TN-2026-021" in projetos and "TN-2026-027" in projetos
    assert "R$ 44.600.000,00" in projetos, "o BAC do projeto 1"


@pytest.mark.usefixtures("carga")
def test_sem_o_cabecalho_do_alpine_os_kpis_mandam_para_o_shell() -> None:
    response = routes.custos_kpis_content(_request(PATH, target=None, projeto="1"))

    assert response.status_code == 302


def test_escopo_sem_curva_mostra_o_vazio_de_origem(cenario: EacScenario) -> None:
    response = routes.custos_kpis_content(
        _request(PATH, email=MEMBER_EMAIL, projeto=str(cenario.first.id))
    )

    assert response.status_code == 200
    body = _body(response)
    assert 'data-situacao="vazio-origem"' in body
    assert "Este projeto ainda não tem curva física e financeira para o valor agregado." in body
    assert "CPI (desempenho de custo)" not in body


@pytest.mark.usefixtures("cenario")
def test_a_carteira_sem_curva_mostra_o_vazio_de_origem() -> None:
    response = routes.custos_kpis_content(_request(PATH, email=MEMBER_EMAIL, projeto="portfolio"))

    assert response.status_code == 200
    body = _body(response)
    assert 'data-situacao="vazio-origem"' in body
    assert "Nenhum projeto da carteira tem curva física e financeira para o valor agregado." in body


# ── Permissões ───────────────────────────────────────────────────────────


def test_o_visualizador_le_os_kpis(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.custos_kpis_content(
        _request(PATH, email=VIEWER_EMAIL, projeto=str(project.id))
    )

    assert response.status_code == 200
    assert "CPI (desempenho de custo)" in _body(response)


def test_o_fornecedor_nao_le_os_kpis_nem_baixa_o_excel(carga: Session) -> None:
    project = _projeto(carga)

    conteudo = routes.custos_kpis_content(
        _request(PATH, email=SUPPLIER_EMAIL, projeto=str(project.id))
    )
    assert conteudo.status_code == 403

    excel = routes.custos_kpis_excel(
        _request(PATH + "/excel", email=SUPPLIER_EMAIL, target=None, projeto=str(project.id))
    )
    assert excel.status_code == 403


# ── Excel e imprimível ───────────────────────────────────────────────────


def test_o_excel_dos_kpis_traz_o_resumo_o_valor_agregado_e_o_periodo(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.custos_kpis_excel(
        _request(PATH + "/excel", target=None, projeto=str(project.id))
    )

    assert response.status_code == 200
    assert response.headers["Content-Disposition"].startswith(
        'attachment; filename="kpis-de-custo-'
    )
    sheets = _sheets(response)
    assert list(sheets) == ["Resumo", "Valor agregado na data de corte", "Período a período"]

    labels = {
        sheets["Resumo"].cell(row=row, column=1).value
        for row in range(1, sheets["Resumo"].max_row + 1)
    }
    assert set(KPI_LABELS) <= labels

    resumo = sheets["Valor agregado na data de corte"]
    first_row = _header_row(resumo)
    assert [resumo.cell(row=first_row, column=column).value for column in range(1, 6)] == list(
        RESUMO_HEADERS
    )
    rows = {
        str(resumo.cell(row=row, column=1).value): row
        for row in range(first_row + 1, resumo.max_row + 1)
    }
    assert set(rows) == {
        "BAC",
        "PV",
        "EV",
        "AC",
        "CV",
        "SV",
        "EAC (CPI)",
        "Projeção",
        "VAC",
        "TCPI",
    }
    assert resumo.cell(row=rows["BAC"], column=3).value == 44_600_000.0, "o BAC em reais"
    assert resumo.cell(row=rows["Projeção"], column=3).value == 45_255_000.0, (
        "a projeção no término bottom-up em reais"
    )
    assert resumo.cell(row=rows["TCPI"], column=4).value == 1.07, "o TCPI no índice"

    valor = sheets["Período a período"]
    first_row = _header_row(valor)
    assert [valor.cell(row=first_row, column=column).value for column in range(1, 11)] == list(
        VALUE_HEADERS
    )
    corte = next(
        row
        for row in range(first_row + 1, valor.max_row + 1)
        if str(valor.cell(row=row, column=1).value).startswith("set/26")
    )
    assert valor.cell(row=corte, column=2).value == 29_302_200.0, "o PV do corte em reais"


@pytest.mark.usefixtures("carga")
def test_o_excel_da_carteira_traz_o_desempenho_por_projeto() -> None:
    response = routes.custos_kpis_excel(_request(PATH + "/excel", target=None, projeto="portfolio"))

    sheets = _sheets(response)
    assert list(sheets) == [
        "Resumo",
        "Desempenho de custo por projeto",
        "Valor agregado na data de corte",
        "Período a período",
    ]
    projeto = sheets["Desempenho de custo por projeto"]
    first_row = _header_row(projeto)
    assert [projeto.cell(row=first_row, column=column).value for column in range(1, 10)] == list(
        PROJECT_HEADERS
    )
    codes = {
        projeto.cell(row=row, column=1).value for row in range(first_row + 1, projeto.max_row + 1)
    }
    assert any(str(code).startswith("TN-2026-014 · ") for code in codes)


def test_o_imprimivel_dos_kpis_traz_o_mesmo_conteudo_do_excel(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.custos_kpis_printable(
        _request(PATH + "/imprimivel", target=FOLHA_TARGET, projeto=str(project.id))
    )

    assert response.status_code == 200
    body = _body(response)
    for texto in (
        "KPIs de custo",
        *KPI_LABELS,
        "CPI e SPI por mês",
        "Valor agregado acumulado",
        "Valor agregado na data de corte",
        "Projeção no término (bottom-up)",
        "Período a período",
        "R$ 46.440.129,45",
        "R$ 45.255.000,00",
    ):
        assert texto in body
    assert body.count('data-grafico="linhas-multiplas"') == 2
