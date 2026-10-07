"""Rotas dos marcos, claims e EOT da ficha do contrato (ISSUE-033, D12, D14).

Costura HTTP: o handler é chamado com a ``HttpRequest`` que o navegador montaria, como em
``test_contrato_rotas.py``. Cada aba é um fragmento pelo ``?numero=``; as operações devolvem a
ficha atualizada numa resposta multi-alvo (os painéis e o formulário vazio). A evidência do marco
é um ``multipart/form-data`` com o anexo de verdade; o claim notificado fora do prazo aparece com
o alerta; o Acordado e a EOT concedida criam a SM na Governança; cada aba tem o seu Excel e o seu
imprimível.
"""

from __future__ import annotations

from collections.abc import Mapping
from datetime import date
from io import BytesIO
from urllib.parse import unquote, urlencode

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select

from src.core import auth, calendario, file_storage
from src.core.models import Attachment
from src.modulos.financeiro import calculations, routes
from src.modulos.financeiro.models import (
    Contract,
    ContractClaim,
    ContractMeasurement,
    ContractPaymentMilestone,
    ContractTimeExtension,
)
from src.modulos.governanca.models import ChangeRequest
from tests.armazenamento_falso import ArmazenamentoEmMemoria
from tests.financeiro.conftest import (
    MANAGER_EMAIL,
    MEMBER_EMAIL,
    VIEWER_EMAIL,
    EacScenario,
)
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal, criar_empresa

HOJE = date(2026, 10, 6)
NUMERO = "CT-2026-011"
ORIGINAL_CENTS = 100_000_000
CONTRACT_END = date(2026, 10, 30)
EAC_LEAF = "1.1.1"
EVIDENCE_BYTES = b"%PDF-1.7 evidencia do marco"
PARECER = "Parecer registrado no teste das rotas."
MILESTONES_TARGET = "contrato-marcos"
CLAIMS_TARGET = "contrato-claims"
EXTENSIONS_TARGET = "contrato-extensoes"
FORM_TARGET = "contrato-formulario"
BOUNDARY = "fronteira-do-teste-dos-fluxos"


@pytest.fixture(autouse=True)
def data_de_hoje(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(calendario, "today", lambda: HOJE)


@pytest.fixture
def armazenamento(monkeypatch: pytest.MonkeyPatch) -> ArmazenamentoEmMemoria:
    """O arquivo da evidência fica na memória: o teste lê o que foi guardado."""
    storage = ArmazenamentoEmMemoria()
    monkeypatch.setattr(file_storage, "configured_storage", lambda: storage)
    return storage


def _request(
    path: str,
    *,
    email: str = MEMBER_EMAIL,
    route: Mapping[str, str] | None = None,
    payload: Mapping[str, str] | tuple[bytes, str] | None = None,
    target: str | None = "contrato-resumo",
    **params: str,
) -> func.HttpRequest:
    """A request as the browser sends it; ``target=None`` is a plain link (no Alpine).

    ``payload`` is the urlencoded form or a ``(body, content-type)`` pair, which is how the
    multipart of the evidence travels.
    """
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    body = b""
    if isinstance(payload, tuple):
        body, content_type = payload
        headers["Content-Type"] = content_type
    elif payload is not None:
        headers["Content-Type"] = "application/x-www-form-urlencoded"
        body = urlencode(payload).encode()
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="POST" if payload is not None else "GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params=dict(route or {}),
        body=body,
    )


def _multipart(fields: Mapping[str, str], files: tuple[str, bytes]) -> tuple[bytes, str]:
    """A ``multipart/form-data`` body as the browser sends it, and its Content-Type."""
    parts = [
        f'--{BOUNDARY}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{value}\r\n'.encode()
        for name, value in fields.items()
    ]
    name, content = files
    header = (
        f"--{BOUNDARY}\r\n"
        f'Content-Disposition: form-data; name="arquivo"; filename="{name}"\r\n'
        "Content-Type: application/pdf\r\n\r\n"
    )
    parts.append(header.encode() + content + b"\r\n")
    parts.append(f"--{BOUNDARY}--\r\n".encode())
    return b"".join(parts), f"multipart/form-data; boundary={BOUNDARY}"


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _toast(response: func.HttpResponse) -> str:
    return unquote(response.headers["X-TN-Toast"])


def _criar_contrato(
    cenario: EacScenario,
    *,
    number: str = NUMERO,
    original_cents: int = ORIGINAL_CENTS,
    claim_notice_days: int = 30,
) -> Contract:
    session = cenario.session
    empresa = criar_empresa(session, "Contratada de teste")
    contract = Contract(
        project_id=cenario.first.id,
        company_id=empresa.id,
        eac_item_id=cenario.items.get(f"TN-EAC-001:{EAC_LEAF}"),
        manager_id=cenario.responsible.id,
        fiscal_id=cenario.responsible.id,
        number=number,
        object="Obra de teste",
        modality="Preço unitário",
        original_value_cents=original_cents,
        start=date(2026, 1, 5),
        original_end=CONTRACT_END,
        current_end=CONTRACT_END,
        retention_pct=5,
        claim_notice_days=claim_notice_days,
        situation="Em execução",
    )
    session.add(contract)
    session.flush()
    return contract


def _marco(cenario: EacScenario, contract: Contract, **fields: object) -> ContractPaymentMilestone:
    """Um marco direto na tabela, com os padrões do cenário e o que o teste quiser mudar."""
    values: dict[str, object] = {
        "number": "M1",
        "description": "Mobilização e plano de projeto",
        "acceptance_criteria": "Plano de projeto aprovado",
        "pct": 10,
        "expected_on": date(2026, 1, 30),
        "completed_on": None,
        "situation": calculations.MILESTONE_FORECAST,
        "attachment_id": None,
    }
    values.update(fields)
    record = ContractPaymentMilestone(contract_id=contract.id, **values)
    cenario.session.add(record)
    cenario.session.flush()
    return record


def _claim(cenario: EacScenario, contract: Contract, **fields: object) -> ContractClaim:
    """Um claim direto na tabela, com os padrões do cenário e o que o teste quiser mudar."""
    values: dict[str, object] = {
        "code": "CLM-TN-2026-0001",
        "direction": "Da contratada",
        "kind": "Custo",
        "cause": "Outros",
        "description": "Back-charge do retrabalho de solda da linha.",
        "clause": "18.1",
        "event_date": date(2026, 8, 19),
        "notification_date": date(2026, 8, 26),
        "claimed_cents": 3_400_000,
        "claimed_days": 0,
        "recognized_cents": None,
        "recognized_days": None,
        "situation": calculations.CLAIM_NOTIFIED,
    }
    values.update(fields)
    record = ContractClaim(contract_id=contract.id, **values)
    cenario.session.add(record)
    cenario.session.flush()
    return record


def _eot(cenario: EacScenario, contract: Contract, **fields: object) -> ContractTimeExtension:
    """Uma EOT direta na tabela, com os padrões do cenário e o que o teste quiser mudar."""
    values: dict[str, object] = {
        "code": "EOT-TN-2026-0001",
        "claim_id": None,
        "change_id": None,
        "event": "Interferência não mapeada no subsolo da fundação F-12.",
        "requested_days": 30,
        "granted_days": None,
        "classification": None,
        "method": "Análise de impacto no tempo",
        "affected_milestone": "Término das fundações",
        "requested_on": date(2026, 4, 20),
        "decided_on": None,
        "situation": calculations.EXTENSION_REQUESTED,
    }
    values.update(fields)
    record = ContractTimeExtension(contract_id=contract.id, **values)
    cenario.session.add(record)
    cenario.session.flush()
    return record


def _sm(
    cenario: EacScenario, *, code: str = "SM-TN-2026-0001", situation: str = "Aprovada"
) -> ChangeRequest:
    sm = ChangeRequest(
        project_id=cenario.first.id,
        requester_id=cenario.responsible.id,
        code=code,
        title="Mudança de teste",
        kind="Escopo",
        origin="Interna",
        priority="Normal",
        request_date=HOJE,
        description="Mudança criada no teste das rotas dos fluxos do contrato.",
        situation=situation,
    )
    cenario.session.add(sm)
    cenario.session.flush()
    return sm


def _ficha_params(cenario: EacScenario) -> dict[str, str]:
    return {"projeto": str(cenario.first.id), "numero": NUMERO}


# ── Leitura das abas ─────────────────────────────────────────────────────


def test_os_marcos_listam_o_fluxo_o_atraso_e_a_soma(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    _marco(
        cenario,
        contract,
        number="M1",
        pct=60,
        expected_on=date(2026, 9, 1),
        situation=calculations.MILESTONE_FORECAST,
    )
    _marco(
        cenario,
        contract,
        number="M2",
        pct=40,
        expected_on=date(2026, 11, 1),
        completed_on=date(2026, 10, 1),
        situation=calculations.MILESTONE_EVIDENCE_SENT,
    )

    response = routes.contrato_marcos(
        _request(
            "/api/financeiro/contrato/marcos",
            email=MANAGER_EMAIL,
            target=MILESTONES_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), MILESTONES_TARGET) is not None
    body = _body(response)
    assert "M1 Mobilização e plano de projeto" in body
    assert "Soma dos marcos: 100%" in body
    assert "35 dias de atraso" in body
    assert "Enviar evidência" in body
    assert "Aprovar" in body
    assert "R$ 600.000,00" in body  # o valor do M1 sobre o valor atual
    assert "R$ 30.000,00" in body  # a retenção de 5% do M1


def test_os_claims_alertam_a_notificacao_fora_do_prazo(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario, claim_notice_days=30)
    _claim(
        cenario,
        contract,
        code="CLM-TN-2026-0001",
        event_date=date(2026, 8, 1),
        notification_date=date(2026, 8, 31),
    )
    _claim(
        cenario,
        contract,
        code="CLM-TN-2026-0002",
        event_date=date(2026, 8, 1),
        notification_date=date(2026, 9, 1),
    )

    response = routes.contrato_claims(
        _request(
            "/api/financeiro/contrato/claims",
            email=MANAGER_EMAIL,
            target=CLAIMS_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), CLAIMS_TARGET) is not None
    body = _body(response)
    assert "CLM-TN-2026-0001" in body
    assert "CLM-TN-2026-0002" in body
    assert "30 dias após o evento" in body  # o último dia ainda está no prazo
    assert "31 dias após o evento · fora do prazo" in body
    assert "Atualizar" in body


def test_as_extensoes_listam_a_decisao_e_a_sm(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    claim = _claim(cenario, contract, kind="Prazo", claimed_cents=0, claimed_days=30)
    sm = _sm(cenario)
    _eot(
        cenario,
        contract,
        claim_id=claim.id,
        situation=calculations.EXTENSION_PARTIALLY_GRANTED,
        granted_days=21,
        classification="Justificável e compensável",
        change_id=sm.id,
    )

    response = routes.contrato_extensoes(
        _request(
            "/api/financeiro/contrato/extensoes",
            target=EXTENSIONS_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), EXTENSIONS_TARGET) is not None
    body = _body(response)
    assert "EOT-TN-2026-0001" in body
    assert claim.code in body
    assert "Concedida parcialmente" in body
    assert "Justificável e compensável" in body
    assert sm.code in body
    assert "Decidir" not in body  # já decidida


def test_sem_o_cabecalho_do_alpine_a_aba_manda_para_o_shell(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_marcos(
        _request("/api/financeiro/contrato/marcos", target=None, numero=NUMERO)
    )

    assert response.status_code == 302


def test_o_membro_le_as_abas(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    for handler, target, path in (
        (routes.contrato_marcos, MILESTONES_TARGET, "/api/financeiro/contrato/marcos"),
        (routes.contrato_claims, CLAIMS_TARGET, "/api/financeiro/contrato/claims"),
        (routes.contrato_extensoes, EXTENSIONS_TARGET, "/api/financeiro/contrato/extensoes"),
    ):
        response = handler(
            _request(path, email=MEMBER_EMAIL, target=target, **_ficha_params(cenario))
        )
        assert response.status_code == 200


# ── Marcos: formulário, criação, evidência e fluxo ───────────────────────


def test_o_formulario_do_marco_mostra_o_disponivel(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    _marco(cenario, contract, pct=60)

    response = routes.contrato_marco_form(
        _request(
            "/api/financeiro/contrato/marcos/novo",
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), FORM_TARGET) is not None
    names = {tag.attrs.get("name") for tag in _tags(response) if tag.name in {"input", "textarea"}}
    assert {"descricao", "criterio", "pct", "prevista"} <= names
    body = _body(response)
    assert "40% disponíveis" in body
    assert f"/api/financeiro/contrato/marcos?numero={NUMERO}" in body


def test_registrar_marco_devolve_a_ficha_e_fecha_o_formulario(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_marco_create(
        _request(
            "/api/financeiro/contrato/marcos",
            email=MANAGER_EMAIL,
            payload={
                "descricao": "Mobilização e plano de projeto",
                "criterio": "Plano de projeto aprovado",
                "pct": "10",
                "prevista": "2026-01-30",
            },
            target=f"{FORM_TARGET} {MILESTONES_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, FORM_TARGET) is not None
    assert find_by_id(tags, MILESTONES_TARGET) is not None
    assert _toast(response) == "Marco M1 cadastrado."

    (record,) = cenario.session.scalars(select(ContractPaymentMilestone)).all()
    assert (record.number, record.pct, record.situation) == (
        "M1",
        10,
        calculations.MILESTONE_FORECAST,
    )


def test_marco_acima_dos_100_volta_no_formulario_com_422(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    _marco(cenario, contract, pct=60)

    response = routes.contrato_marco_create(
        _request(
            "/api/financeiro/contrato/marcos",
            email=MANAGER_EMAIL,
            payload={
                "descricao": "Segundo marco",
                "criterio": "Entrega aceita",
                "pct": "50",
                "prevista": "2026-02-28",
            },
            target=f"{FORM_TARGET} {MILESTONES_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), FORM_TARGET) is not None
    body = _body(response)
    assert "passaria de 100%" in body
    assert "restam 40%" in body
    assert len(cenario.session.scalars(select(ContractPaymentMilestone)).all()) == 1


def test_enviar_evidencia_e_aprovar_o_marco(
    cenario: EacScenario, armazenamento: ArmazenamentoEmMemoria
) -> None:
    contract = _criar_contrato(cenario)
    marco = _marco(cenario, contract)

    form_response = routes.contrato_marco_evidencia_form(
        _request(
            f"/api/financeiro/contrato/marcos/{marco.id}/evidencia",
            route={"marco_id": str(marco.id)},
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert form_response.status_code == 200
    assert find_by_id(_tags(form_response), FORM_TARGET) is not None
    assert "Enviar evidência do marco M1" in _body(form_response)

    sent = routes.contrato_marco_evidencia(
        _request(
            f"/api/financeiro/contrato/marcos/{marco.id}/evidencia",
            email=MANAGER_EMAIL,
            route={"marco_id": str(marco.id)},
            payload=_multipart(
                {"conclusao": "2026-01-28", "versao": str(marco.version)},
                ("evidencia do marco.pdf", EVIDENCE_BYTES),
            ),
            target=f"{FORM_TARGET} {MILESTONES_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert sent.status_code == 200
    assert _toast(sent) == "Evidência do marco M1 enviada para aprovação."
    cenario.session.refresh(marco)
    assert (marco.situation, marco.completed_on) == (
        calculations.MILESTONE_EVIDENCE_SENT,
        date(2026, 1, 28),
    )
    anexo = cenario.session.get(Attachment, marco.attachment_id)
    assert anexo is not None
    assert (anexo.origin_table, anexo.origin_record_id) == ("contrato_marco_pagamento", marco.id)
    assert armazenamento.arquivos == {f"{cenario.first.id}/{anexo.id}": EVIDENCE_BYTES}

    approved = routes.contrato_marco_approve(
        _request(
            f"/api/financeiro/contrato/marcos/{marco.id}/aprovar",
            email=MANAGER_EMAIL,
            route={"marco_id": str(marco.id)},
            payload={"versao": str(marco.version)},
            target=f"{MILESTONES_TARGET} {FORM_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert approved.status_code == 200
    assert _toast(approved) == "Marco M1: aprovado."
    cenario.session.refresh(marco)
    assert marco.situation == calculations.MILESTONE_APPROVED


def test_aprovar_marco_sem_evidencia_volta_no_painel_com_422(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    marco = _marco(cenario, contract, situation=calculations.MILESTONE_EVIDENCE_SENT)

    response = routes.contrato_marco_approve(
        _request(
            f"/api/financeiro/contrato/marcos/{marco.id}/aprovar",
            email=MANAGER_EMAIL,
            route={"marco_id": str(marco.id)},
            payload={"versao": str(marco.version)},
            target=MILESTONES_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), MILESTONES_TARGET) is not None
    assert "anexo de evidência" in _body(response)
    cenario.session.refresh(marco)
    assert marco.situation == calculations.MILESTONE_EVIDENCE_SENT


def test_faturar_e_pagar_o_marco_seguem_o_fluxo(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    marco = _marco(cenario, contract, situation=calculations.MILESTONE_APPROVED)

    faturado = routes.contrato_marco_invoice(
        _request(
            f"/api/financeiro/contrato/marcos/{marco.id}/faturar",
            email=MANAGER_EMAIL,
            route={"marco_id": str(marco.id)},
            payload={"versao": str(marco.version)},
            target=MILESTONES_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert faturado.status_code == 200
    assert _toast(faturado) == "Marco M1: faturado."
    cenario.session.refresh(marco)
    assert marco.situation == calculations.MILESTONE_INVOICED

    pago = routes.contrato_marco_pay(
        _request(
            f"/api/financeiro/contrato/marcos/{marco.id}/pagar",
            email=MANAGER_EMAIL,
            route={"marco_id": str(marco.id)},
            payload={"versao": str(marco.version)},
            target=MILESTONES_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert pago.status_code == 200
    assert _toast(pago) == "Marco M1: pago."
    cenario.session.refresh(marco)
    assert marco.situation == calculations.MILESTONE_PAID


def test_transicao_fora_do_fluxo_do_marco_volta_no_painel_com_422(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    marco = _marco(cenario, contract)

    response = routes.contrato_marco_pay(
        _request(
            f"/api/financeiro/contrato/marcos/{marco.id}/pagar",
            email=MANAGER_EMAIL,
            route={"marco_id": str(marco.id)},
            payload={"versao": str(marco.version)},
            target=MILESTONES_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), MILESTONES_TARGET) is not None
    assert "não pode ir para" in _body(response)


def test_visualizador_nao_registra_marco(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_marco_create(
        _request(
            "/api/financeiro/contrato/marcos",
            email=VIEWER_EMAIL,
            payload={
                "descricao": "Mobilização",
                "criterio": "Plano aprovado",
                "pct": "10",
                "prevista": "2026-01-30",
            },
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 403


# ── Claims: formulário, criação, atualização e decisão ───────────────────


def test_registrar_claim_devolve_a_ficha(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_claim_create(
        _request(
            "/api/financeiro/contrato/claims",
            email=MANAGER_EMAIL,
            payload={
                "direcao": "Da contratada",
                "tipo": "Custo",
                "causa": "Outros",
                "clausula": "18.1",
                "descricao": "Back-charge do retrabalho de solda da linha.",
                "evento": "2026-08-19",
                "notificacao": "2026-08-26",
                "valor": "34.000,00",
                "dias": "0",
            },
            target=f"{FORM_TARGET} {CLAIMS_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert _toast(response) == "Claim CLM-TN-2026-0001 registrado."
    (record,) = cenario.session.scalars(select(ContractClaim)).all()
    assert (record.code, record.situation, record.claimed_cents) == (
        "CLM-TN-2026-0001",
        calculations.CLAIM_NOTIFIED,
        3_400_000,
    )


def test_claim_com_notificacao_anterior_ao_evento_e_recusado(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_claim_create(
        _request(
            "/api/financeiro/contrato/claims",
            email=MANAGER_EMAIL,
            payload={
                "direcao": "Da contratada",
                "tipo": "Custo",
                "causa": "Outros",
                "descricao": "Back-charge do retrabalho de solda da linha.",
                "evento": "2026-08-19",
                "notificacao": "2026-08-10",
                "valor": "34.000,00",
            },
            target=f"{FORM_TARGET} {CLAIMS_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), FORM_TARGET) is not None
    assert "não pode ser anterior ao evento" in _body(response)
    assert cenario.session.scalars(select(ContractClaim)).all() == []


def test_atualizar_claim_ate_a_negociacao(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    claim = _claim(cenario, contract)

    form_response = routes.contrato_claim_update_form(
        _request(
            f"/api/financeiro/contrato/claims/{claim.id}/atualizar",
            route={"claim_id": str(claim.id)},
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert form_response.status_code == 200
    body = _body(form_response)
    assert "Atualizar claim CLM-TN-2026-0001" in body
    assert "Em análise" in body
    assert "Em negociação" not in body  # o fluxo do Notificado só admite Em análise

    response = routes.contrato_claim_update(
        _request(
            f"/api/financeiro/contrato/claims/{claim.id}/atualizar",
            email=MANAGER_EMAIL,
            route={"claim_id": str(claim.id)},
            payload={"situacao": "Em análise", "parecer": PARECER, "versao": str(claim.version)},
            target=f"{FORM_TARGET} {CLAIMS_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert _toast(response) == "Claim CLM-TN-2026-0001: em análise."
    cenario.session.refresh(claim)
    assert claim.situation == calculations.CLAIM_UNDER_REVIEW


def test_decidir_claim_acordado_cria_sm_na_governanca(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    claim = _claim(cenario, contract, situation=calculations.CLAIM_IN_NEGOTIATION)

    form_response = routes.contrato_claim_decide_form(
        _request(
            f"/api/financeiro/contrato/claims/{claim.id}/decidir",
            route={"claim_id": str(claim.id)},
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert form_response.status_code == 200
    body = _body(form_response)
    assert "Acordado" in body
    assert "Rejeitado" in body
    assert "Em disputa" not in body  # a decisão de mérito só oferece Acordado e Rejeitado

    response = routes.contrato_claim_decide(
        _request(
            f"/api/financeiro/contrato/claims/{claim.id}/decidir",
            email=MANAGER_EMAIL,
            route={"claim_id": str(claim.id)},
            payload={
                "situacao": "Acordado",
                "reconhecido": "30.000,00",
                "dias": "0",
                "parecer": PARECER,
                "versao": str(claim.version),
            },
            target=f"{FORM_TARGET} {CLAIMS_TARGET} {EXTENSIONS_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert _toast(response) == (
        "Claim CLM-TN-2026-0001: acordado. SM SM-TN-2026-0001 aberta em Governança."
    )
    cenario.session.refresh(claim)
    assert (claim.situation, claim.recognized_cents) == (calculations.CLAIM_AGREED, 3_000_000)
    assert claim.change_id is not None
    sm = cenario.session.get(ChangeRequest, claim.change_id)
    assert sm is not None
    assert sm.title == "Claim CLM-TN-2026-0001 acordado"


def test_decisao_restrita_nao_aceita_negociacao(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    claim = _claim(cenario, contract)

    response = routes.contrato_claim_decide(
        _request(
            f"/api/financeiro/contrato/claims/{claim.id}/decidir",
            email=MANAGER_EMAIL,
            route={"claim_id": str(claim.id)},
            payload={"situacao": "Em negociação", "parecer": PARECER, "versao": str(claim.version)},
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    assert "Acordado ou Rejeitado" in _body(response)


# ── Extensões de prazo: criação e decisão ────────────────────────────────


def test_registrar_e_decidir_eot_concedida_prorroga_e_cria_sm(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    claim = _claim(cenario, contract, kind="Prazo", claimed_cents=0, claimed_days=30)

    form_response = routes.contrato_extensao_form(
        _request(
            "/api/financeiro/contrato/extensoes/nova",
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert form_response.status_code == 200
    assert claim.code in _body(form_response)

    created = routes.contrato_extensao_create(
        _request(
            "/api/financeiro/contrato/extensoes",
            email=MANAGER_EMAIL,
            payload={
                "claim": str(claim.id),
                "evento": "Interferência não mapeada no subsolo da fundação F-12.",
                "dias": "30",
                "marco": "Término das fundações",
                "metodo": "Análise de impacto no tempo",
            },
            target=f"{FORM_TARGET} {EXTENSIONS_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert created.status_code == 200
    assert _toast(created) == "Extensão de prazo EOT-TN-2026-0001 registrada."
    (eot,) = cenario.session.scalars(select(ContractTimeExtension)).all()
    assert (eot.situation, eot.claim_id) == (calculations.EXTENSION_REQUESTED, claim.id)

    decision_form = routes.contrato_extensao_decisao_form(
        _request(
            f"/api/financeiro/contrato/extensoes/{eot.id}/decidir",
            route={"extensao_id": str(eot.id)},
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert decision_form.status_code == 200
    body = _body(decision_form)
    assert "Decidir extensão EOT-TN-2026-0001" in body
    assert "30 dia(s) solicitados" in body

    decided = routes.contrato_extensao_decide(
        _request(
            f"/api/financeiro/contrato/extensoes/{eot.id}/decidir",
            email=MANAGER_EMAIL,
            route={"extensao_id": str(eot.id)},
            payload={
                "situacao": "Concedida parcialmente",
                "dias": "21",
                "classificacao": "Justificável e compensável",
                "parecer": PARECER,
                "versao": str(eot.version),
            },
            target=f"{FORM_TARGET} {EXTENSIONS_TARGET} contrato-resumo",
            **_ficha_params(cenario),
        )
    )

    assert decided.status_code == 200
    assert _toast(decided) == (
        "Extensão EOT-TN-2026-0001: concedida parcialmente. Término vigente prorrogado em 21 dias."
    )
    cenario.session.refresh(eot)
    cenario.session.refresh(contract)
    assert (eot.granted_days, eot.classification) == (21, "Justificável e compensável")
    assert contract.current_end == date(2026, 11, 20)
    assert eot.change_id is not None
    sm = cenario.session.get(ChangeRequest, eot.change_id)
    assert sm is not None
    assert sm.title == f"EOT {eot.code}: Concedida parcialmente"


def test_eot_negada_nao_prorroga_nem_cria_sm(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    eot = _eot(cenario, contract)

    response = routes.contrato_extensao_decide(
        _request(
            f"/api/financeiro/contrato/extensoes/{eot.id}/decidir",
            email=MANAGER_EMAIL,
            route={"extensao_id": str(eot.id)},
            payload={
                "situacao": "Negada",
                "parecer": "Sem nexo de causalidade.",
                "versao": str(eot.version),
            },
            target=f"{EXTENSIONS_TARGET} {FORM_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert _toast(response) == "Extensão EOT-TN-2026-0001: negada."
    cenario.session.refresh(eot)
    cenario.session.refresh(contract)
    assert (eot.situation, eot.change_id, eot.granted_days) == (
        calculations.EXTENSION_DENIED,
        None,
        0,
    )
    assert contract.current_end == CONTRACT_END
    assert cenario.session.scalars(select(ChangeRequest)).all() == []


def test_concedida_integral_com_dias_diferentes_volta_no_formulario(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    eot = _eot(cenario, contract)

    response = routes.contrato_extensao_decide(
        _request(
            f"/api/financeiro/contrato/extensoes/{eot.id}/decidir",
            email=MANAGER_EMAIL,
            route={"extensao_id": str(eot.id)},
            payload={
                "situacao": "Concedida",
                "dias": "10",
                "classificacao": "Justificável e compensável",
                "parecer": PARECER,
                "versao": str(eot.version),
            },
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), FORM_TARGET) is not None
    assert "Concedida integralmente: 30 dias" in _body(response)


# ── Excel e imprimível por aba ───────────────────────────────────────────


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> tuple[int, list[object]]:
    first_row = int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    width = sheet.max_column
    return first_row, [
        sheet.cell(row=first_row, column=column).value for column in range(1, width + 1)
    ]


def test_o_excel_dos_marcos_traz_a_tabela_e_a_soma(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    _marco(cenario, contract, pct=60, expected_on=date(2026, 9, 1))
    _marco(cenario, contract, number="M2", pct=40, expected_on=date(2026, 11, 1))

    response = routes.contrato_marcos_excel(
        _request("/api/financeiro/contrato/marcos/excel", target=None, numero=NUMERO)
    )

    assert response.status_code == 200
    sheets = _sheets(response)
    assert list(sheets) == ["Marcos de pagamento"]
    sheet = sheets["Marcos de pagamento"]
    first_row, header = _header_row(sheet)
    assert header == [
        "Nº",
        "Marco",
        "Critério de aceite",
        "%",
        "Valor",
        "Retenção",
        "Prevista",
        "Atraso (dias)",
        "Conclusão",
        "Situação",
        "Evidência",
    ]
    assert sheet.cell(row=first_row + 1, column=1).value == "M1"
    assert sheet.cell(row=first_row + 1, column=4).value == 60
    assert sheet.cell(row=first_row + 1, column=8).value == 35
    assert sheet.cell(row=first_row + 3, column=2).value == "Total"
    assert sheet.cell(row=first_row + 3, column=4).value == 100


def test_o_excel_dos_claims_traz_o_alerta_e_o_excel_das_eot_a_decisao(
    cenario: EacScenario,
) -> None:
    contract = _criar_contrato(cenario)
    _claim(
        cenario,
        contract,
        event_date=date(2026, 8, 1),
        notification_date=date(2026, 9, 1),
    )
    _eot(
        cenario,
        contract,
        situation=calculations.EXTENSION_PARTIALLY_GRANTED,
        granted_days=21,
        classification="Justificável e compensável",
    )

    claims = _sheets(
        routes.contrato_claims_excel(
            _request("/api/financeiro/contrato/claims/excel", target=None, numero=NUMERO)
        )
    )
    assert list(claims) == ["Claims"]
    first_row, header = _header_row(claims["Claims"])
    assert "Fora do prazo" in header
    late_column = header.index("Fora do prazo") + 1
    assert claims["Claims"].cell(row=first_row + 1, column=late_column).value == "Sim"

    eot = _sheets(
        routes.contrato_extensoes_excel(
            _request("/api/financeiro/contrato/extensoes/excel", target=None, numero=NUMERO)
        )
    )
    assert list(eot) == ["Extensões de prazo"]
    first_row, header = _header_row(eot["Extensões de prazo"])
    assert header[:5] == ["Código", "Claim", "Evento causador", "Solicitados", "Concedidos"]
    assert eot["Extensões de prazo"].cell(row=first_row + 1, column=1).value == "EOT-TN-2026-0001"
    assert eot["Extensões de prazo"].cell(row=first_row + 1, column=5).value == 21


def test_a_versao_imprimivel_de_cada_aba_traz_o_mesmo_conteudo(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    _marco(cenario, contract)
    _claim(cenario, contract)
    _eot(cenario, contract)

    for handler, path, titulo in (
        (
            routes.contrato_marcos_printable,
            "/api/financeiro/contrato/marcos/imprimivel",
            "Marcos de pagamento",
        ),
        (routes.contrato_claims_printable, "/api/financeiro/contrato/claims/imprimivel", "Claims"),
        (
            routes.contrato_extensoes_printable,
            "/api/financeiro/contrato/extensoes/imprimivel",
            "Extensões de prazo",
        ),
    ):
        response = handler(_request(path, target="folha-impressao", numero=NUMERO))
        assert response.status_code == 200
        body = _body(response)
        assert f"Contrato {NUMERO}" in body
        assert titulo in body


@pytest.mark.usefixtures("cenario")
def test_contrato_desconhecido_nas_abas_e_recusado() -> None:
    response = routes.contrato_marcos(
        _request("/api/financeiro/contrato/marcos", numero="CT-9999-999")
    )

    assert response.status_code == 422
    assert "Contrato não encontrado." in _body(response)


def test_marco_de_outro_contrato_nao_e_encontrado(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    outro = _criar_contrato(cenario, number="CT-2026-099")
    marco = _marco(cenario, outro)

    response = routes.contrato_marco_approve(
        _request(
            f"/api/financeiro/contrato/marcos/{marco.id}/aprovar",
            email=MANAGER_EMAIL,
            route={"marco_id": str(marco.id)},
            payload={"versao": str(marco.version)},
            target=MILESTONES_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    assert "Marco de pagamento não encontrado." in _body(response)
    assert contract.id != outro.id


def test_medicao_de_marco_aparece_no_vinculo_da_ficha(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    marco = _marco(cenario, contract, situation=calculations.MILESTONE_APPROVED)
    medicao = ContractMeasurement(
        contract_id=contract.id,
        milestone_id=marco.id,
        number="BM-01",
        period=date(2026, 9, 1),
        gross_cents=1_000_000,
        situation=calculations.SITUATION_UNDER_REVIEW,
    )
    cenario.session.add(medicao)
    cenario.session.flush()

    response = routes.contrato_medicoes(
        _request(
            "/api/financeiro/contrato/medicoes",
            target="contrato-medicoes",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert medicao.milestone_id == marco.id
