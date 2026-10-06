"""Rotas das revisões, do remanejamento e da importação da EAC (ISSUE-030, D12, D14).

Costura HTTP: o handler é chamado com a ``HttpRequest`` que o navegador montaria, como em
``test_eac_rotas.py``. O remanejamento abre a SM e deixa o valor reservado; a revisão exige SM
aprovada com custo e congela a linha de base; o item novo entra por remanejamento ou por revisão;
a importação confere linha a linha e só grava na confirmação, criando a revisão vinculada à SM.
Os erros voltam no próprio lugar, com a mensagem e o 422; o 403 é de quem não pode.
"""

from __future__ import annotations

from collections.abc import Mapping
from datetime import date
from io import BytesIO
from urllib.parse import urlencode

import azure.functions as func
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from src.blueprints.importing import import_confirm as generic_confirm
from src.core import auth, calendario, importing
from src.core.rbac import GeneralProfile
from src.core.scope import Scope
from src.modulos.financeiro import importers, routes, service
from src.modulos.financeiro.models import EacItem, EacRevisionItem
from src.modulos.governanca import service as governanca
from src.modulos.governanca.models import ChangeReallocation, ChangeRequest
from tests.financeiro.conftest import (
    MANAGER_EMAIL,
    MEMBER_EMAIL,
    VIEWER_EMAIL,
    EacScenario,
    user_of,
)
from tests.financeiro.test_eac_revisoes_servico import _sm_com_custo
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal
from tests.importadores_de_teste import corpo_multipart, planilha

HOJE = date(2026, 10, 6)
ORCAMENTO = 4_460_000_000
JUSTIFICATIVA = "Remanejamento aprovado na reunião de custos do projeto."
REVISIONS_TARGET = "eac-revisoes"
REALLOCATIONS_TARGET = "eac-remanejamentos"
CONTENT_TARGET = "eac-conteudo"
IMPORT_TARGET = "eac-importacao-resultado"

IMPORT_HEADERS = (
    "Código",
    "Descrição",
    "Tipo de custo",
    "Unidade",
    "Quantidade",
    "Preço unitário (R$)",
    "CAPEX/OPEX",
    "Centro de custo",
    "Responsável",
)
IMPORT_ROW = (
    "1.1.3",
    "Isolamento térmico",
    "Serviço",
    "un",
    10,
    "85,00",
    "CAPEX",
    "CC-1",
    "Maria Membro",
)


@pytest.fixture(autouse=True)
def data_de_hoje(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(calendario, "today", lambda: HOJE)


@pytest.fixture
def orcamento(cenario: EacScenario) -> EacScenario:
    """O projeto 1 com orçamento de cadastro, para a alçada do gerente do projeto."""
    cenario.first.budget_cents = ORCAMENTO
    cenario.session.flush()
    return cenario


def _request(
    path: str,
    *,
    email: str = MEMBER_EMAIL,
    form: Mapping[str, str] | None = None,
    target: str | None = REVISIONS_TARGET,
    file: tuple[str, bytes] | None = None,
    **params: str,
) -> func.HttpRequest:
    """A request as the browser sends it; ``target=None`` is a plain link (no Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    body = b""
    if file is not None:
        body, content_type = corpo_multipart(form or {}, file)
        headers["Content-Type"] = content_type
    elif form is not None:
        headers["Content-Type"] = "application/x-www-form-urlencoded"
        body = urlencode(form).encode()
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="POST" if (form is not None or file is not None) else "GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params={},
        body=body,
    )


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _reallocation_form(**changes: str) -> dict[str, str]:
    values = {
        "origem": "1.1.1",
        "destino": "2.1.1",
        "valor": "300,00",
        "prioridade": "Normal",
        "justificativa": JUSTIFICATIVA,
    }
    values.update(changes)
    return values


def _revision_form(change_id: int, **changes: str) -> dict[str, str]:
    values = {
        "mudanca_id": str(change_id),
        "item": "1.1.1",
        "valor": "1.000,00",
        "justificativa": JUSTIFICATIVA,
    }
    values.update(changes)
    return values


def _new_item_form(**changes: str) -> dict[str, str]:
    values = {
        "recurso": "remanejamento",
        "origem": "1.1.1",
        "codigo": "1.1.3",
        "descricao": "Isolamento térmico",
        "pai": "1.1",
        "tipo_custo": "Serviço",
        "unidade": "un",
        "quantidade": "10",
        "preco": "85,00",
        "classificacao": "CAPEX",
        "centro_custo": "CC-1",
        "responsavel": "Maria Membro",
        "prioridade": "Normal",
        "justificativa": JUSTIFICATIVA,
    }
    values.update(changes)
    return values


def _import_form(change: ChangeRequest, **changes: str) -> dict[str, str]:
    values = {
        "mudanca_id": str(change.id),
        "justificativa": "Itens novos aprovados na SM de custo.",
    }
    values.update(changes)
    return values


def _spreadsheet(rows: tuple[tuple[object, ...], ...] = (IMPORT_ROW,)) -> bytes:
    return planilha(rows, cabecalho=IMPORT_HEADERS)


def _sm_registrada(cenario: EacScenario) -> ChangeRequest:
    """Uma SM só registrada, sem análise: a revisão precisa recusá-la."""
    return governanca.create_change(
        cenario.session,
        user=user_of(cenario.member, cenario.session, GeneralProfile.MEMBER),
        scope=Scope(project_id=cenario.first.id, source="url"),
        form={
            "titulo": "Inclusão de escopo sem análise",
            "tipo": "Custo",
            "origem": "Cliente",
            "prioridade": "Normal",
            "data_solicitacao": HOJE.isoformat(),
            "descricao": JUSTIFICATIVA,
        },
        reference_date=HOJE,
    )


def _pending_transfers(cenario: EacScenario) -> list[ChangeReallocation]:
    return list(
        cenario.session.scalars(
            select(ChangeReallocation).where(ChangeReallocation.applied.is_(False))
        )
    )


# ── Leitura dos painéis ──────────────────────────────────────────────────


def test_revisoes_do_projeto_voltam_na_tabela(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)
    service.create_revision(
        orcamento.session,
        user=user_of(orcamento.manager, orcamento.session, GeneralProfile.MANAGER),
        scope=Scope(project_id=orcamento.first.id, source="url"),
        request=service.RevisionRequest(change_id=change.id, justification=JUSTIFICATIVA),
        reference_date=HOJE,
    )

    response = routes.eac_revisions(
        _request("/api/financeiro/eac/revisoes", projeto=str(orcamento.first.id))
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, REVISIONS_TARGET) is not None
    assert "Rev 0" in _body(response)
    assert "Linha de base" in _body(response)
    assert "Vigente" in _body(response)


def test_remanejamentos_do_projeto_voltam_com_aplicados_e_pendentes(
    cenario: EacScenario,
) -> None:
    criada = service.reallocate(
        cenario.session,
        user=user_of(cenario.member, cenario.session, GeneralProfile.MEMBER),
        scope=Scope(project_id=cenario.first.id, source="url"),
        request=service.ReallocationRequest(
            justification=JUSTIFICATIVA,
            transfers=(
                service.TransferRequest(
                    source_code="1.1.1", target_code="2.1.1", value_cents=300_000
                ),
            ),
        ),
        reference_date=HOJE,
    )

    response = routes.eac_reallocations(
        _request(
            "/api/financeiro/eac/remanejamentos",
            target=REALLOCATIONS_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), REALLOCATIONS_TARGET) is not None
    assert criada.code in _body(response)
    assert "aplica na aprovação" in _body(response)
    assert "R$ 10.000,00" in _body(response)  # o saldo livre da origem


def test_painel_do_portfolio_traz_o_projeto_em_cada_linha(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)
    service.create_revision(
        orcamento.session,
        user=user_of(orcamento.manager, orcamento.session, GeneralProfile.MANAGER),
        scope=Scope(project_id=orcamento.first.id, source="url"),
        request=service.RevisionRequest(change_id=change.id, justification=JUSTIFICATIVA),
        reference_date=HOJE,
    )

    response = routes.eac_revisions(_request("/api/financeiro/eac/revisoes"))

    assert response.status_code == 200
    assert "Projeto" in _body(response)
    assert "TN-EAC-001 · Fábrica de teste" in _body(response)


def test_a_revisao_lista_as_sm_aprovadas_com_custo_pendentes(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)

    response = routes.eac_revisions(
        _request("/api/financeiro/eac/revisoes", projeto=str(orcamento.first.id))
    )

    assert response.status_code == 200
    assert change.code in _body(response)
    assert "R$ 1.000,00" in _body(response)


@pytest.mark.usefixtures("cenario")
def test_sem_o_cabecalho_do_alpine_a_leitura_manda_para_o_shell() -> None:
    response = routes.eac_revisions(_request("/api/financeiro/eac/revisoes", target=None))

    assert response.status_code == 302


# ── Remanejamento ────────────────────────────────────────────────────────


def test_remanejar_abre_a_sm_deixa_pendente_e_avisa(cenario: EacScenario) -> None:
    response = routes.eac_reallocate(
        _request(
            "/api/financeiro/eac/remanejamentos",
            form=_reallocation_form(),
            target=REALLOCATIONS_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    assert response.headers["X-TN-Toast-Tipo"] == "ok"
    assert "SM-" in response.headers["X-TN-Toast"].replace("%20", " ")
    assert find_by_id(_tags(response), REALLOCATIONS_TARGET) is not None
    assert "aplica na aprovação" in _body(response)

    transfer = _pending_transfers(cenario)
    assert len(transfer) == 1
    change = cenario.session.get(ChangeRequest, transfer[0].change_id)
    assert change is not None
    assert change.kind == "Remanejamento de orçamento"
    assert transfer[0].value_cents == 30_000
    assert cenario.items[f"{cenario.first.code}:1.1.1"] == transfer[0].source_item_id


def test_remanejar_acima_do_saldo_livre_e_recusado_no_painel(cenario: EacScenario) -> None:
    response = routes.eac_reallocate(
        _request(
            "/api/financeiro/eac/remanejamentos",
            form=_reallocation_form(valor="1.000.000,00"),
            target=REALLOCATIONS_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 422
    assert "saldo livre" in _body(response)
    assert response.headers["X-TN-Toast-Tipo"] == "erro"
    assert _pending_transfers(cenario) == []


def test_remanejar_sem_justificativa_e_recusado(cenario: EacScenario) -> None:
    response = routes.eac_reallocate(
        _request(
            "/api/financeiro/eac/remanejamentos",
            form=_reallocation_form(justificativa="curta"),
            target=REALLOCATIONS_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 422
    assert "justificativa" in _body(response)


def test_remanejar_como_visualizador_e_recusado_com_403(cenario: EacScenario) -> None:
    response = routes.eac_reallocate(
        _request(
            "/api/financeiro/eac/remanejamentos",
            email=VIEWER_EMAIL,
            form=_reallocation_form(),
            target=REALLOCATIONS_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 403


@pytest.mark.usefixtures("cenario")
def test_remanejar_no_portfolio_e_recusado_com_422() -> None:
    response = routes.eac_reallocate(
        _request(
            "/api/financeiro/eac/remanejamentos",
            form=_reallocation_form(),
            target=REALLOCATIONS_TARGET,
        )
    )

    assert response.status_code == 422
    assert "Escolha um projeto" in _body(response)


# ── Revisão ──────────────────────────────────────────────────────────────


def test_nova_revisao_congela_a_linha_e_incorpora_a_sm(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)

    response = routes.eac_create_revision(
        _request(
            "/api/financeiro/eac/revisoes",
            email=MANAGER_EMAIL,
            form=_revision_form(change.id),
            target=f"{REVISIONS_TARGET} {CONTENT_TARGET}",
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 200
    assert response.headers["X-TN-Toast-Tipo"] == "ok"
    assert "Rev 0" in _body(response)
    assert find_by_id(_tags(response), REVISIONS_TARGET) is not None
    assert find_by_id(_tags(response), CONTENT_TARGET) is not None

    change = orcamento.session.get(ChangeRequest, change.id)
    assert change is not None and change.eac_revision == 0
    frozen = orcamento.session.scalars(
        select(EacRevisionItem).where(
            EacRevisionItem.item_id == orcamento.items[f"{orcamento.first.code}:1.1.1"]
        )
    ).all()
    assert len(frozen) == 1
    assert frozen[0].base_cents == 1_100_000


def test_nova_revisao_sem_sm_aprovada_e_recusada(orcamento: EacScenario) -> None:
    criada = _sm_registrada(orcamento)

    response = routes.eac_create_revision(
        _request(
            "/api/financeiro/eac/revisoes",
            email=MANAGER_EMAIL,
            form=_revision_form(criada.id),
            target=REVISIONS_TARGET,
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 422
    assert "SM aprovada com custo" in _body(response)


def test_nova_revisao_como_membro_e_recusada_com_403(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)

    response = routes.eac_create_revision(
        _request(
            "/api/financeiro/eac/revisoes",
            form=_revision_form(change.id),
            target=REVISIONS_TARGET,
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 403


# ── Item novo ────────────────────────────────────────────────────────────


def test_item_novo_por_remanejamento_nasce_com_base_zero(cenario: EacScenario) -> None:
    response = routes.eac_new_item(
        _request(
            "/api/financeiro/eac/itens",
            form=_new_item_form(),
            target=f"{REALLOCATIONS_TARGET} {CONTENT_TARGET}",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), CONTENT_TARGET) is not None
    assert "1.1.3" in _body(response)

    item = cenario.session.scalars(
        select(EacItem).where(EacItem.project_id == cenario.first.id, EacItem.code == "1.1.3")
    ).one()
    assert item.level == 3
    assert item.quantity == 10
    transfer = _pending_transfers(cenario)
    assert len(transfer) == 1
    assert transfer[0].target_item_id == item.id
    assert transfer[0].new_item is True
    assert transfer[0].value_cents == 85_000


def test_item_novo_por_revisao_entra_na_revisao(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)

    response = routes.eac_new_item(
        _request(
            "/api/financeiro/eac/itens",
            email=MANAGER_EMAIL,
            form=_new_item_form(recurso="revisao", mudanca_id=str(change.id)),
            target=f"{REVISIONS_TARGET} {CONTENT_TARGET}",
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 200
    assert "Rev 0" in _body(response)
    item = orcamento.session.scalars(
        select(EacItem).where(EacItem.project_id == orcamento.first.id, EacItem.code == "1.1.3")
    ).one()
    frozen = orcamento.session.scalars(
        select(EacRevisionItem).where(EacRevisionItem.item_id == item.id)
    ).one()
    assert frozen.base_cents == 85_000


def test_item_novo_com_codigo_repetido_e_recusado(cenario: EacScenario) -> None:
    response = routes.eac_new_item(
        _request(
            "/api/financeiro/eac/itens",
            form=_new_item_form(codigo="1.1.2"),
            target=REALLOCATIONS_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 422
    assert "Já existe item com o código 1.1.2" in _body(response)


# ── Importação ───────────────────────────────────────────────────────────


def test_o_formulario_da_importacao_vem_do_servidor_e_respeita_o_projeto(
    orcamento: EacScenario,
) -> None:
    response = routes.eac_import_form(
        _request(
            "/api/financeiro/eac/revisoes/importar/formulario",
            target="eac-importar-modal",
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 200
    body = _body(response)
    assert find_by_id(_tags(response), "eac-importar-modal") is not None
    assert "Baixar o modelo" in body
    assert 'id="eac-importacao-resultado"' in body
    assert "Escolha a SM que aprovou o acréscimo" in body


@pytest.mark.usefixtures("orcamento")
def test_o_formulario_da_importacao_no_portfolio_e_recusado() -> None:
    response = routes.eac_import_form(
        _request("/api/financeiro/eac/revisoes/importar/formulario", target="eac-importar-modal")
    )

    assert response.status_code == 422
    assert "Escolha um projeto" in _body(response)


def test_a_conferencia_responde_ao_resultado_e_aos_dois_paineis(orcamento: EacScenario) -> None:
    content = _spreadsheet()

    response = routes.eac_import_check(
        _request(
            "/api/financeiro/eac/revisoes/importar/conferir",
            form={},
            file=("itens.xlsx", content),
            target=f"{IMPORT_TARGET} {REVISIONS_TARGET} {CONTENT_TARGET}",
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, IMPORT_TARGET) is not None
    assert find_by_id(tags, REVISIONS_TARGET) is not None
    assert find_by_id(tags, CONTENT_TARGET) is not None
    assert "Vão ser gravadas" in _body(response)


def test_a_conferencia_mostra_as_linhas_e_nao_grava(orcamento: EacScenario) -> None:
    content = _spreadsheet()

    response = routes.eac_import_check(
        _request(
            "/api/financeiro/eac/revisoes/importar/conferir",
            form={},
            file=("itens.xlsx", content),
            target=IMPORT_TARGET,
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 200
    body = _body(response)
    assert "Isolamento térmico" in body
    assert "Vão ser gravadas" in body
    assert orcamento.session.scalars(select(EacItem).where(EacItem.code == "1.1.3")).all() == []


def test_a_conferencia_recusa_linha_com_subpacote_inexistente(orcamento: EacScenario) -> None:
    content = _spreadsheet((("9.9.9", *IMPORT_ROW[1:]),))

    response = routes.eac_import_check(
        _request(
            "/api/financeiro/eac/revisoes/importar/conferir",
            form={},
            file=("itens.xlsx", content),
            target=IMPORT_TARGET,
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 200
    assert "subpacote 9.9 não existe" in _body(response)


def test_a_confirmacao_cria_a_revisao_com_os_itens(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)
    content = _spreadsheet()
    digest = importing.digest_of(importing.Upload(name="itens.xlsx", content=content))

    response = routes.eac_import_confirm(
        _request(
            "/api/financeiro/eac/revisoes/importar/confirmar",
            email=MANAGER_EMAIL,
            form=_import_form(change, conferido=digest),
            file=("itens.xlsx", content),
            target=f"{IMPORT_TARGET} {REVISIONS_TARGET} {CONTENT_TARGET}",
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 200
    assert "1 item novo" in response.headers["X-TN-Toast"].replace("%20", " ")
    assert "Rev 0" in _body(response)
    item = orcamento.session.scalars(
        select(EacItem).where(EacItem.project_id == orcamento.first.id, EacItem.code == "1.1.3")
    ).one()
    assert item.quantity == 10
    frozen = orcamento.session.scalars(
        select(EacRevisionItem).where(EacRevisionItem.item_id == item.id)
    ).one()
    assert frozen.base_cents == 85_000
    change = orcamento.session.get(ChangeRequest, change.id)
    assert change is not None and change.eac_revision == 0


def test_a_confirmacao_sem_o_conferido_e_recusada(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)
    content = _spreadsheet()

    response = routes.eac_import_confirm(
        _request(
            "/api/financeiro/eac/revisoes/importar/confirmar",
            email=MANAGER_EMAIL,
            form=_import_form(change),
            file=("itens.xlsx", content),
            target=IMPORT_TARGET,
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 422
    assert "Confira a planilha antes de importar" in _body(response)
    assert orcamento.session.scalars(select(EacItem).where(EacItem.code == "1.1.3")).all() == []


def test_a_confirmacao_com_erro_sem_o_ciente_e_recusada(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)
    row = ("9.9.9", *IMPORT_ROW[1:])
    content = _spreadsheet((IMPORT_ROW, row))
    digest = importing.digest_of(importing.Upload(name="itens.xlsx", content=content))

    response = routes.eac_import_confirm(
        _request(
            "/api/financeiro/eac/revisoes/importar/confirmar",
            email=MANAGER_EMAIL,
            form=_import_form(change, conferido=digest),
            file=("itens.xlsx", content),
            target=IMPORT_TARGET,
            projeto=str(orcamento.first.id),
        )
    )

    assert response.status_code == 422
    assert "linhas com erro não serão importadas" in _body(response)


def test_a_confirmacao_do_fluxo_generico_recusa_a_eac(cenario: EacScenario) -> None:
    """A confirmação genérica não grava item sem revisão: o importador recusa a linha."""
    content = _spreadsheet()
    digest = importing.digest_of(importing.Upload(name="itens.xlsx", content=content))
    body, content_type = corpo_multipart(
        {"conferido": digest, "ciente": "sim"}, ("itens.xlsx", content)
    )
    request = func.HttpRequest(
        method="POST",
        url="/api/importacao/eac-itens/confirmar",
        headers={
            auth.PRINCIPAL_HEADER: cabecalho_do_principal(MEMBER_EMAIL),
            "Content-Type": content_type,
            "X-Alpine-Request": "true",
            "X-Alpine-Target": IMPORT_TARGET,
        },
        params={"projeto": str(cenario.first.id)},
        route_params={"chave": importers.EAC_ITEMS_KEY},
        body=body,
    )

    response = generic_confirm(request)

    assert response.status_code == 422
    assert "confirmada pela revisão" in _body(response)
    assert cenario.session.scalars(select(EacItem).where(EacItem.code == "1.1.3")).all() == []


# ── Exportação ───────────────────────────────────────────────────────────


def test_o_excel_traz_as_revisoes_e_os_remanejamentos(orcamento: EacScenario) -> None:
    change = _sm_com_custo(orcamento, valor_cents=100_000)
    routes.eac_create_revision(
        _request(
            "/api/financeiro/eac/revisoes",
            email=MANAGER_EMAIL,
            form=_revision_form(change.id),
            target=REVISIONS_TARGET,
            projeto=str(orcamento.first.id),
        )
    )
    routes.eac_reallocate(
        _request(
            "/api/financeiro/eac/remanejamentos",
            form=_reallocation_form(),
            target=REALLOCATIONS_TARGET,
            projeto=str(orcamento.first.id),
        )
    )

    response = routes.eac_excel(
        _request("/api/financeiro/eac/excel", target=None, projeto=str(orcamento.first.id))
    )

    assert response.status_code == 200
    book = load_workbook(BytesIO(response.get_body()))
    sheets = {sheet.title: sheet for sheet in book.worksheets}
    assert "Revisões do orçamento" in sheets
    assert "Remanejamentos" in sheets
    revisions = sheets["Revisões do orçamento"]
    header = int(revisions.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    assert revisions.cell(row=header + 1, column=1).value == "Rev 0"
    assert revisions.cell(row=header + 1, column=3).value == 71_010.0
    reallocations = sheets["Remanejamentos"]
    header = int(reallocations.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    assert reallocations.cell(row=header + 1, column=4).value == 300.0
    assert reallocations.cell(row=header + 1, column=6).value == "Registrada"


def test_quem_nao_esta_no_cadastro_nao_le_o_painel(cenario: EacScenario) -> None:
    response = routes.eac_reallocations(
        _request(
            "/api/financeiro/eac/remanejamentos",
            email="intruso@example.invalid",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 403
