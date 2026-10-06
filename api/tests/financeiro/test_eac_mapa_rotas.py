"""Rotas do mapa de controle, da projeção e da importação dos custos do ERP (ISSUE-031).

Costura HTTP: o handler é chamado com a ``HttpRequest`` que o navegador montaria, como em
``test_eac_rotas.py``. O mapa serve o projeto e a carteira; a projeção exige justificativa, guarda
o fato e devolve o formulário e o mapa no 422 e no sucesso. A importação do ERP confere a planilha
sem gravar nada e só a confirmação escreve o fato por item e mês, substituindo o mês reimportado.
O Excel e o imprimível levam as colunas do mapa em R$ integral, com a coluna Projeto na carteira.
"""

from __future__ import annotations

from collections.abc import Mapping
from datetime import date
from io import BytesIO
from urllib.parse import urlencode

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select

from src.blueprints.importing import import_confirm as generic_confirm
from src.core import auth, calendario, importing
from src.modulos.financeiro import importers, routes, service
from src.modulos.financeiro.models import EacProjection, ErpCost
from tests.financeiro.conftest import MANAGER_EMAIL, MEMBER_EMAIL, VIEWER_EMAIL, EacScenario
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal
from tests.importadores_de_teste import corpo_multipart, planilha

HOJE = date(2026, 10, 6)
MAP_TARGET = "mapa-conteudo"
PROJECTION_TARGET = "mapa-projecao-formulario"
ERP_RESULT_TARGET = "custos-importacao-resultado"
JUSTIFICATIVA = "Premissas revisadas com o fornecedor."

ERP_HEADERS = (
    "Código do item",
    "Mês (1º dia)",
    "Comprometido (R$)",
    "Realizado (R$)",
    "Projeção no término (R$)",
)
ERP_ROW = ("1.1.1", "01/09/2026", "4.000,00", "3.000,00", "12.000,00")


@pytest.fixture(autouse=True)
def data_de_hoje(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(calendario, "today", lambda: HOJE)


def _request(
    path: str,
    *,
    email: str = MEMBER_EMAIL,
    route: Mapping[str, str] | None = None,
    form: Mapping[str, str] | None = None,
    target: str | None = MAP_TARGET,
    **params: str,
) -> func.HttpRequest:
    """A request as the browser sends it; ``target=None`` is a plain link (no Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    body = b""
    if form is not None:
        headers["Content-Type"] = "application/x-www-form-urlencoded"
        body = urlencode(form).encode()
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="POST" if form is not None else "GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params=dict(route or {}),
        body=body,
    )


def _upload_request(
    path: str,
    *,
    file: tuple[str, bytes],
    email: str = MEMBER_EMAIL,
    form: Mapping[str, str] | None = None,
    target: str | None = MAP_TARGET,
    **params: str,
) -> func.HttpRequest:
    """A ``multipart/form-data`` request with the spreadsheet of the import."""
    body, content_type = corpo_multipart(form or {}, file)
    headers = {
        auth.PRINCIPAL_HEADER: cabecalho_do_principal(email),
        "Content-Type": content_type,
    }
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="POST",
        url=path,
        headers=headers,
        params=dict(params),
        route_params={},
        body=body,
    )


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _codes(response: func.HttpResponse) -> list[str]:
    return [tag.attrs["data-codigo"] for tag in _tags(response) if "data-codigo" in tag.attrs]


def _item_id(scenario: EacScenario, code: str = "1.1.1") -> int:
    return scenario.items[f"{scenario.first.code}:{code}"]


def _spreadsheet(rows: tuple[tuple[object, ...], ...] = (ERP_ROW,)) -> bytes:
    return planilha(rows, cabecalho=ERP_HEADERS)


def _import_actual(scenario: EacScenario, *, actual_cents: int) -> None:
    service.import_erp_costs(
        scenario.session,
        projeto_id=scenario.first.id,
        linhas=(
            service.ErpCostLine(
                item_id=_item_id(scenario), month=date(2026, 9, 1), actual_cents=actual_cents
            ),
        ),
        importado_por_id=scenario.member.person_id,
        hoje=HOJE,
    )


# ── Leitura do mapa ──────────────────────────────────────────────────────


def test_o_mapa_do_projeto_traz_as_colunas_e_a_linha_do_projeto(cenario: EacScenario) -> None:
    response = routes.mapa_content(
        _request("/api/financeiro/mapa-de-controle", projeto=str(cenario.first.id))
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), MAP_TARGET) is not None
    assert _codes(response) == ["1", "1.1", "1.1.1", "1.1.2", "2", "2.1", "2.1.1"]
    body = _body(response)
    assert 'data-mapa="projeto"' in body
    for coluna in ("Orçado", "Remanejamentos", "Orçado atual", "Saldo a comprometer", "Desvio %"):
        assert coluna in body
    assert "R$ 70.010,00" in body  # a linha do projeto


@pytest.mark.usefixtures("cenario")
def test_o_mapa_da_carteira_tem_um_registro_por_projeto() -> None:
    response = routes.mapa_content(
        _request("/api/financeiro/mapa-de-controle", projeto="portfolio")
    )

    assert response.status_code == 200
    body = _body(response)
    assert 'data-mapa="carteira"' in body
    assert "TN-EAC-001 · Fábrica de teste" in body
    assert "TN-EAC-002 · Caldeira de teste" in body
    assert "Projeto" in body


@pytest.mark.usefixtures("cenario")
def test_sem_o_cabecalho_do_alpine_o_mapa_manda_para_o_shell() -> None:
    response = routes.mapa_content(
        _request("/api/financeiro/mapa-de-controle", target=None, projeto="1")
    )

    assert response.status_code == 302


# ── Projeção no término ──────────────────────────────────────────────────


def test_o_formulario_da_projecao_traz_os_valores_vigentes(cenario: EacScenario) -> None:
    response = routes.mapa_projection_form(
        _request(
            "/api/financeiro/mapa-de-controle/projecao/x/formulario",
            route={"item_id": str(_item_id(cenario))},
            target=PROJECTION_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, PROJECTION_TARGET) is not None
    names = {tag.attrs.get("name") for tag in tags if tag.name in {"input", "textarea"}}
    assert {"projecao", "justificativa"} <= names
    body = _body(response)
    assert "1.1.1" in body
    assert "R$ 10.000,00" in body  # o orçado atual do item
    assert f"/api/financeiro/mapa-de-controle/projecao/{_item_id(cenario)}" in body


def test_atualizar_projecao_grava_o_fato_e_devolve_o_mapa(cenario: EacScenario) -> None:
    response = routes.mapa_projection_save(
        _request(
            "/api/financeiro/mapa-de-controle/projecao/x",
            route={"item_id": str(_item_id(cenario))},
            form={"projecao": "10.500,00", "justificativa": JUSTIFICATIVA},
            target=f"{PROJECTION_TARGET} {MAP_TARGET}",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, PROJECTION_TARGET) is not None
    assert find_by_id(tags, MAP_TARGET) is not None
    assert response.headers["X-TN-Toast-Tipo"] == "ok"
    assert "1.1.1" in response.headers["X-TN-Toast"].replace("%20", " ")
    assert "R$ 10.500,00" in _body(response)
    (record,) = cenario.session.scalars(select(EacProjection)).all()
    assert record.item_id == _item_id(cenario)
    assert record.value_cents == 1_050_000
    assert record.justification == JUSTIFICATIVA


def test_projecao_sem_justificativa_volta_no_formulario_com_422(cenario: EacScenario) -> None:
    response = routes.mapa_projection_save(
        _request(
            "/api/financeiro/mapa-de-controle/projecao/x",
            route={"item_id": str(_item_id(cenario))},
            form={"projecao": "10.500,00", "justificativa": ""},
            target=f"{PROJECTION_TARGET} {MAP_TARGET}",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 422
    body = _body(response)
    assert "Informe a justificativa da projeção" in body
    assert "10.500,00" in body  # o que a pessoa digitou
    assert cenario.session.scalars(select(EacProjection)).all() == []


def test_projecao_abaixo_do_realizado_do_erp_e_recusada(cenario: EacScenario) -> None:
    _import_actual(cenario, actual_cents=900_000)

    response = routes.mapa_projection_save(
        _request(
            "/api/financeiro/mapa-de-controle/projecao/x",
            route={"item_id": str(_item_id(cenario))},
            form={"projecao": "8.000,00", "justificativa": JUSTIFICATIVA},
            target=PROJECTION_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 422
    assert "menor que o realizado" in _body(response)
    assert cenario.session.scalars(select(EacProjection)).all() == []


def test_projecao_no_portfolio_e_recusada_com_403(cenario: EacScenario) -> None:
    response = routes.mapa_projection_form(
        _request(
            "/api/financeiro/mapa-de-controle/projecao/x/formulario",
            route={"item_id": str(_item_id(cenario))},
            target=PROJECTION_TARGET,
            projeto="portfolio",
        )
    )

    assert response.status_code == 403


def test_projecao_de_item_de_outro_projeto_e_recusada(cenario: EacScenario) -> None:
    other = cenario.items[f"{cenario.second.code}:1.1.1"]

    response = routes.mapa_projection_form(
        _request(
            "/api/financeiro/mapa-de-controle/projecao/x/formulario",
            route={"item_id": str(other)},
            target=PROJECTION_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 422


def test_visualizador_nao_atualiza_projecao(cenario: EacScenario) -> None:
    response = routes.mapa_projection_save(
        _request(
            "/api/financeiro/mapa-de-controle/projecao/x",
            email=VIEWER_EMAIL,
            route={"item_id": str(_item_id(cenario))},
            form={"projecao": "10.500,00", "justificativa": JUSTIFICATIVA},
            target=PROJECTION_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 403


# ── Importação dos custos do ERP ─────────────────────────────────────────


def test_o_formulario_da_importacao_vem_do_servidor_e_traz_o_modelo(cenario: EacScenario) -> None:
    response = routes.erp_import_form(
        _request(
            "/api/financeiro/custos/importar/formulario",
            target="custos-importar-modal",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    body = _body(response)
    assert "/api/importacao/custos-erp/modelo" in body
    assert 'id="custos-importacao-resultado"' in body
    assert "Nada é gravado antes da sua confirmação" in body


@pytest.mark.usefixtures("cenario")
def test_o_formulario_da_importacao_no_portfolio_e_recusado() -> None:
    response = routes.erp_import_form(
        _request("/api/financeiro/custos/importar/formulario", target="custos-importar-modal")
    )

    assert response.status_code == 422
    assert "Escolha um projeto" in _body(response)


def test_a_conferencia_mostra_as_linhas_e_nao_grava(cenario: EacScenario) -> None:
    content = _spreadsheet()

    response = routes.erp_import_check(
        _upload_request(
            "/api/financeiro/custos/importar/conferir",
            form={},
            file=("custos.xlsx", content),
            target=f"{ERP_RESULT_TARGET} {MAP_TARGET}",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, ERP_RESULT_TARGET) is not None
    assert find_by_id(tags, MAP_TARGET) is not None
    body = _body(response)
    assert "Vão ser gravadas" in body
    assert "1.1.1" in body
    assert cenario.session.scalars(select(ErpCost)).all() == []


def test_a_conferencia_recusa_item_que_nao_existe(cenario: EacScenario) -> None:
    content = _spreadsheet((("9.9.9", *ERP_ROW[1:]),))

    response = routes.erp_import_check(
        _upload_request(
            "/api/financeiro/custos/importar/conferir",
            form={},
            file=("custos.xlsx", content),
            target=ERP_RESULT_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    assert "não existe na EAC" in _body(response)


def test_a_confirmacao_sem_o_conferido_e_recusada_e_nao_grava(cenario: EacScenario) -> None:
    content = _spreadsheet()

    response = routes.erp_import_confirm(
        _upload_request(
            "/api/financeiro/custos/importar/confirmar",
            form={},
            file=("custos.xlsx", content),
            target=ERP_RESULT_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 422
    assert "Confira a planilha antes de importar" in _body(response)
    assert cenario.session.scalars(select(ErpCost)).all() == []


def test_a_confirmacao_grava_o_fechamento_do_mes(cenario: EacScenario) -> None:
    content = _spreadsheet()
    digest = importing.digest_of(importing.Upload(name="custos.xlsx", content=content))

    response = routes.erp_import_confirm(
        _upload_request(
            "/api/financeiro/custos/importar/confirmar",
            form={"conferido": digest},
            file=("custos.xlsx", content),
            target=f"{ERP_RESULT_TARGET} {MAP_TARGET}",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), MAP_TARGET) is not None
    assert "1 item atualizado" in response.headers["X-TN-Toast"].replace("%20", " ")
    assert "R$ 12.000,00" in _body(response)
    (record,) = cenario.session.scalars(select(ErpCost)).all()
    assert record.item_id == _item_id(cenario)
    assert record.month == date(2026, 9, 1)
    assert (record.committed_cents, record.actual_cents, record.projection_cents) == (
        400_000,
        300_000,
        1_200_000,
    )


def test_reimportar_o_mesmo_mes_substitui_o_fato(cenario: EacScenario) -> None:
    first = _spreadsheet()
    digest = importing.digest_of(importing.Upload(name="custos.xlsx", content=first))
    routes.erp_import_confirm(
        _request(
            "/api/financeiro/custos/importar/confirmar",
            form={"conferido": digest},
            file=("custos.xlsx", first),
            target=ERP_RESULT_TARGET,
            projeto=str(cenario.first.id),
        )
    )
    second = _spreadsheet((("1.1.1", "01/09/2026", "4.500,00", "3.200,00", "12.500,00"),))
    digest = importing.digest_of(importing.Upload(name="custos.xlsx", content=second))

    response = routes.erp_import_confirm(
        _upload_request(
            "/api/financeiro/custos/importar/confirmar",
            form={"conferido": digest},
            file=("custos.xlsx", second),
            target=ERP_RESULT_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    (record,) = cenario.session.scalars(select(ErpCost)).all()
    assert (record.committed_cents, record.actual_cents, record.projection_cents) == (
        450_000,
        320_000,
        1_250_000,
    )


@pytest.mark.usefixtures("cenario")
def test_o_fluxo_generico_da_importacao_grava_com_o_importador_registrado(
    cenario: EacScenario,
) -> None:
    assert importing.find(importers.ERP_COSTS_KEY) is not None
    content = _spreadsheet()
    digest = importing.digest_of(importing.Upload(name="custos.xlsx", content=content))
    body, content_type = corpo_multipart(
        {"conferido": digest, "ciente": "sim"}, ("custos.xlsx", content)
    )
    request = func.HttpRequest(
        method="POST",
        url="/api/importacao/custos-erp/confirmar",
        headers={
            auth.PRINCIPAL_HEADER: cabecalho_do_principal(MEMBER_EMAIL),
            "Content-Type": content_type,
            "X-Alpine-Request": "true",
            "X-Alpine-Target": ERP_RESULT_TARGET,
        },
        params={"projeto": str(cenario.first.id)},
        route_params={"chave": importers.ERP_COSTS_KEY},
        body=body,
    )

    response = generic_confirm(request)

    assert response.status_code == 200
    (record,) = cenario.session.scalars(select(ErpCost)).all()
    assert record.projection_cents == 1_200_000


@pytest.mark.usefixtures("cenario")
def test_a_importacao_no_portfolio_e_recusada() -> None:
    response = routes.erp_import_check(
        _upload_request(
            "/api/financeiro/custos/importar/conferir",
            form={},
            file=("custos.xlsx", _spreadsheet()),
            target=ERP_RESULT_TARGET,
        )
    )

    assert response.status_code == 422
    assert "Escolha um projeto" in _body(response)


def test_visualizador_nao_importa_custos(cenario: EacScenario) -> None:
    response = routes.erp_import_check(
        _upload_request(
            "/api/financeiro/custos/importar/conferir",
            email=VIEWER_EMAIL,
            form={},
            file=("custos.xlsx", _spreadsheet()),
            target=ERP_RESULT_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 403


# ── Excel e PDF ──────────────────────────────────────────────────────────


def _excel(**params: str) -> func.HttpResponse:
    return routes.mapa_excel(
        _request("/api/financeiro/mapa-de-controle/excel", target=None, **params)
    )


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> tuple[int, list[object]]:
    first_row = int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    width = sheet.max_column
    return first_row, [
        sheet.cell(row=first_row, column=column).value for column in range(1, width + 1)
    ]


MAP_HEADERS = [
    "Código",
    "Descrição",
    "Orçado",
    "Remanejamentos",
    "Orçado atual",
    "Comprometido",
    "Realizado",
    "Saldo a comprometer",
    "Projeção no término",
    "Desvio",
    "Desvio %",
]


def test_excel_do_projeto_tem_as_colunas_do_mapa_em_reais(cenario: EacScenario) -> None:
    service.import_erp_costs(
        cenario.session,
        projeto_id=cenario.first.id,
        linhas=(
            service.ErpCostLine(
                item_id=_item_id(cenario),
                month=date(2026, 9, 1),
                committed_cents=400_000,
                actual_cents=300_000,
                projection_cents=1_200_000,
            ),
        ),
        importado_por_id=cenario.member.person_id,
        hoje=HOJE,
    )

    response = _excel(projeto=str(cenario.first.id))

    assert response.status_code == 200
    sheets = _sheets(response)
    assert list(sheets) == ["Resumo", "Mapa de controle"]
    sheet = sheets["Mapa de controle"]
    first_row, header = _header_row(sheet)
    assert header == MAP_HEADERS
    codes = [sheet.cell(row=first_row + offset, column=1).value for offset in range(1, 8)]
    assert codes == ["1", "1.1", "1.1.1", "1.1.2", "2", "2.1", "2.1.1"]
    row = first_row + 3  # o item 1.1.1
    assert sheet.cell(row=row, column=4).value == 0  # remanejamentos
    assert sheet.cell(row=row, column=5).value == 10_000.0
    assert sheet.cell(row=row, column=6).value == 4_000.0
    assert sheet.cell(row=row, column=7).value == 3_000.0
    assert sheet.cell(row=row, column=8).value == 6_000.0
    assert sheet.cell(row=row, column=9).value == 12_000.0
    assert sheet.cell(row=row, column=10).value == 2_000.0
    assert sheet.cell(row=row, column=11).value == pytest.approx(0.2)
    totals = first_row + 8
    assert sheet.cell(row=totals, column=5).value == 70_010.0


@pytest.mark.usefixtures("cenario")
def test_excel_da_carteira_abre_com_a_coluna_projeto() -> None:
    response = _excel(projeto="portfolio")

    sheet = _sheets(response)["Mapa de controle"]
    first_row, header = _header_row(sheet)
    assert header == ["Projeto", *MAP_HEADERS]
    projects = {sheet.cell(row=first_row + offset, column=1).value for offset in range(1, 3)}
    assert projects == {"TN-EAC-001 · Fábrica de teste", "TN-EAC-002 · Caldeira de teste"}


def test_versao_imprimivel_traz_o_mesmo_conteudo_do_excel(cenario: EacScenario) -> None:
    response = routes.mapa_printable(
        _request(
            "/api/financeiro/mapa-de-controle/imprimivel",
            target="folha-impressao",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    body = _body(response)
    for text in ("Mapa de controle da EAC", "R$ 70.010,00", "Desvio %"):
        assert text in body


def test_gestor_tambem_atualiza_projecao(cenario: EacScenario) -> None:
    response = routes.mapa_projection_save(
        _request(
            "/api/financeiro/mapa-de-controle/projecao/x",
            email=MANAGER_EMAIL,
            route={"item_id": str(_item_id(cenario))},
            form={"projecao": "10.500,00", "justificativa": JUSTIFICATIVA},
            target=PROJECTION_TARGET,
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
