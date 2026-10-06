"""Rotas da EAC (ISSUE-029, D8, D12, D14): conteúdo, cadastro do item e as duas exportações.

Costura HTTP: o handler é chamado com a ``HttpRequest`` montada no teste, como o Azure a
entregaria. A tela é um fragmento (gate do Alpine); o Excel é um download (sem o gate). O
formulário volta preenchido no 422 e no 409; no Portfólio nada se edita e as exportações trazem
a coluna Projeto.
"""

from __future__ import annotations

from collections.abc import Mapping
from io import BytesIO
from urllib.parse import urlencode

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.core import auth
from src.modulos.configuracoes.models import Person
from src.modulos.financeiro import routes
from src.modulos.financeiro.models import EacItem
from tests.financeiro.conftest import MANAGER_EMAIL, MEMBER_EMAIL, VIEWER_EMAIL, EacScenario
from tests.html_tags import Tag, find_all, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal

TARGET = "eac-conteudo"
JUSTIFICATION = "Reclassificação pedida pelo controller."


def _request(
    path: str,
    *,
    email: str = MEMBER_EMAIL,
    route: Mapping[str, str] | None = None,
    form: Mapping[str, str] | None = None,
    target: str | None = TARGET,
    **query: str,
) -> func.HttpRequest:
    """A request as the browser or the shell sends it; ``target=None`` is a plain link (no Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    if form is not None:
        headers["Content-Type"] = "application/x-www-form-urlencoded"
    return func.HttpRequest(
        method="POST" if form is not None else "GET",
        url=path,
        headers=headers,
        params=dict(query),
        route_params=dict(route or {}),
        body=urlencode(form).encode() if form is not None else b"",
    )


def _content(*, email: str = MEMBER_EMAIL, **params: str) -> func.HttpResponse:
    return routes.eac_content(_request("/api/financeiro/eac", email=email, **params))


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(response.get_body().decode())


def _codes(response: func.HttpResponse) -> list[str]:
    return [tag.attrs["data-codigo"] for tag in _tags(response) if "data-codigo" in tag.attrs]


def _item_url_params(scenario: EacScenario, code: str = "1.1.1") -> dict[str, str]:
    return {"item_id": str(scenario.items[f"{scenario.first.code}:{code}"])}


def _form_values(scenario: EacScenario, **changes: str) -> dict[str, str]:
    values = {
        "descricao": "Engenharia civil e estrutural",
        "tipo_custo": "Serviço",
        "classificacao": "CAPEX",
        "centro_custo": "CC-1",
        "responsavel": str(scenario.responsible.id),
        "justificativa": JUSTIFICATION,
        "versao": "1",
    }
    values.update(changes)
    return values


def _save(
    scenario: EacScenario, form: Mapping[str, str], *, email: str = MEMBER_EMAIL
) -> func.HttpResponse:
    return routes.eac_item_save(
        _request(
            "/api/financeiro/eac/itens/x",
            email=email,
            route=_item_url_params(scenario),
            form=form,
            target="eac-formulario eac-conteudo",
            projeto=str(scenario.first.id),
        )
    )


# ── O conteúdo da tela ───────────────────────────────────────────────────


def test_conteudo_do_projeto_traz_a_arvore_com_a_linha_zero_e_os_totais(
    cenario: EacScenario,
) -> None:
    response = _content(projeto=str(cenario.first.id))

    assert response.status_code == 200
    assert _codes(response) == ["0", "1", "1.1", "1.1.1", "1.1.2", "2", "2.1", "2.1.1"]
    tags = _tags(response)
    assert find_by_id(tags, TARGET) is not None
    marker = next(tag for tag in tags if "data-resultado" in tag.attrs)
    assert marker.attrs["data-resultado"] == "conteudo"
    assert marker.attrs["data-carteira"] == "nao"
    body = response.get_body().decode()
    assert "R$ 70.010,00" in body  # a linha 0, o total do projeto
    assert "R$ 10.010,00" in body  # o pacote 1
    assert "Engenharia civil" in body


def test_conteudo_de_membro_tem_o_botao_de_editar_por_item(cenario: EacScenario) -> None:
    tags = _tags(_content(projeto=str(cenario.first.id)))

    edit_links = [tag for tag in tags if "data-editar" in tag.attrs]
    assert len(edit_links) == 3  # só os itens, nunca pacote nem subpacote
    assert all(link.attrs["x-target"] == "eac-formulario" for link in edit_links)
    assert any("/editar" in link.attrs["href"] for link in edit_links)


def test_conteudo_de_visualizador_nao_tem_o_botao_de_editar(cenario: EacScenario) -> None:
    tags = _tags(_content(email=VIEWER_EMAIL, projeto=str(cenario.first.id)))

    assert not [tag for tag in tags if "data-editar" in tag.attrs]


@pytest.mark.usefixtures("cenario")
def test_conteudo_do_portfolio_e_so_leitura_com_projeto_e_pacote_principal() -> None:
    response = _content(projeto="portfolio")

    assert _codes(response) == ["0", "1", "1.1", "1.2", "2", "2.1"]
    tags = _tags(response)
    marker = next(tag for tag in tags if "data-resultado" in tag.attrs)
    assert marker.attrs["data-carteira"] == "sim"
    assert not [tag for tag in tags if "data-editar" in tag.attrs]
    open_links = [
        tag for tag in tags if tag.attrs.get("href", "").startswith("/financeiro/eac?projeto=")
    ]
    assert len(open_links) == 2  # um "Abrir" por projeto
    body = response.get_body().decode()
    assert "Peso na carteira" in body
    assert "58,34%" in body
    assert "TN-EAC-001 · Fábrica de teste" in body


def test_o_escopo_vem_do_cookie_quando_a_url_nao_o_traz(cenario: EacScenario) -> None:
    plain = _request("/api/financeiro/eac")
    request = func.HttpRequest(
        method="GET",
        url="/api/financeiro/eac",
        headers={**dict(plain.headers), "Cookie": f"gestnow_projeto={cenario.second.id}"},
        params={},
        route_params={},
        body=b"",
    )

    assert _codes(routes.eac_content(request)) == ["0", "1", "1.1", "1.1.1"]


def test_filtros_da_url_chegam_a_arvore_e_ao_total_filtrado(cenario: EacScenario) -> None:
    response = _content(projeto=str(cenario.first.id), busca="BOMBAS")

    assert _codes(response) == ["0", "2", "2.1", "2.1.1"]
    assert "Total filtrado" in response.get_body().decode()


def test_nivel_ilegivel_na_url_vale_o_mais_fundo(cenario: EacScenario) -> None:
    response = _content(projeto=str(cenario.first.id), nivel="abc", tipo="Invenção")

    assert len(_codes(response)) == 8


def test_filtro_sem_resultado_e_vazio_por_filtro(cenario: EacScenario) -> None:
    response = _content(projeto=str(cenario.first.id), busca="zzz")

    marker = next(tag for tag in _tags(response) if "data-resultado" in tag.attrs)
    assert marker.attrs["data-resultado"] == "vazio-filtro"


def test_projeto_sem_eac_e_vazio_de_origem(cenario: EacScenario) -> None:
    cenario.session.query(EacItem).filter(EacItem.project_id == cenario.second.id).delete()

    response = _content(projeto=str(cenario.second.id))

    marker = next(tag for tag in _tags(response) if "data-resultado" in tag.attrs)
    assert marker.attrs["data-resultado"] == "vazio-origem"


@pytest.mark.usefixtures("cenario")
def test_sem_o_cabecalho_do_alpine_a_rota_manda_para_o_shell() -> None:
    response = routes.eac_content(_request("/api/financeiro/eac", target=None))

    assert response.status_code == 302


@pytest.mark.usefixtures("cenario")
def test_quem_nao_esta_no_cadastro_recebe_403() -> None:
    response = _content(email="intruso@example.invalid")

    assert response.status_code == 403


# ── O formulário do cadastro ─────────────────────────────────────────────


def test_formulario_traz_os_campos_cadastrais_a_versao_e_o_escopo_dos_filtros(
    cenario: EacScenario,
) -> None:
    request = _request(
        "/api/financeiro/eac/itens/x/editar",
        route=_item_url_params(cenario),
        target="eac-formulario",
        projeto=str(cenario.first.id),
        busca="civil",
        nivel="3",
    )

    response = routes.eac_item_form(request)

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, "eac-formulario") is not None
    names = {tag.attrs.get("name") for tag in tags if tag.name in {"input", "select", "textarea"}}
    assert {
        "descricao",
        "tipo_custo",
        "classificacao",
        "centro_custo",
        "responsavel",
        "justificativa",
        "versao",
        "busca",
        "nivel",
        "tipo",
    } <= names
    hidden = {
        tag.attrs["name"]: tag.attrs["value"]
        for tag in find_all(tags, "input")
        if tag.attrs.get("type") == "hidden"
    }
    assert hidden["versao"] == "1"
    assert hidden["busca"] == "civil"
    form = find_all(tags, "form")[0]
    assert form.attrs["x-target"] == "eac-formulario eac-conteudo"
    assert "Orçado atual R$ 10.000,00" in response.get_body().decode()


def test_formulario_do_portfolio_e_recusado_com_403(cenario: EacScenario) -> None:
    request = _request(
        "/api/financeiro/eac/itens/x/editar",
        route=_item_url_params(cenario),
        target="eac-formulario",
        projeto="portfolio",
    )

    assert routes.eac_item_form(request).status_code == 403


def test_formulario_do_visualizador_e_recusado_com_403(cenario: EacScenario) -> None:
    request = _request(
        "/api/financeiro/eac/itens/x/editar",
        email=VIEWER_EMAIL,
        route=_item_url_params(cenario),
        target="eac-formulario",
        projeto=str(cenario.first.id),
    )

    assert routes.eac_item_form(request).status_code == 403


# ── O cadastro do item ───────────────────────────────────────────────────


def test_salvar_com_justificativa_grava_responde_os_dois_alvos_e_avisa(
    cenario: EacScenario,
) -> None:
    response = _save(cenario, _form_values(cenario))

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, "eac-formulario") is not None  # volta vazio: o formulário fecha
    assert find_by_id(tags, "eac-conteudo") is not None  # e o conteúdo volta atualizado
    assert "Engenharia civil e estrutural" in response.get_body().decode()
    assert response.headers["X-TN-Toast-Tipo"] == "ok"
    assert "1.1.1" in response.headers["X-TN-Toast"].replace("%20", " ")
    item = cenario.session.get(EacItem, cenario.items[f"{cenario.first.code}:1.1.1"])
    assert item is not None
    assert item.description == "Engenharia civil e estrutural"
    assert item.version == 2


def test_salvar_sem_justificativa_volta_o_formulario_preenchido_com_422(
    cenario: EacScenario,
) -> None:
    response = _save(cenario, _form_values(cenario, justificativa=""))

    assert response.status_code == 422
    tags = _tags(response)
    assert find_by_id(tags, "eac-formulario") is not None
    description = next(tag for tag in tags if tag.attrs.get("name") == "descricao")
    assert description.attrs["value"] == "Engenharia civil e estrutural"  # o que a pessoa digitou
    errors = find_all(tags, "span", "field__erro")
    assert len(errors) == 1
    assert "mínimo de 10 caracteres" in errors[0].clean_text()
    item = cenario.session.get(EacItem, cenario.items[f"{cenario.first.code}:1.1.1"])
    assert item is not None
    assert item.description == "Engenharia civil"


def test_salvar_com_versao_antiga_volta_o_formulario_com_409_e_o_nome_de_quem_gravou(
    cenario: EacScenario,
) -> None:
    assert _save(cenario, _form_values(cenario)).status_code == 200

    response = _save(cenario, _form_values(cenario, descricao="Outro nome", versao="1"))

    assert response.status_code == 409
    body = response.get_body().decode()
    assert "Este registro foi alterado por Maria Membro às" in body
    description = next(tag for tag in _tags(response) if tag.attrs.get("name") == "descricao")
    assert description.attrs["value"] == "Outro nome"


def test_salvar_sem_diferenca_avisa_que_nada_mudou(cenario: EacScenario) -> None:
    values = _form_values(cenario, descricao="Engenharia civil")

    response = _save(cenario, values)

    assert response.status_code == 200
    assert response.headers["X-TN-Toast-Tipo"] == "aviso"


def test_salvar_no_portfolio_e_recusado_com_403(cenario: EacScenario) -> None:
    request = _request(
        "/api/financeiro/eac/itens/x",
        route=_item_url_params(cenario),
        form=_form_values(cenario),
        target="eac-formulario eac-conteudo",
        projeto="portfolio",
    )

    assert routes.eac_item_save(request).status_code == 403


def test_salvar_como_visualizador_e_recusado_com_403(cenario: EacScenario) -> None:
    response = _save(cenario, _form_values(cenario), email=VIEWER_EMAIL)

    assert response.status_code == 403


def test_salvar_item_que_nao_existe_e_422(cenario: EacScenario) -> None:
    request = _request(
        "/api/financeiro/eac/itens/x",
        route={"item_id": "987654321"},
        form=_form_values(cenario),
        target="eac-formulario eac-conteudo",
        projeto=str(cenario.first.id),
    )

    assert routes.eac_item_save(request).status_code == 422


def test_gestor_tambem_salva(cenario: EacScenario) -> None:
    response = _save(cenario, _form_values(cenario), email=MANAGER_EMAIL)

    assert response.status_code == 200


def test_outra_pessoa_como_responsavel_e_aceita(cenario: EacScenario) -> None:
    other = Person(name="Outra Pessoa", email="outra-rota@example.invalid")
    cenario.session.add(other)
    cenario.session.flush()

    response = _save(cenario, _form_values(cenario, responsavel=str(other.id)))

    assert response.status_code == 200
    assert "Outra Pessoa" in response.get_body().decode()


# ── Excel e PDF ──────────────────────────────────────────────────────────


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> tuple[int, list[object]]:
    first_row = int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    width = sheet.max_column
    return first_row, [
        sheet.cell(row=first_row, column=column).value for column in range(1, width + 1)
    ]


def _excel(**params: str) -> func.HttpResponse:
    return routes.eac_excel(_request("/api/financeiro/eac/excel", target=None, **params))


def test_excel_do_projeto_tem_as_colunas_da_tela_e_a_arvore_inteira(cenario: EacScenario) -> None:
    response = _excel(projeto=str(cenario.first.id))

    assert response.status_code == 200
    sheets = _sheets(response)
    assert list(sheets) == ["Resumo", "Estrutura analítica de custos"]
    sheet = sheets["Estrutura analítica de custos"]
    first_row, header = _header_row(sheet)
    assert header == [
        "Código",
        "Descrição",
        "Tipo de custo",
        "Un.",
        "Qtd.",
        "Preço unitário",
        "Valor orçado",
        "CAPEX/OPEX",
        "Centro de custo",
        "Responsável",
    ]
    codes = [sheet.cell(row=first_row + offset, column=1).value for offset in range(1, 9)]
    assert codes == ["0", "1", "1.1", "1.1.1", "1.1.2", "2", "2.1", "2.1.1"]
    assert sheet.cell(row=first_row + 1, column=7).value == 70_010.0  # a linha 0: o total
    assert sheet.cell(row=first_row + 4, column=8).value == "CAPEX"
    summary = sheets["Resumo"]
    assert summary["A3"].value == "Escopo: TN-EAC-001 · Fábrica de teste"


@pytest.mark.usefixtures("cenario")
def test_excel_do_portfolio_abre_com_projeto_e_traz_o_peso() -> None:
    response = _excel(projeto="portfolio")

    sheet = _sheets(response)["Estrutura analítica de custos"]
    first_row, header = _header_row(sheet)
    assert header[0] == "Projeto"
    assert header[-1] == "Peso na carteira"
    projects = {sheet.cell(row=first_row + offset, column=1).value for offset in range(2, 7)}
    assert projects == {"TN-EAC-001 · Fábrica de teste", "TN-EAC-002 · Caldeira de teste"}
    weight = sheet.cell(row=first_row + 2, column=header.index("Peso na carteira") + 1).value
    assert weight == pytest.approx(0.5834)


def test_excel_leva_os_filtros_e_o_total_filtrado(cenario: EacScenario) -> None:
    sheet = _sheets(_excel(projeto=str(cenario.first.id), busca="bombas"))[
        "Estrutura analítica de custos"
    ]

    first_row, _ = _header_row(sheet)
    codes = [sheet.cell(row=first_row + offset, column=1).value for offset in range(1, 5)]
    assert codes == ["0", "2", "2.1", "2.1.1"]
    assert sheet.cell(row=first_row + 5, column=1).value == "Total filtrado"
    assert sheet.cell(row=first_row + 5, column=7).value == 60_000.0


@pytest.mark.usefixtures("cenario")
def test_excel_nao_depende_do_cabecalho_do_alpine_e_recusa_quem_nao_esta_no_cadastro() -> None:
    refused = routes.eac_excel(
        _request("/api/financeiro/eac/excel", email="intruso@example.invalid", target=None)
    )

    assert refused.status_code == 403
    assert refused.headers["Content-Type"].startswith("text/plain")


def test_versao_imprimivel_traz_o_mesmo_conteudo_do_excel(cenario: EacScenario) -> None:
    response = routes.eac_printable(
        _request(
            "/api/financeiro/eac/imprimivel",
            target="folha-impressao",
            projeto=str(cenario.first.id),
        )
    )

    assert response.status_code == 200
    body = response.get_body().decode()
    for text in (
        "EAC: estrutura analítica de custos",
        "Engenharia civil",
        "Bombas",
        "R$ 70.010,00",
    ):
        assert text in body
    headers = [tag.clean_text() for tag in find_all(parse_tags(body), "th")]
    assert "Valor orçado" in headers
    assert "Responsável" in headers
