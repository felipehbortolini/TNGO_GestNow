"""Rotas da tela Contingência (ISSUE-041, HU-094, D12, D14).

Costura HTTP: o handler é chamado com a ``HttpRequest`` que o navegador montaria, como em
``test_contratos_consolidado_rotas.py``. A carga de demonstração é a base do escopo, com a âncora
injetada no relógio: o projeto 1 tem os números do mock (saldo de R$ 1,6 mi, consumo de 20,0%,
limite de 71,8% e R$ 480 mil em SMs em análise). O fragmento traz os indicadores, o gráfico do
saldo (com o ``data-dados`` que a biblioteca lê), a composição e o extrato — no Portfólio com a
tabela por projeto e a coluna Projeto. O Excel e o imprimível levam o mesmo conteúdo; o
Visualizador lê e o fornecedor não.
"""

from __future__ import annotations

import json
from io import BytesIO

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.carga import DEMO_ANCHOR, run_demonstration
from src.core import auth, calendario
from src.modulos.configuracoes.models import Project
from src.modulos.financeiro import routes
from tests.financeiro.conftest import MEMBER_EMAIL, EacScenario
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal

ADMIN_EMAIL = "leonardo.gomes@exemplo.com"
VIEWER_EMAIL = "visualizador@cliente.demo"
SUPPLIER_EMAIL = "carlos.nunes@exemplo.com"
PROJECT_CODE = "TN-2026-014"
PATH = "/api/financeiro/contingencia"
CONTENT_TARGET = "contingencia-conteudo"
FOLHA_TARGET = "folha-impressao"

KPI_LABELS = (
    "Saldo da contingência",
    "Consumo da contingência",
    "Pedido em SMs em análise",
    "Cobertura da exposição a riscos",
    "Reserva gerencial",
)
COMPOSITION_ROWS = (
    "Orçado atual da EAC (linha de base de medição)",
    "Saldo da reserva de contingência",
    "Linha de base de custo",
    "Saldo da reserva gerencial",
    "Orçamento do projeto",
)
MOVEMENT_HEADERS = ("Data", "Reserva", "Movimento", "Descrição", "Valor", "Saldo")
BURN_HEADERS = ("Mês", "Saldo previsto", "Saldo real")


@pytest.fixture(autouse=True)
def data_de_referencia(monkeypatch: pytest.MonkeyPatch) -> None:
    """O relógio parado na âncora da demonstração: os números do mock não mudam."""
    monkeypatch.setattr(calendario, "today", lambda: DEMO_ANCHOR)


@pytest.fixture
def carga(rotas_na_transacao: Session) -> Session:
    """A base de demonstração na transação do teste, com as rotas lendo a mesma sessão."""
    run_demonstration(rotas_na_transacao, reference_date=DEMO_ANCHOR)
    return rotas_na_transacao


def _request(
    path: str,
    *,
    email: str = ADMIN_EMAIL,
    target: str | None = CONTENT_TARGET,
    **params: str,
) -> func.HttpRequest:
    """A request como o navegador a envia; ``target=None`` é um link puro (sem Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params={},
        body=b"",
    )


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _table_text(response: func.HttpResponse, attr: str) -> str:
    """O texto de dentro da tabela marcada com o atributo (``data-movimentos``...)."""
    table = next(tag for tag in _tags(response) if tag.name == "table" and attr in tag.attrs)
    return table.clean_text()


def _projeto(session: Session) -> Project:
    return session.scalars(select(Project).where(Project.code == PROJECT_CODE)).one()


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> int:
    return int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))


# ── Leitura do fragmento ─────────────────────────────────────────────────


def test_a_contingencia_do_projeto_traz_kpis_grafico_composicao_e_extrato(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contingencia_content(_request(PATH, projeto=str(project.id)))

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, CONTENT_TARGET) is not None
    body = _body(response)
    assert 'data-situacao="conteudo"' in body
    assert 'data-carteira="nao"' in body
    for rotulo in KPI_LABELS:
        assert rotulo in body
    assert "R$ 1.600.000,00" in body, "o saldo da contingência do projeto 1"
    assert "R$ 2.000.000,00" in body, "a constituída"
    assert "R$ 480.000,00" in body, "o pedido em SMs em análise"
    assert "20,0%" in body, "o consumo"
    assert "71,8%" in body, "o limite do consumo"
    assert "R$ 900.000,00" in body, "a reserva gerencial"

    assert 'data-grafico="linhas-multiplas"' in body
    (host,) = [tag for tag in tags if tag.attrs.get("data-grafico") == "linhas-multiplas"]
    dados = json.loads(host.attrs["data-dados"])
    assert [serie["id"] for serie in dados["series"]] == ["real", "previsto"]
    assert dados["series"][1]["papel"] == "linha-base", "o previsto é a linha tracejada"
    assert dados["periodos"], "o gráfico tem os meses da Curva S congelada"
    assert dados["periodos"][0]["valores"]["previsto"] is not None

    composicao = _table_text(response, "data-composicao")
    for linha in COMPOSITION_ROWS:
        assert linha in composicao
    assert "R$ 1.600.000,00" in composicao

    movimentos = _table_text(response, "data-movimentos")
    for coluna in MOVEMENT_HEADERS:
        assert coluna in movimentos
    assert "SM-TN-2026-0004" in movimentos, "o consumo aprovado com a SM linkada"
    assert "Aguardando comitê" in movimentos, "a situação da linha Em análise"
    assert "Constituição" in movimentos and "Em análise" in movimentos
    assert "Projeto" not in movimentos, "a coluna Projeto só entra no Portfólio"


@pytest.mark.usefixtures("carga")
def test_a_contingencia_da_carteira_traz_a_tabela_por_projeto() -> None:
    response = routes.contingencia_content(_request(PATH, projeto="portfolio"))

    assert response.status_code == 200
    body = _body(response)
    assert 'data-situacao="conteudo"' in body
    assert 'data-carteira="sim"' in body
    assert body.count('<th scope="col">Projeto</th>') == 3, (
        "a tabela por projeto, o extrato e a lista de ameaças ativas"
    )

    projetos = _table_text(response, "data-projetos")
    for coluna in ("Projeto", "Constituída", "Consumida", "Saldo", "Consumo", "Em análise"):
        assert coluna in projetos
    assert "Reserva gerencial" in projetos
    assert "TN-2026-014" in projetos and "TN-2026-021" in projetos and "TN-2026-027" in projetos
    assert "R$ 2.690.000,00" in projetos, "o saldo da carteira é a soma dos projetos"

    movimentos = _table_text(response, "data-movimentos")
    assert "TN-2026-014" in movimentos and "TN-2026-021" in movimentos
    assert "SM-CB-2026-0001" in movimentos, "o consumo do projeto 2"

    ameacas = _table_text(response, "data-ameacas")
    for coluna in ("Risco", "Ameaça", "Severidade", "Situação", "VME", "SMs vinculadas"):
        assert coluna in ameacas
    assert "RSK-TN-2026-0001" in ameacas, "a licença de operação do projeto 1"


@pytest.mark.usefixtures("carga")
def test_sem_o_cabecalho_do_alpine_a_contingencia_manda_para_o_shell() -> None:
    response = routes.contingencia_content(_request(PATH, target=None, projeto="1"))

    assert response.status_code == 302


def test_escopo_sem_reservas_mostra_o_vazio_de_origem(cenario: EacScenario) -> None:
    response = routes.contingencia_content(
        _request(PATH, email=MEMBER_EMAIL, projeto=str(cenario.first.id))
    )

    assert response.status_code == 200
    body = _body(response)
    assert 'data-situacao="vazio-origem"' in body
    assert "Este projeto ainda não tem reservas constituídas." in body
    assert "Saldo da contingência" not in body


# ── Permissões ───────────────────────────────────────────────────────────


def test_o_visualizador_le_a_contingencia(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contingencia_content(
        _request(PATH, email=VIEWER_EMAIL, projeto=str(project.id))
    )

    assert response.status_code == 200
    assert "Saldo da contingência" in _body(response)


def test_o_fornecedor_nao_le_a_contingencia_nem_baixa_o_excel(carga: Session) -> None:
    project = _projeto(carga)

    conteudo = routes.contingencia_content(
        _request(PATH, email=SUPPLIER_EMAIL, projeto=str(project.id))
    )
    assert conteudo.status_code == 403

    excel = routes.contingencia_excel(
        _request(PATH + "/excel", email=SUPPLIER_EMAIL, target=None, projeto=str(project.id))
    )
    assert excel.status_code == 403


# ── Excel e imprimível ───────────────────────────────────────────────────


def test_o_excel_da_contingencia_traz_o_resumo_as_tabelas_e_o_burn(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contingencia_excel(
        _request(PATH + "/excel", target=None, projeto=str(project.id))
    )

    assert response.status_code == 200
    assert response.headers["Content-Disposition"].startswith(
        'attachment; filename="contingencia-e-reservas-'
    )
    sheets = _sheets(response)
    assert "Resumo" in sheets
    assert any(name.startswith("Linha de base de custo") for name in sheets)
    assert "Saldo real x previsto" in sheets
    assert "Extrato de movimentos" in sheets

    labels = {
        sheets["Resumo"].cell(row=row, column=1).value
        for row in range(1, sheets["Resumo"].max_row + 1)
    }
    assert set(KPI_LABELS) <= labels

    composicao = next(sheet for name, sheet in sheets.items() if name.startswith("Linha de base"))
    assert any(
        composicao.cell(row=row, column=1).value == "Linha de base de custo"
        for row in range(1, composicao.max_row + 1)
    )

    burn = sheets["Saldo real x previsto"]
    first_row = _header_row(burn)
    assert [burn.cell(row=first_row, column=column).value for column in range(1, 4)] == list(
        BURN_HEADERS
    )
    assert burn.cell(row=first_row + 1, column=2).value is not None


@pytest.mark.usefixtures("carga")
def test_o_excel_da_carteira_traz_a_tabela_por_projeto() -> None:
    response = routes.contingencia_excel(
        _request(PATH + "/excel", target=None, projeto="portfolio")
    )

    sheets = _sheets(response)
    assert "Reservas por projeto" in sheets
    projeto = sheets["Reservas por projeto"]
    first_row = _header_row(projeto)
    assert projeto.cell(row=first_row, column=1).value == "Projeto"
    codes = {
        projeto.cell(row=row, column=1).value for row in range(first_row + 1, projeto.max_row + 1)
    }
    assert any(str(code).startswith("TN-2026-014 · ") for code in codes)


def test_o_imprimivel_da_contingencia_traz_o_mesmo_conteudo_do_excel(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contingencia_printable(
        _request(PATH + "/imprimivel", target=FOLHA_TARGET, projeto=str(project.id))
    )

    assert response.status_code == 200
    body = _body(response)
    for texto in (
        "Contingência e reservas",
        *KPI_LABELS,
        *COMPOSITION_ROWS,
        "Extrato de movimentos",
        "Saldo real x previsto",
        "R$ 1.600.000,00",
    ):
        assert texto in body
    assert 'data-grafico="linhas-multiplas"' in body
