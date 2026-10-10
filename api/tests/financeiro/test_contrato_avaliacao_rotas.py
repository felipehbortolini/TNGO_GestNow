"""Rotas da avaliação de desempenho da contratada (ISSUE-034, D12, D14).

Costura HTTP: o handler é chamado com a ``HttpRequest`` que o navegador montaria, como em
``test_contrato_fluxos_rotas.py``. A aba lista as avaliações e a série do Gráfico de Áreas de
Avaliação; o formulário traz os critérios vigentes com os pesos que ficam gravados na avaliação; a
gravação é um ``multipart/form-data`` (a evidência viaja como anexo de verdade quando alguma nota
é 1 ou 2) e devolve a aba atualizada com o formulário vazio; a aba tem o seu Excel e o seu
imprimível.
"""

from __future__ import annotations

import json
from collections.abc import Mapping
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any
from urllib.parse import unquote, urlencode

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select

from src.core import auth, calendario, file_storage
from src.core.models import Attachment
from src.core.rbac import GeneralProfile
from src.modulos.central_acoes.models import Action
from src.modulos.configuracoes import service as configuracoes
from src.modulos.financeiro import calculations, routes, service
from src.modulos.financeiro.models import (
    Contract,
    ContractEvaluation,
    ContractEvaluationCriterion,
)
from src.modulos.governanca import lessons_models
from src.modulos.governanca.lessons_models import Lesson
from src.modulos.hse import service as hse_service
from src.modulos.hse.validation import HoursInput, OccurrenceInput
from tests.armazenamento_falso import ArmazenamentoEmMemoria
from tests.financeiro.conftest import (
    MANAGER_EMAIL,
    MEMBER_EMAIL,
    VIEWER_EMAIL,
    EacScenario,
    user_of,
)
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal, criar_empresa

HOJE = date(2026, 10, 6)
NUMERO = "CT-2026-011"
ORIGINAL_CENTS = 100_000_000
CONTRACT_END = date(2026, 10, 30)
PARAMETERS_FROM = date(2026, 1, 1)
EVIDENCE_BYTES = b"%PDF-1.7 evidencia da avaliacao"
NOTES = {"hse": 4, "qualidade": 4, "prazo": 3, "gestao": 3, "recursos": 4, "documentacao": 3}
WEIGHTS = {
    "hse": 25,
    "qualidade": 20,
    "prazo": 20,
    "gestao": 15,
    "recursos": 10,
    "documentacao": 10,
}
EVALUATIONS_TARGET = "contrato-avaliacoes"
FORM_TARGET = "contrato-formulario"
BOUNDARY = "fronteira-do-teste-das-avaliacoes"


@pytest.fixture(autouse=True)
def data_de_hoje(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(calendario, "today", lambda: HOJE)


@pytest.fixture
def armazenamento(monkeypatch: pytest.MonkeyPatch) -> ArmazenamentoEmMemoria:
    """O arquivo da evidência fica na memória: o teste lê o que foi guardado."""
    storage = ArmazenamentoEmMemoria()
    monkeypatch.setattr(file_storage, "configured_storage", lambda: storage)
    return storage


def _request(
    path: str,
    *,
    email: str = MEMBER_EMAIL,
    route: Mapping[str, str] | None = None,
    payload: Mapping[str, str] | tuple[bytes, str] | None = None,
    target: str | None = EVALUATIONS_TARGET,
    **params: str,
) -> func.HttpRequest:
    """A request as the browser sends it; ``target=None`` is a plain link (no Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    body = b""
    if isinstance(payload, tuple):
        body, content_type = payload
        headers["Content-Type"] = content_type
    elif payload is not None:
        headers["Content-Type"] = "application/x-www-form-urlencoded"
        body = urlencode(payload).encode()
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="POST" if payload is not None else "GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params=dict(route or {}),
        body=body,
    )


def _multipart(
    fields: Mapping[str, str], files: tuple[str, bytes] | None = None
) -> tuple[bytes, str]:
    """A ``multipart/form-data`` body as the browser sends it, and its Content-Type."""
    parts = [
        f'--{BOUNDARY}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{value}\r\n'.encode()
        for name, value in fields.items()
    ]
    if files is not None:
        name, content = files
        header = (
            f"--{BOUNDARY}\r\n"
            f'Content-Disposition: form-data; name="arquivo"; filename="{name}"\r\n'
            "Content-Type: application/pdf\r\n\r\n"
        )
        parts.append(header.encode() + content + b"\r\n")
    parts.append(f"--{BOUNDARY}--\r\n".encode())
    return b"".join(parts), f"multipart/form-data; boundary={BOUNDARY}"


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _toast(response: func.HttpResponse) -> str:
    return unquote(response.headers["X-TN-Toast"])


def _criar_contrato(cenario: EacScenario) -> Contract:
    session = cenario.session
    empresa = criar_empresa(session, "Contratada de teste")
    contract = Contract(
        project_id=cenario.first.id,
        company_id=empresa.id,
        manager_id=cenario.responsible.id,
        fiscal_id=cenario.responsible.id,
        number=NUMERO,
        object="Obra de teste",
        modality="Preço unitário",
        original_value_cents=ORIGINAL_CENTS,
        start=date(2026, 1, 5),
        original_end=CONTRACT_END,
        current_end=CONTRACT_END,
        retention_pct=5,
        claim_notice_days=30,
        situation="Em execução",
    )
    session.add(contract)
    session.flush()
    return contract


def _semear_parametros(cenario: EacScenario) -> None:
    configuracoes.seed_initial_parameters(
        cenario.session, author_id=cenario.manager.id, effective_from=PARAMETERS_FROM
    )


def _avaliacao(cenario: EacScenario, contract: Contract, **fields: Any) -> ContractEvaluation:
    """Uma avaliação direta na tabela, com os critérios congelados, como o serviço os grava."""
    session = cenario.session
    version = configuracoes.current_version(
        session, group=service.EVALUATION_GROUP, reference_date=HOJE
    )
    assert version is not None
    notas = fields.get("notas") or NOTES
    record = ContractEvaluation(
        contract_id=contract.id,
        evaluator_id=cenario.responsible.id,
        parameter_version_id=version.id,
        period=date.fromisoformat(f"{fields.get('periodo') or '2026-08'}-01"),
        kind=fields.get("tipo") or calculations.EVALUATION_KIND_MONTHLY,
        comment=fields.get("comentario") or None,
    )
    session.add(record)
    session.flush()
    session.add_all(
        ContractEvaluationCriterion(
            evaluation_id=record.id,
            criterion=criterio,
            weight=peso,
            note=notas[criterio],
        )
        for criterio, peso in WEIGHTS.items()
    )
    session.flush()
    return record


def _ficha_params(cenario: EacScenario) -> dict[str, str]:
    return {"projeto": str(cenario.first.id), "numero": NUMERO}


def _notas_do_form(notas: Mapping[str, int] | None = None) -> dict[str, str]:
    return {criterio: str(nota) for criterio, nota in (notas or NOTES).items()}


# ── Leitura da aba ───────────────────────────────────────────────────────


def test_a_aba_lista_as_avaliacoes_com_notas_classe_e_grafico(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    contract = _criar_contrato(cenario)
    _avaliacao(
        cenario,
        contract,
        periodo="2026-07",
        notas=dict.fromkeys(NOTES, 5),
        comentario="Mobilização dentro do plano.",
    )
    _avaliacao(cenario, contract, periodo="2026-08")

    response = routes.contrato_avaliacoes(
        _request(
            "/api/financeiro/contrato/avaliacoes",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, EVALUATIONS_TARGET) is not None
    body = _body(response)
    assert "07/2026" in body
    assert "08/2026" in body
    assert "Mensal" in body
    assert "Maria Membro" in body
    assert "Mobilização dentro do plano." in body
    assert "5 · 25%" in body
    assert "4 · 25%" in body
    assert "71" in body
    assert "Aprovada" in body

    (host,) = [tag for tag in tags if tag.attrs.get("data-grafico") == "areas-avaliacao"]
    dados = json.loads(host.attrs["data-dados"])
    assert dados["pontos"] == [
        {"ano": 2026, "mes": 7, "valor": 100},
        {"ano": 2026, "mes": 8, "valor": 71},
    ]
    assert dados["zonas"] == [
        {"de": 0, "rotulo": "Classe D", "papel": "erro"},
        {"de": 50, "rotulo": "Classe C", "papel": "alerta"},
        {"de": 70, "rotulo": "Classe B", "papel": "marca"},
        {"de": 85, "rotulo": "Classe A", "papel": "ok"},
    ]
    assert dados["meta"] == 85


def test_a_aba_sem_avaliacoes_traz_o_grafico_vazio(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    _criar_contrato(cenario)

    response = routes.contrato_avaliacoes(
        _request("/api/financeiro/contrato/avaliacoes", **_ficha_params(cenario))
    )

    assert response.status_code == 200
    assert "Nenhuma avaliação registrada." in _body(response)


def test_sem_o_cabecalho_do_alpine_a_aba_manda_para_o_shell(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_avaliacoes(
        _request("/api/financeiro/contrato/avaliacoes", target=None, numero=NUMERO)
    )

    assert response.status_code == 302


# ── Formulário e gravação ────────────────────────────────────────────────


def test_o_formulario_mostra_os_criterios_vigentes_com_os_pesos(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    _criar_contrato(cenario)

    response = routes.contrato_avaliacao_form(
        _request(
            "/api/financeiro/contrato/avaliacoes/nova",
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, FORM_TARGET) is not None
    names = {tag.attrs.get("name") for tag in tags if tag.name in {"input", "select", "textarea"}}
    assert {
        "periodo",
        "tipo",
        "comentario",
        "arquivo",
        "hse",
        "qualidade",
        "prazo",
        "gestao",
        "recursos",
        "documentacao",
    } <= names
    body = _body(response)
    assert "HSE" in body
    assert "peso 25%" in body
    assert "Gestão contratual e comercial" in body
    assert f"/api/financeiro/contrato/avaliacoes?numero={NUMERO}" in body


def test_registrar_avaliacao_mensal_devolve_a_aba_e_fecha_o_formulario(
    cenario: EacScenario,
) -> None:
    _semear_parametros(cenario)
    _criar_contrato(cenario)

    response = routes.contrato_avaliacao_create(
        _request(
            "/api/financeiro/contrato/avaliacoes",
            email=MANAGER_EMAIL,
            payload=_multipart(
                {
                    "periodo": "2026-08",
                    "tipo": calculations.EVALUATION_KIND_MONTHLY,
                    "comentario": "Avaliação do mês.",
                    **_notas_do_form(),
                }
            ),
            target=f"{FORM_TARGET} {EVALUATIONS_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, FORM_TARGET) is not None
    assert find_by_id(tags, EVALUATIONS_TARGET) is not None
    assert _toast(response) == "Avaliação mensal de 08/2026 registrada."
    assert "71" in _body(response)

    (record,) = cenario.session.scalars(select(ContractEvaluation)).all()
    assert (record.period, record.kind, record.comment) == (
        date(2026, 8, 1),
        calculations.EVALUATION_KIND_MONTHLY,
        "Avaliação do mês.",
    )
    criteria = cenario.session.scalars(
        select(ContractEvaluationCriterion)
        .where(ContractEvaluationCriterion.evaluation_id == record.id)
        .order_by(ContractEvaluationCriterion.id)
    ).all()
    assert [(row.criterion, row.weight, row.note) for row in criteria] == [
        ("hse", 25, 4),
        ("qualidade", 20, 4),
        ("prazo", 20, 3),
        ("gestao", 15, 3),
        ("recursos", 10, 4),
        ("documentacao", 10, 3),
    ]


def test_registrar_avaliacao_final_cria_a_licao_em_rascunho(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    contract = _criar_contrato(cenario)

    response = routes.contrato_avaliacao_create(
        _request(
            "/api/financeiro/contrato/avaliacoes",
            email=MANAGER_EMAIL,
            payload=_multipart(
                {
                    "periodo": "2026-09",
                    "tipo": calculations.EVALUATION_KIND_FINAL,
                    **_notas_do_form(dict.fromkeys(NOTES, 4)),
                }
            ),
            target=f"{FORM_TARGET} {EVALUATIONS_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert _toast(response) == (
        "Avaliação final de 09/2026 registrada. Lição de fornecedor criada em rascunho."
    )
    licao = cenario.session.scalars(select(Lesson)).one()
    assert (licao.situation, licao.origin, licao.origin_ref) == (
        lessons_models.SITUATION_DRAFT,
        lessons_models.ORIGIN_CONTRACT,
        contract.number,
    )


def test_avaliacao_com_nota_baixa_sem_evidencia_volta_no_formulario_com_422(
    cenario: EacScenario,
) -> None:
    _semear_parametros(cenario)
    _criar_contrato(cenario)

    response = routes.contrato_avaliacao_create(
        _request(
            "/api/financeiro/contrato/avaliacoes",
            email=MANAGER_EMAIL,
            payload=_multipart(
                {
                    "periodo": "2026-08",
                    "tipo": calculations.EVALUATION_KIND_MONTHLY,
                    **_notas_do_form({**NOTES, "prazo": 2}),
                }
            ),
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), FORM_TARGET) is not None
    body = _body(response)
    assert "exige anexo de evidência" in body
    assert 'name="prazo"' in body
    assert 'value="2"' in body
    assert cenario.session.scalars(select(ContractEvaluation)).all() == []


def test_avaliacao_com_nota_baixa_com_evidencia_cria_o_plano_e_o_anexo(
    cenario: EacScenario, armazenamento: ArmazenamentoEmMemoria
) -> None:
    _semear_parametros(cenario)
    contract = _criar_contrato(cenario)

    response = routes.contrato_avaliacao_create(
        _request(
            "/api/financeiro/contrato/avaliacoes",
            email=MANAGER_EMAIL,
            payload=_multipart(
                {
                    "periodo": "2026-08",
                    "tipo": calculations.EVALUATION_KIND_MONTHLY,
                    **_notas_do_form({**NOTES, "prazo": 2}),
                },
                ("evidencia da avaliacao.pdf", EVIDENCE_BYTES),
            ),
            target=f"{FORM_TARGET} {EVALUATIONS_TARGET}",
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    assert _toast(response) == (
        "Avaliação mensal de 08/2026 registrada. Plano de melhoria aberto na Central."
    )
    avaliacao = cenario.session.scalars(select(ContractEvaluation)).one()
    anexo = cenario.session.scalars(select(Attachment)).one()
    assert (anexo.origin_table, anexo.origin_record_id) == (
        service.EVALUATION_ATTACHMENT_TABLE,
        avaliacao.id,
    )
    assert armazenamento.arquivos == {f"{cenario.first.id}/{anexo.id}": EVIDENCE_BYTES}
    acao = cenario.session.scalars(select(Action)).one()
    assert (acao.origin, acao.origin_ref, acao.project_id) == (
        service.EVALUATION_ACTION_ORIGIN,
        contract.number,
        cenario.first.id,
    )


def test_o_visualizador_nao_registra_avaliacao(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    _criar_contrato(cenario)

    response = routes.contrato_avaliacao_create(
        _request(
            "/api/financeiro/contrato/avaliacoes",
            email=VIEWER_EMAIL,
            payload=_multipart(
                {
                    "periodo": "2026-08",
                    "tipo": calculations.EVALUATION_KIND_MONTHLY,
                    **_notas_do_form(),
                }
            ),
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 403


# ── Excel e imprimível ───────────────────────────────────────────────────


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> tuple[int, list[object]]:
    first_row = int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    width = sheet.max_column
    return first_row, [
        sheet.cell(row=first_row, column=column).value for column in range(1, width + 1)
    ]


def test_o_excel_das_avaliacoes_traz_os_criterios_a_nota_e_a_classe(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    contract = _criar_contrato(cenario)
    _avaliacao(cenario, contract, periodo="2026-08")

    response = routes.contrato_avaliacoes_excel(
        _request("/api/financeiro/contrato/avaliacoes/excel", target=None, numero=NUMERO)
    )

    assert response.status_code == 200
    sheets = _sheets(response)
    assert list(sheets) == ["Avaliações"]
    sheet = sheets["Avaliações"]
    first_row, header = _header_row(sheet)
    assert header == [
        "Período",
        "Tipo",
        "Avaliador",
        "HSE",
        "Qualidade",
        "Prazo",
        "Gestão contratual e comercial",
        "Recursos e mobilização",
        "Documentação e comunicação",
        "Nota",
        "Classe",
    ]
    line = first_row + 1
    assert str(sheet.cell(row=line, column=1).value)[:10] == "2026-08-01"
    assert sheet.cell(row=line, column=2).value == "Mensal"
    assert sheet.cell(row=line, column=3).value == "Maria Membro"
    assert sheet.cell(row=line, column=4).value == "4 · 25%"
    assert sheet.cell(row=line, column=10).value == 71
    assert sheet.cell(row=line, column=11).value == "B"


def test_a_versao_imprimivel_das_avaliacoes_traz_o_mesmo_conteudo(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    contract = _criar_contrato(cenario)
    _avaliacao(cenario, contract, periodo="2026-08")

    response = routes.contrato_avaliacoes_printable(
        _request(
            "/api/financeiro/contrato/avaliacoes/imprimivel",
            target="folha-impressao",
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    body = _body(response)
    assert f"Contrato {NUMERO}" in body
    assert "Avaliações" in body
    assert "HSE" in body
    assert "4 · 25%" in body
    assert "71" in body
    assert "B" in body
    assert "<script" not in body


# ── A sugestão do HSE no critério HSE (ISSUE-075) ────────────────────────────


def _gravar_hht(cenario: EacScenario, contract: Contract) -> None:
    hse_service.save_hours(
        cenario.session,
        user=user_of(cenario.member, cenario.session, GeneralProfile.MEMBER),
        data=HoursInput(
            project_id=cenario.first.id,
            month=date(2026, 9, 1),
            company_id=contract.company_id,
            headcount=10,
            hours=Decimal("300000"),
        ),
        reference_date=HOJE,
    )


def _gravar_lesao_registravel(cenario: EacScenario, contract: Contract) -> None:
    hse_service.save_occurrence(
        cenario.session,
        user=user_of(cenario.member, cenario.session, GeneralProfile.MEMBER),
        data=OccurrenceInput(
            project_id=cenario.first.id,
            happened_on=date(2026, 9, 1),
            happened_time="09:30",
            area="Moagem 210",
            company_id=contract.company_id,
            kind="Trabalho restrito",
            description="Pessoa em trabalho restrito após avaliação médica.",
            role="Operador de moinho",
            potential_p=3,
            potential_i=4,
            immediate_cause="Procedimento não seguido",
            communication_hours=2,
            real_gravity=2,
            medical_notes="Trabalho restrito por 15 dias.",
        ),
        reference_date=HOJE,
    )


def test_o_formulario_pre_seleciona_a_nota_hse_sugerida(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    contract = _criar_contrato(cenario)
    _gravar_hht(cenario, contract)
    _gravar_lesao_registravel(cenario, contract)

    response = routes.contrato_avaliacao_form(
        _request(
            "/api/financeiro/contrato/avaliacoes/nova",
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    campo = find_by_id(_tags(response), "contrato-avaliacao-hse")
    assert campo is not None
    assert campo.attrs["value"] == "4"
    assert "Sugestão do HSE: nota 4 · TRIF da empresa 3,33 contra 3,33 do projeto" in _body(
        response
    )


def test_o_formulario_sem_hht_mostra_o_aviso_e_nao_pre_seleciona(
    cenario: EacScenario,
) -> None:
    _semear_parametros(cenario)
    _criar_contrato(cenario)

    response = routes.contrato_avaliacao_form(
        _request(
            "/api/financeiro/contrato/avaliacoes/nova",
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 200
    campo = find_by_id(_tags(response), "contrato-avaliacao-hse")
    assert campo is not None
    assert campo.attrs["value"] == ""
    assert "Sugestão do HSE: sem HHT registrado para a empresa no projeto." in _body(response)


def test_o_422_preserva_a_nota_digitada_no_criterio_hse(cenario: EacScenario) -> None:
    _semear_parametros(cenario)
    contract = _criar_contrato(cenario)
    _gravar_hht(cenario, contract)
    _gravar_lesao_registravel(cenario, contract)

    response = routes.contrato_avaliacao_create(
        _request(
            "/api/financeiro/contrato/avaliacoes",
            email=MANAGER_EMAIL,
            payload=_multipart(
                {
                    "periodo": "2026-08",
                    "tipo": calculations.EVALUATION_KIND_MONTHLY,
                    **_notas_do_form({**NOTES, "hse": 3, "prazo": 2}),
                }
            ),
            target=FORM_TARGET,
            **_ficha_params(cenario),
        )
    )

    assert response.status_code == 422
    campo = find_by_id(_tags(response), "contrato-avaliacao-hse")
    assert campo is not None
    assert campo.attrs["value"] == "3"
    assert "exige anexo de evidência" in _body(response)
