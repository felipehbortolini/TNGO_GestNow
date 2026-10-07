"""Rotas da ficha do contrato: Resumo, Medições e Aditivos (ISSUE-032, D12, D14).

Costura HTTP: o handler é chamado com a ``HttpRequest`` que o navegador montaria, como em
``test_eac_mapa_rotas.py``. A ficha é aberta pelo número (``?numero=``); o Resumo, as Medições e
os Aditivos são fragmentos por aba, e as operações devolvem a ficha atualizada numa resposta
multi-alvo (os painéis e o formulário vazio). O erro de validação volta no formulário, com o 422;
a transição fora do fluxo volta no painel das medições; quem não pode recebe 403. O Excel e o
imprimível levam os KPIs, a cascata e as tabelas de medições e aditivos.
"""

from __future__ import annotations

from collections.abc import Mapping
from datetime import date
from io import BytesIO
from urllib.parse import unquote, urlencode

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select

from src.core import auth, calendario
from src.core.rbac import GeneralProfile
from src.modulos.financeiro import calculations, routes, service
from src.modulos.financeiro.models import (
    Contract,
    ContractAddendum,
    ContractMeasurement,
    ErpCost,
)
from src.modulos.governanca.models import ChangeRequest
from tests.financeiro.conftest import (
    MANAGER_EMAIL,
    MEMBER_EMAIL,
    VIEWER_EMAIL,
    EacScenario,
    user_of,
)
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal, criar_empresa

HOJE = date(2026, 10, 6)
NUMERO = "CT-2026-011"
ORIGINAL_CENTS = 100_000_000
CONTRACT_END = date(2026, 10, 30)
EAC_LEAF = "1.1.1"
JUSTIFICATIVA = "Engenharia do sistema de tratamento de efluentes."
RESUME_TARGET = "contrato-resumo"
MEASUREMENTS_TARGET = "contrato-medicoes"
ADDENDA_TARGET = "contrato-aditivos"
FORM_TARGET = "contrato-formulario"


@pytest.fixture(autouse=True)
def data_de_hoje(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(calendario, "today", lambda: HOJE)


def _request(
    path: str,
    *,
    email: str = MEMBER_EMAIL,
    route: Mapping[str, str] | None = None,
    form: Mapping[str, str] | None = None,
    target: str | None = RESUME_TARGET,
    **params: str,
) -> func.HttpRequest:
    """A request as the browser sends it; ``target=None`` is a plain link (no Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    body = b""
    if form is not None:
        headers["Content-Type"] = "application/x-www-form-urlencoded"
        body = urlencode(form).encode()
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="POST" if form is not None else "GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params=dict(route or {}),
        body=body,
    )


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _toast(response: func.HttpResponse) -> str:
    return unquote(response.headers["X-TN-Toast"])


def _criar_contrato(
    cenario: EacScenario,
    *,
    number: str = NUMERO,
    original_cents: int = ORIGINAL_CENTS,
    item_code: str | None = EAC_LEAF,
) -> Contract:
    session = cenario.session
    empresa = criar_empresa(session, "Contratada de teste")
    contract = Contract(
        project_id=cenario.first.id,
        company_id=empresa.id,
        eac_item_id=cenario.items.get(f"TN-EAC-001:{item_code}") if item_code else None,
        manager_id=cenario.responsible.id,
        fiscal_id=cenario.responsible.id,
        number=number,
        object="Obra de teste",
        modality="Preço unitário",
        original_value_cents=original_cents,
        start=date(2026, 1, 5),
        original_end=CONTRACT_END,
        current_end=CONTRACT_END,
        retention_pct=5,
        claim_notice_days=30,
        situation="Em execução",
    )
    session.add(contract)
    session.flush()
    return contract


def _medicao(
    cenario: EacScenario,
    contract: Contract,
    *,
    number: str = "BM-01",
    situation: str = calculations.SITUATION_UNDER_REVIEW,
    gross_cents: int = 100_000,
) -> ContractMeasurement:
    record = ContractMeasurement(
        contract_id=contract.id,
        number=number,
        period=date(2026, 9, 1),
        gross_cents=gross_cents,
        situation=situation,
    )
    cenario.session.add(record)
    cenario.session.flush()
    return record


def _aditivo(
    cenario: EacScenario,
    contract: Contract,
    *,
    value_cents: int = 15_000_000,
    days: int = 21,
    change_id: int | None = None,
) -> ContractAddendum:
    record = ContractAddendum(
        contract_id=contract.id,
        change_id=change_id,
        number="AD-01",
        signed_on=date(2026, 5, 18),
        value_cents=value_cents,
        days=days,
        reason=JUSTIFICATIVA,
    )
    cenario.session.add(record)
    cenario.session.flush()
    return record


def _sm(
    cenario: EacScenario, *, code: str = "SM-TN-2026-0001", situation: str = "Aprovada"
) -> ChangeRequest:
    sm = ChangeRequest(
        project_id=cenario.first.id,
        requester_id=cenario.responsible.id,
        code=code,
        title="Mudança de teste",
        kind="Escopo",
        origin="Interna",
        priority="Normal",
        request_date=HOJE,
        description="Mudança criada no teste das rotas da ficha do contrato.",
        situation=situation,
    )
    cenario.session.add(sm)
    cenario.session.flush()
    return sm


# ── Leitura das abas ─────────────────────────────────────────────────────


def test_o_resumo_traz_os_kpis_e_a_cascata(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    _aditivo(cenario, contract)
    _medicao(
        cenario,
        contract,
        situation=calculations.SITUATION_APPROVED,
        gross_cents=20_000_000,
    )

    response = routes.contrato_resumo(
        _request(
            "/api/financeiro/contrato/resumo",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, RESUME_TARGET) is not None
    body = _body(response)
    assert 'data-grafico="cascata-contrato"' in body
    assert "Valor atual" in body
    assert "Saldo a faturar" in body
    assert "R$ 1.150.000,00" in body  # o valor atual
    assert "R$ 950.000,00" in body  # o saldo a faturar
    assert "1.1.1" in body  # o item da EAC


def test_as_medicoes_listam_o_fluxo_e_as_transicoes(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    _medicao(cenario, contract, number="BM-01")
    _medicao(cenario, contract, number="BM-02", situation=calculations.SITUATION_APPROVED)
    devolvida = ContractMeasurement(
        contract_id=contract.id,
        number="BM-03",
        period=date(2026, 9, 1),
        gross_cents=100_000,
        situation=calculations.SITUATION_RETURNED,
        return_reason="Falta a evidência do marco.",
    )
    cenario.session.add(devolvida)
    cenario.session.flush()

    response = routes.contrato_medicoes(
        _request(
            "/api/financeiro/contrato/medicoes",
            email=MANAGER_EMAIL,
            target=MEASUREMENTS_TARGET,
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), MEASUREMENTS_TARGET) is not None
    body = _body(response)
    assert [code for code in ("BM-01", "BM-02", "BM-03") if code in body] == [
        "BM-01",
        "BM-02",
        "BM-03",
    ]
    assert "Aprovar" in body
    assert "Faturar" in body
    assert "Devolver" in body
    assert "Falta a evidência do marco." in body


def test_os_aditivos_listam_valor_dias_e_a_sm(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    sm = _sm(cenario)
    _aditivo(cenario, contract, change_id=sm.id)

    response = routes.contrato_aditivos(
        _request(
            "/api/financeiro/contrato/aditivos",
            target=ADDENDA_TARGET,
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    assert find_by_id(_tags(response), ADDENDA_TARGET) is not None
    body = _body(response)
    assert "AD-01" in body
    assert "R$ 150.000,00" in body
    assert sm.code in body
    assert f"/governanca/mudanca?codigo={sm.code}" in body


@pytest.mark.usefixtures("cenario")
def test_sem_o_cabecalho_do_alpine_a_ficha_manda_para_o_shell() -> None:
    response = routes.contrato_resumo(
        _request("/api/financeiro/contrato/resumo", target=None, numero=NUMERO)
    )

    assert response.status_code == 302


@pytest.mark.usefixtures("cenario")
def test_contrato_desconhecido_e_recusado_com_422() -> None:
    response = routes.contrato_resumo(
        _request("/api/financeiro/contrato/resumo", numero="CT-9999-999")
    )

    assert response.status_code == 422
    assert "Contrato não encontrado." in _body(response)


# ── Medições: formulário, criação e fluxo ────────────────────────────────


def test_o_formulario_da_medicao_traz_o_saldo_e_os_campos(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_medicao_form(
        _request(
            "/api/financeiro/contrato/medicoes/nova",
            target=FORM_TARGET,
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, FORM_TARGET) is not None
    names = {tag.attrs.get("name") for tag in tags if tag.name in {"input", "textarea"}}
    assert {"periodo", "bruto", "marco"} <= names
    body = _body(response)
    assert f"/api/financeiro/contrato/medicoes?numero={NUMERO}" in body
    assert "disponível a medir R$ 1.000.000,00" in body


def test_registrar_medicao_devolve_a_ficha_e_fecha_o_formulario(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_medicao_create(
        _request(
            "/api/financeiro/contrato/medicoes",
            email=MANAGER_EMAIL,
            form={"periodo": "2026-09", "bruto": "1.000,00"},
            target=f"{FORM_TARGET} {MEASUREMENTS_TARGET} {RESUME_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, FORM_TARGET) is not None
    assert find_by_id(tags, MEASUREMENTS_TARGET) is not None
    assert find_by_id(tags, RESUME_TARGET) is not None
    assert _toast(response) == "Medição BM-01 registrada em análise."
    assert "BM-01" in _body(response)

    (record,) = cenario.session.scalars(select(ContractMeasurement)).all()
    assert (record.number, record.situation, record.gross_cents) == (
        "BM-01",
        calculations.SITUATION_UNDER_REVIEW,
        100_000,
    )
    assert record.period == date(2026, 9, 1)


def test_medicao_acima_do_saldo_volta_no_formulario_com_422(cenario: EacScenario) -> None:
    _criar_contrato(cenario, original_cents=100_000)

    response = routes.contrato_medicao_create(
        _request(
            "/api/financeiro/contrato/medicoes",
            email=MANAGER_EMAIL,
            form={"periodo": "2026-09", "bruto": "2.000,00"},
            target=f"{FORM_TARGET} {MEASUREMENTS_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), FORM_TARGET) is not None
    body = _body(response)
    assert "supera o saldo disponível" in body
    assert "2.000,00" in body
    assert cenario.session.scalars(select(ContractMeasurement)).all() == []


def test_aprovar_medicao_alimenta_a_eac_e_o_resumo(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    medicao = _medicao(cenario, contract)

    response = routes.contrato_medicao_approve(
        _request(
            f"/api/financeiro/contrato/medicoes/{medicao.id}/aprovar",
            email=MANAGER_EMAIL,
            route={"medicao_id": str(medicao.id)},
            form={"versao": str(medicao.version)},
            target=f"{MEASUREMENTS_TARGET} {RESUME_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    assert _toast(response) == "Medição BM-01: aprovada."
    body = _body(response)
    assert "Aprovada" in body
    assert "R$ 1.000,00" in body

    cenario.session.refresh(medicao)
    assert (medicao.situation, medicao.version) == (calculations.SITUATION_APPROVED, 2)
    item_id = cenario.items[f"TN-EAC-001:{EAC_LEAF}"]
    (fact,) = cenario.session.scalars(select(ErpCost).where(ErpCost.item_id == item_id)).all()
    assert (fact.month, fact.actual_cents) == (date(2026, 9, 1), 100_000)


def test_faturar_e_pagar_seguem_o_fluxo(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    medicao = _medicao(cenario, contract, situation=calculations.SITUATION_APPROVED)

    faturada = routes.contrato_medicao_invoice(
        _request(
            f"/api/financeiro/contrato/medicoes/{medicao.id}/faturar",
            email=MANAGER_EMAIL,
            route={"medicao_id": str(medicao.id)},
            form={"versao": str(medicao.version)},
            target=f"{MEASUREMENTS_TARGET} {RESUME_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert faturada.status_code == 200
    assert _toast(faturada) == "Medição BM-01: faturada."
    cenario.session.refresh(medicao)
    assert medicao.situation == calculations.SITUATION_INVOICED

    paga = routes.contrato_medicao_pay(
        _request(
            f"/api/financeiro/contrato/medicoes/{medicao.id}/pagar",
            email=MANAGER_EMAIL,
            route={"medicao_id": str(medicao.id)},
            form={"versao": str(medicao.version)},
            target=f"{MEASUREMENTS_TARGET} {RESUME_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert paga.status_code == 200
    assert _toast(paga) == "Medição BM-01: paga."
    cenario.session.refresh(medicao)
    assert medicao.situation == calculations.SITUATION_PAID


def test_transicao_fora_do_fluxo_volta_no_painel_com_422(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    medicao = _medicao(cenario, contract)

    response = routes.contrato_medicao_pay(
        _request(
            f"/api/financeiro/contrato/medicoes/{medicao.id}/pagar",
            email=MANAGER_EMAIL,
            route={"medicao_id": str(medicao.id)},
            form={"versao": str(medicao.version)},
            target=f"{MEASUREMENTS_TARGET} {RESUME_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), MEASUREMENTS_TARGET) is not None
    assert "não pode ir para" in _body(response)
    cenario.session.refresh(medicao)
    assert medicao.situation == calculations.SITUATION_UNDER_REVIEW


def test_devolver_exige_motivo_e_devolve_a_medicao(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    medicao = _medicao(cenario, contract)

    form_response = routes.contrato_medicao_return_form(
        _request(
            f"/api/financeiro/contrato/medicoes/{medicao.id}/devolucao",
            route={"medicao_id": str(medicao.id)},
            target=FORM_TARGET,
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert form_response.status_code == 200
    assert find_by_id(_tags(form_response), FORM_TARGET) is not None
    assert "Motivo da devolução" in _body(form_response)

    sem_motivo = routes.contrato_medicao_return(
        _request(
            f"/api/financeiro/contrato/medicoes/{medicao.id}/devolver",
            email=MANAGER_EMAIL,
            route={"medicao_id": str(medicao.id)},
            form={"versao": str(medicao.version), "motivo": "   "},
            target=f"{FORM_TARGET} {MEASUREMENTS_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert sem_motivo.status_code == 422
    assert find_by_id(_tags(sem_motivo), FORM_TARGET) is not None
    assert "motivo da devolução" in _body(sem_motivo)

    devolvida = routes.contrato_medicao_return(
        _request(
            f"/api/financeiro/contrato/medicoes/{medicao.id}/devolver",
            email=MANAGER_EMAIL,
            route={"medicao_id": str(medicao.id)},
            form={"versao": str(medicao.version), "motivo": "Falta a evidência do marco."},
            target=f"{FORM_TARGET} {MEASUREMENTS_TARGET} {RESUME_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert devolvida.status_code == 200
    assert _toast(devolvida) == "Medição BM-01 devolvida à contratada."
    cenario.session.refresh(medicao)
    assert (medicao.situation, medicao.return_reason) == (
        calculations.SITUATION_RETURNED,
        "Falta a evidência do marco.",
    )


def test_visualizador_nao_move_a_medicao(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    medicao = _medicao(cenario, contract)

    response = routes.contrato_medicao_approve(
        _request(
            f"/api/financeiro/contrato/medicoes/{medicao.id}/aprovar",
            email=VIEWER_EMAIL,
            route={"medicao_id": str(medicao.id)},
            form={"versao": str(medicao.version)},
            target=MEASUREMENTS_TARGET,
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 403


def test_o_portfolio_nao_cria_medicao(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_medicao_create(
        _request(
            "/api/financeiro/contrato/medicoes",
            email=MANAGER_EMAIL,
            form={"periodo": "2026-09", "bruto": "1.000,00"},
            target=FORM_TARGET,
            numero=NUMERO,
        )
    )

    assert response.status_code == 422
    assert "Escolha um projeto" in _body(response)


# ── Aditivos: formulário e criação ───────────────────────────────────────


def test_o_formulario_do_aditivo_lista_as_sm_aprovadas(cenario: EacScenario) -> None:
    _criar_contrato(cenario)
    sm = _sm(cenario)
    _sm(cenario, code="SM-TN-2026-0002", situation="Registrada")

    response = routes.contrato_aditivo_form(
        _request(
            "/api/financeiro/contrato/aditivos/novo",
            target=FORM_TARGET,
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, FORM_TARGET) is not None
    names = {tag.attrs.get("name") for tag in tags if tag.name in {"input", "select", "textarea"}}
    assert {"data", "valor", "dias", "sm", "motivo"} <= names
    body = _body(response)
    assert sm.code in body
    assert "SM-TN-2026-0002" not in body  # só as aprovadas entram no select


def test_registrar_aditivo_muda_o_valor_e_prorroga_o_termino(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    sm = _sm(cenario)

    response = routes.contrato_aditivo_create(
        _request(
            "/api/financeiro/contrato/aditivos",
            email=MANAGER_EMAIL,
            form={
                "data": "2026-05-18",
                "valor": "150.000,00",
                "dias": "21",
                "sm": str(sm.id),
                "motivo": JUSTIFICATIVA,
            },
            target=f"{FORM_TARGET} {ADDENDA_TARGET} {RESUME_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, FORM_TARGET) is not None
    assert find_by_id(tags, ADDENDA_TARGET) is not None
    assert find_by_id(tags, RESUME_TARGET) is not None
    assert _toast(response) == "Aditivo AD-01 registrado. Término vigente prorrogado em 21 dias."
    assert "AD-01" in _body(response)

    (record,) = cenario.session.scalars(select(ContractAddendum)).all()
    assert (record.change_id, record.value_cents, record.days) == (sm.id, 15_000_000, 21)
    cenario.session.refresh(contract)
    assert contract.current_end == date(2026, 11, 20)


def test_aditivo_sem_valor_nem_dias_volta_no_formulario_com_422(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_aditivo_create(
        _request(
            "/api/financeiro/contrato/aditivos",
            email=MANAGER_EMAIL,
            form={
                "data": "2026-05-18",
                "valor": "0,00",
                "dias": "0",
                "motivo": JUSTIFICATIVA,
            },
            target=f"{FORM_TARGET} {ADDENDA_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 422
    assert find_by_id(_tags(response), FORM_TARGET) is not None
    assert "valor, prorrogação ou ambos" in _body(response)
    assert cenario.session.scalars(select(ContractAddendum)).all() == []


# ── Excel e PDF ──────────────────────────────────────────────────────────


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> tuple[int, list[object]]:
    first_row = int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    width = sheet.max_column
    return first_row, [
        sheet.cell(row=first_row, column=column).value for column in range(1, width + 1)
    ]


def _ficha_completa(cenario: EacScenario) -> Contract:
    contract = _criar_contrato(cenario)
    _aditivo(cenario, contract)
    _medicao(
        cenario,
        contract,
        situation=calculations.SITUATION_PAID,
        gross_cents=10_000_000,
    )
    return contract


def test_o_excel_traz_o_resumo_a_cascata_e_as_tabelas(cenario: EacScenario) -> None:
    _ficha_completa(cenario)

    response = routes.contrato_excel(
        _request("/api/financeiro/contrato/excel", target=None, numero=NUMERO)
    )

    assert response.status_code == 200
    sheets = _sheets(response)
    assert list(sheets) == ["Resumo", "Composição do valor", "Medições", "Aditivos"]
    summary = sheets["Resumo"]
    labels = {summary.cell(row=row, column=1).value for row in range(1, summary.max_row + 1)}
    assert {"Valor atual", "Medido aprovado", "Saldo a faturar", "Término vigente"} <= labels

    cascade = sheets["Composição do valor"]
    first_row, header = _header_row(cascade)
    assert header == ["Etapa", "Valor"]
    assert cascade.cell(row=first_row + 1, column=1).value == "Valor original"
    assert cascade.cell(row=first_row + 1, column=2).value == 1_000_000.0
    assert cascade.cell(row=first_row + 3, column=2).value == 1_150_000.0
    assert cascade.cell(row=first_row + 5, column=2).value == 1_050_000.0

    measurements = sheets["Medições"]
    first_row, header = _header_row(measurements)
    assert header == [
        "Nº",
        "Período",
        "Valor bruto",
        "Retenção",
        "Valor líquido",
        "Situação",
        "Motivo da devolução",
    ]
    assert measurements.cell(row=first_row + 1, column=1).value == "BM-01"
    assert measurements.cell(row=first_row + 1, column=3).value == 100_000.0
    assert measurements.cell(row=first_row + 1, column=4).value == 5_000.0
    assert measurements.cell(row=first_row + 1, column=6).value == calculations.SITUATION_PAID

    addenda = sheets["Aditivos"]
    first_row, header = _header_row(addenda)
    assert header == ["Nº", "Data", "Valor", "Dias", "Motivo", "SM"]
    assert addenda.cell(row=first_row + 1, column=1).value == "AD-01"
    assert addenda.cell(row=first_row + 1, column=3).value == 150_000.0
    assert addenda.cell(row=first_row + 1, column=4).value == 21


def test_a_versao_imprimivel_traz_o_mesmo_conteudo(cenario: EacScenario) -> None:
    _ficha_completa(cenario)

    response = routes.contrato_printable(
        _request(
            "/api/financeiro/contrato/imprimivel",
            target="folha-impressao",
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    body = _body(response)
    assert f"Contrato {NUMERO}" in body
    assert "R$ 1.150.000,00" in body
    assert "Composição do valor" in body
    assert "Medições" in body
    assert "Aditivos" in body
    assert 'data-grafico="cascata-contrato"' in body


def test_gestor_tambem_registra_aditivo(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_aditivo_create(
        _request(
            "/api/financeiro/contrato/aditivos",
            email=MANAGER_EMAIL,
            form={
                "data": "2026-05-18",
                "valor": "150.000,00",
                "dias": "0",
                "motivo": JUSTIFICATIVA,
            },
            target=f"{ADDENDA_TARGET} {RESUME_TARGET}",
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200
    assert _toast(response) == "Aditivo AD-01 registrado."


def test_membro_nao_registra_aditivo(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_aditivo_create(
        _request(
            "/api/financeiro/contrato/aditivos",
            email=MEMBER_EMAIL,
            form={
                "data": "2026-05-18",
                "valor": "150.000,00",
                "dias": "0",
                "motivo": JUSTIFICATIVA,
            },
            target=FORM_TARGET,
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 403


def test_o_membro_le_a_ficha(cenario: EacScenario) -> None:
    _criar_contrato(cenario)

    response = routes.contrato_resumo(
        _request(
            "/api/financeiro/contrato/resumo",
            email=MEMBER_EMAIL,
            projeto=str(cenario.first.id),
            numero=NUMERO,
        )
    )

    assert response.status_code == 200


def test_o_servico_resolve_o_contrato_pelo_numero(cenario: EacScenario) -> None:
    contract = _criar_contrato(cenario)
    usuario = user_of(cenario.member, cenario.session, GeneralProfile.MEMBER)

    ficha = service.contract_by_number(cenario.session, user=usuario, number=NUMERO)

    assert ficha.id == contract.id
