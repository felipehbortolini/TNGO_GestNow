"""Rotas da Curva S financeira (ISSUE-042, HU-093, D12, D14).

Costura HTTP: o handler é chamado com a ``HttpRequest`` que o navegador montaria, como em
``test_contingencia_rotas.py``. A carga de demonstração é a base do escopo, com a âncora injetada
no relógio: o projeto 1 tem os números do mock (planejado de jan/26 = R$ 936.600,00, realizado de
ago/26 = R$ 22.800.000,00, comprometido de R$ 41.130.000,00 e a tendência chegando aos R$ 46,4 mi
do BAC ÷ CPI). O fragmento traz os indicadores, o gráfico das quatro séries (com o ``data-dados``
que a biblioteca lê) e a tabela período a período — na carteira, a soma em reais. O Excel e o
imprimível levam o mesmo conteúdo; o Visualizador lê e o fornecedor não.
"""

from __future__ import annotations

import json
from io import BytesIO

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.carga import DEMO_ANCHOR, run_demonstration
from src.core import auth, calendario
from src.modulos.configuracoes.models import Project
from src.modulos.financeiro import routes
from tests.financeiro.conftest import MEMBER_EMAIL, EacScenario
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal

ADMIN_EMAIL = "leonardo.gomes@exemplo.com"
VIEWER_EMAIL = "visualizador@cliente.demo"
SUPPLIER_EMAIL = "carlos.nunes@exemplo.com"
PROJECT_CODE = "TN-2026-014"
PATH = "/api/financeiro/curva-s"
CONTENT_TARGET = "curva-s-conteudo"
FOLHA_TARGET = "folha-impressao"

KPI_LABELS = (
    "Planejado acumulado",
    "Realizado acumulado",
    "Comprometido acumulado",
    "Projeção no término",
)
POINTS_HEADERS = (
    "Mês",
    "Planejado no mês",
    "Planejado acumulado",
    "Comprometido acumulado",
    "Realizado no mês",
    "Realizado acumulado",
    "Projeção acumulada",
    "Realizado menos planejado",
)
SERIES = ["planejado", "comprometido", "realizado", "tendencia"]


@pytest.fixture(autouse=True)
def data_de_referencia(monkeypatch: pytest.MonkeyPatch) -> None:
    """O relógio parado na âncora da demonstração: os números do mock não mudam."""
    monkeypatch.setattr(calendario, "today", lambda: DEMO_ANCHOR)


@pytest.fixture
def carga(rotas_na_transacao: Session) -> Session:
    """A base de demonstração na transação do teste, com as rotas lendo a mesma sessão."""
    run_demonstration(rotas_na_transacao, reference_date=DEMO_ANCHOR)
    return rotas_na_transacao


def _request(
    path: str,
    *,
    email: str = ADMIN_EMAIL,
    target: str | None = CONTENT_TARGET,
    **params: str,
) -> func.HttpRequest:
    """A request como o navegador a envia; ``target=None`` é um link puro (sem Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params={},
        body=b"",
    )


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _table_text(response: func.HttpResponse, attr: str) -> str:
    """O texto de dentro da tabela marcada com o atributo (``data-periodos``...)."""
    table = next(tag for tag in _tags(response) if tag.name == "table" and attr in tag.attrs)
    return table.clean_text()


def _projeto(session: Session) -> Project:
    return session.scalars(select(Project).where(Project.code == PROJECT_CODE)).one()


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> int:
    return int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))


# ── Leitura do fragmento ─────────────────────────────────────────────────


def test_a_curva_do_projeto_traz_kpis_grafico_e_tabela(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.curva_financeira_content(_request(PATH, projeto=str(project.id)))

    assert response.status_code == 200
    body = _body(response)
    assert find_by_id(_tags(response), CONTENT_TARGET) is not None
    assert 'data-situacao="conteudo"' in body
    assert 'data-carteira="nao"' in body
    for rotulo in KPI_LABELS:
        assert rotulo in body
    assert "R$ 29.302.200,00" in body, "o planejado acumulado do corte"
    assert "R$ 28.700.000,00" in body, "o realizado acumulado do corte"
    assert "R$ 41.130.000,00" in body, "o comprometido acumulado do corte"
    assert "R$ 46.440.129,45" in body, "a tendência chegando à EAC por desempenho"
    assert "Orçado: R$ 44.600.000,00" in body, "a referência de gestão"
    assert "92,2% do orçamento" in body
    assert "linha de base de custo em set/26" in body

    assert 'data-grafico="curva-s-linha"' in body
    (host,) = [tag for tag in _tags(response) if tag.attrs.get("data-grafico") == "curva-s-linha"]
    dados = json.loads(host.attrs["data-dados"])
    assert [serie["id"] for serie in dados["series"]] == SERIES
    assert [serie["papel"] for serie in dados["series"]] == [
        "linha-base",
        "comprometido",
        "realizado",
        "tendencia",
    ]
    periodos = dados["periodos"]
    assert len(periodos) == 16, "os meses do mock convertido"
    assert periodos[0]["mes"] == 1 and periodos[0]["valores"]["planejado"] == 936.6
    agosto = next(item for item in periodos if item["ano"] == 2026 and item["mes"] == 8)
    assert agosto["valores"]["realizado"] == 22800.0, "o realizado de 2026-08 do mock"
    assert agosto["valores"]["tendencia"] is None, "a tendência só existe a partir do corte"
    setembro = next(item for item in periodos if item["ano"] == 2026 and item["mes"] == 9)
    assert setembro["valores"]["tendencia"] is not None

    pontos = _table_text(response, "data-periodos")
    for coluna in POINTS_HEADERS:
        assert coluna in pontos
    assert "R$ 936.600,00" in pontos, "o planejado de jan/26"
    assert "R$ 22.800.000,00" in pontos, "o realizado acumulado de ago/26"
    assert "corte" in pontos


@pytest.mark.usefixtures("carga")
def test_a_curva_da_carteira_soma_os_projetos() -> None:
    response = routes.curva_financeira_content(_request(PATH, projeto="portfolio"))

    assert response.status_code == 200
    body = _body(response)
    assert 'data-carteira="sim"' in body
    assert "soma dos valores em R$ dos projetos (sem ponderação)" in body
    (host,) = [tag for tag in _tags(response) if tag.attrs.get("data-grafico") == "curva-s-linha"]
    dados = json.loads(host.attrs["data-dados"])
    setembro = next(item for item in dados["periodos"] if item["ano"] == 2026 and item["mes"] == 9)
    assert setembro["valores"]["planejado"] == 38387.2, "2.930.220.000 + 664.300.000 + 244.200.000"
    assert setembro["valores"]["realizado"] == 37170.0, "2.870.000.000 + 575.000.000 + 272.000.000"


@pytest.mark.usefixtures("carga")
def test_sem_o_cabecalho_do_alpine_a_curva_manda_para_o_shell() -> None:
    response = routes.curva_financeira_content(_request(PATH, target=None, projeto="1"))

    assert response.status_code == 302


def test_escopo_sem_curva_mostra_o_vazio_de_origem(cenario: EacScenario) -> None:
    response = routes.curva_financeira_content(
        _request(PATH, email=MEMBER_EMAIL, projeto=str(cenario.first.id))
    )

    assert response.status_code == 200
    body = _body(response)
    assert 'data-situacao="vazio-origem"' in body
    assert "Este projeto ainda não tem curva financeira cadastrada." in body
    assert "Planejado acumulado" not in body


# ── Permissões ───────────────────────────────────────────────────────────


def test_o_visualizador_le_a_curva(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.curva_financeira_content(
        _request(PATH, email=VIEWER_EMAIL, projeto=str(project.id))
    )

    assert response.status_code == 200
    assert "Planejado acumulado" in _body(response)


def test_o_fornecedor_nao_le_a_curva_nem_baixa_o_excel(carga: Session) -> None:
    project = _projeto(carga)

    conteudo = routes.curva_financeira_content(
        _request(PATH, email=SUPPLIER_EMAIL, projeto=str(project.id))
    )
    assert conteudo.status_code == 403

    excel = routes.curva_financeira_excel(
        _request(PATH + "/excel", email=SUPPLIER_EMAIL, target=None, projeto=str(project.id))
    )
    assert excel.status_code == 403


# ── Excel e imprimível ───────────────────────────────────────────────────


def test_o_excel_da_curva_traz_o_resumo_e_a_tabela(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.curva_financeira_excel(
        _request(PATH + "/excel", target=None, projeto=str(project.id))
    )

    assert response.status_code == 200
    assert response.headers["Content-Disposition"].startswith(
        'attachment; filename="curva-s-financeira-'
    )
    sheets = _sheets(response)
    assert list(sheets) == ["Resumo", "Período a período"]

    labels = {
        sheets["Resumo"].cell(row=row, column=1).value
        for row in range(1, sheets["Resumo"].max_row + 1)
    }
    assert set(KPI_LABELS) <= labels

    pontos = sheets["Período a período"]
    first_row = _header_row(pontos)
    assert [pontos.cell(row=first_row, column=column).value for column in range(1, 9)] == list(
        POINTS_HEADERS
    )
    assert pontos.cell(row=first_row + 1, column=3).value == 936_600.0, "R$ 936.600,00 em reais"


def test_o_imprimivel_da_curva_traz_o_mesmo_conteudo_do_excel(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.curva_financeira_printable(
        _request(PATH + "/imprimivel", target=FOLHA_TARGET, projeto=str(project.id))
    )

    assert response.status_code == 200
    body = _body(response)
    for texto in (
        "Curva S financeira",
        *KPI_LABELS,
        "R$ 29.302.200,00",
        "Período a período",
    ):
        assert texto in body
    assert 'data-grafico="curva-s-linha"' in body
