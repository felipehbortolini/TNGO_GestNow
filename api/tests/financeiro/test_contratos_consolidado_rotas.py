"""Rotas da visão consolidada da administração contratual (ISSUE-035, D12, D14).

Costura HTTP: o handler é chamado com a ``HttpRequest`` que o navegador montaria, como em
``test_eac_mapa_rotas.py``. A carga de demonstração é a base do escopo, com a âncora injetada no
relógio: o projeto 1 e o Portfólio têm os números do mock. O fragmento traz os KPIs do escopo
inteiro e as cinco abas filtradas; o Excel e o imprimível levam o mesmo conteúdo, e o nome do
arquivo é ``contratos-<código do projeto|portfolio>``; o Visualizador lê e o fornecedor não.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

import azure.functions as func
import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.carga import DEMO_ANCHOR, run_demonstration
from src.core import auth, calendario
from src.modulos.configuracoes.models import Project
from src.modulos.financeiro import routes
from tests.financeiro.conftest import MEMBER_EMAIL, EacScenario
from tests.html_tags import Tag, find_by_id, parse_tags
from tests.identidades import cabecalho_do_principal

ADMIN_EMAIL = "leonardo.gomes@exemplo.com"
VIEWER_EMAIL = "visualizador@cliente.demo"
SUPPLIER_EMAIL = "carlos.nunes@exemplo.com"
PROJECT_CODE = "TN-2026-014"
PATH = "/api/financeiro/contratos/consolidado"
CONTENT_TARGET = "contratos-conteudo"
KPIS_TARGET = "contratos-kpis"
TABELAS_TARGET = "contratos-tabelas"
FOLHA_TARGET = "folha-impressao"

KPI_LABELS = (
    "Contratos",
    "Exposição de claims",
    "Notificados fora do prazo",
    "Marcos atrasados",
    "Extensão de prazo",
    "Contratadas C ou D",
)
SHEETS = [
    "Resumo",
    "Contratos",
    "Claims",
    "Extensões de prazo",
    "Marcos de pagamento",
    "Avaliações",
]
CONTRACTS_HEADERS = [
    "Nº",
    "Contratada",
    "Objeto",
    "Modalidade",
    "Valor atual",
    "Medido",
    "Saldo a faturar",
    "Término vigente",
    "Dias aditados",
    "Claims abertos",
    "Última avaliação",
]
FRAGMENT_HEADERS = (
    "Nº",
    "Contratada",
    "Valor atual",
    "Medido",
    "Saldo a faturar",
    "Término vigente",
    "Claims abertos",
    "Última avaliação",
    "Direção / tipo",
    "Descrição",
    "Notificação",
    "Pleiteado",
    "Reconhecido",
    "Evento causador",
    "Dias solicitados",
    "Dias concedidos",
    "Classificação",
    "Marco afetado",
    "Marco",
    "Prevista",
    "Conclusão",
    "Período",
    "Classe",
    "Plano de melhoria",
    "Comentário",
    "Avaliador",
)


@pytest.fixture(autouse=True)
def data_de_referencia(monkeypatch: pytest.MonkeyPatch) -> None:
    """O relógio parado na âncora da demonstração: os números do mock não mudam."""
    monkeypatch.setattr(calendario, "today", lambda: DEMO_ANCHOR)


@pytest.fixture
def carga(rotas_na_transacao: Session) -> Session:
    """A base de demonstração na transação do teste, com as rotas lendo a mesma sessão."""
    run_demonstration(rotas_na_transacao, reference_date=DEMO_ANCHOR)
    return rotas_na_transacao


def _request(
    path: str,
    *,
    email: str = ADMIN_EMAIL,
    target: str | None = CONTENT_TARGET,
    **params: str,
) -> func.HttpRequest:
    """A request as the browser sends it; ``target=None`` is a plain link (no Alpine)."""
    headers = {auth.PRINCIPAL_HEADER: cabecalho_do_principal(email)}
    if target is not None:
        headers["X-Alpine-Request"] = "true"
        headers["X-Alpine-Target"] = target
    return func.HttpRequest(
        method="GET",
        url=path,
        headers=headers,
        params=dict(params),
        route_params={},
        body=b"",
    )


def _body(response: func.HttpResponse) -> str:
    return response.get_body().decode()


def _tags(response: func.HttpResponse) -> list[Tag]:
    return parse_tags(_body(response))


def _table_text(response: func.HttpResponse, attr: str) -> str:
    """O texto de dentro da tabela marcada com o atributo (``data-contratos``...)."""
    table = next(tag for tag in _tags(response) if tag.name == "table" and attr in tag.attrs)
    return table.clean_text()


def _cell_value(value: object) -> object:
    """O valor da célula, com a data de volta como ``date`` (o openpyxl a lê como ``datetime``)."""
    return value.date() if isinstance(value, datetime) else value


def _projeto(session: Session) -> Project:
    return session.scalars(select(Project).where(Project.code == PROJECT_CODE)).one()


def _sheets(response: func.HttpResponse) -> dict[str, Worksheet]:
    book = load_workbook(BytesIO(response.get_body()))
    return {sheet.title: sheet for sheet in book.worksheets}


def _header_row(sheet: Worksheet) -> tuple[int, list[object]]:
    first_row = int(sheet.auto_filter.ref.split(":")[0].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    width = sheet.max_column
    return first_row, [
        sheet.cell(row=first_row, column=column).value for column in range(1, width + 1)
    ]


# ── Leitura do fragmento ─────────────────────────────────────────────────


def test_o_consolidado_do_projeto_traz_kpis_abas_e_colunas(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contratos_consolidado(_request(PATH, projeto=str(project.id)))

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, CONTENT_TARGET) is not None
    body = _body(response)
    assert 'data-contratos="projeto"' in body
    for rotulo in KPI_LABELS:
        assert rotulo in body
    for titulo in SHEETS[1:]:
        assert titulo in body
    for coluna in FRAGMENT_HEADERS:
        assert coluna in body
    assert "CT-2026-011" in _table_text(response, "data-contratos")
    assert "CLM-TN-2026-0005" in _table_text(response, "data-claims")
    assert "EOT-TN-2026-0003" in _table_text(response, "data-eot")
    assert "5 dias de atraso" in _table_text(response, "data-marcos")
    assert "B (83)" in _table_text(response, "data-contratos")
    assert "R$ 605.000,00" in body  # exposição da contratada no projeto
    assert "2 claims abertos" in body
    assert "56 dias após o evento · fora do prazo" in _table_text(response, "data-claims")
    assert '<th scope="col">Projeto</th>' not in body


def test_o_consolidado_do_projeto_nao_tem_botoes_de_inclusao(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contratos_consolidado(_request(PATH, projeto=str(project.id)))

    body = _body(response)
    for botao in ("Novo contrato", "Novo claim", "Nova extensão", "Novo marco", "Nova avaliação"):
        assert botao not in body


@pytest.mark.usefixtures("carga")
def test_o_consolidado_da_carteira_traz_a_coluna_projeto_em_toda_aba() -> None:
    response = routes.contratos_consolidado(_request(PATH, projeto="portfolio"))

    assert response.status_code == 200
    body = _body(response)
    assert 'data-contratos="carteira"' in body
    assert body.count('<th scope="col">Projeto</th>') == 5
    assert "TN-2026-014 · " in _table_text(response, "data-contratos")
    assert "CT-2026-020" in _table_text(response, "data-contratos")


def test_os_filtros_do_consolidado_estreitam_as_abas(carga: Session) -> None:
    project = _projeto(carga)

    por_contrato = routes.contratos_consolidado(
        _request(PATH, projeto=str(project.id), contrato="CT-2026-011")
    )
    contratos = _table_text(por_contrato, "data-contratos")
    assert "CT-2026-011" in contratos
    assert "CT-2026-012" not in contratos
    assert "CT-2026-014" not in contratos
    assert "R$ 605.000,00" in _body(por_contrato)  # os KPIs seguem do escopo inteiro

    por_busca = routes.contratos_consolidado(
        _request(PATH, projeto=str(project.id), busca="clm-tn-2026-0003")
    )
    claims = _table_text(por_busca, "data-claims")
    assert "CLM-TN-2026-0003" in claims
    assert "CLM-TN-2026-0001" not in claims
    assert "Nenhum contrato encontrado" in _table_text(por_busca, "data-contratos")


def test_o_formulario_de_filtros_recebe_os_dois_blocos(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contratos_consolidado(
        _request(
            PATH,
            target=f"{KPIS_TARGET} {TABELAS_TARGET}",
            projeto=str(project.id),
            busca="clm-tn-2026-0003",
        )
    )

    assert response.status_code == 200
    tags = _tags(response)
    assert find_by_id(tags, KPIS_TARGET) is not None
    assert find_by_id(tags, TABELAS_TARGET) is not None
    body = _body(response)
    assert "<form" not in body  # o formulário fica fora dos blocos substituídos
    assert "CLM-TN-2026-0003" in body
    assert "CLM-TN-2026-0001" not in body


@pytest.mark.usefixtures("carga")
def test_sem_o_cabecalho_do_alpine_o_consolidado_manda_para_o_shell() -> None:
    response = routes.contratos_consolidado(_request(PATH, target=None, projeto="1"))

    assert response.status_code == 302


def test_escopo_sem_contratos_mostra_o_estado_de_origem(cenario: EacScenario) -> None:
    response = routes.contratos_consolidado(
        _request(PATH, email=MEMBER_EMAIL, projeto=str(cenario.first.id))
    )

    assert response.status_code == 200
    body = _body(response)
    assert 'data-estado="vazio-origem"' in body
    assert "Nenhum contrato no escopo" in body
    assert "Exposição de claims" not in body


# ── Permissões ───────────────────────────────────────────────────────────


def test_o_visualizador_le_o_consolidado(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contratos_consolidado(
        _request(PATH, email=VIEWER_EMAIL, projeto=str(project.id))
    )

    assert response.status_code == 200
    assert "Contratos" in _body(response)


def test_o_fornecedor_nao_le_o_consolidado_nem_baixa_o_excel(carga: Session) -> None:
    project = _projeto(carga)

    conteudo = routes.contratos_consolidado(
        _request(PATH, email=SUPPLIER_EMAIL, projeto=str(project.id))
    )
    assert conteudo.status_code == 403

    excel = routes.contratos_consolidado_excel(
        _request(PATH + "/excel", email=SUPPLIER_EMAIL, target=None, projeto=str(project.id))
    )
    assert excel.status_code == 403


# ── Excel e imprimível ───────────────────────────────────────────────────


def test_o_excel_do_consolidado_traz_o_resumo_e_as_cinco_tabelas(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contratos_consolidado_excel(
        _request(PATH + "/excel", target=None, projeto=str(project.id))
    )

    assert response.status_code == 200
    assert response.headers["Content-Disposition"].startswith(
        'attachment; filename="contratos-TN-2026-014.xlsx"'
    )
    sheets = _sheets(response)
    assert list(sheets) == SHEETS
    labels = {
        sheets["Resumo"].cell(row=row, column=1).value
        for row in range(1, sheets["Resumo"].max_row + 1)
    }
    assert set(KPI_LABELS) <= labels
    contratos = sheets["Contratos"]
    first_row, header = _header_row(contratos)
    assert header == CONTRACTS_HEADERS
    assert contratos.cell(row=first_row + 1, column=1).value == "CT-2026-011"
    assert contratos.cell(row=first_row + 1, column=5).value == 3_450_000.0
    assert _cell_value(contratos.cell(row=first_row + 1, column=8).value) == date(2026, 11, 20)


def test_o_excel_do_consolidado_respeita_os_filtros(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contratos_consolidado_excel(
        _request(PATH + "/excel", target=None, projeto=str(project.id), contrato="CT-2026-011")
    )

    sheet = _sheets(response)["Contratos"]
    first_row, _header = _header_row(sheet)
    assert sheet.cell(row=first_row + 1, column=1).value == "CT-2026-011"
    assert sheet.cell(row=first_row + 2, column=1).value in (None, "")
    assert "CT-2026-012" not in {
        sheet.cell(row=row, column=1).value for row in range(1, sheet.max_row + 1)
    }
    assert "Filtro: contrato: CT-2026-011" in {
        sheet.cell(row=row, column=1).value for row in range(1, sheet.max_row + 1)
    }


@pytest.mark.usefixtures("carga")
def test_o_excel_da_carteira_abre_com_a_coluna_projeto() -> None:
    response = routes.contratos_consolidado_excel(
        _request(PATH + "/excel", target=None, projeto="portfolio")
    )

    assert response.headers["Content-Disposition"].startswith(
        'attachment; filename="contratos-portfolio.xlsx"'
    )
    sheets = _sheets(response)
    first_row, header = _header_row(sheets["Contratos"])
    assert header[:2] == ["Projeto", "Nº"]
    projects = [
        sheets["Contratos"].cell(row=first_row + offset, column=1).value for offset in range(1, 10)
    ]
    assert all(projects)
    assert any(str(value).startswith("TN-2026-014 · ") for value in projects)


def test_o_imprimivel_do_consolidado_traz_o_mesmo_conteudo_do_excel(carga: Session) -> None:
    project = _projeto(carga)

    response = routes.contratos_consolidado_printable(
        _request(PATH + "/imprimivel", target=FOLHA_TARGET, projeto=str(project.id))
    )

    assert response.status_code == 200
    body = _body(response)
    for texto in (
        "Contratos: administração contratual",
        *KPI_LABELS,
        "Extensões de prazo",
        "Marcos de pagamento",
        "R$ 605.000,00",
    ):
        assert texto in body


@pytest.mark.usefixtures("carga")
def test_o_imprimivel_da_carteira_traz_a_coluna_projeto() -> None:
    response = routes.contratos_consolidado_printable(
        _request(PATH + "/imprimivel", target=FOLHA_TARGET, projeto="portfolio")
    )

    body = _body(response)
    assert ">Projeto<" in body
    assert "CT-2026-020" in body
