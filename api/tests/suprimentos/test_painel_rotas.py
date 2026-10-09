"""As rotas do painel de suprimentos (ISSUE-063): a tela, o Excel e o imprimível.

A rota roda na transação do teste (fixture ``sessao``), no modo demonstração e com o relógio em
25/09/2026. A requisição é montada como o shell a envia: cabeçalho do Alpine AJAX, o principal do
colaborador e o escopo no parâmetro ``projeto``. O cenário é o do diligenciamento: o PC-01 tem
pedido emitido e o PC-02 é planejado, sem pedido.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from src.modulos.suprimentos import routes
from tests.apoio_diligenciamento import (
    CenarioDiligenciamento,
    montar_cenario_diligenciamento,
)
from tests.identidades import requisicao

ANA = "ana.admin@plano.example.invalid"
VITOR = "vitor.visualizador@plano.example.invalid"
CABECALHO_AUSENTE = "cabeçalho não encontrado: {title}"


@pytest.fixture
def cenario(sessao: Session) -> CenarioDiligenciamento:
    return montar_cenario_diligenciamento(sessao)


def _params(cenario: CenarioDiligenciamento) -> dict[str, str]:
    return {"projeto": str(cenario.projeto.id)}


def _header_row(sheet: Worksheet, title: str) -> int:
    for row in sheet.iter_rows(min_row=1, max_row=20):
        if any(cell.value == title for cell in row):
            return row[0].row
    raise AssertionError(CABECALHO_AUSENTE.format(title=title))


# ── A tela ──────────────────────────────────────────────────────────────────────────────────


def test_a_rota_da_tela_responde_o_fragmento_do_painel(
    cenario: CenarioDiligenciamento,
) -> None:
    resposta = routes.supply_panel_view(
        requisicao(
            "/api/suprimentos/painel",
            email=ANA,
            params=_params(cenario),
            alvo="painel-conteudo",
        )
    )

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert 'id="painel-conteudo"' in corpo
    assert "Aderência ao plano de compras" in corpo
    assert "Saving sobre estimativa" in corpo
    assert "Saving de negociação" in corpo
    assert "Entrega no prazo (OTD)" in corpo
    assert "Pedidos críticos" in corpo
    assert 'data-grafico="curva-s-linha"' in corpo, "a contratação e o avanço do MAS"
    assert 'data-grafico="linhas-multiplas"' in corpo, "o saving acumulado"
    assert 'data-grafico="kpi-status"' in corpo, "os pacotes por etapa"
    assert "Pedidos por folga em relação ao ROS" in corpo
    assert "PC-01" in corpo
    assert "/suprimentos/plano-de-compras" in corpo, "o KPI da aderência abre o plano"


def test_a_rota_da_tela_no_portfolio_traz_a_coluna_projeto(
    cenario: CenarioDiligenciamento,
) -> None:
    del cenario
    resposta = routes.supply_panel_view(
        requisicao("/api/suprimentos/painel", email=ANA, alvo="painel-conteudo")
    )

    assert resposta.status_code == 200
    assert '<th scope="col">Projeto</th>' in resposta.get_body().decode()


def test_o_visualizador_le_o_painel(cenario: CenarioDiligenciamento) -> None:
    resposta = routes.supply_panel_view(
        requisicao(
            "/api/suprimentos/painel",
            email=VITOR,
            params=_params(cenario),
            alvo="painel-conteudo",
        )
    )

    assert resposta.status_code == 200


# ── O Excel e o imprimível ──────────────────────────────────────────────────────────────────


def test_o_excel_do_painel_traz_o_resumo_e_a_tabela(cenario: CenarioDiligenciamento) -> None:
    resposta = routes.supply_panel_excel(
        requisicao("/api/suprimentos/painel/excel", email=ANA, params=_params(cenario))
    )

    assert resposta.status_code == 200
    assert resposta.headers.get("Content-Type", "").startswith("application/vnd.openxmlformats")
    pasta = load_workbook(BytesIO(resposta.get_body()))
    resumo = pasta["Resumo"]
    indicadores = [cell.value for row in resumo.iter_rows() for cell in row]
    assert "Aderência ao plano de compras" in indicadores
    assert "Saving sobre estimativa" in indicadores
    aba = pasta["Pedidos por folga"]
    header = _header_row(aba, "Pacote")
    colunas = [cell.value for cell in aba[header]]
    assert colunas[:4] == ["Pacote", "Escopo", "Fornecedor", "LLI"]
    assert "Folga (dias)" in colunas
    assert "Situação" in colunas
    assert aba.page_setup.orientation == "landscape"


def test_o_excel_do_painel_no_portfolio_traz_a_coluna_projeto(
    cenario: CenarioDiligenciamento,
) -> None:
    del cenario
    resposta = routes.supply_panel_excel(requisicao("/api/suprimentos/painel/excel", email=ANA))

    assert resposta.status_code == 200
    pasta = load_workbook(BytesIO(resposta.get_body()))
    aba = pasta["Pedidos por folga"]
    header = _header_row(aba, "Projeto")
    colunas = [cell.value for cell in aba[header]]
    assert colunas[:2] == ["Projeto", "Pacote"]


def test_o_imprimivel_sai_em_a4_com_os_graficos(cenario: CenarioDiligenciamento) -> None:
    resposta = routes.supply_panel_printable(
        requisicao(
            "/api/suprimentos/painel/imprimivel",
            email=ANA,
            params=_params(cenario),
            alvo="folha-impressao",
        )
    )

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert "folha--a4" in corpo
    assert "Painel de suprimentos" in corpo
    assert "Aderência ao plano de compras" in corpo
    assert 'data-grafico="curva-s-linha"' in corpo
    assert "Pedidos por folga" in corpo
