"""As rotas do MAS (ISSUE-062): a tela, o detalhe no modal, o Excel e o imprimível A3.

A rota roda na transação do teste (fixture ``sessao``), no modo demonstração e com o relógio em
25/09/2026. A requisição é montada como o shell a envia: cabeçalho do Alpine AJAX, o principal do
colaborador e o escopo no parâmetro ``projeto``. O cenário é o do diligenciamento: o PC-01 tem
pedido emitido e o PC-02 é planejado, sem pedido.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from src.modulos.suprimentos import routes
from src.modulos.suprimentos.models import PackageMilestone
from tests.apoio_diligenciamento import (
    CenarioDiligenciamento,
    montar_cenario_diligenciamento,
)
from tests.identidades import requisicao

ANA = "ana.admin@plano.example.invalid"
VITOR = "vitor.visualizador@plano.example.invalid"
CABECALHO_AUSENTE = "cabeçalho não encontrado: {title}"


@pytest.fixture
def cenario(sessao: Session) -> CenarioDiligenciamento:
    return montar_cenario_diligenciamento(sessao)


def _params(cenario: CenarioDiligenciamento) -> dict[str, str]:
    return {"projeto": str(cenario.projeto.id)}


def _atrasar(sessao: Session, cenario: CenarioDiligenciamento) -> None:
    """A RFx do PC-01 realizada depois da LB: a célula do marco sai com atraso."""
    row = sessao.scalars(
        select(PackageMilestone).where(
            PackageMilestone.package_id == cenario.pacote.id,
            PackageMilestone.milestone == "rfx",
        )
    ).one()
    row.real_date = date(2026, 8, 20)
    sessao.flush()


def _header_row(sheet: Worksheet, title: str) -> int:
    for row in sheet.iter_rows(min_row=1, max_row=20):
        if any(cell.value == title for cell in row):
            return row[0].row
    raise AssertionError(CABECALHO_AUSENTE.format(title=title))


# ── A tela ──────────────────────────────────────────────────────────────────────────────────


def test_a_rota_da_tela_responde_o_fragmento_do_mapa(
    cenario: CenarioDiligenciamento,
) -> None:
    resposta = routes.mas_list(
        requisicao(
            "/api/suprimentos/mas",
            email=ANA,
            params=_params(cenario),
            alvo="mas-conteudo",
        )
    )

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert 'id="mas-conteudo"' in corpo
    assert "PC-01" in corpo
    assert "Aquisição" in corpo
    assert "Fabricação e entrega" in corpo
    assert "Documentos aprovados" in corpo
    assert "Realizados / previstos na LB até hoje" in corpo
    assert "mas__celula--estimada" in corpo, "o PC-02 sem pedido leva a LB estimada em itálico"


def test_a_rota_da_tela_no_portfolio_traz_a_coluna_projeto(
    cenario: CenarioDiligenciamento,
) -> None:
    del cenario
    resposta = routes.mas_list(requisicao("/api/suprimentos/mas", email=ANA, alvo="mas-conteudo"))

    assert resposta.status_code == 200
    assert "Projeto" in resposta.get_body().decode()


def test_o_modo_desvio_troca_o_conteudo_das_celulas(
    cenario: CenarioDiligenciamento,
) -> None:
    resposta = routes.mas_list(
        requisicao(
            "/api/suprimentos/mas",
            email=ANA,
            params={**_params(cenario), "modo": "desvio"},
            alvo="mas-conteudo",
        )
    )

    corpo = resposta.get_body().decode()
    assert "desvio em dias em relação à LB" in corpo
    assert ">vencido<" in corpo, "os marcos do plano do PC-01 já passaram da LB"


def test_a_fase_corta_os_grupos_da_tabela(cenario: CenarioDiligenciamento) -> None:
    resposta = routes.mas_list(
        requisicao(
            "/api/suprimentos/mas",
            email=ANA,
            params={**_params(cenario), "fase": "fabricacao"},
            alvo="mas-conteudo",
        )
    )

    corpo = resposta.get_body().decode()
    assert "mas__grupo--fabricacao" in corpo
    assert "mas__grupo--aquisicao" not in corpo


def test_o_visualizador_le_o_mapa(cenario: CenarioDiligenciamento) -> None:
    resposta = routes.mas_list(
        requisicao(
            "/api/suprimentos/mas",
            email=VITOR,
            params=_params(cenario),
            alvo="mas-conteudo",
        )
    )

    assert resposta.status_code == 200


# ── O detalhe do pacote ─────────────────────────────────────────────────────────────────────


def _detalhe(cenario: CenarioDiligenciamento, pacote_id: int, *, email: str = ANA):
    return routes.mas_package_detail(
        requisicao(
            f"/api/suprimentos/mas/pacote/{pacote_id}",
            email=email,
            params=_params(cenario),
            rota={"pacote_id": str(pacote_id)},
            alvo="mas-modal-corpo",
        )
    )


def test_o_detalhe_do_pacote_com_pedido_traz_os_doze_marcos(
    cenario: CenarioDiligenciamento,
) -> None:
    resposta = _detalhe(cenario, cenario.pacote.id)

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert 'id="mas-modal-corpo"' in corpo
    assert "Documentos aprovados" in corpo
    assert "Entrega" in corpo
    assert "PED-2026-0001" in corpo
    assert "(LB estimada)" not in corpo


def test_o_detalhe_do_pacote_sem_pedido_avisa_a_lb_estimada(
    cenario: CenarioDiligenciamento,
) -> None:
    resposta = _detalhe(cenario, cenario.planejado.id)

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert "(LB estimada)" in corpo


def test_o_pacote_fora_do_escopo_nao_abre(cenario: CenarioDiligenciamento) -> None:
    resposta = routes.mas_package_detail(
        requisicao(
            f"/api/suprimentos/mas/pacote/{cenario.pacote.id}",
            email=ANA,
            params={"projeto": str(cenario.outro.id)},
            rota={"pacote_id": str(cenario.pacote.id)},
            alvo="mas-modal-corpo",
        )
    )

    assert resposta.status_code == 422


# ── O Excel e o imprimível ──────────────────────────────────────────────────────────────────


def test_o_excel_do_mapa_tem_as_colunas_do_prototipo_em_a3(
    cenario: CenarioDiligenciamento,
) -> None:
    resposta = routes.mas_excel(
        requisicao("/api/suprimentos/mas/excel", email=ANA, params=_params(cenario))
    )

    assert resposta.status_code == 200
    assert resposta.headers.get("Content-Type", "").startswith("application/vnd.openxmlformats")
    pasta = load_workbook(BytesIO(resposta.get_body()))
    aba = pasta["Mapa de Suprimentos"]
    header = _header_row(aba, "Pacote")
    colunas = [cell.value for cell in aba[header]]
    assert colunas[:3] == ["Pacote", "Escopo", "Tipo"]
    assert "Documentos aprovados LB" in colunas
    assert "Documentos aprovados previsão ou real" in colunas
    assert "Documentos aprovados situação" in colunas
    assert "Avanço real (%)" in colunas
    assert aba.page_setup.orientation == "landscape"
    assert aba.page_setup.paperSize == 8, "A3"


def test_o_imprimivel_sai_em_a3_com_as_celulas_coloridas(
    sessao: Session, cenario: CenarioDiligenciamento
) -> None:
    _atrasar(sessao, cenario)
    resposta = routes.mas_printable(
        requisicao(
            "/api/suprimentos/mas/imprimivel",
            email=ANA,
            params=_params(cenario),
            alvo="folha-impressao",
        )
    )

    assert resposta.status_code == 200
    corpo = resposta.get_body().decode()
    assert "folha--a3" in corpo
    assert "folha__td--erro" in corpo
    assert "com atraso +10 d" in corpo
    assert "Mapa de Suprimentos (MAS)" in corpo
