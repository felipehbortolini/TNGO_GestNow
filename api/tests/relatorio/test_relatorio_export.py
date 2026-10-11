"""O Excel do relatório gerencial e o contrato da impressão (ISSUE-086).

O download chama a rota de arquivo como o link do navegador a chama — a resposta é o
arquivo, não um fragmento — e abre o ``.xlsx`` com openpyxl para conferir as abas: o
Resumo com os indicadores de todas as folhas, uma aba por tabela exibida e as duas abas
de fechamento do protótipo (Análises do período e Desvios negativos e comentários).
Tipo ou período inválidos recusam em texto simples com 422. A impressão do relatório é
a do navegador, pela tela (ISSUE-083): o teste guarda o contrato — o botão nas folhas e
o CSS da página — lendo os arquivos do app.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import azure.functions as func
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from src.core.auth import DEMO_COOKIE
from src.modulos.planejamento.validation import WEEKLY
from src.modulos.relatorio import routes, validation
from tests.identidades import requisicao
from tests.relatorio.conftest import Demonstracao

RAIZ_DO_REPOSITORIO = Path(__file__).resolve().parents[3]

VIEW = RAIZ_DO_REPOSITORIO / "app" / "_views" / "relatorio" / "gerencial.html"
FOLHAS = RAIZ_DO_REPOSITORIO / "api" / "src" / "templates" / "relatorio" / "folhas.html"
CSS = RAIZ_DO_REPOSITORIO / "app" / "paginas" / "relatorio" / "gerencial.css"

PERIODO = "2026-S38"
CABECALHO_AUSENTE = "cabeçalho não encontrado: {title}"


def _excel(
    demonstracao: Demonstracao, *, projeto_id: int | None, params: dict[str, str]
) -> func.HttpResponse:
    """O download como o navegador o pede: o pedido do link, com o cookie do seletor de perfil."""
    parametros = dict(params)
    if projeto_id is not None:
        parametros["projeto"] = str(projeto_id)
    return routes.report_sheets_excel(
        requisicao(
            "/api/relatorio/gerencial/excel",
            cookies={DEMO_COOKIE: str(demonstracao.admin.id)},
            params=parametros,
        )
    )


def _header_row(sheet: Worksheet, title: str) -> int:
    for row in sheet.iter_rows(min_row=1, max_row=30):
        if any(cell.value == title for cell in row):
            return row[0].row
    raise AssertionError(CABECALHO_AUSENTE.format(title=title))


def test_o_excel_do_portfolio_tem_o_resumo_as_tabelas_e_as_analises(
    demonstracao: Demonstracao, sessao_das_rotas: Session
) -> None:
    del sessao_das_rotas
    resposta = _excel(demonstracao, projeto_id=None, params={"tipo": WEEKLY, "periodo": PERIODO})

    assert resposta.status_code == 200
    assert resposta.headers.get("Content-Type", "").startswith("application/vnd.openxmlformats")
    pasta = load_workbook(BytesIO(resposta.get_body()))
    assert set(pasta.sheetnames) >= {
        "Resumo",
        "Projetos da carteira",
        "Critérios de ponderação",
        "Avanço por projeto",
        "Relatos dos projetos",
        "Principais riscos",
        "Análises do período",
        "Desvios negativos e comentários",
    }

    resumo = pasta["Resumo"]
    indicadores = [cell.value for row in resumo.iter_rows() for cell in row]
    assert "Carteira · Projetos na carteira" in indicadores
    assert "Planejamento · Avanço real (acumulado)" in indicadores
    assert any(str(value).startswith("Riscos · ") for value in indicadores if value)

    carteira = pasta["Projetos da carteira"]
    header = _header_row(carteira, "Projeto")
    assert [cell.value for cell in carteira[header]][:3] == ["Projeto", "Peso", "BAC"]
    assert carteira.page_setup.orientation == "landscape"


def test_o_excel_recusa_tipo_ou_periodo_invalido_em_texto_simples(
    demonstracao: Demonstracao, sessao_das_rotas: Session
) -> None:
    del sessao_das_rotas
    resposta = _excel(demonstracao, projeto_id=demonstracao.project_id, params={})

    assert resposta.status_code == 422
    assert resposta.headers.get("Content-Type", "").startswith("text/plain")
    assert validation.CHOOSE_MESSAGE in resposta.get_body().decode()


def test_o_contrato_da_impressao_esta_na_tela_na_barra_e_no_estilo() -> None:
    view = VIEW.read_text(encoding="utf-8")
    folhas = FOLHAS.read_text(encoding="utf-8")
    css = CSS.read_text(encoding="utf-8")

    assert "data-relatorio-imprimir" in view
    assert "data-tn-excel" in folhas
    assert "data-relatorio-alterar" in folhas
    assert "/api/relatorio/gerencial/excel?tipo=" in folhas
    assert "@media print" in css
    assert "break-after: page" in css
    assert "print-color-adjust" in css
    assert "@page relatorio-gerencial" in css
